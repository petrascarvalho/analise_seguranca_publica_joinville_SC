from __future__ import annotations

import json
import hashlib
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scripts.transform import CITY_CRIME_URL, CITY_CRIME_VALUES, YEARS
from scripts.utils import project_path


SINESP_URL = (
    "https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/"
    "estatistica/dados-nacionais-1/base-de-dados-e-notas-metodologicas-"
    "dos-gestores-estaduais-sinesp-vde-2022-e-2023"
)
SSP_URL = "https://ssp.sc.gov.br/segurancaemnumeros/"
SSP_HOMICIDE = "Nº de vítimas de homicídio"
ANALYSIS_DIR = project_path("dados", "processados", "analises")
TABLE_DIR = project_path("documentacao", "analises", "tabelas")
CHART_DIR = project_path("documentacao", "analises", "graficos")
SINESP_EXPECTED_CITY_ROWS = {2020: 132, 2021: 132, 2022: 132, 2023: 132, 2024: 132, 2025: 144, 2026: 96}
SSP_ANNUAL_CONTROL = {2020: 66, 2021: 62, 2022: 49, 2023: 56, 2024: 37, 2025: 32}
SSP_PARTIAL_CONTROL = {2023: 36, 2024: 20, 2025: 17, 2026: 22}


def csv_quality(path: Path) -> dict:
    frame = pd.read_csv(path, encoding="utf-8-sig", dtype="string", keep_default_na=False)
    return {
        "arquivo": str(path.relative_to(project_path())),
        "tipo": "CSV",
        "registros": len(frame),
        "campos": len(frame.columns),
        "municipios": frame["municipio"].nunique() if "municipio" in frame else None,
        "campos_com_vazios": int(frame.eq("").any().sum()),
        "status": "disponível",
    }


def xlsx_quality(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        fields = list(next(rows, ()))
        count = 0
        cities = set()
        for row in rows:
            count += 1
            if len(row) > 1 and row[1]:
                cities.add(row[1])
    finally:
        workbook.close()
    return {
        "arquivo": str(path.relative_to(project_path())),
        "tipo": "XLSX tratado",
        "registros": count,
        "campos": len(fields),
        "municipios": len(cities),
        "campos_com_vazios": "não calculado",
        "status": "disponível",
    }


def quality_inventory() -> list[dict]:
    results = []
    for path in sorted(project_path("dados", "processados", "sinesp").glob("bancovde-*-joinville-sc.xlsx")):
        results.append(xlsx_quality(path))
    for directory in (
        project_path("dados", "processados", "sinesp"),
        project_path("dados", "processados", "ssp_sc"),
        project_path("dados", "processados", "prefeitura"),
        ANALYSIS_DIR,
    ):
        for path in sorted(directory.glob("*.csv")):
            results.append(csv_quality(path))
    return results


def _read_sinesp() -> pd.DataFrame:
    frames = []
    for year in YEARS:
        path = project_path("dados", "processados", "sinesp", f"bancovde-{year}-joinville-sc.xlsx")
        if not path.exists():
            raise FileNotFoundError(f"Recorte Sinesp municipal ausente: {path}")
        frame = pd.read_excel(path, dtype={"uf": "string", "municipio": "string"})
        frame["data_referencia"] = pd.to_datetime(frame["data_referencia"], errors="raise")
        if not frame["data_referencia"].dt.year.eq(year).all():
            raise ValueError(f"data_referencia fora do ano no arquivo {path.name}")
        frames.append(frame)
    return pd.concat(frames, ignore_index=True)


def _read_ssp() -> pd.DataFrame:
    paths = sorted(project_path("dados", "processados", "ssp_sc").glob("ssp_joinville_*.csv"))
    if len(paths) != 4:
        raise ValueError(f"Esperados quatro CSVs SSP/SC, encontrados {len(paths)}")
    frames = [pd.read_csv(path, encoding="utf-8-sig") for path in paths]
    frame = pd.concat(frames, ignore_index=True)
    frame["valor"] = pd.to_numeric(frame["valor"], errors="raise").astype("int64")
    frame["ano_referencia"] = pd.to_numeric(frame["ano_referencia"], errors="raise").astype("int64")
    frame["ano_relatorio"] = pd.to_numeric(frame["ano_relatorio"], errors="raise").astype("int64")
    frame["periodo_referencia"] = frame["periodo_referencia"].astype("string")
    frame["indicador_normalizado"] = frame["indicador_normalizado"].fillna(frame["indicador"])
    for column in ("pagina_pdf", "arquivo_pdf", "indicador", "indicador_normalizado", "periodo_tipo", "municipio", "data_atualizacao"):
        frame[column] = frame[column].astype("string").fillna("")
    return frame


def _read_city_crime() -> pd.DataFrame:
    paths = sorted(project_path("dados", "processados", "prefeitura").glob("furto_roubo_joinville_*.csv"))
    if len(paths) != 3:
        raise ValueError(f"Esperados três CSVs da Prefeitura, encontrados {len(paths)}")
    frame = pd.concat([pd.read_csv(p, encoding="utf-8-sig") for p in paths], ignore_index=True)
    frame["ano"] = pd.to_numeric(frame["ano"], errors="raise").astype("int64")
    frame["quantidade"] = pd.to_numeric(frame["quantidade"], errors="raise").astype("int64")
    if set(frame["url_fonte"]) != {CITY_CRIME_URL}:
        raise ValueError("A URL de origem da série municipal de furto/roubo divergiu")
    found = {
        (int(row.ano), row.indicador): int(row.quantidade)
        for row in frame.itertuples(index=False)
    }
    expected = {
        (year, indicator): value
        for year, indicators in CITY_CRIME_VALUES.items()
        for indicator, value in indicators.items()
    }
    if found != expected:
        raise ValueError("Totais de furto/roubo não conferem com a fonte da Prefeitura")
    return frame


def _write_frame(frame: pd.DataFrame, stem: str, source_fields: list[str]) -> dict:
    csv_path = ANALYSIS_DIR / f"{stem}.csv"
    parquet_path = ANALYSIS_DIR / f"{stem}.parquet"
    frame.to_csv(csv_path, index=False, encoding="utf-8-sig", date_format="%Y-%m-%d")
    frame.to_parquet(parquet_path, index=False, engine="pyarrow")
    csv_frame = pd.read_csv(csv_path, encoding="utf-8-sig", dtype="string", keep_default_na=False)
    parquet_frame = pd.read_parquet(parquet_path, engine="pyarrow")
    if len(csv_frame) != len(frame) or len(parquet_frame) != len(frame):
        raise ValueError(f"Falha na contagem após exportar {stem}")
    if list(csv_frame.columns) != list(frame.columns) or list(parquet_frame.columns) != list(frame.columns):
        raise ValueError(f"Falha no esquema após exportar {stem}")
    if not set(source_fields).issubset(frame.columns):
        raise ValueError(f"Campos originais ausentes em {stem}: {source_fields}")
    from pandas.testing import assert_frame_equal
    assert_frame_equal(frame.reset_index(drop=True), parquet_frame.reset_index(drop=True), check_dtype=False)
    expected_csv = frame.copy()
    for column in expected_csv.columns:
        if pd.api.types.is_datetime64_any_dtype(expected_csv[column]):
            expected_csv[column] = expected_csv[column].dt.strftime("%Y-%m-%d")
    expected_csv = expected_csv.astype("string").fillna("")
    assert_frame_equal(csv_frame.reset_index(drop=True), expected_csv.reset_index(drop=True), check_dtype=False)
    return {
        "conjunto": stem,
        "registros": len(frame),
        "campos": len(frame.columns),
        "csv": str(csv_path.relative_to(project_path())),
        "parquet": str(parquet_path.relative_to(project_path())),
        "csv_releitura": len(csv_frame),
        "parquet_releitura": len(parquet_frame),
        "sha256_csv": hashlib.sha256(csv_path.read_bytes()).hexdigest(),
        "sha256_parquet": hashlib.sha256(parquet_path.read_bytes()).hexdigest(),
        "status": "validado",
    }


def _source_url_for_ssp(source_names: str) -> str:
    return f"{SSP_URL} — documentos: {source_names}"


def _summarize_sinesp(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame]:
    data = frame.copy()
    data["ano"] = data["data_referencia"].dt.year.astype(int)
    data["mes"] = data["data_referencia"].dt.to_period("M").astype(str)
    measures = [column for column in ("total_vitima", "total") if column in data]
    long = data.melt(
        id_vars=["ano", "mes", "evento", "abrangencia"],
        value_vars=measures,
        var_name="campo_medida",
        value_name="valor_mensal",
    ).dropna(subset=["valor_mensal"])
    # As observações repetidas no mesmo mês/categoria seguem a nota metodológica
    # oficial do Sinesp: tentativa de homicídio e estupro devem ser somados.
    monthly = (
        long.groupby(["ano", "mes", "evento", "abrangencia", "campo_medida"], dropna=False)
        .agg(valor_mensal=("valor_mensal", "sum"), linhas_fonte=("valor_mensal", "size"))
        .reset_index()
    )
    summary = (
        monthly.groupby(["ano", "evento", "abrangencia", "campo_medida"], dropna=False)
        .agg(
            meses_com_dado=("mes", "nunique"),
            soma_periodo=("valor_mensal", "sum"),
            media_mensal=("valor_mensal", "mean"),
            minimo_mensal=("valor_mensal", "min"),
            maximo_mensal=("valor_mensal", "max"),
            linhas_fonte=("linhas_fonte", "sum"),
        )
        .reset_index()
    )
    summary["cobertura"] = summary["ano"].map(lambda year: "parcial" if year == 2026 else "ano completo")
    summary["unidade"] = summary["campo_medida"].map(
        {"total_vitima": "vítimas informadas", "total": "ocorrências/mandados informados"}
    )
    summary["fonte"] = "MJSP / Sinesp VDE"
    summary["url_fonte"] = SINESP_URL
    return monthly, summary


def _summarize_ssp(frame: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    homicide = frame[frame["indicador_normalizado"].eq(SSP_HOMICIDE)].copy()
    homicide["data_ordem"] = pd.to_datetime(homicide["data_atualizacao"], format="%d/%m/%Y", errors="coerce")
    complete = homicide[homicide["periodo_tipo"].eq("ano completo")].sort_values("data_ordem")
    annual = complete.drop_duplicates("ano_referencia", keep="last").copy()
    annual = annual[annual["ano_referencia"].between(2020, 2025)]
    annual["fonte"] = "SSP/SC — boletim mensal; último valor anual publicado para cada referência"
    annual["url_fonte"] = SSP_URL
    annual = annual[[
        "ano_referencia", "valor", "data_atualizacao", "arquivo_pdf", "pagina_pdf", "fonte", "url_fonte"
    ]].rename(columns={"ano_referencia": "ano", "valor": "vitimas_homicidio"})

    accumulated = homicide[homicide["periodo_tipo"].eq("período acumulado")].copy()
    if accumulated.empty:
        raise ValueError("Não foram encontradas observações acumuladas de homicídio")
    latest_date = accumulated["data_ordem"].max()
    latest = accumulated[accumulated["data_ordem"].eq(latest_date)].copy()
    partial = latest[latest["ano_referencia"].between(2023, 2026)].sort_values("ano_referencia")
    if partial.empty or partial["ano_referencia"].max() != 2026:
        raise ValueError("Não há comparação parcial que inclua 2026")
    partial["periodo_comparavel"] = partial["periodo_referencia"]
    partial["cobertura"] = "parcial acumulado, conforme data do boletim"
    partial["fonte"] = "SSP/SC — mesmo boletim para os quatro anos comparados"
    partial["url_fonte"] = SSP_URL
    partial = partial[[
        "ano_referencia", "valor", "periodo_comparavel", "data_atualizacao",
        "arquivo_pdf", "pagina_pdf", "cobertura", "fonte", "url_fonte",
    ]].rename(columns={"ano_referencia": "ano", "valor": "vitimas_homicidio"})

    domestic = frame[frame["periodo_tipo"].eq("mensal")].copy()
    domestic["mes"] = domestic["periodo_referencia"].astype("string")
    key = ["ano_referencia", "mes", "indicador_normalizado"]
    if domestic.duplicated(key).any():
        duplicates = domestic.loc[domestic.duplicated(key, keep=False), key]
        raise ValueError(f"Violência doméstica tem mês/categoria duplicado após normalizar: {duplicates.head()}")
    domestic_stats = (
        domestic.groupby(["ano_referencia", "indicador_normalizado"], dropna=False)
        .agg(
            meses_com_dado=("mes", "nunique"),
            total_periodo=("valor", "sum"),
            media_mensal=("valor", "mean"),
            minimo_mensal=("valor", "min"),
            maximo_mensal=("valor", "max"),
            rotulos_originais=("indicador", lambda values: " | ".join(sorted(set(values)))),
            documentos=("arquivo_pdf", lambda values: " | ".join(sorted(set(values)))),
        )
        .reset_index()
    )
    domestic_stats["fonte"] = "SSP/SC — relatórios temáticos de violência doméstica"
    domestic_stats["url_fonte"] = SSP_URL
    domestic_stats["pagina_pdf"] = "ver arquivo_pdf nos registros mensais"
    return annual, partial, domestic_stats


def _summarize_city_crime(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame[["ano", "indicador", "periodo", "quantidade", "fonte", "url_fonte"]].sort_values(
        ["indicador", "ano"]
    )
    result["cobertura"] = "janeiro-dezembro"
    return result


def _compare_homicide_sources(sinesp: pd.DataFrame, ssp_annual: pd.DataFrame,
                              ssp_partial: pd.DataFrame) -> pd.DataFrame:
    sinesp_rows = sinesp[
        sinesp["evento"].eq("Homicídio doloso")
        & sinesp["total_vitima"].notna()
    ].copy()
    sinesp_rows["ano"] = sinesp_rows["data_referencia"].dt.year
    sinesp_annual = sinesp_rows.groupby("ano")["total_vitima"].sum().rename("sinesp_vitimas_homicidio")
    result = pd.DataFrame({"ano": range(2020, 2027)}).merge(
        sinesp_annual, left_on="ano", right_index=True, how="left"
    )
    result = result.merge(
        ssp_annual[["ano", "vitimas_homicidio"]].rename(columns={"vitimas_homicidio": "ssp_vitimas_homicidio_ano_completo"}),
        on="ano", how="left",
    )
    result["diferenca_sinesp_menos_ssp"] = (
        result["sinesp_vitimas_homicidio"] - result["ssp_vitimas_homicidio_ano_completo"]
    )
    partial = ssp_partial.set_index("ano")[["vitimas_homicidio", "periodo_comparavel"]]
    result["ssp_vitimas_homicidio_parcial"] = result["ano"].map(partial["vitimas_homicidio"])
    result["ssp_periodo_parcial"] = result["ano"].map(partial["periodo_comparavel"])
    result["sinesp_periodo"] = result["ano"].map(lambda year: "janeiro–agosto" if year == 2026 else "janeiro–dezembro")
    result["comparacao_anual_valida"] = result["ano"].lt(2026) & result["ssp_vitimas_homicidio_ano_completo"].notna()
    result["observacao"] = result["comparacao_anual_valida"].map(
        {True: "fontes mantidas separadas; diferenças refletem publicação/revisão própria de cada série",
         False: "sem comparação anual equivalente; Sinesp cobre jan–ago/2026 e SSP jan–jul/2026"}
    )
    result["fonte_sinesp"] = SINESP_URL
    result["fonte_ssp_sc"] = SSP_URL
    return result


def _save_charts(sinesp_monthly: pd.DataFrame, ssp_annual: pd.DataFrame,
                 ssp_partial: pd.DataFrame, domestic: pd.DataFrame,
                 city: pd.DataFrame) -> list[str]:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import numpy as np

    CHART_DIR.mkdir(parents=True, exist_ok=True)
    paths = []

    crime = sinesp_monthly[
        sinesp_monthly["campo_medida"].eq("total_vitima")
        & sinesp_monthly["evento"].isin([
            "Homicídio doloso", "Tentativa de homicídio", "Feminicídio",
            "Tentativa de feminicídio", "Lesão corporal seguida de morte",
        ])
    ]
    if not crime.empty:
        annual = crime.groupby(["ano", "evento"], as_index=False)["valor_mensal"].sum()
        years = sorted(annual["ano"].unique())
        fig, ax = plt.subplots(figsize=(10, 5.5), constrained_layout=True)
        for label, group in annual.groupby("evento"):
            group = group.sort_values("ano")
            ax.plot(group["ano"], group["valor_mensal"], marker="o", linewidth=2, label=label)
            if 2026 in group["ano"].values:
                partial = group[group["ano"].eq(2026)].iloc[0]
                ax.scatter([2026], [partial["valor_mensal"]], marker="D", s=55, zorder=4)
        ax.set_title("Sinesp VDE — indicadores de vítimas por ano")
        ax.set_xlabel("Ano de referência")
        ax.set_ylabel("Vítimas informadas; 2026 parcial (jan–ago)")
        ax.set_xticks(years)
        ax.grid(axis="y", alpha=.25)
        ax.legend(frameon=False, ncol=2, loc="upper left")
        ax.text(0, -0.2, f"Fonte: MJSP / Sinesp VDE. {SINESP_URL}", transform=ax.transAxes, fontsize=7, wrap=True)
        out = CHART_DIR / "sinesp_vitimas_por_indicador_ano.png"
        fig.savefig(out, dpi=180)
        plt.close(fig)
        paths.append(str(out.relative_to(project_path())))

    fig, ax = plt.subplots(figsize=(8, 4.5), constrained_layout=True)
    annual = ssp_annual.sort_values("ano")
    ax.bar(annual["ano"].astype(str), annual["vitimas_homicidio"], color="#335f8a")
    ax.set_title("Joinville — vítimas de homicídio por ano completo")
    ax.set_xlabel("Ano de referência")
    ax.set_ylabel("Vítimas")
    ax.grid(axis="y", alpha=.25)
    ax.text(0, -0.24, f"Fonte: SSP/SC, último valor anual publicado por ano. {SSP_URL}", transform=ax.transAxes, fontsize=7, wrap=True)
    out = CHART_DIR / "ssp_homicidios_ano_completo.png"
    fig.savefig(out, dpi=180)
    plt.close(fig)
    paths.append(str(out.relative_to(project_path())))

    fig, ax = plt.subplots(figsize=(8, 4.5), constrained_layout=True)
    partial = ssp_partial.sort_values("ano")
    ax.bar(partial["ano"].astype(str), partial["vitimas_homicidio"], color=["#8fa8be"] * (len(partial)-1) + ["#cf5c36"])
    ax.set_title("Joinville — vítimas de homicídio no período jan–jul")
    ax.set_xlabel("Ano de referência")
    ax.set_ylabel("Vítimas acumuladas no mesmo período")
    ax.grid(axis="y", alpha=.25)
    ax.text(0, -0.24, f"2026 é parcial. Um único boletim SSP/SC (atualizado em {partial['data_atualizacao'].iloc[-1]}). {SSP_URL}", transform=ax.transAxes, fontsize=7, wrap=True)
    out = CHART_DIR / "ssp_homicidios_comparacao_jan_jul.png"
    fig.savefig(out, dpi=180)
    plt.close(fig)
    paths.append(str(out.relative_to(project_path())))

    pivot = domestic.pivot(index="indicador_normalizado", columns="ano_referencia", values="total_periodo").fillna(0)
    if not pivot.empty:
        fig, ax = plt.subplots(figsize=(10, max(4.5, .42 * len(pivot))), constrained_layout=True)
        y = np.arange(len(pivot))
        years = list(pivot.columns)
        width = .36 if len(years) == 2 else .7
        for idx, year in enumerate(years):
            ax.barh(y + (idx - (len(years)-1)/2) * width, pivot[year], height=width, label=str(year))
        ax.set_yticks(y, labels=pivot.index)
        ax.invert_yaxis()
        ax.set_title("Violência doméstica em Joinville — soma dos valores mensais por indicador")
        ax.set_xlabel("Registros informados nos relatórios mensais")
        ax.legend(title="Ano", frameon=False)
        ax.grid(axis="x", alpha=.25)
        ax.text(0, -0.14, f"Fonte: SSP/SC; soma de meses distintos; rótulos normalizados com original preservado. {SSP_URL}", transform=ax.transAxes, fontsize=7, wrap=True)
        out = CHART_DIR / "ssp_violencia_domestica_por_indicador.png"
        fig.savefig(out, dpi=180)
        plt.close(fig)
        paths.append(str(out.relative_to(project_path())))

    pivot = city.pivot(index="ano", columns="indicador", values="quantidade")
    if not pivot.empty:
        fig, axes = plt.subplots(1, len(pivot.columns), figsize=(10, 4.5), constrained_layout=True)
        if len(pivot.columns) == 1:
            axes = [axes]
        for ax, indicator in zip(axes, pivot.columns):
            ax.bar(pivot.index.astype(str), pivot[indicator], color="#598a72")
            ax.set_title(indicator)
            ax.set_xlabel("Ano")
            ax.set_ylabel("Quantidade")
            ax.grid(axis="y", alpha=.25)
        fig.suptitle("Joinville — totais anuais divulgados pela Prefeitura")
        fig.text(.01, -.03, f"Fonte: Prefeitura de Joinville / SEPROT. {CITY_CRIME_URL}", fontsize=7, wrap=True)
        out = CHART_DIR / "prefeitura_furtos_roubos_ano.png"
        fig.savefig(out, dpi=180, bbox_inches="tight")
        plt.close(fig)
        paths.append(str(out.relative_to(project_path())))
    return paths


def run_eda_and_validate() -> dict:
    ANALYSIS_DIR.mkdir(parents=True, exist_ok=True)
    TABLE_DIR.mkdir(parents=True, exist_ok=True)
    sinesp = _read_sinesp()
    ssp = _read_ssp()
    city = _read_city_crime()
    if not sinesp["uf"].eq("SC").all() or not sinesp["municipio"].eq("JOINVILLE").all():
        raise ValueError("O recorte Sinesp contém linhas fora de Joinville/SC")
    actual_sinesp_counts = sinesp.groupby(sinesp["data_referencia"].dt.year).size().to_dict()
    if actual_sinesp_counts != SINESP_EXPECTED_CITY_ROWS:
        raise ValueError(f"Contagens Sinesp municipais divergentes: {actual_sinesp_counts}")

    exported = []
    exported.append(_write_frame(
        sinesp, "sinesp_joinville", ["uf", "municipio", "evento", "data_referencia"]
    ))
    exported.append(_write_frame(
        ssp, "ssp_joinville", ["indicador", "indicador_normalizado", "arquivo_pdf", "pagina_pdf"]
    ))
    exported.append(_write_frame(
        city, "prefeitura_furto_roubo_joinville", ["indicador", "quantidade", "url_fonte"]
    ))

    sinesp_monthly, sinesp_stats = _summarize_sinesp(sinesp)
    ssp_annual, ssp_partial, dv_stats = _summarize_ssp(ssp)
    city_stats = _summarize_city_crime(city)
    homicide_comparison = _compare_homicide_sources(sinesp, ssp_annual, ssp_partial)
    if set(ssp_annual["ano"]) != {2020, 2021, 2022, 2023, 2024, 2025}:
        raise ValueError(f"Cobertura anual SSP inesperada: {sorted(set(ssp_annual['ano']))}")
    if set(ssp_partial["ano"]) != {2023, 2024, 2025, 2026}:
        raise ValueError(f"Comparação parcial SSP inesperada: {sorted(set(ssp_partial['ano']))}")
    if ssp_partial["periodo_comparavel"].nunique() != 1:
        raise ValueError("Períodos acumulados SSP não são comparáveis entre si")
    actual_annual = {
        int(row.ano): int(row.vitimas_homicidio)
        for row in ssp_annual.itertuples(index=False)
    }
    actual_partial = {
        int(row.ano): int(row.vitimas_homicidio)
        for row in ssp_partial.itertuples(index=False)
    }
    if actual_annual != SSP_ANNUAL_CONTROL:
        raise ValueError(f"Totais anuais SSP/SC divergentes dos valores de controle dos PDFs: {actual_annual}")
    if actual_partial != SSP_PARTIAL_CONTROL:
        raise ValueError(f"Totais parciais SSP/SC divergentes do boletim de julho/2026: {actual_partial}")

    tables = {
        "estatisticas_sinesp_indicador_ano": sinesp_stats,
        "estatisticas_ssp_homicidio_ano": ssp_annual,
        "comparacao_ssp_homicidio_parcial": ssp_partial,
        "estatisticas_ssp_violencia_domestica": dv_stats,
        "estatisticas_prefeitura_furto_roubo": city_stats,
        "comparacao_fontes_homicidio": homicide_comparison,
        "sinesp_observacoes_mensais_agregadas": sinesp_monthly,
    }
    for name, table in tables.items():
        table.to_csv(TABLE_DIR / f"{name}.csv", index=False, encoding="utf-8-sig")

    chart_paths = _save_charts(sinesp_monthly, ssp_annual, ssp_partial, dv_stats, city_stats)

    quality = []
    field_quality = []
    for dataset, frame in (("sinesp", sinesp), ("ssp_sc", ssp), ("prefeitura", city)):
        blanks = frame.astype("string").eq("")
        missing = frame.isna() | blanks
        quality.append({
            "base": dataset,
            "registros": len(frame),
            "campos": len(frame.columns),
            "municipios": int(frame["municipio"].nunique()) if "municipio" in frame else None,
            "valores_vazios": int(missing.sum().sum()),
            "fonte": {"sinesp": SINESP_URL, "ssp_sc": SSP_URL, "prefeitura": CITY_CRIME_URL}[dataset],
        })
        for column in frame.columns:
            count = int(missing[column].sum())
            field_quality.append({
                "base": dataset,
                "campo": column,
                "vazios": count,
                "percentual_vazio": round(count / len(frame) * 100, 2) if len(frame) else 0,
                "registros": len(frame),
                "interpretação": "não preenchido na fonte/extração; não convertido para zero",
            })
    pd.DataFrame(quality).to_csv(TABLE_DIR / "auditoria_qualidade_bases.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(field_quality).to_csv(TABLE_DIR / "auditoria_campos_vazios.csv", index=False, encoding="utf-8-sig")
    pd.DataFrame(exported).to_csv(TABLE_DIR / "validacao_csv_parquet.csv", index=False, encoding="utf-8-sig")
    manifest = {
        "fontes": {"sinesp": SINESP_URL, "ssp_sc": SSP_URL, "prefeitura": CITY_CRIME_URL},
        "arquivos": exported,
        "tabelas": {name: len(frame) for name, frame in tables.items()},
        "graficos": chart_paths,
        "sinesp_anos": sorted(int(year) for year in sinesp["data_referencia"].dt.year.unique()),
        "ssp_2026_periodo_parcial": sorted(set(ssp_partial["periodo_comparavel"])),
        "limite_guarda_municipal": "sem série estruturada validada no conjunto coletado",
    }
    manifest_path = TABLE_DIR / "manifesto_analise.json"
    manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return {
        "exported": exported,
        "tables": tables,
        "charts": chart_paths,
        "quality": quality,
        "manifest": manifest_path,
        "sinesp": sinesp,
        "ssp": ssp,
        "city": city,
    }
