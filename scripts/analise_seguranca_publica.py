"""Inspeção acadêmica, em células interativas, das fontes de Joinville/SC.

Abra a pasta do projeto no VS Code. Execute cada célula com "Run Cell".
As variáveis DataFrame ficam disponíveis no Interactive Window para abrir no
Data Viewer. As etapas 1–4 inspecionam as fontes; a etapa 5 valida e concatena
somente em memória os recortes anuais do Sinesp. A etapa 6 diagnostica a
qualidade desse DataFrame; a etapa 7 investiga os achados e propõe colunas e
tipos. A etapa 8 prepara uma cópia analítica em memória, com datas e inteiros
anuláveis validados. A etapa 9 calcula tabelas descritivas por indicador e
abrangência, somente em memória. A etapa 10 exibe quatro gráficos exploratórios
validados contra essas tabelas. Não há exportação.
"""

# %% [markdown]
# # ETAPA 1 — Importação das bibliotecas
# Objetivo: importar Pandas para leitura de planilhas e CSVs e pathlib para
# localizar as fontes. Nenhuma transformação é aplicada nesta etapa.

# %%
from pathlib import Path
import sys

import pandas as pd
import pdfplumber
from scripts import transform

# Mantém os acentos legíveis no terminal quando o fluxo não está no Interactive Window.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

print(f"Pandas: {pd.__version__}")


# %% [markdown]
# # ETAPA 2 — Definição dos diretórios
# Objetivo: localizar as fontes a partir da pasta do projeto aberta no VS Code.
# Procedimento: em execução como módulo, usa a pasta acima de scripts/; no
# Interactive Window, usa a pasta de trabalho (abra a raiz do projeto no VS Code).
# Validação: os caminhos são exibidos junto com seu estado de disponibilidade.

# %%
RAIZ_PROJETO = (
    Path(__file__).resolve().parents[1] if "__file__" in globals() else Path.cwd()
)
PASTA_SINESP = RAIZ_PROJETO / "dados" / "brutos" / "sinesp"
ARQUIVO_SINESP_2023 = PASTA_SINESP / "bancovde-2023.xlsx"
PASTA_CSV_SSP = RAIZ_PROJETO / "dados" / "brutos" / "ssp_sc" / "csv_extraidos"
PASTA_PDF_SSP = RAIZ_PROJETO / "dados" / "brutos" / "ssp_sc" / "pdfs"
PASTA_PREFEITURA = RAIZ_PROJETO / "dados" / "processados" / "prefeitura"

# Altere este nome para inspecionar outro CSV existente da SSP/SC.
CSV_SSP_SELECIONADO = "2026_08_2026-julho.csv"

print(f"Raiz do projeto: {RAIZ_PROJETO}")
for nome, caminho in (
    ("XLSX Sinesp 2023", ARQUIVO_SINESP_2023),
    ("CSV extraídos SSP/SC", PASTA_CSV_SSP),
    ("CSV Prefeitura/SEPROT", PASTA_PREFEITURA),
):
    print(f"{nome}: {'disponível' if caminho.exists() else 'não encontrado'} — {caminho}")

# %% [markdown]
# # ETAPA 3 — Leitura dos arquivos brutos
# Objetivo: ler uma amostra nacional do Sinesp, selecionar Joinville/SC em
# memória, ler os CSVs SSP já extraídos e abrir separadamente os CSVs SEPROT.
# Nenhuma fonte é sobrescrita, exportada ou consolidada.

# %% [markdown]
# ## Sinesp — amostra nacional e recorte em memória
# read_excel(nrows=20) lê a amostra nacional. O leitor XML especializado já
# existente seleciona as linhas SC/JOINVILLE do XLSX e as devolve em memória.
# As datas permanecem como armazenadas no arquivo, inclusive seriais do Excel.

# %%
sinesp_2023 = pd.read_excel(
    ARQUIVO_SINESP_2023,
    nrows=20,
    engine="openpyxl",
)
joinville_sc = pd.DataFrame.from_records(
    transform._joinville_rows(ARQUIVO_SINESP_2023),
    columns=transform.HEADERS,
)
print(f"Amostra nacional: {len(sinesp_2023)} linhas; recorte Joinville/SC: {len(joinville_sc)} linhas.")

# %% [markdown]
# ## SSP/SC — inventário e cabeçalhos dos CSVs já extraídos
# O inventário consulta os cabeçalhos com nrows=0. As colunas celula_* variam
# entre relatórios porque as tabelas dos PDFs têm larguras diferentes.

# %%
arquivos_ssp = sorted(PASTA_CSV_SSP.glob("*.csv"))
inventario_linhas_ssp = []
esquemas_ssp = {}
for arquivo in arquivos_ssp:
    colunas_arquivo = pd.read_csv(
        arquivo,
        sep=";",
        encoding="utf-8-sig",
        nrows=0,
    ).columns.tolist()
    esquemas_ssp.setdefault(tuple(colunas_arquivo), []).append(arquivo.name)
    inventario_linhas_ssp.append(
        {
            "arquivo_csv": arquivo.name,
            "formato": "CSV",
            "tamanho_bytes": arquivo.stat().st_size,
            "ano_no_nome": arquivo.name[:4],
            "numero_colunas": len(colunas_arquivo),
        }
    )
inventario_ssp = pd.DataFrame(inventario_linhas_ssp)

arquivo_ssp_selecionado = PASTA_CSV_SSP / CSV_SSP_SELECIONADO
ssp_selecionado = pd.read_csv(
    arquivo_ssp_selecionado,
    sep=";",
    encoding="utf-8-sig",
    dtype="string",
    keep_default_na=False,
)
print(f"SSP/SC: {len(arquivos_ssp)} CSVs; arquivo selecionado: {arquivo_ssp_selecionado.name}.")

# %% [markdown]
# ## Prefeitura/SEPROT — arquivos transcritos
# Os três CSVs são lidos separadamente. Cada um contém duas linhas; não são
# empilhados em uma tabela única.

# %%
arquivos_prefeitura = sorted(PASTA_PREFEITURA.glob("furto_roubo_joinville_*.csv"))
tabelas_prefeitura = [
    pd.read_csv(arquivo, encoding="utf-8-sig") for arquivo in arquivos_prefeitura
]
prefeitura_2023 = tabelas_prefeitura[0]
prefeitura_2024 = tabelas_prefeitura[1]
prefeitura_2025 = tabelas_prefeitura[2]
print(f"Prefeitura/SEPROT: {len(arquivos_prefeitura)} arquivos, {sum(map(len, tabelas_prefeitura))} linhas.")

# %% [markdown]
# # ETAPA 4 — Inspeção e visualização dos dados
# Execute as células abaixo separadamente. No Interactive Window, cada expressão
# final apresenta a tabela; os nomes das variáveis ficam disponíveis no painel
# Variables para abrir no Data Viewer.
# Como abrir uma tabela: execute as células anteriores e a célula da fonte;
# no Interactive Window, clique em Variables. Na lista, localize a variável e
# selecione o botão Show variable in data viewer (ícone de tabela).
# Variáveis: sinesp_2023, joinville_sc, data_referencia_nacional_20,
# data_referencia_joinville_20, inventario_ssp, ssp_selecionado,
# ssp_com_joinville, ssp_tabela_pagina_2, ssp_tabela_pagina_8,
# ssp_contexto_pagina_2, ssp_contexto_pagina_8 e prefeitura_2023–2025.

# %% [markdown]
# ## Sinesp — colunas, tipos e primeiras 20 linhas nacionais
# As 14 colunas e seus tipos são mostrados conforme a amostra lida. Os tipos
# inferidos a partir de 20 linhas não são apresentados como tipos de toda a base.

# %%
print("Arquivo nacional Sinesp:", ARQUIVO_SINESP_2023.name)
print("Dimensões da amostra:", sinesp_2023.shape)
print("Colunas:", sinesp_2023.columns.tolist())
print("Tipos Pandas observados na amostra:")
print(sinesp_2023.dtypes.to_string())
sinesp_2023.head(20)

# %% [markdown]
# ## Sinesp — comparação de data_referencia sem conversão
# Objetivo: observar os primeiros 20 valores e os tipos inferidos por cada
# caminho de leitura. read_excel interpreta as células XLSX formatadas como
# data; o leitor XML auxiliar retorna o conteúdo numérico serial da célula.
# A diferença descreve os leitores, não comprova erro no arquivo. Nenhum valor
# é convertido nesta inspeção.

# %%
data_referencia_nacional_20 = sinesp_2023["data_referencia"].head(20)
data_referencia_joinville_20 = joinville_sc["data_referencia"].head(20)
print("Tipo original na amostra nacional:", sinesp_2023["data_referencia"].dtype)
print("Tipo original no recorte Joinville:", joinville_sc["data_referencia"].dtype)
print("Valores nacionais — data_referencia:")
data_referencia_nacional_20

# %%
print("Valores Joinville/SC — data_referencia (sem conversão):")
data_referencia_joinville_20

# %% [markdown]
# ## Sinesp — primeiras 20 linhas do recorte Joinville/SC
# É uma seleção em memória do arquivo nacional, sem exportação. O leitor
# especializado preserva a data como número serial do XLSX quando assim gravada.

# %%
print("Dimensões do recorte Joinville/SC:", joinville_sc.shape)
print("Tipos Pandas observados no recorte:")
print(joinville_sc.dtypes.to_string())
joinville_sc.head(20)

# %% [markdown]
# ### Significado dos campos Sinesp
# Baseado no catálogo e nas notas oficiais do MJSP e nos nomes originais. O PDF
# do dicionário não abriu nesta consulta; por isso, campos sem definição confirmada
# são identificados como tal, sem atribuir significado por suposição.
# Referência: https://dados.mj.gov.br/dataset/sistema-nacional-de-estatisticas-de-seguranca-publica

# %%
print(
    "uf e municipio: referência geográfica; evento: indicador; "
    "data_referencia: data/período de referência registrado na base."
)
print(
    "feminino, masculino e nao_informado: nomes de campos de sexo; as notas do "
    "MJSP descrevem totais por sexo para indicadores enviados como vítimas. "
    "total_vitima: campo declarado para total de vítimas."
)
print(
    "agente, arma e faixa_etaria: campos presentes na planilha; sem confirmação "
    "do dicionário, não se atribui definição mais específica a cada um."
)
print(
    "total e total_peso: nomes de campos de medida; unidade e interpretação "
    "dependem do indicador e não são inferidas apenas pelos nomes."
)
print(
    "abrangencia: campo assim nomeado na base; não identifica, por si só, qual "
    "corporação registrou os dados municipais."
)

# %% [markdown]
# ## SSP/SC — inventário e esquemas de colunas dos CSVs
# arquivo_pdf, pagina_pdf, tabela_pdf e linha_tabela são metadados do extrator.
# celula_* contém o texto extraído e não representa um esquema estatístico comum.

# %%
print(f"CSV encontrados: {len(inventario_ssp)} (esperados: 45)")
inventario_ssp

# %%
for numero, (colunas, arquivos_do_esquema) in enumerate(esquemas_ssp.items(), start=1):
    print(f"Esquema {numero}: {len(colunas)} colunas; arquivos: {arquivos_do_esquema}")
    print(list(colunas))

# %% [markdown]
# ## SSP/SC — colunas e primeiras 20 linhas do CSV selecionado
# O nome abaixo pode ser trocado em CSV_SSP_SELECIONADO, na etapa 2.

# %%
print("Arquivo selecionado:", arquivo_ssp_selecionado.name)
print("Dimensões:", ssp_selecionado.shape)
print("Colunas:", ssp_selecionado.columns.tolist())
print(
    "Rastreabilidade:",
    [nome for nome in ("arquivo_pdf", "pagina_pdf", "tabela_pdf", "linha_tabela")
     if nome in ssp_selecionado.columns],
)
ssp_selecionado.head(20)

# %% [markdown]
# ## SSP/SC — todas as linhas que contêm JOINVILLE
# A máscara procura o texto em qualquer coluna e apenas seleciona linhas para
# visualização. O CSV original permanece intacto.

# %%
linhas_joinville_ssp = ssp_selecionado.apply(
    lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
).any(axis=1)
ssp_com_joinville = ssp_selecionado.loc[linhas_joinville_ssp]
print(f"Linhas encontradas: {len(ssp_com_joinville)}")
ssp_com_joinville

# %% [markdown]
# ## SSP/SC — cabeçalho e contexto no PDF original
# Objetivo: conferir visualmente as duas linhas encontradas no CSV, diretamente
# nas páginas originais. pdfplumber apenas lê as tabelas do PDF em memória; as
# linhas e cabeçalhos não são limpos, renomeados no arquivo, nem exportados.
# tabela_pdf=1 é o número registrado no CSV pelo extrator. Os nomes temporários
# coluna_XX permitem visualizar as células em DataFrame sem presumir cabeçalhos.

# %%
ARQUIVO_PDF_SSP_SELECIONADO = (
    PASTA_PDF_SSP / "2026_08_2026-julho.pdf"
)
with pdfplumber.open(ARQUIVO_PDF_SSP_SELECIONADO) as documento_ssp:
    tabela_pdf_pagina_2 = documento_ssp.pages[1].extract_tables()[0]
    tabela_pdf_pagina_8 = documento_ssp.pages[7].extract_tables()[0]

ssp_tabela_pagina_2 = pd.DataFrame(
    tabela_pdf_pagina_2,
    columns=[f"coluna_{numero:02d}" for numero in range(1, len(tabela_pdf_pagina_2[0]) + 1)],
)
ssp_tabela_pagina_8 = pd.DataFrame(
    tabela_pdf_pagina_8,
    columns=[f"coluna_{numero:02d}" for numero in range(1, len(tabela_pdf_pagina_8[0]) + 1)],
)
linha_joinville_pagina_2 = ssp_tabela_pagina_2.index[
    ssp_tabela_pagina_2.astype("string").apply(
        lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
    ).any(axis=1)
][0]
linha_joinville_pagina_8 = ssp_tabela_pagina_8.index[
    ssp_tabela_pagina_8.astype("string").apply(
        lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
    ).any(axis=1)
][0]
ssp_contexto_pagina_2 = ssp_tabela_pagina_2.iloc[
    max(0, linha_joinville_pagina_2 - 3):linha_joinville_pagina_2 + 4
]
ssp_contexto_pagina_8 = ssp_tabela_pagina_8.iloc[
    max(0, linha_joinville_pagina_8 - 3):linha_joinville_pagina_8 + 4
]
print(f"PDF: {ARQUIVO_PDF_SSP_SELECIONADO.name} | página 2 | tabela 1")
print("Cabeçalho e contexto: tabela de variações de homicídios por município, período 01/01–31/07, comparação 2025–2026.")
ssp_contexto_pagina_2

# %%
print(f"PDF: {ARQUIVO_PDF_SSP_SELECIONADO.name} | página 8 | tabela 1")
print("Cabeçalho da tabela municipal: vítimas por ano (2023–2025) e no período 01/01–31/07 (2023–2026).")
print("Primeiras linhas da tabela, incluindo cabeçalhos originais:")
ssp_tabela_pagina_8.head(5)

# %%
print("Linha Joinville e municípios vizinhos na tabela original da página 8:")
ssp_contexto_pagina_8

# %%
print("As duas linhas completas do CSV, com metadados de rastreabilidade:")
ssp_com_joinville

# %% [markdown]
# ### Significado dos campos SSP/SC
# A página oficial informa que Segurança em Números é elaborado pela Gerência
# de Estatística e Análise Criminal. As quatro colunas de rastreabilidade e as
# celula_* são estrutura do extrator; os cabeçalhos/rótulos estatísticos ficam
# dentro das células conforme cada PDF, sem normalização.
# Referência: https://ssp.sc.gov.br/segurancaemnumeros/

# %%
print(
    "arquivo_pdf: nome do PDF de origem; pagina_pdf: página; tabela_pdf: número "
    "da tabela detectada; linha_tabela: posição da linha extraída. "
    "celula_*: conteúdo da linha conforme o extrator."
)

# %% [markdown]
# ## Prefeitura/SEPROT — seis linhas, em três DataFrames separados
# Os CSVs são transcrições anteriores da publicação municipal. A publicação
# informa os totais anuais; não fornece dicionário para o esquema dos CSVs.
# Referência: https://www.joinville.sc.gov.br/noticias/indices-de-criminalidade-tem-reducao-em-joinville-pelo-terceiro-ano-consecutivo/

# %%
print("Prefeitura/SEPROT — 2023 (transcrição anterior)")
prefeitura_2023

# %%
print("Prefeitura/SEPROT — 2024 (transcrição anterior)")
prefeitura_2024

# %%
print("Prefeitura/SEPROT — 2025 (transcrição anterior)")
prefeitura_2025

# %% [markdown]
# ### Significado dos campos Prefeitura/SEPROT
# uf, municipio e ano identificam o recorte transcrito; indicador nomeia Furto
# ou Roubo; periodo registra janeiro a dezembro; quantidade reproduz o valor
# publicado; fonte e url_fonte registram a proveniência da transcrição.

# %%
print(
    "A Prefeitura publicou totais anuais para Furto e Roubo em 2023–2025. "
    "Os campos de organização e proveniência pertencem à transcrição dos CSVs."
)

# %% [markdown]
# # ETAPA 5 — Validação e prévia da concatenação Sinesp
# Objetivo: empilhar os sete recortes municipais compatíveis somente em memória.
# Critério: uf == "SC" e municipio == "JOINVILLE", sem correspondência parcial.
# O leitor XML existente preserva os valores e os seriais Excel; Pandas infere
# os tipos dos recortes. A amostra read_excel(nrows=20) é apenas uma comparação
# dos leitores, não um perfil de tipos do arquivo nacional inteiro.
# Nenhum tipo ou valor é corrigido. Metadados ficam em tabelas auxiliares.
# Execute as células da etapa 5 após a etapa 2. As tabelas abaixo e
# df_sinesp_joinville_concatenado ficam disponíveis em Variables > Data Viewer.

# %%
import re
from zipfile import ZipFile

from openpyxl import load_workbook

anos_sinesp = tuple(range(2020, 2027))
arquivos_sinesp_por_ano = {}
for ano_sinesp in anos_sinesp:
    candidatos_sinesp = sorted(PASTA_SINESP.rglob(f"bancovde-{ano_sinesp}.xlsx"))
    if len(candidatos_sinesp) != 1:
        raise ValueError(
            f"Esperado um original de {ano_sinesp}; encontrados: {candidatos_sinesp}"
        )
    arquivos_sinesp_por_ano[ano_sinesp] = candidatos_sinesp[0]


def inspecionar_original_sinesp(caminho, ano):
    """Confere a estrutura e lê Joinville sem tipar datas nem gravar arquivos."""
    livro = load_workbook(caminho, read_only=True, data_only=False)
    try:
        abas = livro.sheetnames
        if abas != [str(ano)]:
            raise ValueError(f"Planilhas inesperadas em {caminho.name}: {abas}")
        planilha = livro[abas[0]]
        cabecalho = list(next(planilha.iter_rows(min_row=1, max_row=1, values_only=True)))
        if cabecalho != transform.HEADERS:
            raise ValueError(f"Colunas incompatíveis em {caminho.name}: {cabecalho}")
    finally:
        livro.close()

    registros_municipais = []
    linhas_origem = []
    linhas_nacionais = 0
    datas_municipais = set()
    with ZipFile(caminho) as pacote:
        textos = transform._shared_strings(pacote)
        ids_joinville = [str(i).encode("ascii") for i, texto in enumerate(textos)
                        if texto == "JOINVILLE"]
        padrao_municipio = (
            re.compile(rb'<c\b[^>]*\br="B\d+"[^>]*>\s*<v>(?:'
                       + b"|".join(ids_joinville) + rb')</v>\s*</c>')
            if ids_joinville else None
        )
        xml = pacote.read("xl/worksheets/sheet1.xml")
        if re.search(rb'<c\b[^>]*\br="(?:[O-Z]|[A-Z]{2,})\d+"', xml):
            raise ValueError(f"Células além das 14 colunas em {caminho.name}")
        for correspondencia in re.finditer(
            rb'<row\b[^>]*\br="(\d+)"[^>]*>.*?</row>', xml, re.DOTALL
        ):
            numero_linha = int(correspondencia.group(1))
            if numero_linha == 1:
                continue
            linhas_nacionais += 1
            linha_xml = correspondencia.group()
            if b"JOINVILLE" not in linha_xml and not (
                padrao_municipio is not None and padrao_municipio.search(linha_xml)
            ):
                continue
            elemento = transform._parse_sheet_row(linha_xml)
            celulas = {
                celula.get("r").rstrip("0123456789"): transform._cell_value(celula, textos)
                for celula in elemento.findall(transform.NS + "c")
            }
            if celulas.get("A") != "SC" or celulas.get("B") != "JOINVILLE":
                continue
            valores = [celulas.get(chr(65 + i)) for i in range(len(cabecalho))]
            registros_municipais.append(valores)
            linhas_origem.append(numero_linha)
            datas_municipais.add((type(valores[3]).__name__, repr(valores[3])))

    recorte = pd.DataFrame.from_records(registros_municipais, columns=cabecalho)
    amostra = pd.read_excel(caminho, sheet_name=abas[0], nrows=20, engine="openpyxl")
    if amostra.columns.tolist() != cabecalho or recorte.empty:
        raise ValueError(f"Cabeçalhos divergentes ou recorte vazio: {caminho.name}")
    return recorte, amostra, {
        "ano": ano, "arquivo": str(caminho.relative_to(RAIZ_PROJETO)),
        "tamanho_bytes": caminho.stat().st_size, "extensao": caminho.suffix,
        "planilha": abas[0], "colunas": tuple(cabecalho),
        "linhas_nacionais_sem_cabecalho": linhas_nacionais,
        "linhas_joinville": len(recorte),
        "tipo_data_xml": sorted({tipo for tipo, _ in datas_municipais}),
    }, linhas_origem


recortes_sinesp_por_ano = {}
amostras_sinesp_por_ano = {}
metadados_sinesp_por_ano = {}
inventario_sinesp_linhas = []
for ano_sinesp, caminho_sinesp in arquivos_sinesp_por_ano.items():
    recorte_sinesp, amostra_sinesp, ficha_sinesp, linhas_sinesp = inspecionar_original_sinesp(
        caminho_sinesp, ano_sinesp
    )
    recortes_sinesp_por_ano[ano_sinesp] = recorte_sinesp
    amostras_sinesp_por_ano[ano_sinesp] = amostra_sinesp
    inventario_sinesp_linhas.append(ficha_sinesp)
    metadados_sinesp_por_ano[ano_sinesp] = pd.DataFrame({
        "ano_arquivo": ano_sinesp, "arquivo_origem": ficha_sinesp["arquivo"],
        "planilha_origem": ficha_sinesp["planilha"], "linha_xlsx": linhas_sinesp,
    })
    print(f"Sinesp {ano_sinesp}: estrutura compatível; {len(recorte_sinesp)} linhas de SC/JOINVILLE.")

inventario_sinesp = pd.DataFrame(inventario_sinesp_linhas)
comparacao_esquemas_sinesp = inventario_sinesp[["ano", "planilha", "colunas"]]
tipos_recortes_sinesp = pd.DataFrame({
    ano: quadro.dtypes for ano, quadro in recortes_sinesp_por_ano.items()
})
tipos_amostras_sinesp = pd.DataFrame({
    ano: quadro.dtypes for ano, quadro in amostras_sinesp_por_ano.items()
})
contagens_sinesp = inventario_sinesp[[
    "ano", "linhas_nacionais_sem_cabecalho", "linhas_joinville"
]]

# %%
inventario_sinesp

# %%
comparacao_esquemas_sinesp

# %%
contagens_sinesp

# %%
tipos_amostras_sinesp

# %%
tipos_recortes_sinesp

# %% [markdown]
# ## Concatenação autorizada somente em memória
# Granularidade: indicador, competência, município e abrangência, com dimensões
# agente/arma/faixa_etaria preservadas. A chave não é apenas ano/município.
# Os 900 registros anteriores são referência de conferência, não regra de filtro.
# Diferenças de tipos inferidos em campos com nulos serão mostradas, sem cast.
# Se houver esquema incompatível, a célula anterior interrompe antes de concat.

# %%
if any(quadro.columns.tolist() != transform.HEADERS
       for quadro in recortes_sinesp_por_ano.values()):
    raise ValueError("Esquemas incompatíveis: concatenação interrompida.")

df_sinesp_joinville_concatenado = pd.concat(
    list(recortes_sinesp_por_ano.values()), ignore_index=True, sort=False
)
metadados_sinesp_concatenado = pd.concat(
    list(metadados_sinesp_por_ano.values()), ignore_index=True, sort=False
)
linhas_esperadas_sinesp = sum(len(quadro) for quadro in recortes_sinesp_por_ano.values())
assert df_sinesp_joinville_concatenado.shape == (linhas_esperadas_sinesp, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
assert len(metadados_sinesp_concatenado) == linhas_esperadas_sinesp
assert df_sinesp_joinville_concatenado["uf"].eq("SC").all()
assert df_sinesp_joinville_concatenado["municipio"].eq("JOINVILLE").all()

# Confere valores, ausências e ordem de cada bloco. Inferências de dtype são
# relatadas separadamente, sem conversão manual ou substituição de ausências.
inicio_bloco_sinesp = 0
for ano_sinesp, recorte_sinesp in recortes_sinesp_por_ano.items():
    fim_bloco_sinesp = inicio_bloco_sinesp + len(recorte_sinesp)
    pd.testing.assert_frame_equal(
        df_sinesp_joinville_concatenado.iloc[inicio_bloco_sinesp:fim_bloco_sinesp].reset_index(drop=True),
        recorte_sinesp, check_dtype=False,
    )
    inicio_bloco_sinesp = fim_bloco_sinesp

referencia_contagens_sinesp = {2020: 132, 2021: 132, 2022: 132, 2023: 132,
                             2024: 132, 2025: 144, 2026: 96}
for ano_sinesp, quadro_sinesp in recortes_sinesp_por_ano.items():
    if len(quadro_sinesp) != referencia_contagens_sinesp[ano_sinesp]:
        print(f"ATENÇÃO: {ano_sinesp} diverge do inventário anterior; investigar sem ajustar linhas.")

print("Dimensões concatenadas:", df_sinesp_joinville_concatenado.shape)
print("Conferência: soma dos recortes, ordem das colunas e valores de cada bloco preservados.")

# %%
df_sinesp_joinville_concatenado.head(10)

# %%
df_sinesp_joinville_concatenado.tail(10)

# %%
tipos_concatenado_sinesp = df_sinesp_joinville_concatenado.dtypes.to_frame("tipo_pandas")
tipos_concatenado_sinesp

# %%
ausencias_concatenado_sinesp = df_sinesp_joinville_concatenado.isna().sum().to_frame("valores_ausentes")
ausencias_concatenado_sinesp

# %%
# Frequência de linhas tabulares por indicador; não é soma de vítimas ou crimes.
frequencia_eventos_sinesp = df_sinesp_joinville_concatenado["evento"].value_counts(dropna=False).to_frame("linhas_tabulares")
frequencia_eventos_sinesp

# %% [markdown]
# ## Limitações e próxima decisão
# As datas dos recortes continuam seriais Excel; não se convertem nesta etapa.
# Amostras nacionais de 20 linhas não representam os tipos de todo o arquivo.
# Ausência não significa zero; metadados auxiliares preservam ano/arquivo/linha.
# A etapa 6 abaixo foi autorizada para diagnóstico. Tratamento de ausências,
# tipagem, agregação, exportação e painel dependem de nova autorização.

# %% [markdown]
# # ETAPA 6 — Diagnóstico de duplicidades e qualidade Sinesp
# Pré-requisito: df_sinesp_joinville_concatenado já produzido pela etapa 5 no
# mesmo kernel. Esta etapa não lê arquivos nem repete a concatenação.
# As contagens de linhas e ausências são diagnósticas; não se somam medidas dos
# indicadores. A base é apenas consultada; quadros e máscaras são auxiliares.
# Execute cada célula e abra as variáveis em Variables > Data Viewer.

# %%
if "df_sinesp_joinville_concatenado" not in globals():
    raise RuntimeError("A etapa 6 requer o DataFrame da etapa 5 no mesmo kernel.")
assert df_sinesp_joinville_concatenado.shape == (900, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
retrato_antes_diagnostico_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)

# %% [markdown]
# ## Duplicidades exatas e possíveis repetições de competência
# data_referencia é a competência mensal original, incluindo o ano no serial.
# Usá-la diretamente evita converter datas ou criar mês/ano por suposição.
# A chave também mantém município/UF. Datas distintas não são agrupadas.
# Linhas repetidas na chave resumida são apenas candidatas: conferir agente,
# arma e faixa_etaria antes de interpretar. Valores ausentes entram na análise.

# %%
chave_competencia_sinesp = ["uf", "municipio", "data_referencia", "evento", "abrangencia"]
chave_dimensoes_sinesp = chave_competencia_sinesp + ["agente", "arma", "faixa_etaria"]
chave_sem_abrangencia_sinesp = ["uf", "municipio", "data_referencia", "evento"]
mascara_exatas_sinesp = df_sinesp_joinville_concatenado.duplicated(keep=False)
mascara_competencia_sinesp = df_sinesp_joinville_concatenado.duplicated(
    subset=chave_competencia_sinesp, keep=False
)
mascara_dimensoes_sinesp = df_sinesp_joinville_concatenado.duplicated(
    subset=chave_dimensoes_sinesp, keep=False
)
mascara_sem_abrangencia_sinesp = df_sinesp_joinville_concatenado.duplicated(
    subset=chave_sem_abrangencia_sinesp, keep=False
)
resumo_duplicidades_sinesp = pd.DataFrame([
    {
        "criterio": criterio,
        "linhas_envolvidas": int(mascara.sum()),
        "repeticoes_apos_primeira": int(df_sinesp_joinville_concatenado.duplicated(subset=chave).sum()),
        "grupos_repetidos": int((df_sinesp_joinville_concatenado.value_counts(subset=chave, dropna=False) > 1).sum()),
    }
    for criterio, chave, mascara in [
        ("14 colunas originais", transform.HEADERS, mascara_exatas_sinesp),
        ("competencia/evento/abrangencia/municipio", chave_competencia_sinesp, mascara_competencia_sinesp),
        ("chave com todas as dimensoes", chave_dimensoes_sinesp, mascara_dimensoes_sinesp),
        ("controle sem abrangencia", chave_sem_abrangencia_sinesp, mascara_sem_abrangencia_sinesp),
    ]
])
resumo_duplicidades_sinesp

# %%
linhas_duplicadas_exatas_sinesp = df_sinesp_joinville_concatenado.loc[mascara_exatas_sinesp]
linhas_duplicadas_exatas_sinesp

# %%
linhas_suspeitas_chave_sinesp = df_sinesp_joinville_concatenado.loc[mascara_competencia_sinesp]
linhas_suspeitas_chave_sinesp

# %%
# Controle: evidencia por que abrangencia precisa permanecer na chave.
linhas_repetidas_sem_abrangencia_sinesp = df_sinesp_joinville_concatenado.loc[mascara_sem_abrangencia_sinesp]
linhas_repetidas_sem_abrangencia_sinesp

# %%
# Metadados de ano/arquivo/linha ficam fora das 14 colunas, alinhados pelo índice.
# A identificação do mês permanece na competência original, sem conversão.
if "metadados_sinesp_concatenado" in globals():
    assert metadados_sinesp_concatenado.index.equals(df_sinesp_joinville_concatenado.index)
    origem_linhas_repetidas_sinesp = metadados_sinesp_concatenado.loc[mascara_sem_abrangencia_sinesp]
    origem_linhas_repetidas_sinesp

# %% [markdown]
# ## Ausências gerais e por indicador
# Percentuais usam a quantidade de linhas da base ou do evento correspondente.
# Nenhuma ausência é preenchida. Linhas com ausência são mantidas integralmente.

# %%
ausencias_colunas_sinesp = pd.DataFrame({
    "quantidade": df_sinesp_joinville_concatenado.isna().sum(),
    "percentual": df_sinesp_joinville_concatenado.isna().mean() * 100,
})
ausencias_colunas_sinesp

# %%
ausencias_eventos_linhas_sinesp = []
for evento_sinesp in df_sinesp_joinville_concatenado["evento"].unique():
    linhas_evento_sinesp = df_sinesp_joinville_concatenado.loc[
        df_sinesp_joinville_concatenado["evento"].eq(evento_sinesp)
    ]
    for campo_sinesp in df_sinesp_joinville_concatenado.columns:
        quantidade_ausente_sinesp = int(linhas_evento_sinesp[campo_sinesp].isna().sum())
        ausencias_eventos_linhas_sinesp.append({
            "evento": evento_sinesp, "campo": campo_sinesp,
            "linhas_evento": len(linhas_evento_sinesp),
            "ausencias": quantidade_ausente_sinesp,
            "percentual": quantidade_ausente_sinesp / len(linhas_evento_sinesp) * 100,
        })
ausencias_por_evento_sinesp = pd.DataFrame(ausencias_eventos_linhas_sinesp)
ausencias_por_evento_sinesp

# %%
# Matriz apenas para visualização das contagens, sem somar valores dos crimes.
matriz_ausencias_eventos_sinesp = ausencias_por_evento_sinesp.pivot(
    index="evento", columns="campo", values="ausencias"
)
matriz_ausencias_eventos_sinesp

# %%
# Recorte para inspecionar ausência de sexo em indicadores com total de vítimas.
# Trata-se de suspeita de detalhamento incompleto, não de erro comprovado.
campos_sexo_sinesp = ["feminino", "masculino", "nao_informado"]
mascara_sexo_ausente_sinesp = (
    df_sinesp_joinville_concatenado["total_vitima"].notna()
    & df_sinesp_joinville_concatenado[campos_sexo_sinesp].isna().any(axis=1)
)
linhas_sexo_ausente_sinesp = df_sinesp_joinville_concatenado.loc[mascara_sexo_ausente_sinesp]
linhas_sexo_ausente_sinesp

# %% [markdown]
# ## Tipos e necessidade futura de interpretar datas
# No recorte da etapa 5, data_referencia foi lida do XML como serial inteiro.
# A coluna permanece intacta. Não se chama to_datetime nem se calcula calendário.
# Futuramente: conferir o sistema de datas do XLSX, validar a regra de conversão
# e amostras por ano, mantendo referência ao valor original e à origem.
# float64 nas contagens pode decorrer de ausências, não de medidas fracionárias.

# %%
diagnostico_tipos_sinesp = df_sinesp_joinville_concatenado.dtypes.to_frame("tipo_pandas")
diagnostico_tipos_sinesp

# %%
amostra_datas_originais_sinesp = df_sinesp_joinville_concatenado[["data_referencia"]].head(12)
amostra_datas_originais_sinesp

# %% [markdown]
# ## Interpretação cautelosa e integridade
# Campos total_vitima e total atendem a indicadores diferentes; ausência de uma
# medida quando a outra se aplica é compatível com estrutura, sem autorizar zero.
# total_peso não é preenchido neste recorte; agente/arma/faixa_etaria não trazem
# detalhamento municipal. Campo inteiramente vazio não comprova defeito.
# Sexo ausente com total_vitima informado requer conferir original e definição;
# não inferir sexo, não calcular parcela residual e não declarar ausência = zero.
# A próxima etapa depende de autorização específica, após avaliar os quadros.

# %%
pd.testing.assert_frame_equal(
    df_sinesp_joinville_concatenado, retrato_antes_diagnostico_sinesp,
    check_dtype=True, check_exact=True,
)
assert df_sinesp_joinville_concatenado.isna().equals(retrato_antes_diagnostico_sinesp.isna())
assert df_sinesp_joinville_concatenado.shape == (900, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
print("Etapa 6: 900 linhas e 14 colunas preservadas; valores, tipos, índices e ausências intactos.")

# %% [markdown]
# # ETAPA 7 — Investigação e propostas de colunas e tipos
# Execute no mesmo kernel das etapas 5–6. Não há leitura de CSV histórico,
# concatenação, limpeza, seleção definitiva de colunas ou conversão da base.
# As tabelas auxiliares apenas apresentam evidências e recomendações.

# %%
for requisito_sinesp in (
    "df_sinesp_joinville_concatenado", "metadados_sinesp_concatenado",
    "arquivos_sinesp_por_ano", "amostras_sinesp_por_ano",
    "linhas_sexo_ausente_sinesp", "linhas_repetidas_sem_abrangencia_sinesp",
    "retrato_antes_diagnostico_sinesp", "ausencias_colunas_sinesp", "campos_sexo_sinesp",
):
    if requisito_sinesp not in globals():
        raise RuntimeError(f"Execute as etapas 5–6 no mesmo kernel: falta {requisito_sinesp}.")
assert df_sinesp_joinville_concatenado.shape == (900, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
assert pd.api.types.is_integer_dtype(df_sinesp_joinville_concatenado["data_referencia"])
assert metadados_sinesp_concatenado.index.equals(df_sinesp_joinville_concatenado.index)
pd.testing.assert_frame_equal(
    df_sinesp_joinville_concatenado, retrato_antes_diagnostico_sinesp,
    check_dtype=True, check_exact=True,
)
pd.testing.assert_frame_equal(
    linhas_sexo_ausente_sinesp,
    df_sinesp_joinville_concatenado.loc[
        df_sinesp_joinville_concatenado["total_vitima"].notna()
        & df_sinesp_joinville_concatenado[campos_sexo_sinesp].isna().any(axis=1)
    ], check_exact=True,
)
pd.testing.assert_frame_equal(
    linhas_repetidas_sem_abrangencia_sinesp,
    df_sinesp_joinville_concatenado.loc[
        df_sinesp_joinville_concatenado.duplicated(
            subset=["uf", "municipio", "data_referencia", "evento"], keep=False
        )
    ], check_exact=True,
)
retrato_antes_investigacao_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)
ausencias_confirmadas_etapa7_sinesp = pd.DataFrame({
    "quantidade": df_sinesp_joinville_concatenado.isna().sum(),
    "percentual": df_sinesp_joinville_concatenado.isna().mean() * 100,
})
pd.testing.assert_frame_equal(
    ausencias_confirmadas_etapa7_sinesp, ausencias_colunas_sinesp,
    check_exact=True,
)
ausencias_confirmadas_etapa7_sinesp

# %% [markdown]
# ## Rótulos de competência para inspeção, sem converter seriais da base
# A amostra read_excel da etapa 5 já contém datas interpretadas pelo leitor.
# Confrontamos cada uma com o serial da MESMA célula original (D2 a D21).
# Isso fornece um rótulo de apresentação; nenhum serial é convertido ou trocado.
# Meses sem correspondência comprovada ficam identificados para investigação.

# %%
rotulos_competencia_sinesp = {}
sistemas_datas_originais_sinesp = []
for ano_rotulo_sinesp, arquivo_rotulo_sinesp in arquivos_sinesp_por_ano.items():
    livro_rotulo_sinesp = load_workbook(arquivo_rotulo_sinesp, read_only=True)
    try:
        sistemas_datas_originais_sinesp.append({
            "ano": ano_rotulo_sinesp, "epoch_xlsx": repr(livro_rotulo_sinesp.epoch),
        })
    finally:
        livro_rotulo_sinesp.close()
    with ZipFile(arquivo_rotulo_sinesp) as pacote_rotulo_sinesp:
        with pacote_rotulo_sinesp.open("xl/worksheets/sheet1.xml") as fluxo_rotulo_sinesp:
            prefixo_rotulo_sinesp = b""
            while prefixo_rotulo_sinesp.count(b"</row>") < 21:
                trecho_rotulo_sinesp = fluxo_rotulo_sinesp.read(65536)
                if not trecho_rotulo_sinesp:
                    break
                prefixo_rotulo_sinesp += trecho_rotulo_sinesp
    for correspondencia_rotulo_sinesp in re.finditer(
        rb'<row\b[^>]*\br="(\d+)"[^>]*>.*?</row>', prefixo_rotulo_sinesp, re.DOTALL
    ):
        linha_rotulo_sinesp = int(correspondencia_rotulo_sinesp.group(1))
        if not 2 <= linha_rotulo_sinesp <= 21:
            continue
        elemento_rotulo_sinesp = transform._parse_sheet_row(correspondencia_rotulo_sinesp.group())
        celula_data_rotulo_sinesp = next(
            c for c in elemento_rotulo_sinesp.findall(transform.NS + "c")
            if c.get("r") == f"D{linha_rotulo_sinesp}"
        )
        serial_rotulo_sinesp = transform._cell_value(celula_data_rotulo_sinesp, [])
        data_amostra_rotulo_sinesp = amostras_sinesp_por_ano[ano_rotulo_sinesp].iloc[
            linha_rotulo_sinesp - 2
        ]["data_referencia"]
        if not isinstance(data_amostra_rotulo_sinesp, pd.Timestamp):
            raise ValueError("Data da amostra não interpretada como Timestamp; revisar o leitor.")
        chave_rotulo_sinesp = (ano_rotulo_sinesp, serial_rotulo_sinesp)
        rotulo_competencia_sinesp = data_amostra_rotulo_sinesp.strftime("%Y-%m")
        if chave_rotulo_sinesp in rotulos_competencia_sinesp:
            assert rotulos_competencia_sinesp[chave_rotulo_sinesp] == rotulo_competencia_sinesp
        rotulos_competencia_sinesp[chave_rotulo_sinesp] = rotulo_competencia_sinesp
sistemas_datas_sinesp = pd.DataFrame(sistemas_datas_originais_sinesp)
sistemas_datas_sinesp

# %%
investigacao_sexo_linhas_sinesp = []
for indice_sexo_sinesp, linha_sexo_sinesp in linhas_sexo_ausente_sinesp.iterrows():
    origem_sexo_sinesp = metadados_sinesp_concatenado.loc[indice_sexo_sinesp]
    ano_sexo_sinesp = origem_sexo_sinesp["ano_arquivo"]
    parcelas_informadas_sinesp = [
        linha_sexo_sinesp[campo] for campo in campos_sexo_sinesp
        if pd.notna(linha_sexo_sinesp[campo])
    ]
    soma_informada_sinesp = sum(parcelas_informadas_sinesp) if parcelas_informadas_sinesp else None
    if not parcelas_informadas_sinesp:
        relacao_sexo_sinesp = "Sem parcelas de sexo; reconciliação indisponível."
    elif soma_informada_sinesp == linha_sexo_sinesp["total_vitima"]:
        relacao_sexo_sinesp = "Parcelas informadas coincidem com total; não inferir ausentes."
    else:
        relacao_sexo_sinesp = "Parcelas informadas diferem do total; investigar definição e preenchimento."
    investigacao_sexo_linhas_sinesp.append({
        "indice_base": indice_sexo_sinesp, "ano": ano_sexo_sinesp,
        "competencia": rotulos_competencia_sinesp.get(
            (ano_sexo_sinesp, linha_sexo_sinesp["data_referencia"]), "SEM RÓTULO VALIDADO"
        ),
        "serial_original": linha_sexo_sinesp["data_referencia"],
        "evento": linha_sexo_sinesp["evento"], "abrangencia": linha_sexo_sinesp["abrangencia"],
        **{campo: linha_sexo_sinesp[campo] for campo in campos_sexo_sinesp},
        "total_vitima": linha_sexo_sinesp["total_vitima"],
        "soma_apenas_informados": soma_informada_sinesp,
        "relacao": relacao_sexo_sinesp,
        "arquivo_origem": origem_sexo_sinesp["arquivo_origem"],
        "linha_xlsx": origem_sexo_sinesp["linha_xlsx"],
    })
investigacao_sexo_sinesp = pd.DataFrame(investigacao_sexo_linhas_sinesp)
investigacao_sexo_sinesp

# %% [markdown]
# ## Mandados estaduais e Polícia Federal: comparação sem somar
# Mantém os 40 registros em um quadro e apresenta cada par em colunas separadas.
# A abrangência distingue as séries; exclusividade estatística não é presumida.

# %%
mandados_investigacao_sinesp = pd.DataFrame([
    {
        "indice_base": indice,
        "ano": metadados_sinesp_concatenado.loc[indice, "ano_arquivo"],
        "competencia": rotulos_competencia_sinesp.get(
            (metadados_sinesp_concatenado.loc[indice, "ano_arquivo"], linha["data_referencia"]),
            "SEM RÓTULO VALIDADO",
        ),
        "serial_original": linha["data_referencia"], "evento": linha["evento"],
        "abrangencia": linha["abrangencia"], "total": linha["total"],
        "agente": linha["agente"], "arma": linha["arma"], "faixa_etaria": linha["faixa_etaria"],
        "arquivo_origem": metadados_sinesp_concatenado.loc[indice, "arquivo_origem"],
        "linha_xlsx": metadados_sinesp_concatenado.loc[indice, "linha_xlsx"],
    }
    for indice, linha in linhas_repetidas_sem_abrangencia_sinesp.iterrows()
])
mandados_investigacao_sinesp

# %%
pares_mandados_sinesp = mandados_investigacao_sinesp.pivot(
    index=["ano", "competencia", "serial_original", "evento"],
    columns="abrangencia", values="total",
)
pares_mandados_sinesp

# %% [markdown]
# ## Avaliação das 14 colunas e proposta de tipagem (sem aplicação)
# A seleção futura será uma visão analítica, mantendo sempre a camada original.
# Os campos inteiramente vazios são candidatos apenas nessa visão e dependem
# de verificar outros indicadores/atualizações. Nenhuma coluna é removida aqui.

# %%
propostas_campos_sinesp = {
    "uf": ("Manter", "Território estadual e filtro exato", "municipio", "string", "Impede misturar municípios homônimos."),
    "municipio": ("Manter", "Território municipal", "uf", "string", "Identifica o recorte; não autoriza desagregar bairros."),
    "evento": ("Manter", "Identificador original do indicador", "abrangencia e campos de medida", "string; category apenas após validar domínio", "Preservar rótulos e distinguir unidades e sobreposições."),
    "data_referencia": ("Manter", "Competência mensal", "sistema de datas e arquivo/ano de origem", "datetime64[ns], após validar conversão", "Manter referência ao serial original para rastreabilidade."),
    "agente": ("Candidata à exclusão da visão futura", "Possível dimensão de agentes; dicionário pendente", "evento e abrangencia", "string anulável se mantida", "100% ausente neste recorte; não remover da camada original."),
    "arma": ("Candidata à exclusão da visão futura", "Possível dimensão de arma; dicionário pendente", "evento", "string anulável se mantida", "100% ausente; avaliar novos indicadores antes de selecionar colunas."),
    "faixa_etaria": ("Candidata à exclusão da visão futura", "Possível dimensão etária; dicionário pendente", "evento e medida de vítimas", "string anulável se mantida", "Não oferece detalhamento neste recorte; preservar no original."),
    "feminino": ("Manter e avaliar ausências", "Parcela de vítimas por sexo", "evento e total_vitima", "Int64 anulável", "Validar integralidade, não negatividade e significado das ausências."),
    "masculino": ("Manter e avaliar ausências", "Parcela de vítimas por sexo", "evento e total_vitima", "Int64 anulável", "Não inferir zero a partir da soma das parcelas informadas."),
    "nao_informado": ("Manter e avaliar ausências", "Parcela de vítimas com sexo não informado", "evento e total_vitima", "Int64 anulável", "Ausência da célula difere da categoria sexo não informado."),
    "total_vitima": ("Manter", "Total publicado de vítimas", "evento, sexo e abrangencia", "Int64 anulável", "Não substituir por total nem somar novamente às parcelas."),
    "total": ("Manter", "Contagem publicada de mandados no recorte", "evento e abrangencia", "Int64 anulável", "Manter estadual e PF separados; não representa vítimas únicas."),
    "total_peso": ("Candidata à exclusão da visão futura", "Medida de peso; unidade pendente", "evento e unidade oficial", "decimal anulável se houver uso e unidade confirmada", "100% ausente no recorte; não presumir quilogramas ou preencher."),
    "abrangencia": ("Manter", "Distingue séries estaduais, PF e PRF", "evento e competência", "string; category após validar domínio", "É necessária à chave; não identifica corporação individual em toda série."),
}
proposta_colunas_sinesp = pd.DataFrame([
    {
        "coluna": campo, "tipo_atual": str(df_sinesp_joinville_concatenado[campo].dtype),
        "percentual_ausente": ausencias_confirmadas_etapa7_sinesp.loc[campo, "percentual"],
        "exemplos_originais": repr(df_sinesp_joinville_concatenado[campo].unique()[:3].tolist()),
        "recomendacao": proposta[0], "finalidade": proposta[1],
        "dependencias": proposta[2], "tipo_proposto": proposta[3], "justificativa": proposta[4],
    }
    for campo, proposta in propostas_campos_sinesp.items()
])
proposta_colunas_sinesp

# %% [markdown]
# ## Validações necessárias antes de qualquer conversão futura
# Datas: validar epoch 1900/1904 em todos os XLSX, regra de dias seriais e
# particularidade do calendário Excel; conferir ausência, frações, limites,
# correspondência ano/mês e primeiro dia do mês. Confrontar o resultado futuro
# com células lidas como datas pelo Excel/openpyxl e preservar a origem.
# Contagens: conferir valores inteiros, finitos e não negativos; preservar nulos
# e categorias por unidade. Só depois propor Int64, sem preencher as ausências.
# Texto: manter grafia original; eventual categoria exige domínio validado.
# Esta etapa não aplica nenhuma dessas conversões ou seleções.

# %%
pd.testing.assert_frame_equal(
    df_sinesp_joinville_concatenado, retrato_antes_investigacao_sinesp,
    check_dtype=True, check_exact=True,
)
assert df_sinesp_joinville_concatenado.isna().equals(retrato_antes_investigacao_sinesp.isna())
assert df_sinesp_joinville_concatenado.shape == (900, 14)
print("Etapa 7: investigação concluída em memória; 900 linhas, 14 colunas e valores originais preservados.")

# %% [markdown]
# # ETAPA 8 — Preparação controlada da cópia analítica Sinesp
# Execute depois das etapas 5–7 no mesmo processo/kernel. A etapa 5 reconstrói
# os recortes pelos sete XLSX, sem usar CSV histórico. Nesta etapa, qualquer
# divergência interrompe a execução antes de disponibilizar a base analítica.
# A seleção usa .loc[..., colunas].copy(deep=True); a base original de 14
# colunas e seus metadados nunca recebem atribuições de valores ou tipos.

# %%
import hashlib
import math


def exigir_sinesp_etapa8(condicao, mensagem):
    if not condicao:
        raise ValueError(f"Etapa 8 interrompida: {mensagem}")


def sha256_original_sinesp(caminho):
    resumo = hashlib.sha256()
    with caminho.open("rb") as fluxo:
        for bloco in iter(lambda: fluxo.read(1024 * 1024), b""):
            resumo.update(bloco)
    return resumo.hexdigest()


for requisito_etapa8 in (
    "df_sinesp_joinville_concatenado", "metadados_sinesp_concatenado",
    "recortes_sinesp_por_ano", "arquivos_sinesp_por_ano",
    "amostras_sinesp_por_ano", "rotulos_competencia_sinesp",
    "linhas_sexo_ausente_sinesp", "linhas_repetidas_sem_abrangencia_sinesp",
):
    exigir_sinesp_etapa8(requisito_etapa8 in globals(), f"falta {requisito_etapa8}; execute etapas 5–7.")
exigir_sinesp_etapa8(df_sinesp_joinville_concatenado.shape == (900, 14), "dimensões originais divergentes.")
exigir_sinesp_etapa8(df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS, "esquema original divergente.")
original_preservado_etapa8 = df_sinesp_joinville_concatenado.copy(deep=True)
metadados_preservados_etapa8 = metadados_sinesp_concatenado.copy(deep=True)

# Confere hashes contra o inventário anterior, não apenas antes/depois desta sessão.
inventario_hash_sinesp_etapa8 = pd.read_csv(
    RAIZ_PROJETO / "documentacao/auditoria/inventario_sha256_pre_reorganizacao.csv"
)
hashes_originais_etapa8 = {}
epocas_originais_etapa8 = {}
for ano_etapa8, caminho_etapa8 in arquivos_sinesp_por_ano.items():
    relativo_etapa8 = caminho_etapa8.relative_to(RAIZ_PROJETO).as_posix()
    evidencia_hash_etapa8 = inventario_hash_sinesp_etapa8.loc[
        inventario_hash_sinesp_etapa8["caminho_destino_planejado"].eq(relativo_etapa8)
    ]
    exigir_sinesp_etapa8(len(evidencia_hash_etapa8) == 1, f"evidência de hash ausente/ambígua: {relativo_etapa8}")
    hash_etapa8 = sha256_original_sinesp(caminho_etapa8)
    exigir_sinesp_etapa8(hash_etapa8 == evidencia_hash_etapa8.iloc[0]["sha256"], f"hash divergente: {relativo_etapa8}")
    exigir_sinesp_etapa8(caminho_etapa8.stat().st_size == evidencia_hash_etapa8.iloc[0]["tamanho_bytes"], f"tamanho divergente: {relativo_etapa8}")
    hashes_originais_etapa8[ano_etapa8] = hash_etapa8
    livro_epoca_etapa8 = load_workbook(caminho_etapa8, read_only=True)
    try:
        epocas_originais_etapa8[ano_etapa8] = livro_epoca_etapa8.epoch
    finally:
        livro_epoca_etapa8.close()
exigir_sinesp_etapa8(
    all(epoca == pd.Timestamp("1899-12-30") for epoca in epocas_originais_etapa8.values()),
    "sistema de datas diferente do Excel 1900 previamente identificado.",
)

# Valida o alinhamento de origem e cada bloco da reconstrução da etapa 5.
exigir_sinesp_etapa8(metadados_sinesp_concatenado.shape == (900, 4), "dimensões dos metadados divergentes.")
exigir_sinesp_etapa8(metadados_sinesp_concatenado.index.equals(original_preservado_etapa8.index), "índices de origem desalinhados.")
exigir_sinesp_etapa8(not metadados_sinesp_concatenado.isna().any().any(), "metadados incompletos.")
exigir_sinesp_etapa8(not metadados_sinesp_concatenado.duplicated(["arquivo_origem", "planilha_origem", "linha_xlsx"]).any(), "linhas de origem repetidas.")
contagens_referencia_etapa8 = {2020: 132, 2021: 132, 2022: 132, 2023: 132, 2024: 132, 2025: 144, 2026: 96}
eventos_referencia_etapa8 = {
    "Feminicídio": 80, "Homicídio doloso": 80, "Lesão corporal seguida de morte": 80,
    "Mandado de prisão cumprido": 100,
    "Morte no trânsito ou em decorrência dele (exceto homicídio doloso)": 80,
    "Mortes a esclarecer (sem indício de crime)": 80, "Mortes no trânsito": 80,
    "Roubo seguido de morte (latrocínio)": 80, "Suicídio": 80,
    "Tentativa de feminicídio": 80, "Tentativa de homicídio": 80,
}
exigir_sinesp_etapa8(original_preservado_etapa8["evento"].value_counts().to_dict() == eventos_referencia_etapa8, "indicadores/frequências diferentes do inventário.")
exigir_sinesp_etapa8(original_preservado_etapa8["uf"].eq("SC").all() and original_preservado_etapa8["municipio"].eq("JOINVILLE").all(), "recorte municipal divergente.")
inicio_etapa8 = 0
for ano_etapa8, quantidade_etapa8 in contagens_referencia_etapa8.items():
    bloco_etapa8 = original_preservado_etapa8.iloc[inicio_etapa8:inicio_etapa8 + quantidade_etapa8]
    origem_bloco_etapa8 = metadados_sinesp_concatenado.iloc[inicio_etapa8:inicio_etapa8 + quantidade_etapa8]
    pd.testing.assert_frame_equal(bloco_etapa8.reset_index(drop=True), recortes_sinesp_por_ano[ano_etapa8], check_exact=True)
    exigir_sinesp_etapa8(origem_bloco_etapa8["ano_arquivo"].eq(ano_etapa8).all(), "ano de origem desalinhado.")
    exigir_sinesp_etapa8(origem_bloco_etapa8["planilha_origem"].eq(str(ano_etapa8)).all(), "planilha de origem divergente.")
    exigir_sinesp_etapa8(origem_bloco_etapa8["arquivo_origem"].eq(str(arquivos_sinesp_por_ano[ano_etapa8].relative_to(RAIZ_PROJETO))).all(), "arquivo de origem divergente.")
    exigir_sinesp_etapa8(origem_bloco_etapa8["linha_xlsx"].ge(2).all(), "linha de origem inválida.")
    inicio_etapa8 += quantidade_etapa8
print("Etapa 8: pré-condições aprovadas; sete originais íntegros, 900 linhas municipais e origem alinhada.")

# %% [markdown]
# ## Validar datas antes de atribuí-las à cópia
# pd.to_datetime(unit="D", origin=..., errors="raise") interpreta o serial.
# A série temporária precisa passar por TODAS as verificações abaixo. Só então
# será atribuída à cópia. Não se usa errors="coerce" ou correção de datas.

# %%
seriais_etapa8 = original_preservado_etapa8["data_referencia"]
exigir_sinesp_etapa8(seriais_etapa8.notna().all(), "serial de data ausente.")
exigir_sinesp_etapa8(all(
    isinstance(valor, (int, float)) and not isinstance(valor, bool)
    and math.isfinite(valor) and valor == int(valor) and 0 < valor < 100000
    for valor in seriais_etapa8.tolist()
), "serial inválido, fracionário ou fora dos limites previstos.")
datas_validadas_etapa8 = pd.to_datetime(seriais_etapa8, unit="D", origin="1899-12-30", errors="raise")
exigir_sinesp_etapa8(datas_validadas_etapa8.dt.year.eq(metadados_sinesp_concatenado["ano_arquivo"]).all(), "data não corresponde ao ano do arquivo.")
exigir_sinesp_etapa8(datas_validadas_etapa8.dt.day.eq(1).all(), "competência fora do primeiro dia do mês.")
exigir_sinesp_etapa8(datas_validadas_etapa8.eq(datas_validadas_etapa8.dt.normalize()).all(), "horário inesperado na competência.")
datas_leitor_etapa8 = {}
for ano_etapa8, amostra_etapa8 in amostras_sinesp_por_ano.items():
    for data_leitor_etapa8 in amostra_etapa8["data_referencia"]:
        chave_leitor_etapa8 = (ano_etapa8, data_leitor_etapa8.strftime("%Y-%m"))
        if chave_leitor_etapa8 in datas_leitor_etapa8:
            exigir_sinesp_etapa8(datas_leitor_etapa8[chave_leitor_etapa8] == data_leitor_etapa8, "amostra Excel tem datas conflitantes.")
        datas_leitor_etapa8[chave_leitor_etapa8] = data_leitor_etapa8
for indice_etapa8, data_etapa8 in datas_validadas_etapa8.items():
    ano_etapa8 = metadados_sinesp_concatenado.loc[indice_etapa8, "ano_arquivo"]
    competencia_etapa8 = data_etapa8.strftime("%Y-%m")
    exigir_sinesp_etapa8(rotulos_competencia_sinesp.get((ano_etapa8, seriais_etapa8.loc[indice_etapa8])) == competencia_etapa8, "serial não corresponde à célula Excel amostrada.")
    exigir_sinesp_etapa8(datas_leitor_etapa8.get((ano_etapa8, competencia_etapa8)) == data_etapa8, "data difere da interpretação do leitor Excel.")
for ano_etapa8 in contagens_referencia_etapa8:
    meses_etapa8 = sorted(datas_validadas_etapa8.loc[datas_validadas_etapa8.dt.year.eq(ano_etapa8)].dt.month.unique().tolist())
    exigir_sinesp_etapa8(meses_etapa8 == list(range(1, 9 if ano_etapa8 == 2026 else 13)), f"cobertura mensal divergente em {ano_etapa8}.")

# %% [markdown]
# ## Validar contagens antes da tipagem anulável
# Apenas os valores preenchidos são examinados. Não há arredondamento, soma de
# indicadores ou substituição de ausência. Int64 mantém NA em cada posição.

# %%
campos_contagens_etapa8 = ["feminino", "masculino", "nao_informado", "total_vitima", "total"]
for campo_etapa8 in campos_contagens_etapa8:
    valores_etapa8 = original_preservado_etapa8.loc[original_preservado_etapa8[campo_etapa8].notna(), campo_etapa8]
    exigir_sinesp_etapa8(all(
        math.isfinite(valor) and valor == int(valor) and 0 <= valor <= 2**63 - 1
        for valor in valores_etapa8
    ), f"contagem não finita, fracionária, negativa ou fora de Int64: {campo_etapa8}")

# %% [markdown]
# ## Criar a candidata de 900 × 10, sem atribuir à base original
# A variável definitiva só é disponibilizada depois das conferências finais.
# Metadados ficam separados, conservando o serial e o hash do arquivo de origem.

# %%
colunas_analiticas_sinesp = [
    "uf", "municipio", "evento", "data_referencia", "feminino", "masculino",
    "nao_informado", "total_vitima", "total", "abrangencia",
]
candidata_analitica_sinesp = df_sinesp_joinville_concatenado.loc[:, colunas_analiticas_sinesp].copy(deep=True)
candidata_analitica_sinesp["data_referencia"] = datas_validadas_etapa8
for campo_etapa8 in campos_contagens_etapa8:
    candidata_analitica_sinesp[campo_etapa8] = candidata_analitica_sinesp[campo_etapa8].astype("Int64")
metadados_candidatos_sinesp = metadados_sinesp_concatenado.copy(deep=True)
metadados_candidatos_sinesp["serial_original"] = seriais_etapa8.copy(deep=True)
metadados_candidatos_sinesp["sha256_original"] = metadados_candidatos_sinesp["ano_arquivo"].map(hashes_originais_etapa8)

# %% [markdown]
# ## Catálogo auxiliar: evento e abrangência definem medida e limitações
# Um catálogo não agrega valores. Feminicídios e tentativas mantêm rótulos
# próprios. Mandados usam total; os demais eventos deste recorte usam vítimas.

# %%
naturezas_eventos_etapa8 = {
    "Feminicídio": "criminal", "Homicídio doloso": "criminal",
    "Lesão corporal seguida de morte": "criminal", "Mandado de prisão cumprido": "atividade policial",
    "Morte no trânsito ou em decorrência dele (exceto homicídio doloso)": "complementar: mortalidade no trânsito",
    "Mortes a esclarecer (sem indício de crime)": "complementar: mortes a esclarecer",
    "Mortes no trânsito": "complementar: mortalidade no trânsito",
    "Roubo seguido de morte (latrocínio)": "criminal", "Suicídio": "complementar: suicídio",
    "Tentativa de feminicídio": "criminal", "Tentativa de homicídio": "criminal",
}
catalogo_linhas_etapa8 = []
for evento_etapa8 in original_preservado_etapa8["evento"].unique():
    for abrangencia_etapa8 in original_preservado_etapa8.loc[original_preservado_etapa8["evento"].eq(evento_etapa8), "abrangencia"].unique():
        mandado_etapa8 = evento_etapa8 == "Mandado de prisão cumprido"
        catalogo_linhas_etapa8.append({
            "evento_original": evento_etapa8, "abrangencia": abrangencia_etapa8,
            "medida_utilizada": "total" if mandado_etapa8 else "total_vitima",
            "unidade_medida": "mandados cumpridos (contagem publicada)" if mandado_etapa8 else "vítimas",
            "natureza": naturezas_eventos_etapa8[evento_etapa8],
            "limitacoes": (
                "Não representa pessoas únicas; estadual e PF separados; exclusividade não demonstrada."
                if mandado_etapa8 else
                "Não somar às parcelas de sexo; manter categoria e abrangência; sobreposição entre indicadores/fontes requer definição."
            ) + " 2026 parcial; linha tabular não equivale a ocorrência individual.",
        })
catalogo_candidato_sinesp = pd.DataFrame(catalogo_linhas_etapa8)
exigir_sinesp_etapa8(len(catalogo_candidato_sinesp) == 12, "combinações evento/abrangência inesperadas.")
mascara_mandados_etapa8 = candidata_analitica_sinesp["evento"].eq("Mandado de prisão cumprido")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[mascara_mandados_etapa8, "total"].notna().all(), "mandados sem medida total.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[mascara_mandados_etapa8, "total_vitima"].isna().all(), "medidas de mandados/vítimas inesperadas.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[~mascara_mandados_etapa8, "total_vitima"].notna().all(), "indicadores de vítimas sem total.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[~mascara_mandados_etapa8, "total"].isna().all(), "medida total fora dos mandados.")

# %% [markdown]
# ## Conferências finais: valores, nulos, linhas, tipos e origens
# Compara máscaras de ausência e valores preenchidos, sem somar medidas.
# As 20 linhas com sexo incompleto e os 20 pares estaduais/PF são preservados.

# %%
exigir_sinesp_etapa8(candidata_analitica_sinesp.shape == (900, 10), "dimensões analíticas divergentes.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.index.equals(original_preservado_etapa8.index), "índices analíticos divergentes.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.isna().equals(original_preservado_etapa8[colunas_analiticas_sinesp].isna()), "posições de ausência alteradas.")
for campo_etapa8 in campos_contagens_etapa8:
    preenchidos_etapa8 = original_preservado_etapa8[campo_etapa8].notna()
    exigir_sinesp_etapa8(str(candidata_analitica_sinesp[campo_etapa8].dtype) == "Int64", "tipagem de contagem divergente.")
    exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[preenchidos_etapa8, campo_etapa8].eq(original_preservado_etapa8.loc[preenchidos_etapa8, campo_etapa8]).all(), f"valores alterados: {campo_etapa8}")
for campo_etapa8 in ["uf", "municipio", "evento", "abrangencia"]:
    pd.testing.assert_series_equal(candidata_analitica_sinesp[campo_etapa8], original_preservado_etapa8[campo_etapa8], check_exact=True)
mascara_sexo_etapa8 = candidata_analitica_sinesp["total_vitima"].notna() & candidata_analitica_sinesp[campos_sexo_sinesp].isna().any(axis=1)
exigir_sinesp_etapa8(int(mascara_sexo_etapa8.sum()) == 20, "quantidade de linhas com sexo incompleto divergente.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.index[mascara_sexo_etapa8].equals(linhas_sexo_ausente_sinesp.index), "linhas com sexo incompleto alteradas.")
mascara_pares_etapa8 = candidata_analitica_sinesp.duplicated(["uf", "municipio", "data_referencia", "evento"], keep=False)
exigir_sinesp_etapa8(int(mascara_pares_etapa8.sum()) == 40, "pares de mandados divergentes.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.index[mascara_pares_etapa8].equals(linhas_repetidas_sem_abrangencia_sinesp.index), "pares de mandados alterados.")
for ano_etapa8, caminho_etapa8 in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa8(sha256_original_sinesp(caminho_etapa8) == hashes_originais_etapa8[ano_etapa8], "arquivo original alterado durante a preparação.")
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_preservado_etapa8, check_dtype=True, check_exact=True)
pd.testing.assert_frame_equal(metadados_sinesp_concatenado, metadados_preservados_etapa8, check_dtype=True, check_exact=True)
pd.testing.assert_frame_equal(metadados_candidatos_sinesp[metadados_preservados_etapa8.columns], metadados_preservados_etapa8, check_exact=True)
exigir_sinesp_etapa8(metadados_candidatos_sinesp.index.equals(candidata_analitica_sinesp.index), "metadados analíticos desalinhados.")
df_sinesp_joinville_analitico = candidata_analitica_sinesp
metadados_sinesp_analitico = metadados_candidatos_sinesp
catalogo_indicadores_sinesp = catalogo_candidato_sinesp
print("Etapa 8 aprovada: original 900 × 14 intacto; analítico 900 × 10; sete hashes preservados.")

# %%
comparacao_tipos_etapa8_sinesp = pd.DataFrame({
    "antes": original_preservado_etapa8[colunas_analiticas_sinesp].dtypes,
    "depois": df_sinesp_joinville_analitico.dtypes,
})
comparacao_tipos_etapa8_sinesp

# %%
comparacao_ausencias_etapa8_sinesp = pd.DataFrame({
    "antes": original_preservado_etapa8[colunas_analiticas_sinesp].isna().sum(),
    "depois": df_sinesp_joinville_analitico.isna().sum(),
})
comparacao_ausencias_etapa8_sinesp

# %%
conferencia_datas_etapa8_sinesp = pd.DataFrame({
    "ano_arquivo": metadados_sinesp_analitico["ano_arquivo"],
    "serial_original": metadados_sinesp_analitico["serial_original"],
    "data_analitica": df_sinesp_joinville_analitico["data_referencia"],
})
conferencia_datas_etapa8_sinesp

# %%
sexo_incompleto_analitico_sinesp = df_sinesp_joinville_analitico.loc[mascara_sexo_etapa8]
sexo_incompleto_analitico_sinesp

# %%
mandados_pares_analitico_sinesp = df_sinesp_joinville_analitico.loc[mascara_pares_etapa8]
mandados_pares_analitico_sinesp

# %%
df_sinesp_joinville_analitico

# %%
catalogo_indicadores_sinesp

# %%
metadados_sinesp_analitico

# %% [markdown]
# # ETAPA 9 — Análise exploratória descritiva Sinesp
# Reconstrua as etapas 5–8 no mesmo processo/kernel. A análise consulta a base
# analítica sem novos tratamentos. Fonte, indicador, abrangência e unidade
# acompanham cada resultado. Não há total geral de crimes nem taxas populacionais.

# %%
def exigir_sinesp_etapa9(condicao, mensagem):
    if not condicao:
        raise ValueError(f"Etapa 9 interrompida: {mensagem}")


for requisito_etapa9 in (
    "df_sinesp_joinville_concatenado", "df_sinesp_joinville_analitico",
    "catalogo_indicadores_sinesp", "metadados_sinesp_analitico",
    "hashes_originais_etapa8", "original_preservado_etapa8",
):
    exigir_sinesp_etapa9(requisito_etapa9 in globals(), f"execute etapas 5–8: falta {requisito_etapa9}.")
exigir_sinesp_etapa9(df_sinesp_joinville_concatenado.shape == (900, 14), "dimensões originais divergentes.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico.shape == (900, 10), "dimensões analíticas divergentes.")
exigir_sinesp_etapa9(pd.api.types.is_datetime64_any_dtype(df_sinesp_joinville_analitico["data_referencia"]), "datas não tipadas.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico["data_referencia"].notna().all(), "datas ausentes.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico["data_referencia"].dt.day.eq(1).all(), "competência mensal inesperada.")
exigir_sinesp_etapa9(all(str(df_sinesp_joinville_analitico[c].dtype) == "Int64" for c in campos_contagens_etapa8), "medidas não anuláveis Int64.")
exigir_sinesp_etapa9(not df_sinesp_joinville_analitico.duplicated(["uf", "municipio", "evento", "abrangencia", "data_referencia"]).any(), "repetição de competência na mesma série.")
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_preservado_etapa8, check_exact=True)
pd.testing.assert_frame_equal(catalogo_indicadores_sinesp, catalogo_candidato_sinesp, check_exact=True)
for ano_etapa9, caminho_etapa9 in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa9(sha256_original_sinesp(caminho_etapa9) == hashes_originais_etapa8[ano_etapa9], "arquivo original com hash divergente.")
original_antes_eda_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)
analitico_antes_eda_sinesp = df_sinesp_joinville_analitico.copy(deep=True)
catalogo_antes_eda_sinesp = catalogo_indicadores_sinesp.copy(deep=True)
metadados_antes_eda_sinesp = metadados_sinesp_analitico.copy(deep=True)

# %% [markdown]
# ## Tabela mensal de observações publicadas
# Consulta o campo de medida definido pelo catálogo: vítimas em total_vitima;
# mandados em total. Cada linha continua sendo uma observação mensal agregada,
# não uma ocorrência individual. Os campos ano/mês pertencem só à apresentação.

# %%
catalogo_consulta_eda_sinesp = {
    (linha.evento_original, linha.abrangencia): linha
    for linha in catalogo_indicadores_sinesp.itertuples(index=False)
}
linhas_mensais_eda_sinesp = []
for linha_eda_sinesp in df_sinesp_joinville_analitico.itertuples(index=False):
    chave_eda_sinesp = (linha_eda_sinesp.evento, linha_eda_sinesp.abrangencia)
    exigir_sinesp_etapa9(chave_eda_sinesp in catalogo_consulta_eda_sinesp, "indicador sem catálogo.")
    definicao_eda_sinesp = catalogo_consulta_eda_sinesp[chave_eda_sinesp]
    linhas_mensais_eda_sinesp.append({
        "fonte": "MJSP / Sinesp VDE", "uf": linha_eda_sinesp.uf,
        "municipio": linha_eda_sinesp.municipio, "evento": linha_eda_sinesp.evento,
        "abrangencia": linha_eda_sinesp.abrangencia,
        "medida": definicao_eda_sinesp.medida_utilizada,
        "unidade": definicao_eda_sinesp.unidade_medida, "natureza": definicao_eda_sinesp.natureza,
        "competencia": linha_eda_sinesp.data_referencia,
        "ano": linha_eda_sinesp.data_referencia.year, "mes": linha_eda_sinesp.data_referencia.month,
        "valor_publicado": getattr(linha_eda_sinesp, definicao_eda_sinesp.medida_utilizada),
    })
tabela_mensal_sinesp = pd.DataFrame(linhas_mensais_eda_sinesp)
exigir_sinesp_etapa9(len(tabela_mensal_sinesp) == 900, "observações perdidas na apresentação mensal.")
chaves_series_eda_sinesp = ["fonte", "uf", "municipio", "evento", "abrangencia", "medida", "unidade", "natureza"]
tabela_mensal_sinesp

# %% [markdown]
# ## Estatísticas mensais por série e ano
# Soma com min_count=1 preserva ausência quando não há medida preenchida.
# Média/mediana usam meses com medida; os dois denominadores ficam explícitos.
# Desvio-padrão descritivo com ddof=0 representa a dispersão dos meses observados
# no recorte, sem inferência amostral; só é calculado com pelo menos dois valores.
# Meses de valor zero são observações publicadas, não ausências preenchidas.

# %%
def descrever_meses_sinesp(grupo):
    valores = grupo["valor_publicado"]
    n_validos = int(valores.count())
    return {
        "meses_disponiveis": int(grupo["competencia"].nunique()),
        "meses_com_medida": n_validos,
        "soma_medida_publicada": valores.sum(min_count=1),
        "media_mensal": valores.mean() if n_validos else pd.NA,
        "mediana_mensal": valores.median() if n_validos else pd.NA,
        "minimo_mensal": valores.min() if n_validos else pd.NA,
        "maximo_mensal": valores.max() if n_validos else pd.NA,
        "desvio_padrao_mensal_ddof0": valores.std(ddof=0) if n_validos >= 2 else pd.NA,
        "meses_zero_publicado": int(valores.eq(0).sum()),
    }


linhas_anuais_eda_sinesp = []
linhas_equivalentes_eda_sinesp = []
for chave_ano_eda_sinesp, grupo_ano_eda_sinesp in tabela_mensal_sinesp.groupby(chaves_series_eda_sinesp + ["ano"], sort=False, dropna=False):
    identificacao_ano_eda_sinesp = dict(zip(chaves_series_eda_sinesp + ["ano"], chave_ano_eda_sinesp))
    meses_ano_eda_sinesp = tuple(sorted(grupo_ano_eda_sinesp["mes"].unique().tolist()))
    cobertura_ano_eda_sinesp = "ano completo" if meses_ano_eda_sinesp == tuple(range(1, 13)) else "parcial janeiro–agosto" if meses_ano_eda_sinesp == tuple(range(1, 9)) else "cobertura incompleta: investigar"
    linhas_anuais_eda_sinesp.append({
        **identificacao_ano_eda_sinesp, "cobertura": cobertura_ano_eda_sinesp,
        "meses_presentes": meses_ano_eda_sinesp, **descrever_meses_sinesp(grupo_ano_eda_sinesp),
    })
    grupo_equivalente_eda_sinesp = grupo_ano_eda_sinesp.loc[grupo_ano_eda_sinesp["mes"].le(8)]
    exigir_sinesp_etapa9(tuple(sorted(grupo_equivalente_eda_sinesp["mes"].unique())) == tuple(range(1, 9)), "janeiro–agosto incompleto; não comparar.")
    linhas_equivalentes_eda_sinesp.append({
        **identificacao_ano_eda_sinesp, "periodo": "janeiro–agosto",
        **descrever_meses_sinesp(grupo_equivalente_eda_sinesp),
    })
tabela_anual_sinesp = pd.DataFrame(linhas_anuais_eda_sinesp)
comparacao_jan_ago_sinesp = pd.DataFrame(linhas_equivalentes_eda_sinesp)
estatisticas_descritivas_sinesp = pd.DataFrame([
    {
        **dict(zip(chaves_series_eda_sinesp, chave)),
        "inicio_observado": grupo["competencia"].min(), "fim_observado": grupo["competencia"].max(),
        "escopo": "todo o período observado; inclui 2026 parcial",
        **descrever_meses_sinesp(grupo),
    }
    for chave, grupo in tabela_mensal_sinesp.groupby(chaves_series_eda_sinesp, sort=False, dropna=False)
])
exigir_sinesp_etapa9(len(tabela_anual_sinesp) == 79 and len(comparacao_jan_ago_sinesp) == 79, "quantidade de séries/anos divergente.")
tabela_anual_sinesp

# %%
comparacao_jan_ago_sinesp

# %%
estatisticas_descritivas_sinesp

# %% [markdown]
# ## Comparação equivalente 2026 versus 2025
# Todos os totais abaixo têm oito meses. Percentual só é calculado com base
# anterior positiva; zero anterior não é substituído por outro denominador.
# Variação descritiva não demonstra efeito de políticas ou ações policiais.

# %%
linhas_comparacao_2026_sinesp = []
for chave_2026_sinesp, grupo_2026_sinesp in comparacao_jan_ago_sinesp.groupby(chaves_series_eda_sinesp, sort=False, dropna=False):
    anteriores_2026_sinesp = grupo_2026_sinesp.loc[grupo_2026_sinesp["ano"].eq(2025)]
    atuais_2026_sinesp = grupo_2026_sinesp.loc[grupo_2026_sinesp["ano"].eq(2026)]
    exigir_sinesp_etapa9(len(anteriores_2026_sinesp) == 1 and len(atuais_2026_sinesp) == 1, "comparação 2025/2026 indisponível.")
    anterior_2026_sinesp = anteriores_2026_sinesp.iloc[0]["soma_medida_publicada"]
    atual_2026_sinesp = atuais_2026_sinesp.iloc[0]["soma_medida_publicada"]
    ambos_informados_2026_sinesp = pd.notna(anterior_2026_sinesp) and pd.notna(atual_2026_sinesp)
    linhas_comparacao_2026_sinesp.append({
        **dict(zip(chaves_series_eda_sinesp, chave_2026_sinesp)), "periodo": "janeiro–agosto",
        "valor_2025": anterior_2026_sinesp, "valor_2026": atual_2026_sinesp,
        "diferenca_absoluta": atual_2026_sinesp - anterior_2026_sinesp if ambos_informados_2026_sinesp else pd.NA,
        "variacao_percentual": (atual_2026_sinesp / anterior_2026_sinesp - 1) * 100 if ambos_informados_2026_sinesp and anterior_2026_sinesp > 0 else pd.NA,
        "ressalva": "percentual não calculável: base zero ou ausente" if not ambos_informados_2026_sinesp or anterior_2026_sinesp == 0 else "variação descritiva; observar magnitude absoluta",
    })
comparacao_2026_2025_sinesp = pd.DataFrame(linhas_comparacao_2026_sinesp)
comparacao_2026_2025_sinesp

# %% [markdown]
# ## Tabelas separadas por indicador para Data Viewer
# O dicionário disponibiliza tabelas anuais, mensais e janeiro–agosto de cada
# evento. Abrangências continuam identificadas dentro de cada tabela.

# %%
tabelas_por_indicador_sinesp = {
    evento: {
        "anual": tabela_anual_sinesp.loc[tabela_anual_sinesp["evento"].eq(evento)],
        "mensal": tabela_mensal_sinesp.loc[tabela_mensal_sinesp["evento"].eq(evento)],
        "janeiro_agosto": comparacao_jan_ago_sinesp.loc[comparacao_jan_ago_sinesp["evento"].eq(evento)],
    }
    for evento in df_sinesp_joinville_analitico["evento"].unique()
}
feminicidio_anual_sinesp = tabelas_por_indicador_sinesp["Feminicídio"]["anual"]
tentativa_feminicidio_anual_sinesp = tabelas_por_indicador_sinesp["Tentativa de feminicídio"]["anual"]
homicidio_doloso_anual_sinesp = tabelas_por_indicador_sinesp["Homicídio doloso"]["anual"]
tentativa_homicidio_anual_sinesp = tabelas_por_indicador_sinesp["Tentativa de homicídio"]["anual"]
latrocinio_anual_sinesp = tabelas_por_indicador_sinesp["Roubo seguido de morte (latrocínio)"]["anual"]
lesao_seguida_morte_anual_sinesp = tabelas_por_indicador_sinesp["Lesão corporal seguida de morte"]["anual"]
complementares_anuais_sinesp = tabela_anual_sinesp.loc[~tabela_anual_sinesp["natureza"].eq("criminal")]
violencia_mulher_anual_sinesp = tabela_anual_sinesp.loc[tabela_anual_sinesp["evento"].isin(["Feminicídio", "Tentativa de feminicídio"])]
violencia_mulher_jan_ago_sinesp = comparacao_jan_ago_sinesp.loc[comparacao_jan_ago_sinesp["evento"].isin(["Feminicídio", "Tentativa de feminicídio"])]
feminicidio_anual_sinesp

# %%
tentativa_feminicidio_anual_sinesp

# %%
homicidio_doloso_anual_sinesp

# %%
tentativa_homicidio_anual_sinesp

# %%
latrocinio_anual_sinesp

# %%
lesao_seguida_morte_anual_sinesp

# %%
complementares_anuais_sinesp

# %%
violencia_mulher_jan_ago_sinesp

# %% [markdown]
# ## Limites da interpretação e integridade final
# Os totais de 2026 representam apenas janeiro–agosto. Indicadores raros, com
# muitos meses de zero ou poucos casos, devem ser interpretados pela magnitude
# absoluta; percentuais podem parecer grandes com bases pequenas.
# Não somar feminicídio a homicídio ou as duas tentativas. Trânsito, suicídio,
# mortes a esclarecer e mandados são complementares. Não há explicação causal,
# taxa populacional, gráfico ou exportação nesta etapa.

# %%
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_antes_eda_sinesp, check_exact=True)
pd.testing.assert_frame_equal(df_sinesp_joinville_analitico, analitico_antes_eda_sinesp, check_exact=True)
pd.testing.assert_frame_equal(catalogo_indicadores_sinesp, catalogo_antes_eda_sinesp, check_exact=True)
pd.testing.assert_frame_equal(metadados_sinesp_analitico, metadados_antes_eda_sinesp, check_exact=True)
for ano_etapa9, caminho_etapa9 in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa9(sha256_original_sinesp(caminho_etapa9) == hashes_originais_etapa8[ano_etapa9], "arquivo original alterado durante a EDA.")
print("Etapa 9: 900 observações mensais, 79 séries/anos, 79 comparações janeiro–agosto e 12 resumos descritivos; bases e arquivos preservados.")

# %% [markdown]
# # ETAPA 10 — Visualização exploratória Sinesp
# Pré-requisito: tabelas validadas da etapa 9 no mesmo processo/kernel.
# As figuras usam apenas valores publicados dessas tabelas. Não se recalculam
# medidas, não se alteram as bases e não se exportam imagens. No Python
# Interactive, execute cada célula com plt.show() para visualizar o gráfico.

# %%
import textwrap

import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator

exigir_sinesp_etapa9(df_sinesp_joinville_concatenado.shape == (900, 14), "base original divergente antes dos gráficos.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico.shape == (900, 10), "base analítica divergente antes dos gráficos.")
original_antes_graficos_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)
analitico_antes_graficos_sinesp = df_sinesp_joinville_analitico.copy(deep=True)
anual_antes_graficos_sinesp = tabela_anual_sinesp.copy(deep=True)
jan_ago_antes_graficos_sinesp = comparacao_jan_ago_sinesp.copy(deep=True)
cores_graficos_sinesp = {
    "Homicídio doloso": "#2563A6", "Tentativa de homicídio": "#188977",
    "Feminicídio": "#7A5195", "Tentativa de feminicídio": "#D97706",
}


def selecionar_tabela_grafico_sinesp(tabela, eventos, anos, meses):
    """Seleciona e ordena observações já agregadas na etapa 9, sem preencher."""
    dados = tabela.loc[
        tabela["evento"].isin(eventos) & tabela["ano"].isin(anos)
        & tabela["abrangencia"].eq("Estadual")
    ].copy(deep=True)
    exigir_sinesp_etapa9(len(dados) == len(eventos) * len(anos), "observações inesperadas para o gráfico.")
    exigir_sinesp_etapa9(not dados.duplicated(["ano", "evento"]).any(), "ano/indicador repetido no gráfico.")
    exigir_sinesp_etapa9(dados["fonte"].eq("MJSP / Sinesp VDE").all(), "fonte divergente no gráfico.")
    exigir_sinesp_etapa9(dados["unidade"].eq("vítimas").all() and dados["medida"].eq("total_vitima").all(), "unidades incompatíveis no gráfico.")
    exigir_sinesp_etapa9(dados["meses_disponiveis"].eq(meses).all() and dados["meses_com_medida"].eq(meses).all(), "cobertura incompleta para o gráfico.")
    exigir_sinesp_etapa9(dados["soma_medida_publicada"].notna().all(), "valor ausente; gráfico interrompido sem preenchimento.")
    if meses == 12:
        exigir_sinesp_etapa9(dados["cobertura"].eq("ano completo").all(), "período parcial na série anual.")
    else:
        exigir_sinesp_etapa9(dados["periodo"].eq("janeiro–agosto").all(), "períodos não equivalentes.")
    quadro = dados.pivot(index="ano", columns="evento", values="soma_medida_publicada").reindex(index=anos, columns=eventos)
    exigir_sinesp_etapa9(not quadro.isna().any().any(), "célula ausente na tabela visual.")
    for linha in dados.itertuples(index=False):
        exigir_sinesp_etapa9(quadro.loc[linha.ano, linha.evento] == linha.soma_medida_publicada, "valor visual difere da etapa 9.")
    return quadro


def preparar_figura_sinesp(titulo, periodo, interpretacao, limite_y):
    fig, ax = plt.subplots(figsize=(10, 6.2), dpi=120)
    fig.subplots_adjust(left=0.09, right=0.97, bottom=0.27, top=0.80)
    fig.suptitle(titulo, x=0.08, y=0.97, ha="left", fontsize=15, fontweight="semibold")
    fig.text(0.08, 0.91, periodo, fontsize=11, color="#444444")
    ax.set_xlabel("Ano", fontsize=11)
    ax.set_ylabel("Vítimas", fontsize=11)
    ax.set_ylim(0, limite_y)
    ax.yaxis.set_major_locator(MaxNLocator(nbins=6, integer=True))
    ax.set_axisbelow(True)
    ax.grid(axis="y", color="#DDDDDD", linewidth=0.7)
    ax.spines[["top", "right"]].set_visible(False)
    fig.text(0.08, 0.125, textwrap.fill(interpretacao, width=108), va="top", fontsize=9, color="#333333")
    fig.text(0.08, 0.035, "Fonte: MJSP / Sinesp VDE | Abrangência: Estadual | Unidade: vítimas", fontsize=9, color="#555555")
    return fig, ax


def desenhar_linhas_sinesp(ax, dados):
    for numero, evento in enumerate(dados.columns):
        valores = dados[evento].tolist()
        artista, = ax.plot(
            dados.index, valores, label=evento, color=cores_graficos_sinesp[evento],
            marker="o" if numero == 0 else "s", linestyle="-" if numero == 0 else "--",
            linewidth=2, markersize=6,
        )
        exigir_sinesp_etapa9(list(artista.get_ydata()) == valores, "coordenadas da linha divergentes.")
        for ano, valor in zip(dados.index, valores):
            ax.annotate(f"{valor:g}", (ano, valor), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=10)
    ax.set_xticks(dados.index.tolist())
    if len(dados.columns) > 1:
        ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=2, frameon=False, fontsize=10)


def desenhar_barras_sinesp(ax, dados):
    posicoes = list(range(len(dados)))
    largura = 0.34
    for numero, evento in enumerate(dados.columns):
        valores = dados[evento].tolist()
        barras = ax.bar(
            [posicao + (numero - 0.5) * largura for posicao in posicoes], valores,
            width=largura, label=evento, color=cores_graficos_sinesp[evento],
            hatch="//" if numero == 0 else None,
        )
        exigir_sinesp_etapa9([barra.get_height() for barra in barras] == valores, "altura de barra divergente.")
        ax.bar_label(barras, labels=[f"{valor:g}" for valor in valores], padding=4, fontsize=10)
    ax.set_xticks(posicoes, dados.index.tolist())
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=2, frameon=False, fontsize=10)


dados_grafico_homicidio_anual_sinesp = selecionar_tabela_grafico_sinesp(
    tabela_anual_sinesp, ["Homicídio doloso"], list(range(2020, 2026)), 12
)
dados_grafico_mulher_anual_sinesp = selecionar_tabela_grafico_sinesp(
    tabela_anual_sinesp, ["Feminicídio", "Tentativa de feminicídio"], list(range(2020, 2026)), 12
)
dados_grafico_homicidios_jan_ago_sinesp = selecionar_tabela_grafico_sinesp(
    comparacao_jan_ago_sinesp, ["Homicídio doloso", "Tentativa de homicídio"], list(range(2020, 2027)), 8
)
dados_grafico_mulher_jan_ago_sinesp = selecionar_tabela_grafico_sinesp(
    comparacao_jan_ago_sinesp, ["Feminicídio", "Tentativa de feminicídio"], [2025, 2026], 8
)
interpretacoes_graficos_sinesp = {
    "homicidio_anual": f"O total publicado passou de {dados_grafico_homicidio_anual_sinesp.loc[2020, 'Homicídio doloso']} vítimas em 2020 para {dados_grafico_homicidio_anual_sinesp.loc[2025, 'Homicídio doloso']} em 2025, com oscilações intermediárias. A série não demonstra causas.",
    "mulher_anual": f"Feminicídios variaram de {dados_grafico_mulher_anual_sinesp['Feminicídio'].min()} a {dados_grafico_mulher_anual_sinesp['Feminicídio'].max()} vítimas anuais; tentativas, de {dados_grafico_mulher_anual_sinesp['Tentativa de feminicídio'].min()} a {dados_grafico_mulher_anual_sinesp['Tentativa de feminicídio'].max()}. São pequenas contagens de categorias distintas, sem soma entre elas.",
    "homicidios_jan_ago": f"De janeiro a agosto, homicídios passaram de {dados_grafico_homicidios_jan_ago_sinesp.loc[2025, 'Homicídio doloso']} em 2025 para {dados_grafico_homicidios_jan_ago_sinesp.loc[2026, 'Homicídio doloso']} em 2026; tentativas, de {dados_grafico_homicidios_jan_ago_sinesp.loc[2025, 'Tentativa de homicídio']} para {dados_grafico_homicidios_jan_ago_sinesp.loc[2026, 'Tentativa de homicídio']}. Os períodos têm a mesma duração.",
    "mulher_jan_ago": f"Tentativas passaram de {dados_grafico_mulher_jan_ago_sinesp.loc[2025, 'Tentativa de feminicídio']} para {dados_grafico_mulher_jan_ago_sinesp.loc[2026, 'Tentativa de feminicídio']} vítimas. Feminicídio tem zero publicado nos dois recortes; isso não comprova eliminação do fenômeno.",
}

# %% [markdown]
# ## Gráfico 1 — Homicídio doloso: anos completos
# Linhas com marcadores e valores explícitos; eixo iniciado em zero. 2026 não
# integra esta série anual porque sua cobertura termina em agosto.

# %%
fig_homicidio_anual_sinesp, ax_homicidio_anual_sinesp = preparar_figura_sinesp(
    "Homicídio doloso em Joinville/SC", "Anos completos: 2020–2025",
    interpretacoes_graficos_sinesp["homicidio_anual"], 70,
)
desenhar_linhas_sinesp(ax_homicidio_anual_sinesp, dados_grafico_homicidio_anual_sinesp)
plt.show()

# %% [markdown]
# ## Gráfico 2 — Feminicídio e tentativa: séries anuais distintas
# Barras lado a lado, cores e hachura distintas. Escala única, inteira, iniciada
# em zero; não há eixo secundário, empilhamento ou soma de categorias.

# %%
fig_mulher_anual_sinesp, ax_mulher_anual_sinesp = preparar_figura_sinesp(
    "Feminicídio e tentativa de feminicídio em Joinville/SC", "Anos completos: 2020–2025",
    interpretacoes_graficos_sinesp["mulher_anual"], 12,
)
desenhar_barras_sinesp(ax_mulher_anual_sinesp, dados_grafico_mulher_anual_sinesp)
plt.show()

# %% [markdown]
# ## Gráfico 3 — Homicídio e tentativa: janeiro–agosto
# Sete recortes com oito meses cada. Marcadores e estilos distinguem categorias;
# nenhuma linha usa o total anual completo dos anos anteriores.

# %%
fig_homicidios_jan_ago_sinesp, ax_homicidios_jan_ago_sinesp = preparar_figura_sinesp(
    "Homicídio doloso e tentativa de homicídio em Joinville/SC", "Períodos equivalentes: janeiro–agosto de 2020 a 2026",
    interpretacoes_graficos_sinesp["homicidios_jan_ago"], 90,
)
desenhar_linhas_sinesp(ax_homicidios_jan_ago_sinesp, dados_grafico_homicidios_jan_ago_sinesp)
plt.show()

# %% [markdown]
# ## Gráfico 4 — Violência contra a mulher: janeiro–agosto de 2025 e 2026
# Mesmas cores e escala do gráfico 2. Os rótulos zero permanecem visíveis;
# pequenos números são apresentados sem ampliar percentuais no gráfico.

# %%
fig_mulher_jan_ago_sinesp, ax_mulher_jan_ago_sinesp = preparar_figura_sinesp(
    "Feminicídio e tentativa de feminicídio em Joinville/SC", "Períodos equivalentes: janeiro–agosto de 2025 e 2026",
    interpretacoes_graficos_sinesp["mulher_jan_ago"], 12,
)
desenhar_barras_sinesp(ax_mulher_jan_ago_sinesp, dados_grafico_mulher_jan_ago_sinesp)
plt.show()

# %%
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_antes_graficos_sinesp, check_exact=True)
pd.testing.assert_frame_equal(df_sinesp_joinville_analitico, analitico_antes_graficos_sinesp, check_exact=True)
pd.testing.assert_frame_equal(tabela_anual_sinesp, anual_antes_graficos_sinesp, check_exact=True)
pd.testing.assert_frame_equal(comparacao_jan_ago_sinesp, jan_ago_antes_graficos_sinesp, check_exact=True)
for ano_grafico_sinesp, caminho_grafico_sinesp in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa9(sha256_original_sinesp(caminho_grafico_sinesp) == hashes_originais_etapa8[ano_grafico_sinesp], "original alterado durante as visualizações.")
print("Etapa 10: quatro figuras em memória; valores plotados iguais às tabelas da etapa 9; bases 900 × 14 e 900 × 10 preservadas.")

# ETAPA 11 — Planejamento: exportar somente com autorização específica.
# ETAPA 12 — Planejamento: reler e validar arquivos exportados após autorização.
