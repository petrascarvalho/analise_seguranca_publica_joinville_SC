from __future__ import annotations

import csv
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "joinville" / "csv_complementares"
OUT.mkdir(parents=True, exist_ok=True)

CITY_SOURCE = "Prefeitura de Joinville / SEPROT"
CITY_URL = "https://www.joinville.sc.gov.br/noticias/indices-de-criminalidade-tem-reducao-em-joinville-pelo-terceiro-ano-consecutivo/"
SINESP_URL = "https://www.gov.br/mj/pt-br/assuntos/sua-seguranca/seguranca-publica/estatistica/dados-nacionais-1/base-de-dados-e-notas-metodologicas-dos-gestores-estaduais-sinesp-vde-2022-e-2023"


def export_theft_robbery():
    values = {
        2023: {"Furto": 8865, "Roubo": 933},
        2024: {"Furto": 8730, "Roubo": 593},
        2025: {"Furto": 8372, "Roubo": 476},
    }
    for year, metrics in values.items():
        target = OUT / f"furto_roubo_joinville_{year}.csv"
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(["uf", "municipio", "ano", "indicador", "periodo", "quantidade", "fonte", "url_fonte"])
            for metric, value in metrics.items():
                writer.writerow(["SC", "JOINVILLE", year, metric, "janeiro-dezembro", value, CITY_SOURCE, CITY_URL])
        print(f"{target.name}: {len(metrics)} registros")


def excel_value(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value


def export_femicide():
    source_dir = ROOT / "data" / "joinville"
    for source in sorted(source_dir.glob("bancovde-*-joinville-sc.xlsx")):
        year = int(source.stem.split("-")[1])
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        headers = list(next(values))
        selected = []
        for row in values:
            if row[0] != "SC" or row[1] != "JOINVILLE":
                continue
            event = str(row[2] or "")
            if event not in {"Feminicídio", "Tentativa de feminicídio"}:
                continue
            selected.append([excel_value(v) for v in row])
        workbook.close()
        target = OUT / f"feminicidio_sinesp_joinville_{year}.csv"
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow(headers)
            writer.writerows(selected)
        print(f"{target.name}: {len(selected)} registros")


def main():
    export_theft_robbery()
    export_femicide()


if __name__ == "__main__":
    main()
