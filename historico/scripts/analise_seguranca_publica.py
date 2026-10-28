"""Inspeção acadêmica, em células interativas, das fontes de Joinville/SC.

Abra a pasta do projeto no VS Code. Execute cada célula com "Run Cell".
As variáveis DataFrame ficam disponíveis no Interactive Window para abrir no
Data Viewer. As etapas 1–4 inspecionam as fontes; a etapa 5 valida e concatena
somente em memória os recortes anuais do Sinesp. A etapa 6 diagnostica a
qualidade desse DataFrame; a etapa 7 investiga os achados e propõe colunas e
tipos. A etapa 8 prepara uma cópia analítica em memória, com datas e inteiros
anuláveis validados. A etapa 9 calcula tabelas descritivas por indicador e
abrangência, somente em memória. A etapa 10 exibe quatro gráficos exploratórios
validados contra essas tabelas. Não há exportação.
"""

# %% [markdown]
# # ETAPA 1 — Importação das bibliotecas
# Objetivo: importar Pandas para leitura de planilhas e CSVs e pathlib para
# localizar as fontes. Nenhuma transformação é aplicada nesta etapa.

# %%
from pathlib import Path
import sys

import pandas as pd
import pdfplumber
from src.pipeline import transform

# Mantém os acentos legíveis no terminal quando o fluxo não está no Interactive Window.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

print(f"Pandas: {pd.__version__}")


# %% [markdown]
# # ETAPA 2 — Definição dos diretórios
# Objetivo: localizar as fontes a partir da pasta do projeto aberta no VS Code.
# Procedimento: em execução como módulo, usa a raiz do projeto (dois níveis acima de historico/scripts/); no
# Interactive Window, usa a pasta de trabalho (abra a raiz do projeto no VS Code).
# Validação: os caminhos são exibidos junto com seu estado de disponibilidade.

# %%
RAIZ_PROJETO = (
    Path(__file__).resolve().parents[2] if "__file__" in globals() else Path.cwd()
)
PASTA_SINESP = RAIZ_PROJETO / "dados" / "brutos" / "sinesp"
ARQUIVO_SINESP_2023 = PASTA_SINESP / "bancovde-2023.xlsx"
PASTA_CSV_SSP = RAIZ_PROJETO / "dados" / "brutos" / "ssp_sc" / "csv_extraidos"
PASTA_PDF_SSP = RAIZ_PROJETO / "dados" / "brutos" / "ssp_sc" / "pdfs"
PASTA_PREFEITURA = RAIZ_PROJETO / "dados" / "processados" / "prefeitura"

# Altere este nome para inspecionar outro CSV existente da SSP/SC.
CSV_SSP_SELECIONADO = "2026_08_2026-julho.csv"

print(f"Raiz do projeto: {RAIZ_PROJETO}")
for nome, caminho in (
    ("XLSX Sinesp 2023", ARQUIVO_SINESP_2023),
    ("CSV extraídos SSP/SC", PASTA_CSV_SSP),
    ("CSV Prefeitura/SEPROT", PASTA_PREFEITURA),
):
    print(f"{nome}: {'disponível' if caminho.exists() else 'não encontrado'} — {caminho}")

# %% [markdown]
# # ETAPA 3 — Leitura dos arquivos brutos
# Objetivo: ler uma amostra nacional do Sinesp, selecionar Joinville/SC em
# memória, ler os CSVs SSP já extraídos e abrir separadamente os CSVs SEPROT.
# Nenhuma fonte é sobrescrita, exportada ou consolidada.

# %% [markdown]
# ## Sinesp — amostra nacional e recorte em memória
# read_excel(nrows=20) lê a amostra nacional. O leitor XML especializado já
# existente seleciona as linhas SC/JOINVILLE do XLSX e as devolve em memória.
# As datas permanecem como armazenadas no arquivo, inclusive seriais do Excel.

# %%
sinesp_2023 = pd.read_excel(
    ARQUIVO_SINESP_2023,
    nrows=20,
    engine="openpyxl",
)
joinville_sc = pd.DataFrame.from_records(
    transform._joinville_rows(ARQUIVO_SINESP_2023),
    columns=transform.HEADERS,
)
print(f"Amostra nacional: {len(sinesp_2023)} linhas; recorte Joinville/SC: {len(joinville_sc)} linhas.")

# %% [markdown]
# ## SSP/SC — inventário e cabeçalhos dos CSVs já extraídos
# O inventário consulta os cabeçalhos com nrows=0. As colunas celula_* variam
# entre relatórios porque as tabelas dos PDFs têm larguras diferentes.

# %%
arquivos_ssp = sorted(PASTA_CSV_SSP.glob("*.csv"))
inventario_linhas_ssp = []
esquemas_ssp = {}
for arquivo in arquivos_ssp:
    colunas_arquivo = pd.read_csv(
        arquivo,
        sep=";",
        encoding="utf-8-sig",
        nrows=0,
    ).columns.tolist()
    esquemas_ssp.setdefault(tuple(colunas_arquivo), []).append(arquivo.name)
    inventario_linhas_ssp.append(
        {
            "arquivo_csv": arquivo.name,
            "formato": "CSV",
            "tamanho_bytes": arquivo.stat().st_size,
            "ano_no_nome": arquivo.name[:4],
            "numero_colunas": len(colunas_arquivo),
        }
    )
inventario_ssp = pd.DataFrame(inventario_linhas_ssp)

arquivo_ssp_selecionado = PASTA_CSV_SSP / CSV_SSP_SELECIONADO
ssp_selecionado = pd.read_csv(
    arquivo_ssp_selecionado,
    sep=";",
    encoding="utf-8-sig",
    dtype="string",
    keep_default_na=False,
)
print(f"SSP/SC: {len(arquivos_ssp)} CSVs; arquivo selecionado: {arquivo_ssp_selecionado.name}.")

# %% [markdown]
# ## Prefeitura/SEPROT — arquivos transcritos
# Os três CSVs são lidos separadamente. Cada um contém duas linhas; não são
# empilhados em uma tabela única.

# %%
arquivos_prefeitura = sorted(PASTA_PREFEITURA.glob("furto_roubo_joinville_*.csv"))
tabelas_prefeitura = [
    pd.read_csv(arquivo, encoding="utf-8-sig") for arquivo in arquivos_prefeitura
]
prefeitura_2023 = tabelas_prefeitura[0]
prefeitura_2024 = tabelas_prefeitura[1]
prefeitura_2025 = tabelas_prefeitura[2]
print(f"Prefeitura/SEPROT: {len(arquivos_prefeitura)} arquivos, {sum(map(len, tabelas_prefeitura))} linhas.")

# %% [markdown]
# # ETAPA 4 — Inspeção e visualização dos dados
# Execute as células abaixo separadamente. No Interactive Window, cada expressão
# final apresenta a tabela; os nomes das variáveis ficam disponíveis no painel
# Variables para abrir no Data Viewer.
# Como abrir uma tabela: execute as células anteriores e a célula da fonte;
# no Interactive Window, clique em Variables. Na lista, localize a variável e
# selecione o botão Show variable in data viewer (ícone de tabela).
# Variáveis: sinesp_2023, joinville_sc, data_referencia_nacional_20,
# data_referencia_joinville_20, inventario_ssp, ssp_selecionado,
# ssp_com_joinville, ssp_tabela_pagina_2, ssp_tabela_pagina_8,
# ssp_contexto_pagina_2, ssp_contexto_pagina_8 e prefeitura_2023–2025.

# %% [markdown]
# ## Sinesp — colunas, tipos e primeiras 20 linhas nacionais
# As 14 colunas e seus tipos são mostrados conforme a amostra lida. Os tipos
# inferidos a partir de 20 linhas não são apresentados como tipos de toda a base.

# %%
print("Arquivo nacional Sinesp:", ARQUIVO_SINESP_2023.name)
print("Dimensões da amostra:", sinesp_2023.shape)
print("Colunas:", sinesp_2023.columns.tolist())
print("Tipos Pandas observados na amostra:")
print(sinesp_2023.dtypes.to_string())
sinesp_2023.head(20)

# %% [markdown]
# ## Sinesp — comparação de data_referencia sem conversão
# Objetivo: observar os primeiros 20 valores e os tipos inferidos por cada
# caminho de leitura. read_excel interpreta as células XLSX formatadas como
# data; o leitor XML auxiliar retorna o conteúdo numérico serial da célula.
# A diferença descreve os leitores, não comprova erro no arquivo. Nenhum valor
# é convertido nesta inspeção.

# %%
data_referencia_nacional_20 = sinesp_2023["data_referencia"].head(20)
data_referencia_joinville_20 = joinville_sc["data_referencia"].head(20)
print("Tipo original na amostra nacional:", sinesp_2023["data_referencia"].dtype)
print("Tipo original no recorte Joinville:", joinville_sc["data_referencia"].dtype)
print("Valores nacionais — data_referencia:")
data_referencia_nacional_20

# %%
print("Valores Joinville/SC — data_referencia (sem conversão):")
data_referencia_joinville_20

# %% [markdown]
# ## Sinesp — primeiras 20 linhas do recorte Joinville/SC
# É uma seleção em memória do arquivo nacional, sem exportação. O leitor
# especializado preserva a data como número serial do XLSX quando assim gravada.

# %%
print("Dimensões do recorte Joinville/SC:", joinville_sc.shape)
print("Tipos Pandas observados no recorte:")
print(joinville_sc.dtypes.to_string())
joinville_sc.head(20)

# %% [markdown]
# ### Significado dos campos Sinesp
# Baseado no catálogo e nas notas oficiais do MJSP e nos nomes originais. O PDF
# do dicionário não abriu nesta consulta; por isso, campos sem definição confirmada
# são identificados como tal, sem atribuir significado por suposição.
# Referência: https://dados.mj.gov.br/dataset/sistema-nacional-de-estatisticas-de-seguranca-publica

# %%
print(
    "uf e municipio: referência geográfica; evento: indicador; "
    "data_referencia: data/período de referência registrado na base."
)
print(
    "feminino, masculino e nao_informado: nomes de campos de sexo; as notas do "
    "MJSP descrevem totais por sexo para indicadores enviados como vítimas. "
    "total_vitima: campo declarado para total de vítimas."
)
print(
    "agente, arma e faixa_etaria: campos presentes na planilha; sem confirmação "
    "do dicionário, não se atribui definição mais específica a cada um."
)
print(
    "total e total_peso: nomes de campos de medida; unidade e interpretação "
    "dependem do indicador e não são inferidas apenas pelos nomes."
)
print(
    "abrangencia: campo assim nomeado na base; não identifica, por si só, qual "
    "corporação registrou os dados municipais."
)

# %% [markdown]
# ## SSP/SC — inventário e esquemas de colunas dos CSVs
# arquivo_pdf, pagina_pdf, tabela_pdf e linha_tabela são metadados do extrator.
# celula_* contém o texto extraído e não representa um esquema estatístico comum.

# %%
print(f"CSV encontrados: {len(inventario_ssp)} (esperados: 45)")
inventario_ssp

# %%
for numero, (colunas, arquivos_do_esquema) in enumerate(esquemas_ssp.items(), start=1):
    print(f"Esquema {numero}: {len(colunas)} colunas; arquivos: {arquivos_do_esquema}")
    print(list(colunas))

# %% [markdown]
# ## SSP/SC — colunas e primeiras 20 linhas do CSV selecionado
# O nome abaixo pode ser trocado em CSV_SSP_SELECIONADO, na etapa 2.

# %%
print("Arquivo selecionado:", arquivo_ssp_selecionado.name)
print("Dimensões:", ssp_selecionado.shape)
print("Colunas:", ssp_selecionado.columns.tolist())
print(
    "Rastreabilidade:",
    [nome for nome in ("arquivo_pdf", "pagina_pdf", "tabela_pdf", "linha_tabela")
     if nome in ssp_selecionado.columns],
)
ssp_selecionado.head(20)

# %% [markdown]
# ## SSP/SC — todas as linhas que contêm JOINVILLE
# A máscara procura o texto em qualquer coluna e apenas seleciona linhas para
# visualização. O CSV original permanece intacto.

# %%
linhas_joinville_ssp = ssp_selecionado.apply(
    lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
).any(axis=1)
ssp_com_joinville = ssp_selecionado.loc[linhas_joinville_ssp]
print(f"Linhas encontradas: {len(ssp_com_joinville)}")
ssp_com_joinville

# %% [markdown]
# ## SSP/SC — cabeçalho e contexto no PDF original
# Objetivo: conferir visualmente as duas linhas encontradas no CSV, diretamente
# nas páginas originais. pdfplumber apenas lê as tabelas do PDF em memória; as
# linhas e cabeçalhos não são limpos, renomeados no arquivo, nem exportados.
# tabela_pdf=1 é o número registrado no CSV pelo extrator. Os nomes temporários
# coluna_XX permitem visualizar as células em DataFrame sem presumir cabeçalhos.

# %%
ARQUIVO_PDF_SSP_SELECIONADO = (
    PASTA_PDF_SSP / "2026_08_2026-julho.pdf"
)
with pdfplumber.open(ARQUIVO_PDF_SSP_SELECIONADO) as documento_ssp:
    tabela_pdf_pagina_2 = documento_ssp.pages[1].extract_tables()[0]
    tabela_pdf_pagina_8 = documento_ssp.pages[7].extract_tables()[0]

ssp_tabela_pagina_2 = pd.DataFrame(
    tabela_pdf_pagina_2,
    columns=[f"coluna_{numero:02d}" for numero in range(1, len(tabela_pdf_pagina_2[0]) + 1)],
)
ssp_tabela_pagina_8 = pd.DataFrame(
    tabela_pdf_pagina_8,
    columns=[f"coluna_{numero:02d}" for numero in range(1, len(tabela_pdf_pagina_8[0]) + 1)],
)
linha_joinville_pagina_2 = ssp_tabela_pagina_2.index[
    ssp_tabela_pagina_2.astype("string").apply(
        lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
    ).any(axis=1)
][0]
linha_joinville_pagina_8 = ssp_tabela_pagina_8.index[
    ssp_tabela_pagina_8.astype("string").apply(
        lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
    ).any(axis=1)
][0]
ssp_contexto_pagina_2 = ssp_tabela_pagina_2.iloc[
    max(0, linha_joinville_pagina_2 - 3):linha_joinville_pagina_2 + 4
]
ssp_contexto_pagina_8 = ssp_tabela_pagina_8.iloc[
    max(0, linha_joinville_pagina_8 - 3):linha_joinville_pagina_8 + 4
]
print(f"PDF: {ARQUIVO_PDF_SSP_SELECIONADO.name} | página 2 | tabela 1")
print("Cabeçalho e contexto: tabela de variações de homicídios por município, período 01/01–31/07, comparação 2025–2026.")
ssp_contexto_pagina_2

# %%
print(f"PDF: {ARQUIVO_PDF_SSP_SELECIONADO.name} | página 8 | tabela 1")
print("Cabeçalho da tabela municipal: vítimas por ano (2023–2025) e no período 01/01–31/07 (2023–2026).")
print("Primeiras linhas da tabela, incluindo cabeçalhos originais:")
ssp_tabela_pagina_8.head(5)

# %%
print("Linha Joinville e municípios vizinhos na tabela original da página 8:")
ssp_contexto_pagina_8

# %%
print("As duas linhas completas do CSV, com metadados de rastreabilidade:")
ssp_com_joinville

# %% [markdown]
# ### Significado dos campos SSP/SC
# A página oficial informa que Segurança em Números é elaborado pela Gerência
# de Estatística e Análise Criminal. As quatro colunas de rastreabilidade e as
# celula_* são estrutura do extrator; os cabeçalhos/rótulos estatísticos ficam
# dentro das células conforme cada PDF, sem normalização.
# Referência: https://ssp.sc.gov.br/segurancaemnumeros/

# %%
print(
    "arquivo_pdf: nome do PDF de origem; pagina_pdf: página; tabela_pdf: número "
    "da tabela detectada; linha_tabela: posição da linha extraída. "
    "celula_*: conteúdo da linha conforme o extrator."
)

# %% [markdown]
# ## Prefeitura/SEPROT — seis linhas, em três DataFrames separados
# Os CSVs são transcrições anteriores da publicação municipal. A publicação
# informa os totais anuais; não fornece dicionário para o esquema dos CSVs.
# Referência: https://www.joinville.sc.gov.br/noticias/indices-de-criminalidade-tem-reducao-em-joinville-pelo-terceiro-ano-consecutivo/

# %%
print("Prefeitura/SEPROT — 2023 (transcrição anterior)")
prefeitura_2023

# %%
print("Prefeitura/SEPROT — 2024 (transcrição anterior)")
prefeitura_2024

# %%
print("Prefeitura/SEPROT — 2025 (transcrição anterior)")
prefeitura_2025

# %% [markdown]
# ### Significado dos campos Prefeitura/SEPROT
# uf, municipio e ano identificam o recorte transcrito; indicador nomeia Furto
# ou Roubo; periodo registra janeiro a dezembro; quantidade reproduz o valor
# publicado; fonte e url_fonte registram a proveniência da transcrição.

# %%
print(
    "A Prefeitura publicou totais anuais para Furto e Roubo em 2023–2025. "
    "Os campos de organização e proveniência pertencem à transcrição dos CSVs."
)

# %% [markdown]
# # ETAPA 5 — Validação e prévia da concatenação Sinesp
# Objetivo: empilhar os sete recortes municipais compatíveis somente em memória.
# Critério: uf == "SC" e municipio == "JOINVILLE", sem correspondência parcial.
# O leitor XML existente preserva os valores e os seriais Excel; Pandas infere
# os tipos dos recortes. A amostra read_excel(nrows=20) é apenas uma comparação
# dos leitores, não um perfil de tipos do arquivo nacional inteiro.
# Nenhum tipo ou valor é corrigido. Metadados ficam em tabelas auxiliares.
# Execute as células da etapa 5 após a etapa 2. As tabelas abaixo e
# df_sinesp_joinville_concatenado ficam disponíveis em Variables > Data Viewer.

# %%
import re
from zipfile import ZipFile

from openpyxl import load_workbook

anos_sinesp = tuple(range(2020, 2027))
arquivos_sinesp_por_ano = {}
for ano_sinesp in anos_sinesp:
    candidatos_sinesp = sorted(PASTA_SINESP.rglob(f"bancovde-{ano_sinesp}.xlsx"))
    if len(candidatos_sinesp) != 1:
        raise ValueError(
            f"Esperado um original de {ano_sinesp}; encontrados: {candidatos_sinesp}"
        )
    arquivos_sinesp_por_ano[ano_sinesp] = candidatos_sinesp[0]


def inspecionar_original_sinesp(caminho, ano):
    """Confere a estrutura e lê Joinville sem tipar datas nem gravar arquivos."""
    livro = load_workbook(caminho, read_only=True, data_only=False)
    try:
        abas = livro.sheetnames
        if abas != [str(ano)]:
            raise ValueError(f"Planilhas inesperadas em {caminho.name}: {abas}")
        planilha = livro[abas[0]]
        cabecalho = list(next(planilha.iter_rows(min_row=1, max_row=1, values_only=True)))
        if cabecalho != transform.HEADERS:
            raise ValueError(f"Colunas incompatíveis em {caminho.name}: {cabecalho}")
    finally:
        livro.close()

    registros_municipais = []
    linhas_origem = []
    linhas_nacionais = 0
    datas_municipais = set()
    with ZipFile(caminho) as pacote:
        textos = transform._shared_strings(pacote)
        ids_joinville = [str(i).encode("ascii") for i, texto in enumerate(textos)
                        if texto == "JOINVILLE"]
        padrao_municipio = (
            re.compile(rb'<c\b[^>]*\br="B\d+"[^>]*>\s*<v>(?:'
                       + b"|".join(ids_joinville) + rb')</v>\s*</c>')
            if ids_joinville else None
        )
        xml = pacote.read("xl/worksheets/sheet1.xml")
        if re.search(rb'<c\b[^>]*\br="(?:[O-Z]|[A-Z]{2,})\d+"', xml):
            raise ValueError(f"Células além das 14 colunas em {caminho.name}")
        for correspondencia in re.finditer(
            rb'<row\b[^>]*\br="(\d+)"[^>]*>.*?</row>', xml, re.DOTALL
        ):
            numero_linha = int(correspondencia.group(1))
            if numero_linha == 1:
                continue
            linhas_nacionais += 1
            linha_xml = correspondencia.group()
            if b"JOINVILLE" not in linha_xml and not (
                padrao_municipio is not None and padrao_municipio.search(linha_xml)
            ):
                continue
            elemento = transform._parse_sheet_row(linha_xml)
            celulas = {
                celula.get("r").rstrip("0123456789"): transform._cell_value(celula, textos)
                for celula in elemento.findall(transform.NS + "c")
            }
            if celulas.get("A") != "SC" or celulas.get("B") != "JOINVILLE":
                continue
            valores = [celulas.get(chr(65 + i)) for i in range(len(cabecalho))]
            registros_municipais.append(valores)
            linhas_origem.append(numero_linha)
            datas_municipais.add((type(valores[3]).__name__, repr(valores[3])))

    recorte = pd.DataFrame.from_records(registros_municipais, columns=cabecalho)
    amostra = pd.read_excel(caminho, sheet_name=abas[0], nrows=20, engine="openpyxl")
    if amostra.columns.tolist() != cabecalho or recorte.empty:
        raise ValueError(f"Cabeçalhos divergentes ou recorte vazio: {caminho.name}")
    return recorte, amostra, {
        "ano": ano, "arquivo": str(caminho.relative_to(RAIZ_PROJETO)),
        "tamanho_bytes": caminho.stat().st_size, "extensao": caminho.suffix,
        "planilha": abas[0], "colunas": tuple(cabecalho),
        "linhas_nacionais_sem_cabecalho": linhas_nacionais,
        "linhas_joinville": len(recorte),
        "tipo_data_xml": sorted({tipo for tipo, _ in datas_municipais}),
    }, linhas_origem


recortes_sinesp_por_ano = {}
amostras_sinesp_por_ano = {}
metadados_sinesp_por_ano = {}
inventario_sinesp_linhas = []
for ano_sinesp, caminho_sinesp in arquivos_sinesp_por_ano.items():
    recorte_sinesp, amostra_sinesp, ficha_sinesp, linhas_sinesp = inspecionar_original_sinesp(
        caminho_sinesp, ano_sinesp
    )
    recortes_sinesp_por_ano[ano_sinesp] = recorte_sinesp
    amostras_sinesp_por_ano[ano_sinesp] = amostra_sinesp
    inventario_sinesp_linhas.append(ficha_sinesp)
    metadados_sinesp_por_ano[ano_sinesp] = pd.DataFrame({
        "ano_arquivo": ano_sinesp, "arquivo_origem": ficha_sinesp["arquivo"],
        "planilha_origem": ficha_sinesp["planilha"], "linha_xlsx": linhas_sinesp,
    })
    print(f"Sinesp {ano_sinesp}: estrutura compatível; {len(recorte_sinesp)} linhas de SC/JOINVILLE.")

inventario_sinesp = pd.DataFrame(inventario_sinesp_linhas)
comparacao_esquemas_sinesp = inventario_sinesp[["ano", "planilha", "colunas"]]
tipos_recortes_sinesp = pd.DataFrame({
    ano: quadro.dtypes for ano, quadro in recortes_sinesp_por_ano.items()
})
tipos_amostras_sinesp = pd.DataFrame({
    ano: quadro.dtypes for ano, quadro in amostras_sinesp_por_ano.items()
})
contagens_sinesp = inventario_sinesp[[
    "ano", "linhas_nacionais_sem_cabecalho", "linhas_joinville"
]]

# %%
inventario_sinesp

# %%
comparacao_esquemas_sinesp

# %%
contagens_sinesp

# %%
tipos_amostras_sinesp

# %%
tipos_recortes_sinesp

# %% [markdown]
# ## Concatenação autorizada somente em memória
# Granularidade: indicador, competência, município e abrangência, com dimensões
# agente/arma/faixa_etaria preservadas. A chave não é apenas ano/município.
# Os 900 registros anteriores são referência de conferência, não regra de filtro.
# Diferenças de tipos inferidos em campos com nulos serão mostradas, sem cast.
# Se houver esquema incompatível, a célula anterior interrompe antes de concat.

# %%
if any(quadro.columns.tolist() != transform.HEADERS
       for quadro in recortes_sinesp_por_ano.values()):
    raise ValueError("Esquemas incompatíveis: concatenação interrompida.")

df_sinesp_joinville_concatenado = pd.concat(
    list(recortes_sinesp_por_ano.values()), ignore_index=True, sort=False
)
metadados_sinesp_concatenado = pd.concat(
    list(metadados_sinesp_por_ano.values()), ignore_index=True, sort=False
)
linhas_esperadas_sinesp = sum(len(quadro) for quadro in recortes_sinesp_por_ano.values())
assert df_sinesp_joinville_concatenado.shape == (linhas_esperadas_sinesp, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
assert len(metadados_sinesp_concatenado) == linhas_esperadas_sinesp
assert df_sinesp_joinville_concatenado["uf"].eq("SC").all()
assert df_sinesp_joinville_concatenado["municipio"].eq("JOINVILLE").all()

# Confere valores, ausências e ordem de cada bloco. Inferências de dtype são
# relatadas separadamente, sem conversão manual ou substituição de ausências.
inicio_bloco_sinesp = 0
for ano_sinesp, recorte_sinesp in recortes_sinesp_por_ano.items():
    fim_bloco_sinesp = inicio_bloco_sinesp + len(recorte_sinesp)
    pd.testing.assert_frame_equal(
        df_sinesp_joinville_concatenado.iloc[inicio_bloco_sinesp:fim_bloco_sinesp].reset_index(drop=True),
        recorte_sinesp, check_dtype=False,
    )
    inicio_bloco_sinesp = fim_bloco_sinesp

referencia_contagens_sinesp = {2020: 132, 2021: 132, 2022: 132, 2023: 132,
                             2024: 132, 2025: 144, 2026: 96}
for ano_sinesp, quadro_sinesp in recortes_sinesp_por_ano.items():
    if len(quadro_sinesp) != referencia_contagens_sinesp[ano_sinesp]:
        print(f"ATENÇÃO: {ano_sinesp} diverge do inventário anterior; investigar sem ajustar linhas.")

print("Dimensões concatenadas:", df_sinesp_joinville_concatenado.shape)
print("Conferência: soma dos recortes, ordem das colunas e valores de cada bloco preservados.")

# %%
df_sinesp_joinville_concatenado.head(10)

# %%
df_sinesp_joinville_concatenado.tail(10)

# %%
tipos_concatenado_sinesp = df_sinesp_joinville_concatenado.dtypes.to_frame("tipo_pandas")
tipos_concatenado_sinesp

# %%
ausencias_concatenado_sinesp = df_sinesp_joinville_concatenado.isna().sum().to_frame("valores_ausentes")
ausencias_concatenado_sinesp

# %%
# Frequência de linhas tabulares por indicador; não é soma de vítimas ou crimes.
frequencia_eventos_sinesp = df_sinesp_joinville_concatenado["evento"].value_counts(dropna=False).to_frame("linhas_tabulares")
frequencia_eventos_sinesp

# %% [markdown]
# ## Limitações e próxima decisão
# As datas dos recortes continuam seriais Excel; não se convertem nesta etapa.
# Amostras nacionais de 20 linhas não representam os tipos de todo o arquivo.
# Ausência não significa zero; metadados auxiliares preservam ano/arquivo/linha.
# A etapa 6 abaixo foi autorizada para diagnóstico. Tratamento de ausências,
# tipagem, agregação, exportação e painel dependem de nova autorização.

# %% [markdown]
# # ETAPA 6 — Diagnóstico de duplicidades e qualidade Sinesp
# Pré-requisito: df_sinesp_joinville_concatenado já produzido pela etapa 5 no
# mesmo kernel. Esta etapa não lê arquivos nem repete a concatenação.
# As contagens de linhas e ausências são diagnósticas; não se somam medidas dos
# indicadores. A base é apenas consultada; quadros e máscaras são auxiliares.
# Execute cada célula e abra as variáveis em Variables > Data Viewer.

# %%
if "df_sinesp_joinville_concatenado" not in globals():
    raise RuntimeError("A etapa 6 requer o DataFrame da etapa 5 no mesmo kernel.")
assert df_sinesp_joinville_concatenado.shape == (900, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
retrato_antes_diagnostico_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)

# %% [markdown]
# ## Duplicidades exatas e possíveis repetições de competência
# data_referencia é a competência mensal original, incluindo o ano no serial.
# Usá-la diretamente evita converter datas ou criar mês/ano por suposição.
# A chave também mantém município/UF. Datas distintas não são agrupadas.
# Linhas repetidas na chave resumida são apenas candidatas: conferir agente,
# arma e faixa_etaria antes de interpretar. Valores ausentes entram na análise.

# %%
chave_competencia_sinesp = ["uf", "municipio", "data_referencia", "evento", "abrangencia"]
chave_dimensoes_sinesp = chave_competencia_sinesp + ["agente", "arma", "faixa_etaria"]
chave_sem_abrangencia_sinesp = ["uf", "municipio", "data_referencia", "evento"]
mascara_exatas_sinesp = df_sinesp_joinville_concatenado.duplicated(keep=False)
mascara_competencia_sinesp = df_sinesp_joinville_concatenado.duplicated(
    subset=chave_competencia_sinesp, keep=False
)
mascara_dimensoes_sinesp = df_sinesp_joinville_concatenado.duplicated(
    subset=chave_dimensoes_sinesp, keep=False
)
mascara_sem_abrangencia_sinesp = df_sinesp_joinville_concatenado.duplicated(
    subset=chave_sem_abrangencia_sinesp, keep=False
)
resumo_duplicidades_sinesp = pd.DataFrame([
    {
        "criterio": criterio,
        "linhas_envolvidas": int(mascara.sum()),
        "repeticoes_apos_primeira": int(df_sinesp_joinville_concatenado.duplicated(subset=chave).sum()),
        "grupos_repetidos": int((df_sinesp_joinville_concatenado.value_counts(subset=chave, dropna=False) > 1).sum()),
    }
    for criterio, chave, mascara in [
        ("14 colunas originais", transform.HEADERS, mascara_exatas_sinesp),
        ("competencia/evento/abrangencia/municipio", chave_competencia_sinesp, mascara_competencia_sinesp),
        ("chave com todas as dimensoes", chave_dimensoes_sinesp, mascara_dimensoes_sinesp),
        ("controle sem abrangencia", chave_sem_abrangencia_sinesp, mascara_sem_abrangencia_sinesp),
    ]
])
resumo_duplicidades_sinesp

# %%
linhas_duplicadas_exatas_sinesp = df_sinesp_joinville_concatenado.loc[mascara_exatas_sinesp]
linhas_duplicadas_exatas_sinesp

# %%
linhas_suspeitas_chave_sinesp = df_sinesp_joinville_concatenado.loc[mascara_competencia_sinesp]
linhas_suspeitas_chave_sinesp

# %%
# Controle: evidencia por que abrangencia precisa permanecer na chave.
linhas_repetidas_sem_abrangencia_sinesp = df_sinesp_joinville_concatenado.loc[mascara_sem_abrangencia_sinesp]
linhas_repetidas_sem_abrangencia_sinesp

# %%
# Metadados de ano/arquivo/linha ficam fora das 14 colunas, alinhados pelo índice.
# A identificação do mês permanece na competência original, sem conversão.
if "metadados_sinesp_concatenado" in globals():
    assert metadados_sinesp_concatenado.index.equals(df_sinesp_joinville_concatenado.index)
    origem_linhas_repetidas_sinesp = metadados_sinesp_concatenado.loc[mascara_sem_abrangencia_sinesp]
    origem_linhas_repetidas_sinesp

# %% [markdown]
# ## Ausências gerais e por indicador
# Percentuais usam a quantidade de linhas da base ou do evento correspondente.
# Nenhuma ausência é preenchida. Linhas com ausência são mantidas integralmente.

# %%
ausencias_colunas_sinesp = pd.DataFrame({
    "quantidade": df_sinesp_joinville_concatenado.isna().sum(),
    "percentual": df_sinesp_joinville_concatenado.isna().mean() * 100,
})
ausencias_colunas_sinesp

# %%
ausencias_eventos_linhas_sinesp = []
for evento_sinesp in df_sinesp_joinville_concatenado["evento"].unique():
    linhas_evento_sinesp = df_sinesp_joinville_concatenado.loc[
        df_sinesp_joinville_concatenado["evento"].eq(evento_sinesp)
    ]
    for campo_sinesp in df_sinesp_joinville_concatenado.columns:
        quantidade_ausente_sinesp = int(linhas_evento_sinesp[campo_sinesp].isna().sum())
        ausencias_eventos_linhas_sinesp.append({
            "evento": evento_sinesp, "campo": campo_sinesp,
            "linhas_evento": len(linhas_evento_sinesp),
            "ausencias": quantidade_ausente_sinesp,
            "percentual": quantidade_ausente_sinesp / len(linhas_evento_sinesp) * 100,
        })
ausencias_por_evento_sinesp = pd.DataFrame(ausencias_eventos_linhas_sinesp)
ausencias_por_evento_sinesp

# %%
# Matriz apenas para visualização das contagens, sem somar valores dos crimes.
matriz_ausencias_eventos_sinesp = ausencias_por_evento_sinesp.pivot(
    index="evento", columns="campo", values="ausencias"
)
matriz_ausencias_eventos_sinesp

# %%
# Recorte para inspecionar ausência de sexo em indicadores com total de vítimas.
# Trata-se de suspeita de detalhamento incompleto, não de erro comprovado.
campos_sexo_sinesp = ["feminino", "masculino", "nao_informado"]
mascara_sexo_ausente_sinesp = (
    df_sinesp_joinville_concatenado["total_vitima"].notna()
    & df_sinesp_joinville_concatenado[campos_sexo_sinesp].isna().any(axis=1)
)
linhas_sexo_ausente_sinesp = df_sinesp_joinville_concatenado.loc[mascara_sexo_ausente_sinesp]
linhas_sexo_ausente_sinesp

# %% [markdown]
# ## Tipos e necessidade futura de interpretar datas
# No recorte da etapa 5, data_referencia foi lida do XML como serial inteiro.
# A coluna permanece intacta. Não se chama to_datetime nem se calcula calendário.
# Futuramente: conferir o sistema de datas do XLSX, validar a regra de conversão
# e amostras por ano, mantendo referência ao valor original e à origem.
# float64 nas contagens pode decorrer de ausências, não de medidas fracionárias.

# %%
diagnostico_tipos_sinesp = df_sinesp_joinville_concatenado.dtypes.to_frame("tipo_pandas")
diagnostico_tipos_sinesp

# %%
amostra_datas_originais_sinesp = df_sinesp_joinville_concatenado[["data_referencia"]].head(12)
amostra_datas_originais_sinesp

# %% [markdown]
# ## Interpretação cautelosa e integridade
# Campos total_vitima e total atendem a indicadores diferentes; ausência de uma
# medida quando a outra se aplica é compatível com estrutura, sem autorizar zero.
# total_peso não é preenchido neste recorte; agente/arma/faixa_etaria não trazem
# detalhamento municipal. Campo inteiramente vazio não comprova defeito.
# Sexo ausente com total_vitima informado requer conferir original e definição;
# não inferir sexo, não calcular parcela residual e não declarar ausência = zero.
# A próxima etapa depende de autorização específica, após avaliar os quadros.

# %%
pd.testing.assert_frame_equal(
    df_sinesp_joinville_concatenado, retrato_antes_diagnostico_sinesp,
    check_dtype=True, check_exact=True,
)
assert df_sinesp_joinville_concatenado.isna().equals(retrato_antes_diagnostico_sinesp.isna())
assert df_sinesp_joinville_concatenado.shape == (900, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
print("Etapa 6: 900 linhas e 14 colunas preservadas; valores, tipos, índices e ausências intactos.")

# %% [markdown]
# # ETAPA 7 — Investigação e propostas de colunas e tipos
# Execute no mesmo kernel das etapas 5–6. Não há leitura de CSV histórico,
# concatenação, limpeza, seleção definitiva de colunas ou conversão da base.
# As tabelas auxiliares apenas apresentam evidências e recomendações.

# %%
for requisito_sinesp in (
    "df_sinesp_joinville_concatenado", "metadados_sinesp_concatenado",
    "arquivos_sinesp_por_ano", "amostras_sinesp_por_ano",
    "linhas_sexo_ausente_sinesp", "linhas_repetidas_sem_abrangencia_sinesp",
    "retrato_antes_diagnostico_sinesp", "ausencias_colunas_sinesp", "campos_sexo_sinesp",
):
    if requisito_sinesp not in globals():
        raise RuntimeError(f"Execute as etapas 5–6 no mesmo kernel: falta {requisito_sinesp}.")
assert df_sinesp_joinville_concatenado.shape == (900, 14)
assert df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS
assert pd.api.types.is_integer_dtype(df_sinesp_joinville_concatenado["data_referencia"])
assert metadados_sinesp_concatenado.index.equals(df_sinesp_joinville_concatenado.index)
pd.testing.assert_frame_equal(
    df_sinesp_joinville_concatenado, retrato_antes_diagnostico_sinesp,
    check_dtype=True, check_exact=True,
)
pd.testing.assert_frame_equal(
    linhas_sexo_ausente_sinesp,
    df_sinesp_joinville_concatenado.loc[
        df_sinesp_joinville_concatenado["total_vitima"].notna()
        & df_sinesp_joinville_concatenado[campos_sexo_sinesp].isna().any(axis=1)
    ], check_exact=True,
)
pd.testing.assert_frame_equal(
    linhas_repetidas_sem_abrangencia_sinesp,
    df_sinesp_joinville_concatenado.loc[
        df_sinesp_joinville_concatenado.duplicated(
            subset=["uf", "municipio", "data_referencia", "evento"], keep=False
        )
    ], check_exact=True,
)
retrato_antes_investigacao_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)
ausencias_confirmadas_etapa7_sinesp = pd.DataFrame({
    "quantidade": df_sinesp_joinville_concatenado.isna().sum(),
    "percentual": df_sinesp_joinville_concatenado.isna().mean() * 100,
})
pd.testing.assert_frame_equal(
    ausencias_confirmadas_etapa7_sinesp, ausencias_colunas_sinesp,
    check_exact=True,
)
ausencias_confirmadas_etapa7_sinesp

# %% [markdown]
# ## Rótulos de competência para inspeção, sem converter seriais da base
# A amostra read_excel da etapa 5 já contém datas interpretadas pelo leitor.
# Confrontamos cada uma com o serial da MESMA célula original (D2 a D21).
# Isso fornece um rótulo de apresentação; nenhum serial é convertido ou trocado.
# Meses sem correspondência comprovada ficam identificados para investigação.

# %%
rotulos_competencia_sinesp = {}
sistemas_datas_originais_sinesp = []
for ano_rotulo_sinesp, arquivo_rotulo_sinesp in arquivos_sinesp_por_ano.items():
    livro_rotulo_sinesp = load_workbook(arquivo_rotulo_sinesp, read_only=True)
    try:
        sistemas_datas_originais_sinesp.append({
            "ano": ano_rotulo_sinesp, "epoch_xlsx": repr(livro_rotulo_sinesp.epoch),
        })
    finally:
        livro_rotulo_sinesp.close()
    with ZipFile(arquivo_rotulo_sinesp) as pacote_rotulo_sinesp:
        with pacote_rotulo_sinesp.open("xl/worksheets/sheet1.xml") as fluxo_rotulo_sinesp:
            prefixo_rotulo_sinesp = b""
            while prefixo_rotulo_sinesp.count(b"</row>") < 21:
                trecho_rotulo_sinesp = fluxo_rotulo_sinesp.read(65536)
                if not trecho_rotulo_sinesp:
                    break
                prefixo_rotulo_sinesp += trecho_rotulo_sinesp
    for correspondencia_rotulo_sinesp in re.finditer(
        rb'<row\b[^>]*\br="(\d+)"[^>]*>.*?</row>', prefixo_rotulo_sinesp, re.DOTALL
    ):
        linha_rotulo_sinesp = int(correspondencia_rotulo_sinesp.group(1))
        if not 2 <= linha_rotulo_sinesp <= 21:
            continue
        elemento_rotulo_sinesp = transform._parse_sheet_row(correspondencia_rotulo_sinesp.group())
        celula_data_rotulo_sinesp = next(
            c for c in elemento_rotulo_sinesp.findall(transform.NS + "c")
            if c.get("r") == f"D{linha_rotulo_sinesp}"
        )
        serial_rotulo_sinesp = transform._cell_value(celula_data_rotulo_sinesp, [])
        data_amostra_rotulo_sinesp = amostras_sinesp_por_ano[ano_rotulo_sinesp].iloc[
            linha_rotulo_sinesp - 2
        ]["data_referencia"]
        if not isinstance(data_amostra_rotulo_sinesp, pd.Timestamp):
            raise ValueError("Data da amostra não interpretada como Timestamp; revisar o leitor.")
        chave_rotulo_sinesp = (ano_rotulo_sinesp, serial_rotulo_sinesp)
        rotulo_competencia_sinesp = data_amostra_rotulo_sinesp.strftime("%Y-%m")
        if chave_rotulo_sinesp in rotulos_competencia_sinesp:
            assert rotulos_competencia_sinesp[chave_rotulo_sinesp] == rotulo_competencia_sinesp
        rotulos_competencia_sinesp[chave_rotulo_sinesp] = rotulo_competencia_sinesp
sistemas_datas_sinesp = pd.DataFrame(sistemas_datas_originais_sinesp)
sistemas_datas_sinesp

# %%
investigacao_sexo_linhas_sinesp = []
for indice_sexo_sinesp, linha_sexo_sinesp in linhas_sexo_ausente_sinesp.iterrows():
    origem_sexo_sinesp = metadados_sinesp_concatenado.loc[indice_sexo_sinesp]
    ano_sexo_sinesp = origem_sexo_sinesp["ano_arquivo"]
    parcelas_informadas_sinesp = [
        linha_sexo_sinesp[campo] for campo in campos_sexo_sinesp
        if pd.notna(linha_sexo_sinesp[campo])
    ]
    soma_informada_sinesp = sum(parcelas_informadas_sinesp) if parcelas_informadas_sinesp else None
    if not parcelas_informadas_sinesp:
        relacao_sexo_sinesp = "Sem parcelas de sexo; reconciliação indisponível."
    elif soma_informada_sinesp == linha_sexo_sinesp["total_vitima"]:
        relacao_sexo_sinesp = "Parcelas informadas coincidem com total; não inferir ausentes."
    else:
        relacao_sexo_sinesp = "Parcelas informadas diferem do total; investigar definição e preenchimento."
    investigacao_sexo_linhas_sinesp.append({
        "indice_base": indice_sexo_sinesp, "ano": ano_sexo_sinesp,
        "competencia": rotulos_competencia_sinesp.get(
            (ano_sexo_sinesp, linha_sexo_sinesp["data_referencia"]), "SEM RÓTULO VALIDADO"
        ),
        "serial_original": linha_sexo_sinesp["data_referencia"],
        "evento": linha_sexo_sinesp["evento"], "abrangencia": linha_sexo_sinesp["abrangencia"],
        **{campo: linha_sexo_sinesp[campo] for campo in campos_sexo_sinesp},
        "total_vitima": linha_sexo_sinesp["total_vitima"],
        "soma_apenas_informados": soma_informada_sinesp,
        "relacao": relacao_sexo_sinesp,
        "arquivo_origem": origem_sexo_sinesp["arquivo_origem"],
        "linha_xlsx": origem_sexo_sinesp["linha_xlsx"],
    })
investigacao_sexo_sinesp = pd.DataFrame(investigacao_sexo_linhas_sinesp)
investigacao_sexo_sinesp

# %% [markdown]
# ## Mandados estaduais e Polícia Federal: comparação sem somar
# Mantém os 40 registros em um quadro e apresenta cada par em colunas separadas.
# A abrangência distingue as séries; exclusividade estatística não é presumida.

# %%
mandados_investigacao_sinesp = pd.DataFrame([
    {
        "indice_base": indice,
        "ano": metadados_sinesp_concatenado.loc[indice, "ano_arquivo"],
        "competencia": rotulos_competencia_sinesp.get(
            (metadados_sinesp_concatenado.loc[indice, "ano_arquivo"], linha["data_referencia"]),
            "SEM RÓTULO VALIDADO",
        ),
        "serial_original": linha["data_referencia"], "evento": linha["evento"],
        "abrangencia": linha["abrangencia"], "total": linha["total"],
        "agente": linha["agente"], "arma": linha["arma"], "faixa_etaria": linha["faixa_etaria"],
        "arquivo_origem": metadados_sinesp_concatenado.loc[indice, "arquivo_origem"],
        "linha_xlsx": metadados_sinesp_concatenado.loc[indice, "linha_xlsx"],
    }
    for indice, linha in linhas_repetidas_sem_abrangencia_sinesp.iterrows()
])
mandados_investigacao_sinesp

# %%
pares_mandados_sinesp = mandados_investigacao_sinesp.pivot(
    index=["ano", "competencia", "serial_original", "evento"],
    columns="abrangencia", values="total",
)
pares_mandados_sinesp

# %% [markdown]
# ## Avaliação das 14 colunas e proposta de tipagem (sem aplicação)
# A seleção futura será uma visão analítica, mantendo sempre a camada original.
# Os campos inteiramente vazios são candidatos apenas nessa visão e dependem
# de verificar outros indicadores/atualizações. Nenhuma coluna é removida aqui.

# %%
propostas_campos_sinesp = {
    "uf": ("Manter", "Território estadual e filtro exato", "municipio", "string", "Impede misturar municípios homônimos."),
    "municipio": ("Manter", "Território municipal", "uf", "string", "Identifica o recorte; não autoriza desagregar bairros."),
    "evento": ("Manter", "Identificador original do indicador", "abrangencia e campos de medida", "string; category apenas após validar domínio", "Preservar rótulos e distinguir unidades e sobreposições."),
    "data_referencia": ("Manter", "Competência mensal", "sistema de datas e arquivo/ano de origem", "datetime64[ns], após validar conversão", "Manter referência ao serial original para rastreabilidade."),
    "agente": ("Candidata à exclusão da visão futura", "Possível dimensão de agentes; dicionário pendente", "evento e abrangencia", "string anulável se mantida", "100% ausente neste recorte; não remover da camada original."),
    "arma": ("Candidata à exclusão da visão futura", "Possível dimensão de arma; dicionário pendente", "evento", "string anulável se mantida", "100% ausente; avaliar novos indicadores antes de selecionar colunas."),
    "faixa_etaria": ("Candidata à exclusão da visão futura", "Possível dimensão etária; dicionário pendente", "evento e medida de vítimas", "string anulável se mantida", "Não oferece detalhamento neste recorte; preservar no original."),
    "feminino": ("Manter e avaliar ausências", "Parcela de vítimas por sexo", "evento e total_vitima", "Int64 anulável", "Validar integralidade, não negatividade e significado das ausências."),
    "masculino": ("Manter e avaliar ausências", "Parcela de vítimas por sexo", "evento e total_vitima", "Int64 anulável", "Não inferir zero a partir da soma das parcelas informadas."),
    "nao_informado": ("Manter e avaliar ausências", "Parcela de vítimas com sexo não informado", "evento e total_vitima", "Int64 anulável", "Ausência da célula difere da categoria sexo não informado."),
    "total_vitima": ("Manter", "Total publicado de vítimas", "evento, sexo e abrangencia", "Int64 anulável", "Não substituir por total nem somar novamente às parcelas."),
    "total": ("Manter", "Contagem publicada de mandados no recorte", "evento e abrangencia", "Int64 anulável", "Manter estadual e PF separados; não representa vítimas únicas."),
    "total_peso": ("Candidata à exclusão da visão futura", "Medida de peso; unidade pendente", "evento e unidade oficial", "decimal anulável se houver uso e unidade confirmada", "100% ausente no recorte; não presumir quilogramas ou preencher."),
    "abrangencia": ("Manter", "Distingue séries estaduais, PF e PRF", "evento e competência", "string; category após validar domínio", "É necessária à chave; não identifica corporação individual em toda série."),
}
proposta_colunas_sinesp = pd.DataFrame([
    {
        "coluna": campo, "tipo_atual": str(df_sinesp_joinville_concatenado[campo].dtype),
        "percentual_ausente": ausencias_confirmadas_etapa7_sinesp.loc[campo, "percentual"],
        "exemplos_originais": repr(df_sinesp_joinville_concatenado[campo].unique()[:3].tolist()),
        "recomendacao": proposta[0], "finalidade": proposta[1],
        "dependencias": proposta[2], "tipo_proposto": proposta[3], "justificativa": proposta[4],
    }
    for campo, proposta in propostas_campos_sinesp.items()
])
proposta_colunas_sinesp

# %% [markdown]
# ## Validações necessárias antes de qualquer conversão futura
# Datas: validar epoch 1900/1904 em todos os XLSX, regra de dias seriais e
# particularidade do calendário Excel; conferir ausência, frações, limites,
# correspondência ano/mês e primeiro dia do mês. Confrontar o resultado futuro
# com células lidas como datas pelo Excel/openpyxl e preservar a origem.
# Contagens: conferir valores inteiros, finitos e não negativos; preservar nulos
# e categorias por unidade. Só depois propor Int64, sem preencher as ausências.
# Texto: manter grafia original; eventual categoria exige domínio validado.
# Esta etapa não aplica nenhuma dessas conversões ou seleções.

# %%
pd.testing.assert_frame_equal(
    df_sinesp_joinville_concatenado, retrato_antes_investigacao_sinesp,
    check_dtype=True, check_exact=True,
)
assert df_sinesp_joinville_concatenado.isna().equals(retrato_antes_investigacao_sinesp.isna())
assert df_sinesp_joinville_concatenado.shape == (900, 14)
print("Etapa 7: investigação concluída em memória; 900 linhas, 14 colunas e valores originais preservados.")

# %% [markdown]
# # ETAPA 8 — Preparação controlada da cópia analítica Sinesp
# Execute depois das etapas 5–7 no mesmo processo/kernel. A etapa 5 reconstrói
# os recortes pelos sete XLSX, sem usar CSV histórico. Nesta etapa, qualquer
# divergência interrompe a execução antes de disponibilizar a base analítica.
# A seleção usa .loc[..., colunas].copy(deep=True); a base original de 14
# colunas e seus metadados nunca recebem atribuições de valores ou tipos.

# %%
import hashlib
import math


def exigir_sinesp_etapa8(condicao, mensagem):
    if not condicao:
        raise ValueError(f"Etapa 8 interrompida: {mensagem}")


def sha256_original_sinesp(caminho):
    resumo = hashlib.sha256()
    with caminho.open("rb") as fluxo:
        for bloco in iter(lambda: fluxo.read(1024 * 1024), b""):
            resumo.update(bloco)
    return resumo.hexdigest()


for requisito_etapa8 in (
    "df_sinesp_joinville_concatenado", "metadados_sinesp_concatenado",
    "recortes_sinesp_por_ano", "arquivos_sinesp_por_ano",
    "amostras_sinesp_por_ano", "rotulos_competencia_sinesp",
    "linhas_sexo_ausente_sinesp", "linhas_repetidas_sem_abrangencia_sinesp",
):
    exigir_sinesp_etapa8(requisito_etapa8 in globals(), f"falta {requisito_etapa8}; execute etapas 5–7.")
exigir_sinesp_etapa8(df_sinesp_joinville_concatenado.shape == (900, 14), "dimensões originais divergentes.")
exigir_sinesp_etapa8(df_sinesp_joinville_concatenado.columns.tolist() == transform.HEADERS, "esquema original divergente.")
original_preservado_etapa8 = df_sinesp_joinville_concatenado.copy(deep=True)
metadados_preservados_etapa8 = metadados_sinesp_concatenado.copy(deep=True)

# Confere hashes contra o inventário anterior, não apenas antes/depois desta sessão.
inventario_hash_sinesp_etapa8 = pd.read_csv(
    RAIZ_PROJETO / "historico/documentacao_anterior/auditoria/inventario_sha256_pre_reorganizacao.csv"
)
hashes_originais_etapa8 = {}
epocas_originais_etapa8 = {}
for ano_etapa8, caminho_etapa8 in arquivos_sinesp_por_ano.items():
    relativo_etapa8 = caminho_etapa8.relative_to(RAIZ_PROJETO).as_posix()
    evidencia_hash_etapa8 = inventario_hash_sinesp_etapa8.loc[
        inventario_hash_sinesp_etapa8["caminho_destino_planejado"].eq(relativo_etapa8)
    ]
    exigir_sinesp_etapa8(len(evidencia_hash_etapa8) == 1, f"evidência de hash ausente/ambígua: {relativo_etapa8}")
    hash_etapa8 = sha256_original_sinesp(caminho_etapa8)
    exigir_sinesp_etapa8(hash_etapa8 == evidencia_hash_etapa8.iloc[0]["sha256"], f"hash divergente: {relativo_etapa8}")
    exigir_sinesp_etapa8(caminho_etapa8.stat().st_size == evidencia_hash_etapa8.iloc[0]["tamanho_bytes"], f"tamanho divergente: {relativo_etapa8}")
    hashes_originais_etapa8[ano_etapa8] = hash_etapa8
    livro_epoca_etapa8 = load_workbook(caminho_etapa8, read_only=True)
    try:
        epocas_originais_etapa8[ano_etapa8] = livro_epoca_etapa8.epoch
    finally:
        livro_epoca_etapa8.close()
exigir_sinesp_etapa8(
    all(epoca == pd.Timestamp("1899-12-30") for epoca in epocas_originais_etapa8.values()),
    "sistema de datas diferente do Excel 1900 previamente identificado.",
)

# Valida o alinhamento de origem e cada bloco da reconstrução da etapa 5.
exigir_sinesp_etapa8(metadados_sinesp_concatenado.shape == (900, 4), "dimensões dos metadados divergentes.")
exigir_sinesp_etapa8(metadados_sinesp_concatenado.index.equals(original_preservado_etapa8.index), "índices de origem desalinhados.")
exigir_sinesp_etapa8(not metadados_sinesp_concatenado.isna().any().any(), "metadados incompletos.")
exigir_sinesp_etapa8(not metadados_sinesp_concatenado.duplicated(["arquivo_origem", "planilha_origem", "linha_xlsx"]).any(), "linhas de origem repetidas.")
contagens_referencia_etapa8 = {2020: 132, 2021: 132, 2022: 132, 2023: 132, 2024: 132, 2025: 144, 2026: 96}
eventos_referencia_etapa8 = {
    "Feminicídio": 80, "Homicídio doloso": 80, "Lesão corporal seguida de morte": 80,
    "Mandado de prisão cumprido": 100,
    "Morte no trânsito ou em decorrência dele (exceto homicídio doloso)": 80,
    "Mortes a esclarecer (sem indício de crime)": 80, "Mortes no trânsito": 80,
    "Roubo seguido de morte (latrocínio)": 80, "Suicídio": 80,
    "Tentativa de feminicídio": 80, "Tentativa de homicídio": 80,
}
exigir_sinesp_etapa8(original_preservado_etapa8["evento"].value_counts().to_dict() == eventos_referencia_etapa8, "indicadores/frequências diferentes do inventário.")
exigir_sinesp_etapa8(original_preservado_etapa8["uf"].eq("SC").all() and original_preservado_etapa8["municipio"].eq("JOINVILLE").all(), "recorte municipal divergente.")
inicio_etapa8 = 0
for ano_etapa8, quantidade_etapa8 in contagens_referencia_etapa8.items():
    bloco_etapa8 = original_preservado_etapa8.iloc[inicio_etapa8:inicio_etapa8 + quantidade_etapa8]
    origem_bloco_etapa8 = metadados_sinesp_concatenado.iloc[inicio_etapa8:inicio_etapa8 + quantidade_etapa8]
    pd.testing.assert_frame_equal(bloco_etapa8.reset_index(drop=True), recortes_sinesp_por_ano[ano_etapa8], check_exact=True)
    exigir_sinesp_etapa8(origem_bloco_etapa8["ano_arquivo"].eq(ano_etapa8).all(), "ano de origem desalinhado.")
    exigir_sinesp_etapa8(origem_bloco_etapa8["planilha_origem"].eq(str(ano_etapa8)).all(), "planilha de origem divergente.")
    exigir_sinesp_etapa8(origem_bloco_etapa8["arquivo_origem"].eq(str(arquivos_sinesp_por_ano[ano_etapa8].relative_to(RAIZ_PROJETO))).all(), "arquivo de origem divergente.")
    exigir_sinesp_etapa8(origem_bloco_etapa8["linha_xlsx"].ge(2).all(), "linha de origem inválida.")
    inicio_etapa8 += quantidade_etapa8
print("Etapa 8: pré-condições aprovadas; sete originais íntegros, 900 linhas municipais e origem alinhada.")

# %% [markdown]
# ## Validar datas antes de atribuí-las à cópia
# pd.to_datetime(unit="D", origin=..., errors="raise") interpreta o serial.
# A série temporária precisa passar por TODAS as verificações abaixo. Só então
# será atribuída à cópia. Não se usa errors="coerce" ou correção de datas.

# %%
seriais_etapa8 = original_preservado_etapa8["data_referencia"]
exigir_sinesp_etapa8(seriais_etapa8.notna().all(), "serial de data ausente.")
exigir_sinesp_etapa8(all(
    isinstance(valor, (int, float)) and not isinstance(valor, bool)
    and math.isfinite(valor) and valor == int(valor) and 0 < valor < 100000
    for valor in seriais_etapa8.tolist()
), "serial inválido, fracionário ou fora dos limites previstos.")
datas_validadas_etapa8 = pd.to_datetime(seriais_etapa8, unit="D", origin="1899-12-30", errors="raise")
exigir_sinesp_etapa8(datas_validadas_etapa8.dt.year.eq(metadados_sinesp_concatenado["ano_arquivo"]).all(), "data não corresponde ao ano do arquivo.")
exigir_sinesp_etapa8(datas_validadas_etapa8.dt.day.eq(1).all(), "competência fora do primeiro dia do mês.")
exigir_sinesp_etapa8(datas_validadas_etapa8.eq(datas_validadas_etapa8.dt.normalize()).all(), "horário inesperado na competência.")
datas_leitor_etapa8 = {}
for ano_etapa8, amostra_etapa8 in amostras_sinesp_por_ano.items():
    for data_leitor_etapa8 in amostra_etapa8["data_referencia"]:
        chave_leitor_etapa8 = (ano_etapa8, data_leitor_etapa8.strftime("%Y-%m"))
        if chave_leitor_etapa8 in datas_leitor_etapa8:
            exigir_sinesp_etapa8(datas_leitor_etapa8[chave_leitor_etapa8] == data_leitor_etapa8, "amostra Excel tem datas conflitantes.")
        datas_leitor_etapa8[chave_leitor_etapa8] = data_leitor_etapa8
for indice_etapa8, data_etapa8 in datas_validadas_etapa8.items():
    ano_etapa8 = metadados_sinesp_concatenado.loc[indice_etapa8, "ano_arquivo"]
    competencia_etapa8 = data_etapa8.strftime("%Y-%m")
    exigir_sinesp_etapa8(rotulos_competencia_sinesp.get((ano_etapa8, seriais_etapa8.loc[indice_etapa8])) == competencia_etapa8, "serial não corresponde à célula Excel amostrada.")
    exigir_sinesp_etapa8(datas_leitor_etapa8.get((ano_etapa8, competencia_etapa8)) == data_etapa8, "data difere da interpretação do leitor Excel.")
for ano_etapa8 in contagens_referencia_etapa8:
    meses_etapa8 = sorted(datas_validadas_etapa8.loc[datas_validadas_etapa8.dt.year.eq(ano_etapa8)].dt.month.unique().tolist())
    exigir_sinesp_etapa8(meses_etapa8 == list(range(1, 9 if ano_etapa8 == 2026 else 13)), f"cobertura mensal divergente em {ano_etapa8}.")

# %% [markdown]
# ## Validar contagens antes da tipagem anulável
# Apenas os valores preenchidos são examinados. Não há arredondamento, soma de
# indicadores ou substituição de ausência. Int64 mantém NA em cada posição.

# %%
campos_contagens_etapa8 = ["feminino", "masculino", "nao_informado", "total_vitima", "total"]
for campo_etapa8 in campos_contagens_etapa8:
    valores_etapa8 = original_preservado_etapa8.loc[original_preservado_etapa8[campo_etapa8].notna(), campo_etapa8]
    exigir_sinesp_etapa8(all(
        math.isfinite(valor) and valor == int(valor) and 0 <= valor <= 2**63 - 1
        for valor in valores_etapa8
    ), f"contagem não finita, fracionária, negativa ou fora de Int64: {campo_etapa8}")

# %% [markdown]
# ## Criar a candidata de 900 × 10, sem atribuir à base original
# A variável definitiva só é disponibilizada depois das conferências finais.
# Metadados ficam separados, conservando o serial e o hash do arquivo de origem.

# %%
colunas_analiticas_sinesp = [
    "uf", "municipio", "evento", "data_referencia", "feminino", "masculino",
    "nao_informado", "total_vitima", "total", "abrangencia",
]
candidata_analitica_sinesp = df_sinesp_joinville_concatenado.loc[:, colunas_analiticas_sinesp].copy(deep=True)
candidata_analitica_sinesp["data_referencia"] = datas_validadas_etapa8
for campo_etapa8 in campos_contagens_etapa8:
    candidata_analitica_sinesp[campo_etapa8] = candidata_analitica_sinesp[campo_etapa8].astype("Int64")
metadados_candidatos_sinesp = metadados_sinesp_concatenado.copy(deep=True)
metadados_candidatos_sinesp["serial_original"] = seriais_etapa8.copy(deep=True)
metadados_candidatos_sinesp["sha256_original"] = metadados_candidatos_sinesp["ano_arquivo"].map(hashes_originais_etapa8)

# %% [markdown]
# ## Catálogo auxiliar: evento e abrangência definem medida e limitações
# Um catálogo não agrega valores. Feminicídios e tentativas mantêm rótulos
# próprios. Mandados usam total; os demais eventos deste recorte usam vítimas.

# %%
naturezas_eventos_etapa8 = {
    "Feminicídio": "criminal", "Homicídio doloso": "criminal",
    "Lesão corporal seguida de morte": "criminal", "Mandado de prisão cumprido": "atividade policial",
    "Morte no trânsito ou em decorrência dele (exceto homicídio doloso)": "complementar: mortalidade no trânsito",
    "Mortes a esclarecer (sem indício de crime)": "complementar: mortes a esclarecer",
    "Mortes no trânsito": "complementar: mortalidade no trânsito",
    "Roubo seguido de morte (latrocínio)": "criminal", "Suicídio": "complementar: suicídio",
    "Tentativa de feminicídio": "criminal", "Tentativa de homicídio": "criminal",
}
catalogo_linhas_etapa8 = []
for evento_etapa8 in original_preservado_etapa8["evento"].unique():
    for abrangencia_etapa8 in original_preservado_etapa8.loc[original_preservado_etapa8["evento"].eq(evento_etapa8), "abrangencia"].unique():
        mandado_etapa8 = evento_etapa8 == "Mandado de prisão cumprido"
        catalogo_linhas_etapa8.append({
            "evento_original": evento_etapa8, "abrangencia": abrangencia_etapa8,
            "medida_utilizada": "total" if mandado_etapa8 else "total_vitima",
            "unidade_medida": "mandados cumpridos (contagem publicada)" if mandado_etapa8 else "vítimas",
            "natureza": naturezas_eventos_etapa8[evento_etapa8],
            "limitacoes": (
                "Não representa pessoas únicas; estadual e PF separados; exclusividade não demonstrada."
                if mandado_etapa8 else
                "Não somar às parcelas de sexo; manter categoria e abrangência; sobreposição entre indicadores/fontes requer definição."
            ) + " 2026 parcial; linha tabular não equivale a ocorrência individual.",
        })
catalogo_candidato_sinesp = pd.DataFrame(catalogo_linhas_etapa8)
exigir_sinesp_etapa8(len(catalogo_candidato_sinesp) == 12, "combinações evento/abrangência inesperadas.")
mascara_mandados_etapa8 = candidata_analitica_sinesp["evento"].eq("Mandado de prisão cumprido")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[mascara_mandados_etapa8, "total"].notna().all(), "mandados sem medida total.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[mascara_mandados_etapa8, "total_vitima"].isna().all(), "medidas de mandados/vítimas inesperadas.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[~mascara_mandados_etapa8, "total_vitima"].notna().all(), "indicadores de vítimas sem total.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[~mascara_mandados_etapa8, "total"].isna().all(), "medida total fora dos mandados.")

# %% [markdown]
# ## Conferências finais: valores, nulos, linhas, tipos e origens
# Compara máscaras de ausência e valores preenchidos, sem somar medidas.
# As 20 linhas com sexo incompleto e os 20 pares estaduais/PF são preservados.

# %%
exigir_sinesp_etapa8(candidata_analitica_sinesp.shape == (900, 10), "dimensões analíticas divergentes.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.index.equals(original_preservado_etapa8.index), "índices analíticos divergentes.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.isna().equals(original_preservado_etapa8[colunas_analiticas_sinesp].isna()), "posições de ausência alteradas.")
for campo_etapa8 in campos_contagens_etapa8:
    preenchidos_etapa8 = original_preservado_etapa8[campo_etapa8].notna()
    exigir_sinesp_etapa8(str(candidata_analitica_sinesp[campo_etapa8].dtype) == "Int64", "tipagem de contagem divergente.")
    exigir_sinesp_etapa8(candidata_analitica_sinesp.loc[preenchidos_etapa8, campo_etapa8].eq(original_preservado_etapa8.loc[preenchidos_etapa8, campo_etapa8]).all(), f"valores alterados: {campo_etapa8}")
for campo_etapa8 in ["uf", "municipio", "evento", "abrangencia"]:
    pd.testing.assert_series_equal(candidata_analitica_sinesp[campo_etapa8], original_preservado_etapa8[campo_etapa8], check_exact=True)
mascara_sexo_etapa8 = candidata_analitica_sinesp["total_vitima"].notna() & candidata_analitica_sinesp[campos_sexo_sinesp].isna().any(axis=1)
exigir_sinesp_etapa8(int(mascara_sexo_etapa8.sum()) == 20, "quantidade de linhas com sexo incompleto divergente.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.index[mascara_sexo_etapa8].equals(linhas_sexo_ausente_sinesp.index), "linhas com sexo incompleto alteradas.")
mascara_pares_etapa8 = candidata_analitica_sinesp.duplicated(["uf", "municipio", "data_referencia", "evento"], keep=False)
exigir_sinesp_etapa8(int(mascara_pares_etapa8.sum()) == 40, "pares de mandados divergentes.")
exigir_sinesp_etapa8(candidata_analitica_sinesp.index[mascara_pares_etapa8].equals(linhas_repetidas_sem_abrangencia_sinesp.index), "pares de mandados alterados.")
for ano_etapa8, caminho_etapa8 in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa8(sha256_original_sinesp(caminho_etapa8) == hashes_originais_etapa8[ano_etapa8], "arquivo original alterado durante a preparação.")
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_preservado_etapa8, check_dtype=True, check_exact=True)
pd.testing.assert_frame_equal(metadados_sinesp_concatenado, metadados_preservados_etapa8, check_dtype=True, check_exact=True)
pd.testing.assert_frame_equal(metadados_candidatos_sinesp[metadados_preservados_etapa8.columns], metadados_preservados_etapa8, check_exact=True)
exigir_sinesp_etapa8(metadados_candidatos_sinesp.index.equals(candidata_analitica_sinesp.index), "metadados analíticos desalinhados.")
df_sinesp_joinville_analitico = candidata_analitica_sinesp
metadados_sinesp_analitico = metadados_candidatos_sinesp
catalogo_indicadores_sinesp = catalogo_candidato_sinesp
print("Etapa 8 aprovada: original 900 × 14 intacto; analítico 900 × 10; sete hashes preservados.")

# %%
comparacao_tipos_etapa8_sinesp = pd.DataFrame({
    "antes": original_preservado_etapa8[colunas_analiticas_sinesp].dtypes,
    "depois": df_sinesp_joinville_analitico.dtypes,
})
comparacao_tipos_etapa8_sinesp

# %%
comparacao_ausencias_etapa8_sinesp = pd.DataFrame({
    "antes": original_preservado_etapa8[colunas_analiticas_sinesp].isna().sum(),
    "depois": df_sinesp_joinville_analitico.isna().sum(),
})
comparacao_ausencias_etapa8_sinesp

# %%
conferencia_datas_etapa8_sinesp = pd.DataFrame({
    "ano_arquivo": metadados_sinesp_analitico["ano_arquivo"],
    "serial_original": metadados_sinesp_analitico["serial_original"],
    "data_analitica": df_sinesp_joinville_analitico["data_referencia"],
})
conferencia_datas_etapa8_sinesp

# %%
sexo_incompleto_analitico_sinesp = df_sinesp_joinville_analitico.loc[mascara_sexo_etapa8]
sexo_incompleto_analitico_sinesp

# %%
mandados_pares_analitico_sinesp = df_sinesp_joinville_analitico.loc[mascara_pares_etapa8]
mandados_pares_analitico_sinesp

# %%
df_sinesp_joinville_analitico

# %%
catalogo_indicadores_sinesp

# %%
metadados_sinesp_analitico

# %% [markdown]
# # ETAPA 9 — Análise exploratória descritiva Sinesp
# Reconstrua as etapas 5–8 no mesmo processo/kernel. A análise consulta a base
# analítica sem novos tratamentos. Fonte, indicador, abrangência e unidade
# acompanham cada resultado. Não há total geral de crimes nem taxas populacionais.

# %%
def exigir_sinesp_etapa9(condicao, mensagem):
    if not condicao:
        raise ValueError(f"Etapa 9 interrompida: {mensagem}")


for requisito_etapa9 in (
    "df_sinesp_joinville_concatenado", "df_sinesp_joinville_analitico",
    "catalogo_indicadores_sinesp", "metadados_sinesp_analitico",
    "hashes_originais_etapa8", "original_preservado_etapa8",
):
    exigir_sinesp_etapa9(requisito_etapa9 in globals(), f"execute etapas 5–8: falta {requisito_etapa9}.")
exigir_sinesp_etapa9(df_sinesp_joinville_concatenado.shape == (900, 14), "dimensões originais divergentes.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico.shape == (900, 10), "dimensões analíticas divergentes.")
exigir_sinesp_etapa9(pd.api.types.is_datetime64_any_dtype(df_sinesp_joinville_analitico["data_referencia"]), "datas não tipadas.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico["data_referencia"].notna().all(), "datas ausentes.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico["data_referencia"].dt.day.eq(1).all(), "competência mensal inesperada.")
exigir_sinesp_etapa9(all(str(df_sinesp_joinville_analitico[c].dtype) == "Int64" for c in campos_contagens_etapa8), "medidas não anuláveis Int64.")
exigir_sinesp_etapa9(not df_sinesp_joinville_analitico.duplicated(["uf", "municipio", "evento", "abrangencia", "data_referencia"]).any(), "repetição de competência na mesma série.")
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_preservado_etapa8, check_exact=True)
pd.testing.assert_frame_equal(catalogo_indicadores_sinesp, catalogo_candidato_sinesp, check_exact=True)
for ano_etapa9, caminho_etapa9 in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa9(sha256_original_sinesp(caminho_etapa9) == hashes_originais_etapa8[ano_etapa9], "arquivo original com hash divergente.")
original_antes_eda_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)
analitico_antes_eda_sinesp = df_sinesp_joinville_analitico.copy(deep=True)
catalogo_antes_eda_sinesp = catalogo_indicadores_sinesp.copy(deep=True)
metadados_antes_eda_sinesp = metadados_sinesp_analitico.copy(deep=True)

# %% [markdown]
# ## Tabela mensal de observações publicadas
# Consulta o campo de medida definido pelo catálogo: vítimas em total_vitima;
# mandados em total. Cada linha continua sendo uma observação mensal agregada,
# não uma ocorrência individual. Os campos ano/mês pertencem só à apresentação.

# %%
catalogo_consulta_eda_sinesp = {
    (linha.evento_original, linha.abrangencia): linha
    for linha in catalogo_indicadores_sinesp.itertuples(index=False)
}
linhas_mensais_eda_sinesp = []
for linha_eda_sinesp in df_sinesp_joinville_analitico.itertuples(index=False):
    chave_eda_sinesp = (linha_eda_sinesp.evento, linha_eda_sinesp.abrangencia)
    exigir_sinesp_etapa9(chave_eda_sinesp in catalogo_consulta_eda_sinesp, "indicador sem catálogo.")
    definicao_eda_sinesp = catalogo_consulta_eda_sinesp[chave_eda_sinesp]
    linhas_mensais_eda_sinesp.append({
        "fonte": "MJSP / Sinesp VDE", "uf": linha_eda_sinesp.uf,
        "municipio": linha_eda_sinesp.municipio, "evento": linha_eda_sinesp.evento,
        "abrangencia": linha_eda_sinesp.abrangencia,
        "medida": definicao_eda_sinesp.medida_utilizada,
        "unidade": definicao_eda_sinesp.unidade_medida, "natureza": definicao_eda_sinesp.natureza,
        "competencia": linha_eda_sinesp.data_referencia,
        "ano": linha_eda_sinesp.data_referencia.year, "mes": linha_eda_sinesp.data_referencia.month,
        "valor_publicado": getattr(linha_eda_sinesp, definicao_eda_sinesp.medida_utilizada),
    })
tabela_mensal_sinesp = pd.DataFrame(linhas_mensais_eda_sinesp)
exigir_sinesp_etapa9(len(tabela_mensal_sinesp) == 900, "observações perdidas na apresentação mensal.")
chaves_series_eda_sinesp = ["fonte", "uf", "municipio", "evento", "abrangencia", "medida", "unidade", "natureza"]
tabela_mensal_sinesp

# %% [markdown]
# ## Estatísticas mensais por série e ano
# Soma com min_count=1 preserva ausência quando não há medida preenchida.
# Média/mediana usam meses com medida; os dois denominadores ficam explícitos.
# Desvio-padrão descritivo com ddof=0 representa a dispersão dos meses observados
# no recorte, sem inferência amostral; só é calculado com pelo menos dois valores.
# Meses de valor zero são observações publicadas, não ausências preenchidas.

# %%
def descrever_meses_sinesp(grupo):
    valores = grupo["valor_publicado"]
    n_validos = int(valores.count())
    return {
        "meses_disponiveis": int(grupo["competencia"].nunique()),
        "meses_com_medida": n_validos,
        "soma_medida_publicada": valores.sum(min_count=1),
        "media_mensal": valores.mean() if n_validos else pd.NA,
        "mediana_mensal": valores.median() if n_validos else pd.NA,
        "minimo_mensal": valores.min() if n_validos else pd.NA,
        "maximo_mensal": valores.max() if n_validos else pd.NA,
        "desvio_padrao_mensal_ddof0": valores.std(ddof=0) if n_validos >= 2 else pd.NA,
        "meses_zero_publicado": int(valores.eq(0).sum()),
    }


linhas_anuais_eda_sinesp = []
linhas_equivalentes_eda_sinesp = []
for chave_ano_eda_sinesp, grupo_ano_eda_sinesp in tabela_mensal_sinesp.groupby(chaves_series_eda_sinesp + ["ano"], sort=False, dropna=False):
    identificacao_ano_eda_sinesp = dict(zip(chaves_series_eda_sinesp + ["ano"], chave_ano_eda_sinesp))
    meses_ano_eda_sinesp = tuple(sorted(grupo_ano_eda_sinesp["mes"].unique().tolist()))
    cobertura_ano_eda_sinesp = "ano completo" if meses_ano_eda_sinesp == tuple(range(1, 13)) else "parcial janeiro–agosto" if meses_ano_eda_sinesp == tuple(range(1, 9)) else "cobertura incompleta: investigar"
    linhas_anuais_eda_sinesp.append({
        **identificacao_ano_eda_sinesp, "cobertura": cobertura_ano_eda_sinesp,
        "meses_presentes": meses_ano_eda_sinesp, **descrever_meses_sinesp(grupo_ano_eda_sinesp),
    })
    grupo_equivalente_eda_sinesp = grupo_ano_eda_sinesp.loc[grupo_ano_eda_sinesp["mes"].le(8)]
    exigir_sinesp_etapa9(tuple(sorted(grupo_equivalente_eda_sinesp["mes"].unique())) == tuple(range(1, 9)), "janeiro–agosto incompleto; não comparar.")
    linhas_equivalentes_eda_sinesp.append({
        **identificacao_ano_eda_sinesp, "periodo": "janeiro–agosto",
        **descrever_meses_sinesp(grupo_equivalente_eda_sinesp),
    })
tabela_anual_sinesp = pd.DataFrame(linhas_anuais_eda_sinesp)
comparacao_jan_ago_sinesp = pd.DataFrame(linhas_equivalentes_eda_sinesp)
estatisticas_descritivas_sinesp = pd.DataFrame([
    {
        **dict(zip(chaves_series_eda_sinesp, chave)),
        "inicio_observado": grupo["competencia"].min(), "fim_observado": grupo["competencia"].max(),
        "escopo": "todo o período observado; inclui 2026 parcial",
        **descrever_meses_sinesp(grupo),
    }
    for chave, grupo in tabela_mensal_sinesp.groupby(chaves_series_eda_sinesp, sort=False, dropna=False)
])
exigir_sinesp_etapa9(len(tabela_anual_sinesp) == 79 and len(comparacao_jan_ago_sinesp) == 79, "quantidade de séries/anos divergente.")
tabela_anual_sinesp

# %%
comparacao_jan_ago_sinesp

# %%
estatisticas_descritivas_sinesp

# %% [markdown]
# ## Comparação equivalente 2026 versus 2025
# Todos os totais abaixo têm oito meses. Percentual só é calculado com base
# anterior positiva; zero anterior não é substituído por outro denominador.
# Variação descritiva não demonstra efeito de políticas ou ações policiais.

# %%
linhas_comparacao_2026_sinesp = []
for chave_2026_sinesp, grupo_2026_sinesp in comparacao_jan_ago_sinesp.groupby(chaves_series_eda_sinesp, sort=False, dropna=False):
    anteriores_2026_sinesp = grupo_2026_sinesp.loc[grupo_2026_sinesp["ano"].eq(2025)]
    atuais_2026_sinesp = grupo_2026_sinesp.loc[grupo_2026_sinesp["ano"].eq(2026)]
    exigir_sinesp_etapa9(len(anteriores_2026_sinesp) == 1 and len(atuais_2026_sinesp) == 1, "comparação 2025/2026 indisponível.")
    anterior_2026_sinesp = anteriores_2026_sinesp.iloc[0]["soma_medida_publicada"]
    atual_2026_sinesp = atuais_2026_sinesp.iloc[0]["soma_medida_publicada"]
    ambos_informados_2026_sinesp = pd.notna(anterior_2026_sinesp) and pd.notna(atual_2026_sinesp)
    linhas_comparacao_2026_sinesp.append({
        **dict(zip(chaves_series_eda_sinesp, chave_2026_sinesp)), "periodo": "janeiro–agosto",
        "valor_2025": anterior_2026_sinesp, "valor_2026": atual_2026_sinesp,
        "diferenca_absoluta": atual_2026_sinesp - anterior_2026_sinesp if ambos_informados_2026_sinesp else pd.NA,
        "variacao_percentual": (atual_2026_sinesp / anterior_2026_sinesp - 1) * 100 if ambos_informados_2026_sinesp and anterior_2026_sinesp > 0 else pd.NA,
        "ressalva": "percentual não calculável: base zero ou ausente" if not ambos_informados_2026_sinesp or anterior_2026_sinesp == 0 else "variação descritiva; observar magnitude absoluta",
    })
comparacao_2026_2025_sinesp = pd.DataFrame(linhas_comparacao_2026_sinesp)
comparacao_2026_2025_sinesp

# %% [markdown]
# ## Tabelas separadas por indicador para Data Viewer
# O dicionário disponibiliza tabelas anuais, mensais e janeiro–agosto de cada
# evento. Abrangências continuam identificadas dentro de cada tabela.

# %%
tabelas_por_indicador_sinesp = {
    evento: {
        "anual": tabela_anual_sinesp.loc[tabela_anual_sinesp["evento"].eq(evento)],
        "mensal": tabela_mensal_sinesp.loc[tabela_mensal_sinesp["evento"].eq(evento)],
        "janeiro_agosto": comparacao_jan_ago_sinesp.loc[comparacao_jan_ago_sinesp["evento"].eq(evento)],
    }
    for evento in df_sinesp_joinville_analitico["evento"].unique()
}
feminicidio_anual_sinesp = tabelas_por_indicador_sinesp["Feminicídio"]["anual"]
tentativa_feminicidio_anual_sinesp = tabelas_por_indicador_sinesp["Tentativa de feminicídio"]["anual"]
homicidio_doloso_anual_sinesp = tabelas_por_indicador_sinesp["Homicídio doloso"]["anual"]
tentativa_homicidio_anual_sinesp = tabelas_por_indicador_sinesp["Tentativa de homicídio"]["anual"]
latrocinio_anual_sinesp = tabelas_por_indicador_sinesp["Roubo seguido de morte (latrocínio)"]["anual"]
lesao_seguida_morte_anual_sinesp = tabelas_por_indicador_sinesp["Lesão corporal seguida de morte"]["anual"]
complementares_anuais_sinesp = tabela_anual_sinesp.loc[~tabela_anual_sinesp["natureza"].eq("criminal")]
violencia_mulher_anual_sinesp = tabela_anual_sinesp.loc[tabela_anual_sinesp["evento"].isin(["Feminicídio", "Tentativa de feminicídio"])]
violencia_mulher_jan_ago_sinesp = comparacao_jan_ago_sinesp.loc[comparacao_jan_ago_sinesp["evento"].isin(["Feminicídio", "Tentativa de feminicídio"])]
feminicidio_anual_sinesp

# %%
tentativa_feminicidio_anual_sinesp

# %%
homicidio_doloso_anual_sinesp

# %%
tentativa_homicidio_anual_sinesp

# %%
latrocinio_anual_sinesp

# %%
lesao_seguida_morte_anual_sinesp

# %%
complementares_anuais_sinesp

# %%
violencia_mulher_jan_ago_sinesp

# %% [markdown]
# ## Limites da interpretação e integridade final
# Os totais de 2026 representam apenas janeiro–agosto. Indicadores raros, com
# muitos meses de zero ou poucos casos, devem ser interpretados pela magnitude
# absoluta; percentuais podem parecer grandes com bases pequenas.
# Não somar feminicídio a homicídio ou as duas tentativas. Trânsito, suicídio,
# mortes a esclarecer e mandados são complementares. Não há explicação causal,
# taxa populacional, gráfico ou exportação nesta etapa.

# %%
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_antes_eda_sinesp, check_exact=True)
pd.testing.assert_frame_equal(df_sinesp_joinville_analitico, analitico_antes_eda_sinesp, check_exact=True)
pd.testing.assert_frame_equal(catalogo_indicadores_sinesp, catalogo_antes_eda_sinesp, check_exact=True)
pd.testing.assert_frame_equal(metadados_sinesp_analitico, metadados_antes_eda_sinesp, check_exact=True)
for ano_etapa9, caminho_etapa9 in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa9(sha256_original_sinesp(caminho_etapa9) == hashes_originais_etapa8[ano_etapa9], "arquivo original alterado durante a EDA.")
print("Etapa 9: 900 observações mensais, 79 séries/anos, 79 comparações janeiro–agosto e 12 resumos descritivos; bases e arquivos preservados.")

# %% [markdown]
# # ETAPA 10 — Visualização exploratória Sinesp
# Pré-requisito: tabelas validadas da etapa 9 no mesmo processo/kernel.
# As figuras usam apenas valores publicados dessas tabelas. Não se recalculam
# medidas, não se alteram as bases e não se exportam imagens. No Python
# Interactive, execute cada célula com plt.show() para visualizar o gráfico.

# %%
import textwrap

import matplotlib.pyplot as plt
from matplotlib.ticker import MaxNLocator

exigir_sinesp_etapa9(df_sinesp_joinville_concatenado.shape == (900, 14), "base original divergente antes dos gráficos.")
exigir_sinesp_etapa9(df_sinesp_joinville_analitico.shape == (900, 10), "base analítica divergente antes dos gráficos.")
original_antes_graficos_sinesp = df_sinesp_joinville_concatenado.copy(deep=True)
analitico_antes_graficos_sinesp = df_sinesp_joinville_analitico.copy(deep=True)
anual_antes_graficos_sinesp = tabela_anual_sinesp.copy(deep=True)
jan_ago_antes_graficos_sinesp = comparacao_jan_ago_sinesp.copy(deep=True)
cores_graficos_sinesp = {
    "Homicídio doloso": "#2563A6", "Tentativa de homicídio": "#188977",
    "Feminicídio": "#7A5195", "Tentativa de feminicídio": "#D97706",
}


def selecionar_tabela_grafico_sinesp(tabela, eventos, anos, meses):
    """Seleciona e ordena observações já agregadas na etapa 9, sem preencher."""
    dados = tabela.loc[
        tabela["evento"].isin(eventos) & tabela["ano"].isin(anos)
        & tabela["abrangencia"].eq("Estadual")
    ].copy(deep=True)
    exigir_sinesp_etapa9(len(dados) == len(eventos) * len(anos), "observações inesperadas para o gráfico.")
    exigir_sinesp_etapa9(not dados.duplicated(["ano", "evento"]).any(), "ano/indicador repetido no gráfico.")
    exigir_sinesp_etapa9(dados["fonte"].eq("MJSP / Sinesp VDE").all(), "fonte divergente no gráfico.")
    exigir_sinesp_etapa9(dados["unidade"].eq("vítimas").all() and dados["medida"].eq("total_vitima").all(), "unidades incompatíveis no gráfico.")
    exigir_sinesp_etapa9(dados["meses_disponiveis"].eq(meses).all() and dados["meses_com_medida"].eq(meses).all(), "cobertura incompleta para o gráfico.")
    exigir_sinesp_etapa9(dados["soma_medida_publicada"].notna().all(), "valor ausente; gráfico interrompido sem preenchimento.")
    if meses == 12:
        exigir_sinesp_etapa9(dados["cobertura"].eq("ano completo").all(), "período parcial na série anual.")
    else:
        exigir_sinesp_etapa9(dados["periodo"].eq("janeiro–agosto").all(), "períodos não equivalentes.")
    quadro = dados.pivot(index="ano", columns="evento", values="soma_medida_publicada").reindex(index=anos, columns=eventos)
    exigir_sinesp_etapa9(not quadro.isna().any().any(), "célula ausente na tabela visual.")
    for linha in dados.itertuples(index=False):
        exigir_sinesp_etapa9(quadro.loc[linha.ano, linha.evento] == linha.soma_medida_publicada, "valor visual difere da etapa 9.")
    return quadro


def preparar_figura_sinesp(titulo, periodo, interpretacao, limite_y):
    fig, ax = plt.subplots(figsize=(10, 6.2), dpi=120)
    fig.subplots_adjust(left=0.09, right=0.97, bottom=0.27, top=0.80)
    fig.suptitle(titulo, x=0.08, y=0.97, ha="left", fontsize=15, fontweight="semibold")
    fig.text(0.08, 0.91, periodo, fontsize=11, color="#444444")
    ax.set_xlabel("Ano", fontsize=11)
    ax.set_ylabel("Vítimas", fontsize=11)
    ax.set_ylim(0, limite_y)
    ax.yaxis.set_major_locator(MaxNLocator(nbins=6, integer=True))
    ax.set_axisbelow(True)
    ax.grid(axis="y", color="#DDDDDD", linewidth=0.7)
    ax.spines[["top", "right"]].set_visible(False)
    fig.text(0.08, 0.125, textwrap.fill(interpretacao, width=108), va="top", fontsize=9, color="#333333")
    fig.text(0.08, 0.035, "Fonte: MJSP / Sinesp VDE | Abrangência: Estadual | Unidade: vítimas", fontsize=9, color="#555555")
    return fig, ax


def desenhar_linhas_sinesp(ax, dados):
    for numero, evento in enumerate(dados.columns):
        valores = dados[evento].tolist()
        artista, = ax.plot(
            dados.index, valores, label=evento, color=cores_graficos_sinesp[evento],
            marker="o" if numero == 0 else "s", linestyle="-" if numero == 0 else "--",
            linewidth=2, markersize=6,
        )
        exigir_sinesp_etapa9(list(artista.get_ydata()) == valores, "coordenadas da linha divergentes.")
        for ano, valor in zip(dados.index, valores):
            ax.annotate(f"{valor:g}", (ano, valor), textcoords="offset points", xytext=(0, 8), ha="center", fontsize=10)
    ax.set_xticks(dados.index.tolist())
    if len(dados.columns) > 1:
        ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=2, frameon=False, fontsize=10)


def desenhar_barras_sinesp(ax, dados):
    posicoes = list(range(len(dados)))
    largura = 0.34
    for numero, evento in enumerate(dados.columns):
        valores = dados[evento].tolist()
        barras = ax.bar(
            [posicao + (numero - 0.5) * largura for posicao in posicoes], valores,
            width=largura, label=evento, color=cores_graficos_sinesp[evento],
            hatch="//" if numero == 0 else None,
        )
        exigir_sinesp_etapa9([barra.get_height() for barra in barras] == valores, "altura de barra divergente.")
        ax.bar_label(barras, labels=[f"{valor:g}" for valor in valores], padding=4, fontsize=10)
    ax.set_xticks(posicoes, dados.index.tolist())
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.01), ncol=2, frameon=False, fontsize=10)


dados_grafico_homicidio_anual_sinesp = selecionar_tabela_grafico_sinesp(
    tabela_anual_sinesp, ["Homicídio doloso"], list(range(2020, 2026)), 12
)
dados_grafico_mulher_anual_sinesp = selecionar_tabela_grafico_sinesp(
    tabela_anual_sinesp, ["Feminicídio", "Tentativa de feminicídio"], list(range(2020, 2026)), 12
)
dados_grafico_homicidios_jan_ago_sinesp = selecionar_tabela_grafico_sinesp(
    comparacao_jan_ago_sinesp, ["Homicídio doloso", "Tentativa de homicídio"], list(range(2020, 2027)), 8
)
dados_grafico_mulher_jan_ago_sinesp = selecionar_tabela_grafico_sinesp(
    comparacao_jan_ago_sinesp, ["Feminicídio", "Tentativa de feminicídio"], [2025, 2026], 8
)
interpretacoes_graficos_sinesp = {
    "homicidio_anual": f"O total publicado passou de {dados_grafico_homicidio_anual_sinesp.loc[2020, 'Homicídio doloso']} vítimas em 2020 para {dados_grafico_homicidio_anual_sinesp.loc[2025, 'Homicídio doloso']} em 2025, com oscilações intermediárias. A série não demonstra causas.",
    "mulher_anual": f"Feminicídios variaram de {dados_grafico_mulher_anual_sinesp['Feminicídio'].min()} a {dados_grafico_mulher_anual_sinesp['Feminicídio'].max()} vítimas anuais; tentativas, de {dados_grafico_mulher_anual_sinesp['Tentativa de feminicídio'].min()} a {dados_grafico_mulher_anual_sinesp['Tentativa de feminicídio'].max()}. São pequenas contagens de categorias distintas, sem soma entre elas.",
    "homicidios_jan_ago": f"De janeiro a agosto, homicídios passaram de {dados_grafico_homicidios_jan_ago_sinesp.loc[2025, 'Homicídio doloso']} em 2025 para {dados_grafico_homicidios_jan_ago_sinesp.loc[2026, 'Homicídio doloso']} em 2026; tentativas, de {dados_grafico_homicidios_jan_ago_sinesp.loc[2025, 'Tentativa de homicídio']} para {dados_grafico_homicidios_jan_ago_sinesp.loc[2026, 'Tentativa de homicídio']}. Os períodos têm a mesma duração.",
    "mulher_jan_ago": f"Tentativas passaram de {dados_grafico_mulher_jan_ago_sinesp.loc[2025, 'Tentativa de feminicídio']} para {dados_grafico_mulher_jan_ago_sinesp.loc[2026, 'Tentativa de feminicídio']} vítimas. Feminicídio tem zero publicado nos dois recortes; isso não comprova eliminação do fenômeno.",
}

# %% [markdown]
# ## Gráfico 1 — Homicídio doloso: anos completos
# Linhas com marcadores e valores explícitos; eixo iniciado em zero. 2026 não
# integra esta série anual porque sua cobertura termina em agosto.

# %%
fig_homicidio_anual_sinesp, ax_homicidio_anual_sinesp = preparar_figura_sinesp(
    "Homicídio doloso em Joinville/SC", "Anos completos: 2020–2025",
    interpretacoes_graficos_sinesp["homicidio_anual"], 70,
)
desenhar_linhas_sinesp(ax_homicidio_anual_sinesp, dados_grafico_homicidio_anual_sinesp)
plt.show()

# %% [markdown]
# ## Gráfico 2 — Feminicídio e tentativa: séries anuais distintas
# Barras lado a lado, cores e hachura distintas. Escala única, inteira, iniciada
# em zero; não há eixo secundário, empilhamento ou soma de categorias.

# %%
fig_mulher_anual_sinesp, ax_mulher_anual_sinesp = preparar_figura_sinesp(
    "Feminicídio e tentativa de feminicídio em Joinville/SC", "Anos completos: 2020–2025",
    interpretacoes_graficos_sinesp["mulher_anual"], 12,
)
desenhar_barras_sinesp(ax_mulher_anual_sinesp, dados_grafico_mulher_anual_sinesp)
plt.show()

# %% [markdown]
# ## Gráfico 3 — Homicídio e tentativa: janeiro–agosto
# Sete recortes com oito meses cada. Marcadores e estilos distinguem categorias;
# nenhuma linha usa o total anual completo dos anos anteriores.

# %%
fig_homicidios_jan_ago_sinesp, ax_homicidios_jan_ago_sinesp = preparar_figura_sinesp(
    "Homicídio doloso e tentativa de homicídio em Joinville/SC", "Períodos equivalentes: janeiro–agosto de 2020 a 2026",
    interpretacoes_graficos_sinesp["homicidios_jan_ago"], 90,
)
desenhar_linhas_sinesp(ax_homicidios_jan_ago_sinesp, dados_grafico_homicidios_jan_ago_sinesp)
plt.show()

# %% [markdown]
# ## Gráfico 4 — Violência contra a mulher: janeiro–agosto de 2025 e 2026
# Mesmas cores e escala do gráfico 2. Os rótulos zero permanecem visíveis;
# pequenos números são apresentados sem ampliar percentuais no gráfico.

# %%
fig_mulher_jan_ago_sinesp, ax_mulher_jan_ago_sinesp = preparar_figura_sinesp(
    "Feminicídio e tentativa de feminicídio em Joinville/SC", "Períodos equivalentes: janeiro–agosto de 2025 e 2026",
    interpretacoes_graficos_sinesp["mulher_jan_ago"], 12,
)
desenhar_barras_sinesp(ax_mulher_jan_ago_sinesp, dados_grafico_mulher_jan_ago_sinesp)
plt.show()

# %%
pd.testing.assert_frame_equal(df_sinesp_joinville_concatenado, original_antes_graficos_sinesp, check_exact=True)
pd.testing.assert_frame_equal(df_sinesp_joinville_analitico, analitico_antes_graficos_sinesp, check_exact=True)
pd.testing.assert_frame_equal(tabela_anual_sinesp, anual_antes_graficos_sinesp, check_exact=True)
pd.testing.assert_frame_equal(comparacao_jan_ago_sinesp, jan_ago_antes_graficos_sinesp, check_exact=True)
for ano_grafico_sinesp, caminho_grafico_sinesp in arquivos_sinesp_por_ano.items():
    exigir_sinesp_etapa9(sha256_original_sinesp(caminho_grafico_sinesp) == hashes_originais_etapa8[ano_grafico_sinesp], "original alterado durante as visualizações.")
print("Etapa 10: quatro figuras em memória; valores plotados iguais às tabelas da etapa 9; bases 900 × 14 e 900 × 10 preservadas.")

# %% [markdown]
# # ETAPA 11 — Inspeção dos dados brutos SSP/SC
# Execute a partir desta célula, sem executar as etapas Sinesp. As importações e
# os caminhos abaixo são independentes. Leitura literal: os CSVs permanecem em
# um dicionário por arquivo, sem concatenação, limpeza ou conversão de valores.
# As tabelas auxiliares são evidências de inspeção, não uma base analítica final.

# %%
from pathlib import Path as CaminhoSSP11
from collections import defaultdict as agrupar_ssp11
import hashlib as hash_ssp11
import re as regex_ssp11

import pandas as pd_ssp11
import pdfplumber as pdf_ssp11
from pypdf import PdfReader as LeitorSSP11
from IPython.display import display as mostrar_ssp11

raiz_ssp11 = CaminhoSSP11(__file__).resolve().parents[2] if "__file__" in globals() else CaminhoSSP11.cwd()
pasta_pdf_ssp11 = raiz_ssp11 / "dados/brutos/ssp_sc/pdfs"
pasta_csv_ssp11 = raiz_ssp11 / "dados/brutos/ssp_sc/csv_extraidos"
arquivos_pdf_ssp11 = sorted(pasta_pdf_ssp11.rglob("*.pdf"))
arquivos_csv_ssp11 = sorted(pasta_csv_ssp11.rglob("*.csv"))
if len(arquivos_pdf_ssp11) != 45 or len(arquivos_csv_ssp11) != 45:
    raise ValueError("Etapa 11: esperado acervo de 45 PDFs e 45 CSVs; conferir os caminhos e o inventário.")
if {p.stem for p in arquivos_pdf_ssp11} != {p.stem for p in arquivos_csv_ssp11}:
    raise ValueError("Etapa 11: os nomes-base dos PDFs e CSVs não correspondem.")


def sha256_ssp11(caminho):
    """Verifica bytes originais; não grava nem altera os arquivos."""
    with caminho.open("rb") as fluxo:
        return hash_ssp11.file_digest(fluxo, "sha256").hexdigest()


hashes_antes_ssp11 = {p: sha256_ssp11(p) for p in arquivos_pdf_ssp11 + arquivos_csv_ssp11}
# Se as bases Sinesp estiverem no kernel, conferir sua preservação ao terminar.
# A ausência dessas variáveis não impede a inspeção independente da SSP.
retratos_sinesp_ssp11 = {
    nome: globals()[nome].copy(deep=True)
    for nome in ("df_sinesp_joinville_concatenado", "df_sinesp_joinville_analitico")
    if nome in globals()
}
csvs_brutos_ssp = {}
leitores_ssp11 = {}
textos_pdf_ssp = {}
metadados_pdf_ssp11 = {}
linhas_inventario_ssp11 = []
linhas_amostras_ssp11 = []
linhas_qualidade_ssp11 = []
grupos_esquemas_ssp11 = agrupar_ssp11(list)
problemas_ssp11 = []
meses_ssp11 = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")


def texto_pagina_ssp11(nome, pagina):
    chave = (nome, pagina)
    if chave not in textos_pdf_ssp:
        textos_pdf_ssp[chave] = leitores_ssp11[nome].pages[pagina - 1].extract_text() or ""
    return textos_pdf_ssp[chave]


def celulas_ssp11(linha):
    """Retorna as células em sua ordem original, inclusive strings vazias."""
    return tuple(linha[c] for c in linha.index if c.startswith("celula_"))


def problema_ssp11(arquivo, detalhe):
    problemas_ssp11.append({"arquivo": arquivo, "pendencia": detalhe})


# %% [markdown]
# ## 11.1 — Inventário, esquemas e ausências literais
# dtype=str e keep_default_na=False preservam tokens, inclusive zeros e vazios.
# Vazio em celula_* pode ser preenchimento retangular da extração, célula mesclada
# ou ausência no PDF; não equivale automaticamente a falha nem a contagem zero.
# O período vem do conteúdo do PDF, não do prefixo do nome do arquivo.

# %%
for caminho_pdf_ssp11 in arquivos_pdf_ssp11:
    nome_ssp11 = caminho_pdf_ssp11.name
    caminho_csv_ssp11 = next(p for p in arquivos_csv_ssp11 if p.stem == caminho_pdf_ssp11.stem)
    bruto_ssp11 = pd_ssp11.read_csv(
        caminho_csv_ssp11, sep=";", encoding="utf-8-sig", dtype=str,
        keep_default_na=False, na_filter=False,
    )
    csvs_brutos_ssp[caminho_csv_ssp11.name] = bruto_ssp11
    leitores_ssp11[nome_ssp11] = LeitorSSP11(caminho_pdf_ssp11)
    paginas_ssp11 = len(leitores_ssp11[nome_ssp11].pages)
    domestico_ssp11 = "VIOLENCIA-DOMESTICA" in nome_ssp11
    primeiro_texto_ssp11 = texto_pagina_ssp11(nome_ssp11, 1)
    ano_cabecalho_ssp11 = regex_ssp11.search(r"\b20\d{2}\b", primeiro_texto_ssp11)
    ano_ssp11 = ano_cabecalho_ssp11.group() if ano_cabecalho_ssp11 else "pendente"
    if domestico_ssp11:
        if paginas_ssp11 % 2:
            raise ValueError(f"Etapa 11: páginas temáticas sem par em {nome_ssp11}.")
        metade_ssp11 = paginas_ssp11 // 2
        cabecalho_direito_ssp11 = texto_pagina_ssp11(nome_ssp11, metade_ssp11 + 1)
        # Examinar apenas o cabeçalho anterior à primeira linha numérica.
        meses_esquerda_ssp11 = regex_ssp11.findall(r"\b(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)\b", primeiro_texto_ssp11[:120])
        meses_direita_ssp11 = regex_ssp11.findall(r"\b(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)\b", cabecalho_direito_ssp11[:80])
        periodo_ssp11 = f"{ano_ssp11}: {'/'.join(meses_esquerda_ssp11 + meses_direita_ssp11)}"
        atualizacoes_ssp11 = ()
    else:
        metade_ssp11 = None
        meses_esquerda_ssp11, meses_direita_ssp11 = [], []
        texto_segundo_ssp11 = texto_pagina_ssp11(nome_ssp11, 2)
        # Datas de atualização podem diferir entre a página estadual e homicídios.
        atualizacoes_ssp11 = tuple(regex_ssp11.findall(
            r"ATUALIZAÇÃO DOS DADOS:\s*(\d{2}/\d{2}/\d{4})",
            primeiro_texto_ssp11 + "\n" + texto_segundo_ssp11,
        ))
        limite_ssp11 = regex_ssp11.search(r"PERÍODO\s*\((01/01\s*-\s*\d{2}/\d{2})\)", primeiro_texto_ssp11)
        periodo_ssp11 = f"{ano_ssp11}: {limite_ssp11.group(1) if limite_ssp11 else 'pendente'}"
    metadados_pdf_ssp11[nome_ssp11] = {
        "ano_edicao": ano_ssp11, "periodo_edicao": periodo_ssp11,
        "atualizacoes_paginas_1_2": atualizacoes_ssp11,
        "tipo_relatorio": "Violência doméstica" if domestico_ssp11 else "Boletim mensal de indicadores",
        "paginas": paginas_ssp11, "metade": metade_ssp11,
        "meses_esquerda": meses_esquerda_ssp11, "meses_direita": meses_direita_ssp11,
    }
    colunas_ssp11 = tuple(bruto_ssp11.columns)
    grupos_esquemas_ssp11[colunas_ssp11].append(caminho_csv_ssp11.name)
    for caminho_ssp11, tipo_ssp11 in ((caminho_pdf_ssp11, "PDF"), (caminho_csv_ssp11, "CSV")):
        linhas_inventario_ssp11.append({
            "arquivo": caminho_ssp11.name, "tipo": tipo_ssp11,
            "periodo_referencia": periodo_ssp11,
            "tipo_relatorio": metadados_pdf_ssp11[nome_ssp11]["tipo_relatorio"],
            "paginas": paginas_ssp11 if tipo_ssp11 == "PDF" else None,
            "linhas": len(bruto_ssp11) if tipo_ssp11 == "CSV" else None,
            "colunas": len(colunas_ssp11) if tipo_ssp11 == "CSV" else None,
            "cabecalhos": colunas_ssp11 if tipo_ssp11 == "CSV" else tuple(primeiro_texto_ssp11.splitlines()[:10]),
            "associado": caminho_csv_ssp11.name if tipo_ssp11 == "PDF" else nome_ssp11,
            "bytes": caminho_ssp11.stat().st_size, "sha256": hashes_antes_ssp11[caminho_ssp11],
        })
    for indice_ssp11, linha_ssp11 in bruto_ssp11.head(3).iterrows():
        linhas_amostras_ssp11.append({
            "arquivo_csv": caminho_csv_ssp11.name, "indice_csv": indice_ssp11,
            **{c: linha_ssp11[c] for c in colunas_ssp11[:4]},
            "celulas_originais": celulas_ssp11(linha_ssp11),
        })
    for coluna_ssp11 in colunas_ssp11:
        linhas_qualidade_ssp11.append({
            "arquivo_csv": caminho_csv_ssp11.name, "coluna": coluna_ssp11,
            "tipo": str(bruto_ssp11[coluna_ssp11].dtype),
            "vazios_literais": int(bruto_ssp11[coluna_ssp11].eq("").sum()),
            "nulos_pandas": int(bruto_ssp11[coluna_ssp11].isna().sum()),
        })
    if set(bruto_ssp11.arquivo_pdf) != {nome_ssp11}:
        raise ValueError(f"Etapa 11: procedência divergente em {caminho_csv_ssp11.name}.")
    for coluna_ssp11 in ("pagina_pdf", "tabela_pdf", "linha_tabela"):
        if not bruto_ssp11[coluna_ssp11].map(lambda x: bool(regex_ssp11.fullmatch(r"[1-9]\d*", x))).all():
            raise ValueError(f"Etapa 11: referência inválida: {caminho_csv_ssp11.name}/{coluna_ssp11}.")
    if {int(p) for p in bruto_ssp11.pagina_pdf} != set(range(1, paginas_ssp11 + 1)):
        raise ValueError(f"Etapa 11: cobertura de páginas divergente em {nome_ssp11}.")
    if bruto_ssp11.duplicated(["arquivo_pdf", "pagina_pdf", "tabela_pdf", "linha_tabela"]).any():
        raise ValueError(f"Etapa 11: rastreabilidade repetida dentro de {nome_ssp11}.")

inventario_arquivos_ssp = pd_ssp11.DataFrame(linhas_inventario_ssp11)
amostras_csv_ssp = pd_ssp11.DataFrame(linhas_amostras_ssp11)
qualidade_csv_ssp = pd_ssp11.DataFrame(linhas_qualidade_ssp11)
esquemas_ssp = pd_ssp11.DataFrame([
    {"esquema": numero, "colunas": colunas, "quantidade_colunas": len(colunas),
     "quantidade_arquivos": len(nomes), "quantidade_linhas": sum(len(csvs_brutos_ssp[n]) for n in nomes),
     "arquivos": tuple(nomes), "exemplo_registro": tuple(csvs_brutos_ssp[nomes[0]].loc[csvs_brutos_ssp[nomes[0]].celula_01.eq("JOINVILLE")].iloc[0]),
     "tipo_tabela": "Fragmentos mensais de violência doméstica" if "VIOLENCIA-DOMESTICA" in nomes[0] else "Boletim mensal: tabelas estaduais, municipais e gráficos extraídos"}
    for numero, (colunas, nomes) in enumerate(grupos_esquemas_ssp11.items(), 1)
])
mostrar_ssp11(inventario_arquivos_ssp)
mostrar_ssp11(esquemas_ssp)
mostrar_ssp11(amostras_csv_ssp)

# %% [markdown]
# ## 11.2 — Contexto municipal, página e tabela
# Todas as menções a Joinville são candidatas, inclusive rankings. Só classificar
# contagens após conferir o cabeçalho municipal e a seção de homicídios.
# Na tabela doméstica, examinar as linhas seguintes até o próximo município e
# as páginas de continuação horizontal. O campo municipal vazio permanece vazio:
# a associação contextual é registrada em metadados separados, sem preenchê-lo.

# %%
linhas_candidatas_ssp11 = []
linhas_mencoes_ssp11 = []
linhas_contextos_ssp11 = []
linhas_evidencias_ssp11 = []
verificacoes_tabelas_ssp11 = []
linhas_indicadores_ssp11 = []


def registrar_candidato_ssp11(nome_csv, indice, classificacao, municipio_contexto, indicador, evidencia):
    linha = csvs_brutos_ssp[nome_csv].loc[indice]
    linhas_candidatas_ssp11.append({
        "arquivo_csv": nome_csv, "indice_csv": indice,
        **{c: linha[c] for c in ("arquivo_pdf", "pagina_pdf", "tabela_pdf", "linha_tabela")},
        "municipio_publicado": linha["celula_01"], "municipio_contexto": municipio_contexto,
        "indicador": indicador, "classificacao": classificacao,
        "celulas_originais": celulas_ssp11(linha), "evidencia_contextual": evidencia,
    })


def registrar_celula_ssp11(nome_csv, indice, coluna, indicador, periodo_tipo, ano, periodo, unidade, contexto):
    linha = csvs_brutos_ssp[nome_csv].loc[indice]
    valor = linha[coluna]
    linhas_evidencias_ssp11.append({
        "arquivo_pdf": linha.arquivo_pdf, "arquivo_csv": nome_csv,
        "pagina_pdf": linha.pagina_pdf, "tabela_pdf": linha.tabela_pdf,
        "linha_tabela": linha.linha_tabela, "coluna_csv": coluna,
        "municipio": "JOINVILLE", "indicador_original": indicador,
        "periodo_tipo": periodo_tipo, "ano_referencia": ano,
        "periodo_referencia": periodo, "valor_publicado": valor,
        "unidade": unidade, "edicao": metadados_pdf_ssp11[linha.arquivo_pdf]["periodo_edicao"],
        "valor_inteiro_literal": bool(regex_ssp11.fullmatch(r"\d+", valor)),
        "evidencia_contextual": contexto,
    })


for nome_csv_ssp11, bruto_ssp11 in csvs_brutos_ssp.items():
    nome_pdf_ssp11 = CaminhoSSP11(nome_csv_ssp11).with_suffix(".pdf").name
    meta_ssp11 = metadados_pdf_ssp11[nome_pdf_ssp11]
    colunas_celulas_ssp11 = [c for c in bruto_ssp11 if c.startswith("celula_")]
    mencoes_ssp11 = bruto_ssp11[colunas_celulas_ssp11].apply(
        lambda coluna: coluna.str.contains("JOINVILLE", case=False, regex=False)
    ).any(axis=1)
    for indice_ssp11 in bruto_ssp11.index[mencoes_ssp11]:
        linha_ssp11 = bruto_ssp11.loc[indice_ssp11]
        linhas_mencoes_ssp11.append({"arquivo_csv": nome_csv_ssp11, "indice_csv": indice_ssp11,
            "pagina_pdf": linha_ssp11.pagina_pdf, "tabela_pdf": linha_ssp11.tabela_pdf,
            "linha_tabela": linha_ssp11.linha_tabela, "celulas_originais": celulas_ssp11(linha_ssp11)})
        if linha_ssp11.celula_01 != "JOINVILLE":
            pagina_mencao_ssp11 = bruto_ssp11.loc[bruto_ssp11.pagina_pdf.eq(linha_ssp11.pagina_pdf)]
            ranking_ssp11 = pagina_mencao_ssp11[colunas_celulas_ssp11].apply(
                lambda coluna: coluna.str.contains("VARIAÇÕES DE HOMICÍDIOS POR MUNICÍPIO", regex=False)
            ).any(axis=1).any()
            registrar_candidato_ssp11(nome_csv_ssp11, indice_ssp11,
                "Ranking de variação de homicídios; não somar às contagens" if ranking_ssp11 else "Menção textual; conferir contexto",
                "JOINVILLE" if ranking_ssp11 else "pendente", "Variação de homicídios" if ranking_ssp11 else "pendente",
                "Cabeçalho da página: variação absoluta entre dois anos; pares no rótulo municipal" if ranking_ssp11 else "Conferência pendente")
    if meta_ssp11["metade"] is None:
        # A seção municipal inicia antes da página de Joinville. Procurar o título
        # nos próprios CSVs, preservando a linha textual que estabelece o indicador.
        titulos_ssp11 = bruto_ssp11.loc[bruto_ssp11[colunas_celulas_ssp11].apply(
            lambda coluna: coluna.str.contains("VÍTIMAS DE HOMICÍDIO NOS MUNICÍPIOS", regex=False)
        ).any(axis=1)]
        exatas_ssp11 = bruto_ssp11.loc[bruto_ssp11.celula_01.eq("JOINVILLE")]
        for indice_ssp11, linha_ssp11 in exatas_ssp11.iterrows():
            pagina_ssp11 = int(linha_ssp11.pagina_pdf)
            tabela_ssp11 = bruto_ssp11.loc[(bruto_ssp11.pagina_pdf == linha_ssp11.pagina_pdf) & (bruto_ssp11.tabela_pdf == linha_ssp11.tabela_pdf)]
            texto_ssp11 = texto_pagina_ssp11(nome_pdf_ssp11, pagina_ssp11)
            cabecalhos_ssp11 = tabela_ssp11.loc[tabela_ssp11.celula_01.str.contains("MUNICÍPIO", regex=False)]
            if titulos_ssp11.empty or cabecalhos_ssp11.empty or "VÍTIMAS POR ANO" not in texto_ssp11:
                problema_ssp11(nome_pdf_ssp11, f"Página {pagina_ssp11}: contexto municipal de homicídio não confirmado.")
                continue
            cabecalho_ssp11 = cabecalhos_ssp11.iloc[0]
            colunas_ano_ssp11 = [c for c in colunas_celulas_ssp11 if regex_ssp11.fullmatch(r"20\d{2}", cabecalho_ssp11[c])]
            limite_ssp11 = regex_ssp11.search(r"NO PERÍODO\s*\(([^)]+)\)", texto_ssp11)
            if len(colunas_ano_ssp11) != 7 or limite_ssp11 is None:
                problema_ssp11(nome_pdf_ssp11, "Cabeçalho anual/acumulado não interpretável sem conferência adicional.")
                continue
            # Uma edição pode conter cabeçalhos históricos inconsistentes. Manter
            # os anos literalmente publicados, registrar a suspeita e não corrigi-los.
            anos_cabecalho_ssp11 = [cabecalho_ssp11[c] for c in colunas_ano_ssp11]
            ano_edicao_ssp11 = int(meta_ssp11["ano_edicao"])
            anos_esperados_ssp11 = [str(a) for a in range(ano_edicao_ssp11 - 3, ano_edicao_ssp11)] + [str(a) for a in range(ano_edicao_ssp11 - 3, ano_edicao_ssp11 + 1)]
            if anos_cabecalho_ssp11 != anos_esperados_ssp11:
                problema_ssp11(nome_pdf_ssp11, f"Cabeçalhos municipais {anos_cabecalho_ssp11} não acompanham a edição {ano_edicao_ssp11}; conferir fonte, sem relabelar valores.")
            contexto_ssp11 = f"Seção municipal p.{titulos_ssp11.iloc[0].pagina_pdf}; cabeçalho p.{pagina_ssp11}; {limite_ssp11.group(1)}"
            registrar_candidato_ssp11(nome_csv_ssp11, indice_ssp11, "Linha municipal de homicídios", "JOINVILLE", "Nº de vítimas de homicídio", contexto_ssp11)
            linhas_contextos_ssp11.append({"arquivo_pdf": nome_pdf_ssp11, "pagina_pdf": pagina_ssp11, "tabela_pdf": linha_ssp11.tabela_pdf,
                "cabecalhos_brutos": tuple(celulas_ssp11(r) for _, r in tabela_ssp11.head(4).iterrows()),
                "texto_pagina": texto_ssp11, "contexto": contexto_ssp11})
            for posicao_ssp11, coluna_ssp11 in enumerate(colunas_ano_ssp11):
                ano_valor_ssp11 = cabecalho_ssp11[coluna_ssp11]
                anual_ssp11 = posicao_ssp11 < 3
                registrar_celula_ssp11(nome_csv_ssp11, indice_ssp11, coluna_ssp11, "Nº de vítimas de homicídio",
                    "ano completo" if anual_ssp11 else "período acumulado", ano_valor_ssp11,
                    ano_valor_ssp11 if anual_ssp11 else limite_ssp11.group(1), "vítimas", contexto_ssp11)
    else:
        # Inspeção de fragmentos pareados; não fundir nem reescrever suas células.
        metade_ssp11 = meta_ssp11["metade"]
        for indice_inicio_ssp11 in bruto_ssp11.index[bruto_ssp11.celula_01.eq("JOINVILLE")]:
            inicio_ssp11 = bruto_ssp11.loc[indice_inicio_ssp11]
            pagina_esquerda_ssp11 = int(inicio_ssp11.pagina_pdf)
            if pagina_esquerda_ssp11 > metade_ssp11:
                problema_ssp11(nome_pdf_ssp11, "Município apareceu na metade sem rótulos; conferir layout.")
                continue
            esquerda_ssp11 = bruto_ssp11.loc[(bruto_ssp11.pagina_pdf == inicio_ssp11.pagina_pdf) & (bruto_ssp11.tabela_pdf == inicio_ssp11.tabela_pdf)]
            pagina_direita_ssp11 = pagina_esquerda_ssp11 + metade_ssp11
            direita_ssp11 = bruto_ssp11.loc[(bruto_ssp11.pagina_pdf == str(pagina_direita_ssp11)) & (bruto_ssp11.tabela_pdf == inicio_ssp11.tabela_pdf)]
            if esquerda_ssp11.linha_tabela.tolist() != direita_ssp11.linha_tabela.tolist():
                problema_ssp11(nome_pdf_ssp11, f"Fragmentos p.{pagina_esquerda_ssp11}/{pagina_direita_ssp11} desalinhados.")
                continue
            contexto_ssp11 = f"Bloco JOINVILLE p.{pagina_esquerda_ssp11}; continuação horizontal p.{pagina_direita_ssp11}; mesma linha/tabela"
            linhas_contextos_ssp11.append({"arquivo_pdf": nome_pdf_ssp11, "pagina_pdf": pagina_esquerda_ssp11,
                "tabela_pdf": inicio_ssp11.tabela_pdf, "cabecalhos_brutos": (meta_ssp11["meses_esquerda"], meta_ssp11["meses_direita"]),
                "texto_pagina": texto_pagina_ssp11(nome_pdf_ssp11, pagina_esquerda_ssp11),
                "texto_continuacao": texto_pagina_ssp11(nome_pdf_ssp11, pagina_direita_ssp11), "contexto": contexto_ssp11})
            bloco_ssp11 = esquerda_ssp11.loc[indice_inicio_ssp11:]
            for indice_ssp11, linha_ssp11 in bloco_ssp11.iterrows():
                if indice_ssp11 != indice_inicio_ssp11 and linha_ssp11.celula_01 != "":
                    break
                indicador_ssp11 = linha_ssp11.celula_02
                direita_linha_ssp11 = direita_ssp11.loc[direita_ssp11.linha_tabela.eq(linha_ssp11.linha_tabela)]
                if not indicador_ssp11 or len(direita_linha_ssp11) != 1:
                    problema_ssp11(nome_pdf_ssp11, "Categoria/linha da continuação incerta.")
                    continue
                indice_direita_ssp11 = direita_linha_ssp11.index[0]
                for indice_parte_ssp11, meses_parte_ssp11, deslocamento_ssp11 in (
                    (indice_ssp11, meta_ssp11["meses_esquerda"], 3),
                    (indice_direita_ssp11, meta_ssp11["meses_direita"], 1),
                ):
                    registrar_candidato_ssp11(nome_csv_ssp11, indice_parte_ssp11, "Fragmento municipal de violência doméstica", "JOINVILLE", indicador_ssp11, contexto_ssp11)
                    for posicao_ssp11, mes_ssp11 in enumerate(meses_parte_ssp11):
                        coluna_ssp11 = f"celula_{posicao_ssp11 + deslocamento_ssp11:02d}"
                        registrar_celula_ssp11(nome_csv_ssp11, indice_parte_ssp11, coluna_ssp11, indicador_ssp11,
                            "mensal", meta_ssp11["ano_edicao"], f"{meta_ssp11['ano_edicao']}-{meses_ssp11.index(mes_ssp11)+1:02d}",
                            "Fato Comunicado (unidade operacional a validar)", contexto_ssp11)
    # Conferir no PDF a tabela exata de cada linha municipal usada. Não reexportar
    # a extração; comparar cada célula não vazia e os índices de página/tabela/linha.
    candidatas_arquivo_ssp11 = [r for r in linhas_candidatas_ssp11 if r["arquivo_csv"] == nome_csv_ssp11 and r["classificacao"] in ("Linha municipal de homicídios", "Fragmento municipal de violência doméstica")]
    with pdf_ssp11.open(pasta_pdf_ssp11 / nome_pdf_ssp11) as documento_ssp11:
        tabelas_pagina_ssp11 = {}
        for candidata_ssp11 in candidatas_arquivo_ssp11:
            pagina_ssp11 = int(candidata_ssp11["pagina_pdf"])
            if pagina_ssp11 not in tabelas_pagina_ssp11:
                tabelas_pagina_ssp11[pagina_ssp11] = documento_ssp11.pages[pagina_ssp11 - 1].extract_tables()
            tabelas_ssp11 = tabelas_pagina_ssp11[pagina_ssp11]
            numero_tabela_ssp11 = int(candidata_ssp11["tabela_pdf"])
            numero_linha_ssp11 = int(candidata_ssp11["linha_tabela"])
            linha_pdf_ssp11 = tabelas_ssp11[numero_tabela_ssp11 - 1][numero_linha_ssp11 - 1]
            valores_csv_ssp11 = candidata_ssp11["celulas_originais"]
            # None e string vazia são indistinguíveis no CSV histórico. Comparar
            # a representação serializável somente aqui, sem preencher DataFrames.
            confere_ssp11 = all(valores_csv_ssp11[i] == ("" if v is None else v) for i, v in enumerate(linha_pdf_ssp11)) and all(v == "" for v in valores_csv_ssp11[len(linha_pdf_ssp11):])
            verificacoes_tabelas_ssp11.append({"arquivo_pdf": nome_pdf_ssp11, "pagina_pdf": pagina_ssp11,
                "tabela_pdf": numero_tabela_ssp11, "linha_tabela": numero_linha_ssp11, "celulas_conferem": confere_ssp11})
            if not confere_ssp11:
                problema_ssp11(nome_pdf_ssp11, f"Células divergentes na p.{pagina_ssp11}, tabela {numero_tabela_ssp11}, linha {numero_linha_ssp11}.")

# Menções literais e associação contextual são tabelas separadas. As linhas de
# continuação doméstica não contêm o nome da cidade, mas mantêm sua chave física.
mencoes_textuais_joinville_ssp = pd_ssp11.DataFrame(linhas_mencoes_ssp11)
registros_candidatos_joinville_ssp = pd_ssp11.DataFrame(linhas_candidatas_ssp11)
contextos_tabelas_ssp = pd_ssp11.DataFrame(linhas_contextos_ssp11)
evidencias_celulas_ssp = pd_ssp11.DataFrame(linhas_evidencias_ssp11)
conferencia_tabelas_pdf_ssp = pd_ssp11.DataFrame(verificacoes_tabelas_ssp11)
mostrar_ssp11(registros_candidatos_joinville_ssp)
mostrar_ssp11(contextos_tabelas_ssp)
mostrar_ssp11(evidencias_celulas_ssp)

# %% [markdown]
# ## 11.3 — Cobertura, indicadores e republicações
# Evidências por célula são strings originais anotadas para conferência. Não
# somar categorias, edições ou acumulados. Agrupar abaixo apenas para contar
# cobertura e identificar valores republicados; não gerar totais de crimes.
# Rótulos de lesão com '0 Dolosa' e '- Dolosa' permanecem distintos nesta etapa.
# Os intervalos 28/02 e 29/02 são rótulos literais da edição, compartilhados por
# vários anos históricos; ainda não representam datas civis validadas para cada
# ano. Os grupos abaixo são apenas repetições literais, sem resolver esse caso.

# %%
for chave_ssp11, grupo_ssp11 in evidencias_celulas_ssp.groupby(["indicador_original", "ano_referencia", "periodo_tipo"], sort=False):
    linhas_indicadores_ssp11.append({
        "indicador_original": chave_ssp11[0], "ano_referencia": chave_ssp11[1], "periodo_tipo": chave_ssp11[2],
        "periodos_publicados": tuple(sorted(set(grupo_ssp11.periodo_referencia))),
        "celulas_publicadas": len(grupo_ssp11), "unidade": tuple(set(grupo_ssp11.unidade)),
        "arquivos": tuple(sorted(set(grupo_ssp11.arquivo_pdf))),
        "valores_nao_interpretaveis": int((~grupo_ssp11.valor_inteiro_literal).sum()),
    })
indicadores_identificados_ssp = pd_ssp11.DataFrame(linhas_indicadores_ssp11)
republicacoes_ssp11 = []
for chave_ssp11, grupo_ssp11 in evidencias_celulas_ssp.groupby(["municipio", "indicador_original", "periodo_tipo", "ano_referencia", "periodo_referencia", "unidade"], sort=False):
    if len(grupo_ssp11) > 1:
        republicacoes_ssp11.append({
            "chave_observacao": chave_ssp11, "quantidade_edicoes": len(grupo_ssp11),
            "valores_publicados_distintos": tuple(sorted(set(grupo_ssp11.valor_publicado))),
            "divergencia": grupo_ssp11.valor_publicado.nunique() > 1,
            "referencias": tuple(zip(grupo_ssp11.arquivo_pdf, grupo_ssp11.pagina_pdf, grupo_ssp11.tabela_pdf, grupo_ssp11.linha_tabela, grupo_ssp11.valor_publicado)),
        })
republicacoes_divergencias_ssp = pd_ssp11.DataFrame(republicacoes_ssp11)
# Diagnóstico lógico por célula, sem converter a coluna valor_publicado. Os
# inteiros locais servem apenas à comparação; o token original segue preservado.
inconsistencias_ssp11 = []
for _, acumulado_ssp11 in evidencias_celulas_ssp.loc[evidencias_celulas_ssp.periodo_tipo.eq("período acumulado")].iterrows():
    anual_mesma_edicao_ssp11 = evidencias_celulas_ssp.loc[
        evidencias_celulas_ssp.arquivo_pdf.eq(acumulado_ssp11.arquivo_pdf)
        & evidencias_celulas_ssp.periodo_tipo.eq("ano completo")
        & evidencias_celulas_ssp.ano_referencia.eq(acumulado_ssp11.ano_referencia)
    ]
    if len(anual_mesma_edicao_ssp11) == 1:
        anual_ssp11 = anual_mesma_edicao_ssp11.iloc[0]
        if acumulado_ssp11.valor_inteiro_literal and anual_ssp11.valor_inteiro_literal and int(acumulado_ssp11.valor_publicado) > int(anual_ssp11.valor_publicado):
            inconsistencias_ssp11.append({"arquivo_pdf": acumulado_ssp11.arquivo_pdf,
                "pagina_pdf": acumulado_ssp11.pagina_pdf, "linha_tabela": acumulado_ssp11.linha_tabela,
                "ano_publicado": acumulado_ssp11.ano_referencia, "periodo_publicado": acumulado_ssp11.periodo_referencia,
                "acumulado_publicado": acumulado_ssp11.valor_publicado, "anual_publicado": anual_ssp11.valor_publicado,
                "situacao": "Acumulado superior ao anual na mesma edição; conferir publicação"})
inconsistencias_acumulados_ssp = pd_ssp11.DataFrame(inconsistencias_ssp11)
arquivos_identicos_ssp = pd_ssp11.DataFrame([
    {"sha256": digest, "arquivos": tuple(r.arquivo), "quantidade": len(r)}
    for digest, r in inventario_arquivos_ssp.groupby("sha256") if len(r) > 1
])
mostrar_ssp11(indicadores_identificados_ssp)
mostrar_ssp11(republicacoes_divergencias_ssp)
mostrar_ssp11(arquivos_identicos_ssp)
mostrar_ssp11(inconsistencias_acumulados_ssp)

# %% [markdown]
# ## 11.4 — Integridade e pendências
# Comparar os hashes antes/depois dos 90 insumos. O manifesto antigo pode cobrir
# só os PDFs; ausência de CSV no manifesto não comprova alteração nem validação
# histórica. Nenhuma divergência recebe correção automática.

# %%
manifesto_ssp11 = pd_ssp11.read_csv(
    raiz_ssp11 / "historico/documentacao_anterior/auditoria/inventario_sha256_pre_reorganizacao.csv",
    dtype=str, keep_default_na=False,
)
linhas_integridade_ssp11 = []
for caminho_ssp11, digest_ssp11 in hashes_antes_ssp11.items():
    relativo_ssp11 = caminho_ssp11.relative_to(raiz_ssp11).as_posix()
    historico_ssp11 = manifesto_ssp11.loc[manifesto_ssp11.caminho_destino_planejado.eq(relativo_ssp11)]
    hash_historico_ssp11 = historico_ssp11.iloc[0].sha256 if len(historico_ssp11) == 1 else None
    digest_depois_ssp11 = sha256_ssp11(caminho_ssp11)
    linhas_integridade_ssp11.append({
        "arquivo": caminho_ssp11.name, "tipo": caminho_ssp11.suffix,
        "sha256_antes": digest_ssp11, "sha256_depois": digest_depois_ssp11,
        "preservado": digest_ssp11 == digest_depois_ssp11,
        "historico": "não consta" if hash_historico_ssp11 is None else ("confere" if hash_historico_ssp11 == digest_ssp11 else "diverge"),
    })
integridade_arquivos_ssp = pd_ssp11.DataFrame(linhas_integridade_ssp11)
for nome_ssp11, retrato_ssp11 in retratos_sinesp_ssp11.items():
    pd_ssp11.testing.assert_frame_equal(globals()[nome_ssp11], retrato_ssp11, check_exact=True)
if not integridade_arquivos_ssp.preservado.all() or integridade_arquivos_ssp.historico.eq("diverge").any():
    raise ValueError("Etapa 11: divergência de integridade; interromper e apresentar o inventário.")
pendencias_inspecao_ssp = pd_ssp11.DataFrame(problemas_ssp11, columns=["arquivo", "pendencia"])
mostrar_ssp11(qualidade_csv_ssp)
mostrar_ssp11(conferencia_tabelas_pdf_ssp)
mostrar_ssp11(integridade_arquivos_ssp)
mostrar_ssp11(pendencias_inspecao_ssp)
print(f"Etapa 11: inspeção em memória concluída; {len(pendencias_inspecao_ssp)} pendência(s) de contexto e {len(inconsistencias_acumulados_ssp)} inconsistência(s) de acumulado. Consultar os DataFrames no Data Viewer; tratamentos aguardam validação.")

# %% [markdown]
# # ETAPA 12 — Auditoria das inconsistências SSP/SC
# Execute a partir desta célula, independentemente das etapas anteriores. Esta
# auditoria lê os originais e registra diagnósticos em memória. Não corrige anos,
# rótulos ou valores, não seleciona uma edição definitiva e não exporta arquivos.
# As diferenças numéricas abaixo são diagnósticos entre versões publicadas, não
# contagens mensais obtidas por subtração de acumulados.

# %%
from pathlib import Path as CaminhoSSP12
from calendar import isleap as bissexto_ssp12
import hashlib as hash_ssp12
import re as regex_ssp12

import pandas as pd_ssp12
import pdfplumber as pdf_ssp12
from pypdf import PdfReader as LeitorSSP12
from IPython.display import display as mostrar_ssp12

raiz_ssp12 = CaminhoSSP12(__file__).resolve().parents[2] if "__file__" in globals() else CaminhoSSP12.cwd()
pasta_ssp12 = raiz_ssp12 / "dados/brutos/ssp_sc"
pdfs_ssp12 = sorted((pasta_ssp12 / "pdfs").rglob("*.pdf"))
csvs_ssp12 = sorted((pasta_ssp12 / "csv_extraidos").rglob("*.csv"))
if len(pdfs_ssp12) != 45 or len(csvs_ssp12) != 45 or {p.stem for p in pdfs_ssp12} != {p.stem for p in csvs_ssp12}:
    raise ValueError("Etapa 12: acervo diferente dos 45 pares PDF/CSV; interromper a auditoria.")


def sha256_ssp12(caminho):
    with caminho.open("rb") as fluxo:
        return hash_ssp12.file_digest(fluxo, "sha256").hexdigest()


hashes_entrada_ssp12 = {p: sha256_ssp12(p) for p in pdfs_ssp12 + csvs_ssp12}
retratos_preservacao_ssp12 = {
    nome: globals()[nome].copy(deep=True)
    for nome in ("df_sinesp_joinville_concatenado", "df_sinesp_joinville_analitico", "evidencias_celulas_ssp", "republicacoes_divergencias_ssp")
    if nome in globals()
}
brutos_auditoria_ssp12 = {
    p.with_suffix(".pdf").name: pd_ssp12.read_csv(p, sep=";", encoding="utf-8-sig", dtype=str, keep_default_na=False, na_filter=False)
    for p in csvs_ssp12
}
textos_auditoria_ssp12 = {}
notas_auditoria_ssp12 = []
metadados_auditoria_ssp12 = {}
padrao_nota_ssp12 = r"retific|ratific|revis[aã]|corrig|corre[cç]|provis[oó]|consolid|sujeit|defini[cç]|metodolog|contabiliz|uma v[ií]tima|pessoas [uú]nicas|mais de um fato"
for caminho_ssp12 in pdfs_ssp12:
    leitor_ssp12 = LeitorSSP12(caminho_ssp12)
    textos_auditoria_ssp12[caminho_ssp12.name] = tuple(p.extract_text() or "" for p in leitor_ssp12.pages)
    for numero_pagina_ssp12, texto_ssp12 in enumerate(textos_auditoria_ssp12[caminho_ssp12.name], 1):
        linhas_texto_ssp12 = texto_ssp12.splitlines()
        for numero_linha_ssp12, linha_texto_ssp12 in enumerate(linhas_texto_ssp12):
            if regex_ssp12.search(padrao_nota_ssp12, linha_texto_ssp12, regex_ssp12.I):
                notas_auditoria_ssp12.append({
                    "arquivo_pdf": caminho_ssp12.name, "pagina_pdf": numero_pagina_ssp12,
                    "trecho_literal": "\n".join(linhas_texto_ssp12[max(0, numero_linha_ssp12 - 1):numero_linha_ssp12 + 3]),
                    "explicacao_de_revisao": bool(regex_ssp12.search(r"retific|ratific|revis[aã]|corrig|corre[cç]", linha_texto_ssp12, regex_ssp12.I)),
                })
    if "VIOLENCIA-DOMESTICA" not in caminho_ssp12.name:
        pagina_1_ssp12, pagina_2_ssp12 = textos_auditoria_ssp12[caminho_ssp12.name][:2]
        ano_edicao_ssp12 = regex_ssp12.search(r"\b20\d{2}\b", pagina_1_ssp12).group()
        periodo_edicao_ssp12 = regex_ssp12.search(r"PERÍODO\s*\((01/01\s*-\s*\d{2}/\d{2})\)", pagina_1_ssp12).group(1)
        atualizacao_1_ssp12 = regex_ssp12.search(r"ATUALIZAÇÃO DOS DADOS:\s*(\d{2}/\d{2}/\d{4})", pagina_1_ssp12)
        atualizacao_2_ssp12 = regex_ssp12.search(r"ATUALIZAÇÃO DOS DADOS:\s*(\d{2}/\d{2}/\d{4})", pagina_2_ssp12)
        metadados_auditoria_ssp12[caminho_ssp12.name] = {
            "arquivo_pdf": caminho_ssp12.name, "edicao": f"{ano_edicao_ssp12}-{periodo_edicao_ssp12[-2:]}",
            "ano_edicao": ano_edicao_ssp12, "periodo_edicao": periodo_edicao_ssp12,
            "atualizacao_pagina_1": atualizacao_1_ssp12.group(1) if atualizacao_1_ssp12 else None,
            "atualizacao_homicidios": atualizacao_2_ssp12.group(1) if atualizacao_2_ssp12 else None,
            "pagina_atualizacao_homicidios": 2,
        }
notas_documentais_ssp12 = pd_ssp12.DataFrame(notas_auditoria_ssp12)
edicoes_auditoria_ssp12 = pd_ssp12.DataFrame([
    {"codigo_edicao": f"B{n:02d}", **meta}
    for n, meta in enumerate(sorted(metadados_auditoria_ssp12.values(), key=lambda m: m["edicao"]), 1)
])
codigos_edicao_ssp12 = dict(zip(edicoes_auditoria_ssp12.arquivo_pdf, edicoes_auditoria_ssp12.codigo_edicao))

# %% [markdown]
# ## 12.1 — Leitura literal das células municipais e auditoria de janeiro/2025
# O ano é o escrito na própria coluna, mesmo quando incoerente com a edição.
# As datas de atualização da página 2 identificam a seção de homicídios; não
# substituir por datas da página 1, que se refere aos registros estaduais SISP.

# %%
celulas_homicidios_ssp12 = []
cabecalhos_homicidios_ssp12 = []
for nome_pdf_ssp12, meta_ssp12 in metadados_auditoria_ssp12.items():
    bruto_ssp12 = brutos_auditoria_ssp12[nome_pdf_ssp12]
    colunas_ssp12 = [c for c in bruto_ssp12 if c.startswith("celula_")]
    municipais_ssp12 = bruto_ssp12.loc[bruto_ssp12.celula_01.eq("JOINVILLE")]
    if len(municipais_ssp12) != 1:
        raise ValueError(f"Etapa 12: linha municipal inesperada em {nome_pdf_ssp12}.")
    linha_ssp12 = municipais_ssp12.iloc[0]
    tabela_ssp12 = bruto_ssp12.loc[bruto_ssp12.pagina_pdf.eq(linha_ssp12.pagina_pdf) & bruto_ssp12.tabela_pdf.eq(linha_ssp12.tabela_pdf)]
    cabecalho_ssp12 = tabela_ssp12.loc[tabela_ssp12.celula_01.str.contains("MUNICÍPIO", regex=False)].iloc[0]
    colunas_ano_ssp12 = [c for c in colunas_ssp12 if regex_ssp12.fullmatch(r"20\d{2}", cabecalho_ssp12[c])]
    texto_pagina_ssp12 = textos_auditoria_ssp12[nome_pdf_ssp12][int(linha_ssp12.pagina_pdf) - 1]
    periodo_ssp12 = regex_ssp12.search(r"NO PERÍODO\s*\(([^)]+)\)", texto_pagina_ssp12).group(1)
    if len(colunas_ano_ssp12) != 7:
        raise ValueError(f"Etapa 12: cabeçalho municipal inesperado em {nome_pdf_ssp12}.")
    cabecalhos_homicidios_ssp12.append({
        "arquivo_pdf": nome_pdf_ssp12, "pagina_pdf": linha_ssp12.pagina_pdf,
        "tabela_pdf": linha_ssp12.tabela_pdf, "linha_municipal": linha_ssp12.linha_tabela,
        "anos_literais": tuple(cabecalho_ssp12[c] for c in colunas_ano_ssp12),
        "periodo_literal": periodo_ssp12, "texto_pagina": texto_pagina_ssp12,
    })
    for posicao_ssp12, coluna_ssp12 in enumerate(colunas_ano_ssp12):
        valor_ssp12 = linha_ssp12[coluna_ssp12]
        if not regex_ssp12.fullmatch(r"\d+", valor_ssp12):
            raise ValueError(f"Etapa 12: valor não interpretável em {nome_pdf_ssp12}/{coluna_ssp12}; preservar e investigar.")
        anual_ssp12 = posicao_ssp12 < 3
        celulas_homicidios_ssp12.append({
            "arquivo_pdf": nome_pdf_ssp12, "arquivo_csv": CaminhoSSP12(nome_pdf_ssp12).with_suffix(".csv").name,
            "pagina_pdf": linha_ssp12.pagina_pdf, "tabela_pdf": linha_ssp12.tabela_pdf,
            "linha_tabela": linha_ssp12.linha_tabela, "coluna_csv": coluna_ssp12,
            "indicador": "Nº de vítimas de homicídio", "municipio": "JOINVILLE", "unidade": "vítimas",
            "ano_referencia": cabecalho_ssp12[coluna_ssp12],
            "periodo_tipo": "ano completo" if anual_ssp12 else "período acumulado",
            "periodo_publicado": cabecalho_ssp12[coluna_ssp12] if anual_ssp12 else periodo_ssp12,
            "valor_publicado": valor_ssp12, "codigo_edicao": codigos_edicao_ssp12[nome_pdf_ssp12],
            **{c: meta_ssp12[c] for c in ("edicao", "atualizacao_homicidios", "pagina_atualizacao_homicidios")},
        })
evidencias_homicidios_ssp12 = pd_ssp12.DataFrame(celulas_homicidios_ssp12)
cabecalhos_auditados_ssp12 = pd_ssp12.DataFrame(cabecalhos_homicidios_ssp12)
nome_janeiro_ssp12 = "2025_02_01-janeiro-1.pdf"
auditoria_janeiro_2025_ssp12 = evidencias_homicidios_ssp12.loc[evidencias_homicidios_ssp12.arquivo_pdf.eq(nome_janeiro_ssp12)].copy(deep=True)
with pdf_ssp12.open(pasta_ssp12 / "pdfs" / nome_janeiro_ssp12) as documento_ssp12:
    tabela_janeiro_ssp12 = documento_ssp12.pages[7].find_tables()[0]
    linhas_janeiro_ssp12 = tabela_janeiro_ssp12.extract()
    for indice_ssp12, observacao_ssp12 in auditoria_janeiro_2025_ssp12.iterrows():
        numero_coluna_ssp12 = int(observacao_ssp12.coluna_csv[-2:]) - 1
        valor_pdf_ssp12 = linhas_janeiro_ssp12[int(observacao_ssp12.linha_tabela) - 1][numero_coluna_ssp12]
        auditoria_janeiro_2025_ssp12.loc[indice_ssp12, "valor_pdf"] = valor_pdf_ssp12
        auditoria_janeiro_2025_ssp12.loc[indice_ssp12, "pdf_csv_conferem"] = valor_pdf_ssp12 == observacao_ssp12.valor_publicado
        auditoria_janeiro_2025_ssp12.loc[indice_ssp12, "celula_bbox_pdf"] = str(tabela_janeiro_ssp12.rows[int(observacao_ssp12.linha_tabela) - 1].cells[numero_coluna_ssp12])
casos_acumulado_superior_ssp12 = []
for _, observacao_ssp12 in auditoria_janeiro_2025_ssp12.loc[auditoria_janeiro_2025_ssp12.periodo_tipo.eq("período acumulado")].iterrows():
    anuais_ssp12 = auditoria_janeiro_2025_ssp12.loc[auditoria_janeiro_2025_ssp12.periodo_tipo.eq("ano completo") & auditoria_janeiro_2025_ssp12.ano_referencia.eq(observacao_ssp12.ano_referencia)]
    if len(anuais_ssp12) == 1 and int(observacao_ssp12.valor_publicado) > int(anuais_ssp12.iloc[0].valor_publicado):
        casos_acumulado_superior_ssp12.append({"ano_publicado": observacao_ssp12.ano_referencia,
            "anual_publicado": anuais_ssp12.iloc[0].valor_publicado, "acumulado_publicado": observacao_ssp12.valor_publicado,
            "excesso_diagnostico": int(observacao_ssp12.valor_publicado) - int(anuais_ssp12.iloc[0].valor_publicado),
            "arquivo_pdf": nome_janeiro_ssp12, "pagina_pdf": "8", "tabela_pdf": "1", "linha_tabela": "36"})
acumulados_superiores_ssp12 = pd_ssp12.DataFrame(casos_acumulado_superior_ssp12)
mostrar_ssp12(auditoria_janeiro_2025_ssp12)
mostrar_ssp12(acumulados_superiores_ssp12)

# %% [markdown]
# ## 12.2 — Todas as versões dos 58 grupos, sem selecionar uma publicação
# Diferença de valor não prova revisão documentada. A classificação 'possível
# problema de cabeçalho' é usada quando somente a edição suspeita introduz a
# diferença; não se apaga essa edição. Se as outras edições também divergem, a
# classificação permanece 'inconsistência ainda não explicada', com uma ressalva.
# 'Ratificado' no nome de um arquivo não explica uma alteração municipal específica.

# %%
grupos_auditoria_ssp12 = []
versoes_auditoria_ssp12 = []
campos_chave_ssp12 = ["indicador", "municipio", "unidade", "periodo_tipo", "ano_referencia", "periodo_publicado"]
for chave_ssp12, grupo_ssp12 in evidencias_homicidios_ssp12.groupby(campos_chave_ssp12, sort=True):
    if len(grupo_ssp12) < 2:
        continue
    identificador_ssp12 = f"G{len(grupos_auditoria_ssp12) + 1:02d}"
    divergente_ssp12 = grupo_ssp12.valor_publicado.nunique() > 1
    contem_janeiro_ssp12 = grupo_ssp12.arquivo_pdf.eq(nome_janeiro_ssp12).any()
    outras_edicoes_ssp12 = grupo_ssp12.loc[grupo_ssp12.arquivo_pdf.ne(nome_janeiro_ssp12)]
    if not divergente_ssp12:
        classe_ssp12, explicacao_ssp12 = "Sem divergência de valor", "Republicações literais concordantes; não representam observações independentes."
    elif contem_janeiro_ssp12 and outras_edicoes_ssp12.valor_publicado.nunique() == 1:
        classe_ssp12 = "Possível problema de cabeçalho"
        explicacao_ssp12 = "Diferença restrita à edição jan/2025, cujo cabeçalho municipal não acompanha a edição; causa definitiva pendente."
    else:
        classe_ssp12 = "Inconsistência ainda não explicada"
        explicacao_ssp12 = "Não localizada nota documental que explique a mudança desta observação; não presumir revisão, erro ou mudança de cobertura."
    valores_locais_ssp12 = [int(v) for v in grupo_ssp12.valor_publicado]
    grupos_auditoria_ssp12.append({
        "grupo": identificador_ssp12, **dict(zip(campos_chave_ssp12, chave_ssp12)),
        "quantidade_edicoes": len(grupo_ssp12), "divergente": divergente_ssp12,
        "valores_publicados_distintos": tuple(sorted(set(grupo_ssp12.valor_publicado), key=int)),
        "amplitude_entre_versoes": max(valores_locais_ssp12) - min(valores_locais_ssp12),
        "classificacao": classe_ssp12, "explicacao_documental": explicacao_ssp12,
        "inclui_edicao_com_cabecalho_suspeito": bool(contem_janeiro_ssp12),
        "situacao": "pendente de validação" if divergente_ssp12 else "concordante, sem escolha definitiva",
    })
    anterior_ssp12 = None
    # Ordenar somente as evidências auxiliares por edição; nenhuma base original
    # recebe ordenação, conversão ou mudança de índices.
    for _, versao_ssp12 in grupo_ssp12.sort_values("edicao").iterrows():
        valor_local_ssp12 = int(versao_ssp12.valor_publicado)
        versoes_auditoria_ssp12.append({"grupo": identificador_ssp12, **versao_ssp12.to_dict(),
            "diferenca_versao_anterior": None if anterior_ssp12 is None else valor_local_ssp12 - anterior_ssp12,
            "classificacao_grupo": classe_ssp12, "explicacao_documental": explicacao_ssp12,
            "ressalva_versao": "Cabeçalho/período de jan/2025 pendente" if versao_ssp12.arquivo_pdf == nome_janeiro_ssp12 else "Causa da mudança não documentada"})
        anterior_ssp12 = valor_local_ssp12
matriz_republicacoes_ssp12 = pd_ssp12.DataFrame(grupos_auditoria_ssp12)
versoes_republicacoes_ssp12 = pd_ssp12.DataFrame(versoes_auditoria_ssp12)
matriz_divergencias_ssp12 = matriz_republicacoes_ssp12.loc[matriz_republicacoes_ssp12.divergente].copy(deep=True)
versoes_divergentes_ssp12 = versoes_republicacoes_ssp12.loc[versoes_republicacoes_ssp12.grupo.isin(matriz_divergencias_ssp12.grupo)].copy(deep=True)
mostrar_ssp12(edicoes_auditoria_ssp12)
mostrar_ssp12(matriz_republicacoes_ssp12)
mostrar_ssp12(matriz_divergencias_ssp12)
mostrar_ssp12(versoes_divergentes_ssp12)

# %% [markdown]
# ## 12.3 — Violência doméstica: cabeçalhos, alinhamento e 216 células
# Não preencher o município nas células vazias: associar os fragmentos por
# contexto e manter os valores literais em uma tabela exclusiva de auditoria.
# As grafias '0 Dolosa' e '- Dolosa' não são padronizadas. As nove categorias
# são conferidas em cada relatório; os meses vêm dos cabeçalhos originais.

# %%
meses_ssp12 = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
categorias_comuns_ssp12 = {"Ameaça", "Calúnia", "Difamação", "Estupro", "Injúria", "Injúria real", "Vias de fato"}
fragmentos_domesticos_ssp12 = []
celulas_domesticas_ssp12 = []
estrutura_domestica_ssp12 = []
for nome_pdf_ssp12, bruto_ssp12 in brutos_auditoria_ssp12.items():
    if "VIOLENCIA-DOMESTICA" not in nome_pdf_ssp12:
        continue
    textos_ssp12 = textos_auditoria_ssp12[nome_pdf_ssp12]
    metade_ssp12 = len(textos_ssp12) // 2
    ano_ssp12 = regex_ssp12.search(r"\b20\d{2}\b", textos_ssp12[0]).group()
    meses_esquerda_ssp12 = regex_ssp12.findall(r"\b(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)\b", textos_ssp12[0][:120])
    meses_direita_ssp12 = regex_ssp12.findall(r"\b(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)\b", textos_ssp12[metade_ssp12][:80])
    inicio_ssp12 = bruto_ssp12.loc[bruto_ssp12.celula_01.eq("JOINVILLE")].iloc[0]
    pagina_esquerda_ssp12 = int(inicio_ssp12.pagina_pdf)
    pagina_direita_ssp12 = pagina_esquerda_ssp12 + metade_ssp12
    esquerda_ssp12 = bruto_ssp12.loc[bruto_ssp12.pagina_pdf.eq(str(pagina_esquerda_ssp12)) & bruto_ssp12.tabela_pdf.eq(inicio_ssp12.tabela_pdf)]
    direita_ssp12 = bruto_ssp12.loc[bruto_ssp12.pagina_pdf.eq(str(pagina_direita_ssp12)) & bruto_ssp12.tabela_pdf.eq(inicio_ssp12.tabela_pdf)]
    bloco_ssp12 = []
    iniciou_ssp12 = False
    for _, linha_ssp12 in esquerda_ssp12.iterrows():
        if linha_ssp12.celula_01 == "JOINVILLE":
            iniciou_ssp12 = True
        elif iniciou_ssp12 and linha_ssp12.celula_01 != "":
            break
        if iniciou_ssp12:
            bloco_ssp12.append(linha_ssp12)
    categorias_ssp12 = {r.celula_02 for r in bloco_ssp12}
    variantes_ssp12 = (
        {"Lesão corporal grave ou gravíssima 0 Dolosa", "Lesão corporal leve 0 Dolosa"}
        if "1o-SEMESTRE-2023" in nome_pdf_ssp12 else
        {"Lesão corporal grave ou gravíssima - Dolosa", "Lesão corporal leve - Dolosa"}
    )
    if len(bloco_ssp12) != 9 or categorias_ssp12 != categorias_comuns_ssp12 | variantes_ssp12:
        raise ValueError(f"Etapa 12: categorias domésticas divergentes em {nome_pdf_ssp12}.")
    if esquerda_ssp12.linha_tabela.tolist() != direita_ssp12.linha_tabela.tolist():
        raise ValueError(f"Etapa 12: páginas domésticas desalinhadas em {nome_pdf_ssp12}.")
    with pdf_ssp12.open(pasta_ssp12 / "pdfs" / nome_pdf_ssp12) as documento_ssp12:
        tabelas_ssp12 = {pagina: documento_ssp12.pages[pagina - 1].find_tables()[0] for pagina in (pagina_esquerda_ssp12, pagina_direita_ssp12)}
        extraidas_ssp12 = {pagina: tabela.extract() for pagina, tabela in tabelas_ssp12.items()}
        for linha_esquerda_ssp12 in bloco_ssp12:
            linha_direita_ssp12 = direita_ssp12.loc[direita_ssp12.linha_tabela.eq(linha_esquerda_ssp12.linha_tabela)].iloc[0]
            for pagina_ssp12, linha_ssp12, meses_parte_ssp12, deslocamento_ssp12 in (
                (pagina_esquerda_ssp12, linha_esquerda_ssp12, meses_esquerda_ssp12, 3),
                (pagina_direita_ssp12, linha_direita_ssp12, meses_direita_ssp12, 1),
            ):
                indice_linha_ssp12 = int(linha_ssp12.linha_tabela) - 1
                bbox_linha_ssp12 = tabelas_ssp12[pagina_ssp12].rows[indice_linha_ssp12].bbox
                fragmentos_domesticos_ssp12.append({"arquivo_pdf": nome_pdf_ssp12, "pagina_pdf": pagina_ssp12,
                    "tabela_pdf": linha_ssp12.tabela_pdf, "linha_tabela": linha_ssp12.linha_tabela,
                    "indicador_original": linha_esquerda_ssp12.celula_02, "bbox_linha": bbox_linha_ssp12,
                    "meses_cabecalho": tuple(meses_parte_ssp12)})
                for posicao_ssp12, mes_ssp12 in enumerate(meses_parte_ssp12):
                    numero_coluna_ssp12 = posicao_ssp12 + deslocamento_ssp12
                    coluna_ssp12 = f"celula_{numero_coluna_ssp12:02d}"
                    valor_csv_ssp12 = linha_ssp12[coluna_ssp12]
                    valor_pdf_ssp12 = extraidas_ssp12[pagina_ssp12][indice_linha_ssp12][numero_coluna_ssp12 - 1]
                    celulas_domesticas_ssp12.append({"arquivo_pdf": nome_pdf_ssp12, "pagina_pdf": pagina_ssp12,
                        "tabela_pdf": linha_ssp12.tabela_pdf, "linha_tabela": linha_ssp12.linha_tabela, "coluna_csv": coluna_ssp12,
                        "municipio_contextual": "JOINVILLE", "municipio_celula_original": linha_ssp12.celula_01,
                        "indicador_original": linha_esquerda_ssp12.celula_02, "ano": ano_ssp12,
                        "mes_cabecalho": mes_ssp12, "valor_publicado": valor_csv_ssp12, "valor_pdf": valor_pdf_ssp12,
                        "pdf_csv_conferem": valor_csv_ssp12 == valor_pdf_ssp12,
                        "unidade_cabecalho": "Fato Comunicado", "definicao_operacional": "não explicitada nesses PDFs"})
    estrutura_domestica_ssp12.append({"arquivo_pdf": nome_pdf_ssp12, "paginas": len(textos_ssp12),
        "pagina_esquerda": pagina_esquerda_ssp12, "pagina_direita": pagina_direita_ssp12,
        "cabecalho_original_esquerda": "\n".join(l for l in textos_ssp12[0].splitlines() if regex_ssp12.search(r"Munic[ií]pio|Fato Comunicado|^20\d{2}$|^(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)(?:\s+(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez))*$", l)),
        "cabecalho_original_direita": "\n".join(textos_ssp12[metade_ssp12].splitlines()[:2]),
        "meses": tuple(meses_esquerda_ssp12 + meses_direita_ssp12), "categorias_originais": tuple(sorted(categorias_ssp12)),
        "definicao_oficial_no_documento": "Não localizada explicação operacional da contagem ou de pessoas únicas"})
auditoria_estrutura_domestica_ssp12 = pd_ssp12.DataFrame(estrutura_domestica_ssp12)
auditoria_fragmentos_domesticos_ssp12 = pd_ssp12.DataFrame(fragmentos_domesticos_ssp12)
auditoria_celulas_domesticas_ssp12 = pd_ssp12.DataFrame(celulas_domesticas_ssp12)
mostrar_ssp12(auditoria_estrutura_domestica_ssp12)
mostrar_ssp12(auditoria_fragmentos_domesticos_ssp12)
mostrar_ssp12(auditoria_celulas_domesticas_ssp12)

# %% [markdown]
# ## 12.4 — Calendário, edição e documentação da unidade
# 28/02 e 29/02 permanecem como rótulos originais. Identificar o que seria uma
# data civil inválida não autoriza truncar datas, equiparar períodos ou mudar o
# ano. A edição e a data de atualização são metadados distintos da competência.
# Não calcular meses por diferença entre acumulados de edições diferentes.

# %%
linhas_fevereiro_ssp12 = []
for _, registro_ssp12 in evidencias_homicidios_ssp12.iterrows():
    if registro_ssp12.periodo_tipo == "período acumulado" and registro_ssp12.periodo_publicado.endswith(("28/02", "29/02")):
        linhas_fevereiro_ssp12.append({**registro_ssp12.to_dict(),
            "ano_referencia_bissexto": bissexto_ssp12(int(registro_ssp12.ano_referencia)),
            "29_fevereiro_invalido_se_data_literal": registro_ssp12.periodo_publicado.endswith("29/02") and not bissexto_ssp12(int(registro_ssp12.ano_referencia)),
            "situacao": "Rótulo compartilhado da edição; correspondência ao calendário histórico pendente"})
auditoria_fevereiro_ssp12 = pd_ssp12.DataFrame(linhas_fevereiro_ssp12)
documentacao_unidade_ssp12 = pd_ssp12.DataFrame([
    {"fonte": "Três PDFs domésticos locais", "evidencia": "Cabeçalho Fato Comunicado; meses e categorias; sem definição operacional de contagem", "limite": "Não comprova vítimas, pessoas únicas ou relação um para um com BO"},
    {"fonte": "documentacao/fontes/SSP_SC.md e documentacao/auditoria/ssp_csv_extraction_etapa_1_1.md", "evidencia": "Documentam extração e grafias; não são dicionário oficial da medida", "limite": "Não substituem definição da SSP/SC"},
    {"fonte": "https://ssp.sc.gov.br/segurancaemnumeros/", "evidencia": "Portal oficial identifica os três relatórios como indicadores de violência doméstica", "limite": "A página consultada não apresenta definição operacional de Fato Comunicado"},
    {"fonte": "https://www.sas.sc.gov.br/index.php/conselhos/cedca/diagnostico-da-realidade-social-da-crianca-e-do-adolescente-do-estado-de-santa-catarina/3447-caderno-3-direito-a-liberdade-ao-respeito-e-a-dignidade/file", "evidencia": "Diagnóstico CEDCA/SC, 2018, página PDF 55, nota: uma vítima pode ter mais de um fato comunicado", "limite": "Contexto de crianças/adolescentes e dados SSP/2016; não define o algoritmo dos relatórios domésticos 2023–2024"},
])
mostrar_ssp12(auditoria_fevereiro_ssp12)
mostrar_ssp12(documentacao_unidade_ssp12)

# %% [markdown]
# ## 12.5 — Preservação e limites da auditoria
# Conferir as dimensões diagnósticas esperadas e os 90 hashes contra a auditoria
# histórica e contra a leitura inicial. As bases anteriores existentes no kernel
# também são comparadas, sem exigir sua reconstrução.

# %%
if len(matriz_republicacoes_ssp12) != 58 or len(matriz_divergencias_ssp12) != 41:
    raise ValueError("Etapa 12: os grupos não correspondem à inspeção anterior; apresentar a divergência.")
if len(evidencias_homicidios_ssp12) != 294 or len(acumulados_superiores_ssp12) != 3:
    raise ValueError("Etapa 12: as evidências de homicídios divergiram da Etapa 11.")
if len(auditoria_celulas_domesticas_ssp12) != 216 or not auditoria_celulas_domesticas_ssp12.pdf_csv_conferem.all() or not auditoria_janeiro_2025_ssp12.pdf_csv_conferem.all():
    raise ValueError("Etapa 12: células não correspondem aos PDFs; interromper sem corrigir.")
if auditoria_celulas_domesticas_ssp12.duplicated(["arquivo_pdf", "pagina_pdf", "tabela_pdf", "linha_tabela", "coluna_csv"]).any():
    raise ValueError("Etapa 12: célula doméstica diagnosticada mais de uma vez; investigar sem deduplicar.")
manifesto_ssp12 = pd_ssp12.read_csv(raiz_ssp12 / "historico/documentacao_anterior/auditoria/inventario_sha256_pre_reorganizacao.csv", dtype=str, keep_default_na=False)
integridade_auditoria_ssp12 = pd_ssp12.DataFrame([
    {"arquivo": p.name, "sha256_antes": digest, "sha256_depois": sha256_ssp12(p),
     "sha256_historico": manifesto_ssp12.loc[manifesto_ssp12.caminho_destino_planejado.eq(p.relative_to(raiz_ssp12).as_posix()), "sha256"].item()}
    for p, digest in hashes_entrada_ssp12.items()
])
if not (integridade_auditoria_ssp12.sha256_antes.eq(integridade_auditoria_ssp12.sha256_depois) & integridade_auditoria_ssp12.sha256_antes.eq(integridade_auditoria_ssp12.sha256_historico)).all():
    raise ValueError("Etapa 12: divergência de integridade; interromper a auditoria.")
for nome_ssp12, retrato_ssp12 in retratos_preservacao_ssp12.items():
    pd_ssp12.testing.assert_frame_equal(globals()[nome_ssp12], retrato_ssp12, check_exact=True)
mostrar_ssp12(integridade_auditoria_ssp12)
print("Etapa 12: 58 grupos auditados, 41 divergentes; valores originais preservados; causas pendentes não foram presumidas. Nenhuma série definitiva selecionada.")

# %% [markdown]
# # ETAPA 13.1 — Inspeção e exibição dos dados brutos de violência doméstica (SSP/SC)
# Execute a partir desta célula, sem executar as etapas anteriores nem o Sinesp.
# Esta etapa é somente de leitura e visualização: os três relatórios de violência
# doméstica são lidos literalmente (todas as células como texto), os blocos de
# Joinville das páginas esquerda e direita são exibidos separadamente e as
# estruturas são conferidas contra os PDFs. Nada é limpo, convertido, renomeado,
# concatenado em base analítica ou exportado.
#
# Decisão metodológica registrada: o usuário aprovou, para a FUTURA base analítica
# (Etapa 13.2), a equivalência entre "Lesão corporal leve 0 Dolosa" e
# "Lesão corporal leve - Dolosa" e entre "Lesão corporal grave ou gravíssima 0 Dolosa"
# e "Lesão corporal grave ou gravíssima - Dolosa". É uma decisão do projeto, não
# uma confirmação oficial da SSP/SC. Aqui os quatro textos originais permanecem
# exatamente como publicados e nenhum mapeamento é executado.
# "Fato Comunicado" é a unidade publicada e não é equiparada a vítimas, pessoas
# únicas ou ocorrências únicas.

# %% [markdown]
# ## 13.1.1 — Arquivos, referências dos blocos e verificações
# As referências abaixo (páginas, tabela, linhas e meses) são as identificadas na
# proposta da Etapa 13. Elas são apenas hipóteses de conferência: os blocos são
# localizados nos CSVs pelo texto JOINVILLE e depois comparados com as referências.

# %%
from pathlib import Path as CaminhoSSP131
import hashlib as hash_ssp131
import re as regex_ssp131

import pandas as pd_ssp131
import pdfplumber as pdf_ssp131
from IPython.display import display as mostrar_ssp131

raiz_ssp131 = CaminhoSSP131(__file__).resolve().parents[2] if "__file__" in globals() else CaminhoSSP131.cwd()
pasta_pdf_ssp131 = raiz_ssp131 / "dados/brutos/ssp_sc/pdfs"
pasta_csv_ssp131 = raiz_ssp131 / "dados/brutos/ssp_sc/csv_extraidos"

referencias_ssp131 = {
    "1º semestre de 2023": {
        "arquivo": "2025_06_VIOLENCIA-DOMESTICA-1o-SEMESTRE-2023", "ano": "2023",
        "pagina_esquerda": 16, "pagina_direita": 49, "tabela": "1", "linha_inicial": 42, "linha_final": 50,
        "meses_esquerda": ("jan",), "meses_direita": ("fev", "mar", "abr", "mai", "jun"),
    },
    "2º semestre de 2023": {
        "arquivo": "2025_06_VIOLENCIA-DOMESTICA-2o-SEMESTRE-2023", "ano": "2023",
        "pagina_esquerda": 16, "pagina_direita": 48, "tabela": "1", "linha_inicial": 12, "linha_final": 20,
        "meses_esquerda": ("jul", "ago"), "meses_direita": ("set", "out", "nov", "dez"),
    },
    "2024": {
        "arquivo": "2025_06_VIOLENCIA-DOMESTICA-2024-SC", "ano": "2024",
        "pagina_esquerda": 18, "pagina_direita": 55, "tabela": "1", "linha_inicial": 32, "linha_final": 40,
        "meses_esquerda": ("jan", "fev", "mar"), "meses_direita": ("abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez"),
    },
}
# Textos originais esperados em cada relatório (apenas para conferência, sem substituição).
categorias_comuns_ssp131 = ("Ameaça", "Calúnia", "Difamação", "Estupro", "Injúria", "Injúria real", "Vias de fato")
lesoes_ssp131 = {
    "0": ("Lesão corporal grave ou gravíssima 0 Dolosa", "Lesão corporal leve 0 Dolosa"),
    "-": ("Lesão corporal grave ou gravíssima - Dolosa", "Lesão corporal leve - Dolosa"),
}
todas_lesoes_ssp131 = lesoes_ssp131["0"] + lesoes_ssp131["-"]
grafia_esperada_ssp131 = {"1º semestre de 2023": "0", "2º semestre de 2023": "-", "2024": "-"}

verificacoes_ssp131 = {}


def registrar_ssp131(relatorio, verificacao, esperado, encontrado):
    # Registro por chave: executar novamente uma célula substitui, não duplica.
    verificacoes_ssp131[(relatorio, verificacao)] = {
        "relatorio": relatorio, "verificacao": verificacao,
        "esperado": esperado, "encontrado": encontrado, "ok": esperado == encontrado}


for relatorio_ssp131, ref_ssp131 in referencias_ssp131.items():
    for pasta_ssp131, sufixo_ssp131 in ((pasta_pdf_ssp131, ".pdf"), (pasta_csv_ssp131, ".csv")):
        caminho_ssp131 = pasta_ssp131 / (ref_ssp131["arquivo"] + sufixo_ssp131)
        if not caminho_ssp131.is_file():
            raise FileNotFoundError(f"Etapa 13.1: arquivo não encontrado: {caminho_ssp131}")
        registrar_ssp131(relatorio_ssp131, f"existe {sufixo_ssp131}", True, True)

mostrar_ssp131(pd_ssp131.DataFrame([{"relatorio": r, **ref} for r, ref in referencias_ssp131.items()]))

# %% [markdown]
# ## 13.1.2 — Leitura literal dos três CSVs e inventário
# Os CSVs são lidos com `dtype=str` e sem conversão de ausentes: o que está na
# célula é o que aparece aqui. As colunas `celula_*` são nomes do extrator, não
# cabeçalhos do relatório. Os hashes SHA-256 são comparados com o inventário
# histórico para confirmar que os originais não foram alterados.

# %%
def sha256_ssp131(caminho):
    with caminho.open("rb") as fluxo:
        return hash_ssp131.file_digest(fluxo, "sha256").hexdigest()


manifesto_ssp131 = pd_ssp131.read_csv(
    raiz_ssp131 / "historico/documentacao_anterior/auditoria/inventario_sha256_pre_reorganizacao.csv",
    dtype=str, keep_default_na=False)
brutos_dom_ssp131 = {}
paginas_pdf_ssp131 = {}
hashes_antes_ssp131 = {}
linhas_inventario_ssp131 = []
for relatorio_ssp131, ref_ssp131 in referencias_ssp131.items():
    caminho_pdf_ssp131 = pasta_pdf_ssp131 / (ref_ssp131["arquivo"] + ".pdf")
    caminho_csv_ssp131 = pasta_csv_ssp131 / (ref_ssp131["arquivo"] + ".csv")
    brutos_dom_ssp131[relatorio_ssp131] = pd_ssp131.read_csv(
        caminho_csv_ssp131, sep=";", encoding="utf-8-sig", dtype=str, keep_default_na=False, na_filter=False)
    with pdf_ssp131.open(caminho_pdf_ssp131) as documento_ssp131:
        paginas_pdf_ssp131[relatorio_ssp131] = len(documento_ssp131.pages)
    for caminho_ssp131 in (caminho_pdf_ssp131, caminho_csv_ssp131):
        hashes_antes_ssp131[caminho_ssp131] = sha256_ssp131(caminho_ssp131)
        historico_ssp131 = manifesto_ssp131.loc[
            manifesto_ssp131.caminho_destino_planejado.eq(caminho_ssp131.relative_to(raiz_ssp131).as_posix()), "sha256"].tolist()
        registrar_ssp131(relatorio_ssp131, f"sha256 igual ao inventário histórico ({caminho_ssp131.suffix})",
                         [hashes_antes_ssp131[caminho_ssp131]], historico_ssp131)
    bruto_ssp131 = brutos_dom_ssp131[relatorio_ssp131]
    linhas_inventario_ssp131.append({
        "relatorio": relatorio_ssp131, "arquivo_pdf": caminho_pdf_ssp131.name, "paginas_pdf": paginas_pdf_ssp131[relatorio_ssp131],
        "arquivo_csv": caminho_csv_ssp131.name, "linhas_csv": len(bruto_ssp131), "colunas_csv": tuple(bruto_ssp131.columns),
        "caminho_pdf": caminho_pdf_ssp131.relative_to(raiz_ssp131).as_posix(),
        "caminho_csv": caminho_csv_ssp131.relative_to(raiz_ssp131).as_posix(),
        "sha256_pdf": hashes_antes_ssp131[caminho_pdf_ssp131], "sha256_csv": hashes_antes_ssp131[caminho_csv_ssp131]})
inventario_dom_ssp131 = pd_ssp131.DataFrame(linhas_inventario_ssp131)
mostrar_ssp131(inventario_dom_ssp131)
for relatorio_ssp131, bruto_ssp131 in brutos_dom_ssp131.items():
    print(f"\n{relatorio_ssp131} — primeiras linhas do CSV bruto (valores literais):")
    mostrar_ssp131(bruto_ssp131.head(3))

# %% [markdown]
# ## 13.1.3 — Localização dos blocos de Joinville
# O bloco começa na linha cuja `celula_01` é JOINVILLE e vai até a linha anterior
# à próxima `celula_01` preenchida, na mesma página e tabela. As demais linhas do
# bloco têm `celula_01` vazia no CSV: o município é dado pelo contexto do bloco e a
# célula original continua vazia. A página da direita é a da esquerda mais metade
# do total de páginas (como na Etapa 12). Os blocos abaixo são seleções
# temporárias para exibição; os CSVs não são modificados.

# %%
blocos_esquerda_ssp131 = {}
blocos_direita_ssp131 = {}
linhas_localizacao_ssp131 = []
for relatorio_ssp131, ref_ssp131 in referencias_ssp131.items():
    bruto_ssp131 = brutos_dom_ssp131[relatorio_ssp131]
    marcadores_ssp131 = bruto_ssp131.index[bruto_ssp131.celula_01.eq("JOINVILLE")]
    if len(marcadores_ssp131) == 0:
        raise ValueError(f"Etapa 13.1: JOINVILLE não encontrado em {ref_ssp131['arquivo']}.")
    inicio_ssp131 = bruto_ssp131.loc[marcadores_ssp131[0]]
    pagina_esq_ssp131 = int(inicio_ssp131.pagina_pdf)
    pagina_dir_ssp131 = pagina_esq_ssp131 + paginas_pdf_ssp131[relatorio_ssp131] // 2
    pagina_tabela_ssp131 = bruto_ssp131.loc[bruto_ssp131.pagina_pdf.eq(str(pagina_esq_ssp131)) & bruto_ssp131.tabela_pdf.eq(inicio_ssp131.tabela_pdf)]
    posicao_inicio_ssp131 = pagina_tabela_ssp131.index.get_loc(marcadores_ssp131[0])
    posicao_fim_ssp131 = len(pagina_tabela_ssp131)
    for posicao_ssp131 in range(posicao_inicio_ssp131 + 1, len(pagina_tabela_ssp131)):
        if pagina_tabela_ssp131.iloc[posicao_ssp131].celula_01 != "":
            posicao_fim_ssp131 = posicao_ssp131
            break
    bloco_esq_ssp131 = pagina_tabela_ssp131.iloc[posicao_inicio_ssp131:posicao_fim_ssp131].copy()
    bloco_dir_ssp131 = bruto_ssp131.loc[
        bruto_ssp131.pagina_pdf.eq(str(pagina_dir_ssp131)) & bruto_ssp131.tabela_pdf.eq(inicio_ssp131.tabela_pdf)
        & bruto_ssp131.linha_tabela.isin(bloco_esq_ssp131.linha_tabela)].copy()
    blocos_esquerda_ssp131[relatorio_ssp131] = bloco_esq_ssp131
    blocos_direita_ssp131[relatorio_ssp131] = bloco_dir_ssp131
    linhas_bloco_ssp131 = bloco_esq_ssp131.linha_tabela.astype(int)
    encontrado_ssp131 = (pagina_esq_ssp131, pagina_dir_ssp131, inicio_ssp131.tabela_pdf, int(linhas_bloco_ssp131.min()), int(linhas_bloco_ssp131.max()))
    referencia_ssp131 = (ref_ssp131["pagina_esquerda"], ref_ssp131["pagina_direita"], ref_ssp131["tabela"], ref_ssp131["linha_inicial"], ref_ssp131["linha_final"])
    registrar_ssp131(relatorio_ssp131, "ocorrências de JOINVILLE em celula_01", 1, len(marcadores_ssp131))
    registrar_ssp131(relatorio_ssp131, "páginas, tabela e linhas do bloco = referência", referencia_ssp131, encontrado_ssp131)
    registrar_ssp131(relatorio_ssp131, "linhas do bloco na esquerda", 9, len(bloco_esq_ssp131))
    registrar_ssp131(relatorio_ssp131, "linhas correspondentes na direita", 9, len(bloco_dir_ssp131))
    registrar_ssp131(relatorio_ssp131, "linha_tabela esquerda = direita (mesma ordem)",
                     bloco_esq_ssp131.linha_tabela.tolist(), bloco_dir_ssp131.linha_tabela.tolist())
    registrar_ssp131(relatorio_ssp131, "linhas do bloco consecutivas",
                     list(range(int(linhas_bloco_ssp131.min()), int(linhas_bloco_ssp131.min()) + 9)), linhas_bloco_ssp131.tolist())
    linhas_localizacao_ssp131.append({
        "relatorio": relatorio_ssp131, "paginas_pdf_total": paginas_pdf_ssp131[relatorio_ssp131],
        "pagina_esq_referencia": ref_ssp131["pagina_esquerda"], "pagina_esq_encontrada": pagina_esq_ssp131,
        "pagina_dir_referencia": ref_ssp131["pagina_direita"], "pagina_dir_derivada": pagina_dir_ssp131,
        "tabela_referencia": ref_ssp131["tabela"], "tabela_encontrada": inicio_ssp131.tabela_pdf,
        "linhas_referencia": f"{ref_ssp131['linha_inicial']}–{ref_ssp131['linha_final']}",
        "linhas_encontradas": f"{linhas_bloco_ssp131.min()}–{linhas_bloco_ssp131.max()}",
        "linhas_esquerda": len(bloco_esq_ssp131), "linhas_direita": len(bloco_dir_ssp131),
        "indices_csv_esquerda": f"{bloco_esq_ssp131.index.min()}–{bloco_esq_ssp131.index.max()}",
        "confere_referencia": referencia_ssp131 == encontrado_ssp131})
localizacao_blocos_ssp131 = pd_ssp131.DataFrame(linhas_localizacao_ssp131)
mostrar_ssp131(localizacao_blocos_ssp131)

# %% [markdown]
# ## 13.1.4 — Blocos brutos: página esquerda e página direita
# Para cada relatório, primeiro o bloco da página esquerda (município, categoria e
# os primeiros meses) e depois o da página direita (somente os meses restantes,
# sem município nem categoria). O índice é o do CSV bruto e todas as colunas
# `celula_*` originais são mantidas, inclusive as vazias do extrator.

# %%
for relatorio_ssp131 in referencias_ssp131:
    print(f"\n=== {relatorio_ssp131} — PÁGINA ESQUERDA ({referencias_ssp131[relatorio_ssp131]['arquivo']}.csv) ===")
    mostrar_ssp131(blocos_esquerda_ssp131[relatorio_ssp131])
    print(f"=== {relatorio_ssp131} — PÁGINA DIREITA (continuação horizontal) ===")
    mostrar_ssp131(blocos_direita_ssp131[relatorio_ssp131])

# %% [markdown]
# ## 13.1.5 — Quadro de conferência entre as duas metades
# Este quadro existe só para a verificação visual: cada linha mostra a categoria
# original e os valores da esquerda ao lado dos valores da direita que têm o mesmo
# `linha_tabela`. Não é a base analítica, não converte valores e não identifica
# meses; as colunas continuam sendo `celula_*`. Colunas totalmente vazias nas
# linhas do bloco são omitidas somente da exibição.

# %%
quadros_conferencia_ssp131 = {}
colunas_valor_ssp131 = {}
for relatorio_ssp131 in referencias_ssp131:
    esquerda_ssp131 = blocos_esquerda_ssp131[relatorio_ssp131]
    direita_ssp131 = blocos_direita_ssp131[relatorio_ssp131]
    colunas_celula_ssp131 = [c for c in esquerda_ssp131.columns if c.startswith("celula_")]
    valores_esq_ssp131 = [c for c in colunas_celula_ssp131[2:] if esquerda_ssp131[c].ne("").any()]
    valores_dir_ssp131 = [c for c in colunas_celula_ssp131 if direita_ssp131[c].ne("").any()]
    colunas_valor_ssp131[relatorio_ssp131] = (tuple(valores_esq_ssp131), tuple(valores_dir_ssp131))
    parte_esq_ssp131 = esquerda_ssp131[["linha_tabela", "celula_01", "celula_02"] + valores_esq_ssp131].rename(
        columns={c: f"esq_{c}" for c in ["celula_01", "celula_02"] + valores_esq_ssp131})
    parte_dir_ssp131 = direita_ssp131[["linha_tabela"] + valores_dir_ssp131].rename(columns={c: f"dir_{c}" for c in valores_dir_ssp131})
    quadros_conferencia_ssp131[relatorio_ssp131] = parte_esq_ssp131.merge(
        parte_dir_ssp131, on="linha_tabela", how="left", validate="one_to_one", indicator="presente_na_direita")
    registrar_ssp131(relatorio_ssp131, "todas as linhas da esquerda têm par na direita", 9,
                     int(quadros_conferencia_ssp131[relatorio_ssp131].presente_na_direita.eq("both").sum()))
    print(f"\n=== {relatorio_ssp131}: colunas de valores — esquerda {valores_esq_ssp131}; direita {valores_dir_ssp131} ===")
    mostrar_ssp131(quadros_conferencia_ssp131[relatorio_ssp131])

# %% [markdown]
# ## 13.1.6 — Os nove rótulos originais de cada relatório
# Os rótulos são exibidos exatamente como publicados. A coluna
# `caractere_antes_de_Dolosa` mostra o caractere literal e o código Unicode que
# separa "grave ou gravíssima"/"leve" de "Dolosa": `0` (U+0030) ou `-` (U+002D).
# Registro metodológico: a equivalência dessas quatro grafias foi aprovada para a
# futura transformação (Etapa 13.2), como decisão do projeto, e NÃO é confirmação
# oficial da SSP/SC. Aqui nada é substituído, mapeado ou padronizado.

# %%
linhas_categorias_ssp131 = []
for relatorio_ssp131 in referencias_ssp131:
    for _, linha_ssp131 in blocos_esquerda_ssp131[relatorio_ssp131].iterrows():
        rotulo_ssp131 = linha_ssp131.celula_02
        correspondencia_ssp131 = regex_ssp131.search(r"(\S)\s+Dolosa\s*$", rotulo_ssp131)
        linhas_categorias_ssp131.append({
            "relatorio": relatorio_ssp131, "pagina_pdf": linha_ssp131.pagina_pdf, "tabela_pdf": linha_ssp131.tabela_pdf,
            "linha_tabela": linha_ssp131.linha_tabela, "rotulo_original": rotulo_ssp131,
            "grafia_lesao_corporal": rotulo_ssp131 in todas_lesoes_ssp131,
            "caractere_antes_de_Dolosa": correspondencia_ssp131.group(1) if correspondencia_ssp131 else "",
            "unicode": f"U+{ord(correspondencia_ssp131.group(1)):04X}" if correspondencia_ssp131 else ""})
categorias_originais_ssp131 = pd_ssp131.DataFrame(linhas_categorias_ssp131)
for relatorio_ssp131 in referencias_ssp131:
    rotulos_ssp131 = categorias_originais_ssp131.loc[categorias_originais_ssp131.relatorio.eq(relatorio_ssp131), "rotulo_original"].tolist()
    esperados_ssp131 = set(categorias_comuns_ssp131) | set(lesoes_ssp131[grafia_esperada_ssp131[relatorio_ssp131]])
    registrar_ssp131(relatorio_ssp131, "nove rótulos distintos", 9, len(set(rotulos_ssp131)))
    registrar_ssp131(relatorio_ssp131, "conjunto de rótulos esperado", sorted(esperados_ssp131), sorted(rotulos_ssp131))
    print(f"\n=== {relatorio_ssp131} — rótulos originais ===")
    mostrar_ssp131(categorias_originais_ssp131.loc[categorias_originais_ssp131.relatorio.eq(relatorio_ssp131)].reset_index(drop=True))
print("\nQuatro grafias de lesão corporal por relatório (literais):")
mostrar_ssp131(pd_ssp131.DataFrame([
    {"grafia_original": grafia_ssp131, **{r: bool(categorias_originais_ssp131.loc[categorias_originais_ssp131.relatorio.eq(r)].rotulo_original.eq(grafia_ssp131).any()) for r in referencias_ssp131}}
    for grafia_ssp131 in todas_lesoes_ssp131]))

# %% [markdown]
# ## 13.1.7 — Cabeçalhos dos PDFs e cobertura mensal
# O texto das páginas é lido dos PDFs originais. Os cabeçalhos com ano e meses
# são mostrados literalmente (`jan`, `fev`, …): não há conversão para números,
# competência ou data. A tabela mostra também as primeiras linhas de texto das
# páginas dos blocos, para evidenciar onde o cabeçalho de meses está impresso.

# %%
padrao_mes_ssp131 = r"(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)"
padrao_linha_meses_ssp131 = rf"{padrao_mes_ssp131}(?:\s+{padrao_mes_ssp131})*"
textos_pdf_ssp131 = {}
tabelas_pdf_ssp131 = {}
linhas_cabecalho_ssp131 = []
for relatorio_ssp131, ref_ssp131 in referencias_ssp131.items():
    meio_ssp131 = paginas_pdf_ssp131[relatorio_ssp131] // 2
    pagina_esq_ssp131 = int(blocos_esquerda_ssp131[relatorio_ssp131].pagina_pdf.iloc[0])
    pagina_dir_ssp131 = int(blocos_direita_ssp131[relatorio_ssp131].pagina_pdf.iloc[0])
    papeis_ssp131 = (
        (1, "cabeçalho da metade esquerda (1ª página)"), (meio_ssp131 + 1, "cabeçalho da metade direita (1ª página)"),
        (pagina_esq_ssp131, "página do bloco — esquerda"), (pagina_dir_ssp131, "página do bloco — direita"))
    with pdf_ssp131.open(pasta_pdf_ssp131 / (ref_ssp131["arquivo"] + ".pdf")) as documento_ssp131:
        for pagina_ssp131, papel_ssp131 in papeis_ssp131:
            texto_ssp131 = documento_ssp131.pages[pagina_ssp131 - 1].extract_text() or ""
            textos_pdf_ssp131[(relatorio_ssp131, pagina_ssp131)] = texto_ssp131
            primeiras_ssp131 = texto_ssp131.splitlines()[:4]
            meses_pagina_ssp131 = tuple(m for l in primeiras_ssp131 if regex_ssp131.fullmatch(padrao_linha_meses_ssp131, l.strip()) for m in l.split())
            ano_pagina_ssp131 = next((l.strip() for l in primeiras_ssp131 if regex_ssp131.fullmatch(r"20\d{2}", l.strip())), "")
            linhas_cabecalho_ssp131.append({
                "relatorio": relatorio_ssp131, "pagina_pdf": pagina_ssp131, "papel": papel_ssp131,
                "ano_impresso": ano_pagina_ssp131, "meses_impressos": meses_pagina_ssp131,
                "quatro_primeiras_linhas_literais": " | ".join(primeiras_ssp131)})
        for pagina_ssp131, tabela_pdf_ssp131 in ((pagina_esq_ssp131, blocos_esquerda_ssp131[relatorio_ssp131].tabela_pdf.iloc[0]),
                                                  (pagina_dir_ssp131, blocos_direita_ssp131[relatorio_ssp131].tabela_pdf.iloc[0])):
            tabelas_pdf_ssp131[(relatorio_ssp131, pagina_ssp131)] = documento_ssp131.pages[pagina_ssp131 - 1].find_tables()[int(tabela_pdf_ssp131) - 1].extract()
cabecalhos_pdf_ssp131 = pd_ssp131.DataFrame(linhas_cabecalho_ssp131)
mostrar_ssp131(cabecalhos_pdf_ssp131)

linhas_cobertura_ssp131 = []
for relatorio_ssp131, ref_ssp131 in referencias_ssp131.items():
    parte_ssp131 = cabecalhos_pdf_ssp131.loc[cabecalhos_pdf_ssp131.relatorio.eq(relatorio_ssp131)].set_index("papel")
    meses_esq_ssp131 = parte_ssp131.loc["cabeçalho da metade esquerda (1ª página)", "meses_impressos"]
    meses_dir_ssp131 = parte_ssp131.loc["cabeçalho da metade direita (1ª página)", "meses_impressos"]
    registrar_ssp131(relatorio_ssp131, "meses do cabeçalho esquerdo", ref_ssp131["meses_esquerda"], meses_esq_ssp131)
    registrar_ssp131(relatorio_ssp131, "meses do cabeçalho direito", ref_ssp131["meses_direita"], meses_dir_ssp131)
    registrar_ssp131(relatorio_ssp131, "ano impresso nos cabeçalhos", (ref_ssp131["ano"], ref_ssp131["ano"]),
                     (parte_ssp131.loc["cabeçalho da metade esquerda (1ª página)", "ano_impresso"], parte_ssp131.loc["cabeçalho da metade direita (1ª página)", "ano_impresso"]))
    registrar_ssp131(relatorio_ssp131, "colunas de valor na esquerda = nº de meses", len(ref_ssp131["meses_esquerda"]), len(colunas_valor_ssp131[relatorio_ssp131][0]))
    registrar_ssp131(relatorio_ssp131, "colunas de valor na direita = nº de meses", len(ref_ssp131["meses_direita"]), len(colunas_valor_ssp131[relatorio_ssp131][1]))
    # Observação, não correção: o cabeçalho de meses não é reimpresso nas páginas do bloco.
    registrar_ssp131(relatorio_ssp131, "meses impressos nas páginas dos blocos (esperado: nenhum)", ((), ()),
                     (parte_ssp131.loc["página do bloco — esquerda", "meses_impressos"], parte_ssp131.loc["página do bloco — direita", "meses_impressos"]))
    linhas_cobertura_ssp131.append({
        "relatorio": relatorio_ssp131, "meses_esquerda_cabecalho": meses_esq_ssp131,
        "colunas_csv_esquerda": colunas_valor_ssp131[relatorio_ssp131][0],
        "meses_direita_cabecalho": meses_dir_ssp131, "colunas_csv_direita": colunas_valor_ssp131[relatorio_ssp131][1],
        "cobertura_publicada": meses_esq_ssp131 + meses_dir_ssp131})
cobertura_mensal_ssp131 = pd_ssp131.DataFrame(linhas_cobertura_ssp131)
print("\nCobertura mensal pelos cabeçalhos (rótulos literais; sem conversão):")
mostrar_ssp131(cobertura_mensal_ssp131)
meses_ano_ssp131 = ("jan", "fev", "mar", "abr", "mai", "jun", "jul", "ago", "set", "out", "nov", "dez")
registrar_ssp131("2023 (dois relatórios)", "meses jan–jun + jul–dez sem repetição", meses_ano_ssp131,
                 cobertura_mensal_ssp131.loc[cobertura_mensal_ssp131.relatorio.eq("1º semestre de 2023"), "cobertura_publicada"].iloc[0]
                 + cobertura_mensal_ssp131.loc[cobertura_mensal_ssp131.relatorio.eq("2º semestre de 2023"), "cobertura_publicada"].iloc[0])
registrar_ssp131("2024", "meses jan–dez", meses_ano_ssp131,
                 cobertura_mensal_ssp131.loc[cobertura_mensal_ssp131.relatorio.eq("2024"), "cobertura_publicada"].iloc[0])

# %% [markdown]
# ## 13.1.8 — Conferência dos textos e valores do CSV com a tabela do PDF
# Compara, linha a linha e célula a célula, o texto do CSV com o texto extraído
# da própria tabela da página do PDF. Células do CSV além da largura da tabela do
# PDF devem estar vazias (preenchimento do extrator). Só leitura: nada é corrigido.

# %%
linhas_pdf_csv_ssp131 = []
for relatorio_ssp131 in referencias_ssp131:
    for lado_ssp131, blocos_ssp131 in (("esquerda", blocos_esquerda_ssp131), ("direita", blocos_direita_ssp131)):
        colunas_celula_ssp131 = [c for c in blocos_ssp131[relatorio_ssp131].columns if c.startswith("celula_")]
        for _, linha_ssp131 in blocos_ssp131[relatorio_ssp131].iterrows():
            tabela_pdf_ssp131 = tabelas_pdf_ssp131[(relatorio_ssp131, int(linha_ssp131.pagina_pdf))]
            linha_pdf_ssp131 = [c if c is not None else "" for c in tabela_pdf_ssp131[int(linha_ssp131.linha_tabela) - 1]]
            celulas_csv_ssp131 = [linha_ssp131[c] for c in colunas_celula_ssp131]
            largura_ssp131 = len(linha_pdf_ssp131)
            linhas_pdf_csv_ssp131.append({
                "relatorio": relatorio_ssp131, "lado": lado_ssp131, "pagina_pdf": linha_ssp131.pagina_pdf,
                "linha_tabela": linha_ssp131.linha_tabela, "celulas_pdf": tuple(linha_pdf_ssp131),
                "celulas_csv_na_largura_do_pdf": tuple(celulas_csv_ssp131[:largura_ssp131]),
                "csv_alem_do_pdf_vazio": all(c == "" for c in celulas_csv_ssp131[largura_ssp131:]),
                "pdf_csv_conferem": celulas_csv_ssp131[:largura_ssp131] == linha_pdf_ssp131 and all(c == "" for c in celulas_csv_ssp131[largura_ssp131:])})
conferencia_pdf_csv_ssp131 = pd_ssp131.DataFrame(linhas_pdf_csv_ssp131)
for relatorio_ssp131 in referencias_ssp131:
    parte_ssp131 = conferencia_pdf_csv_ssp131.loc[conferencia_pdf_csv_ssp131.relatorio.eq(relatorio_ssp131)]
    registrar_ssp131(relatorio_ssp131, "18 linhas do bloco conferem com o PDF", 18, int(parte_ssp131.pdf_csv_conferem.sum()))
mostrar_ssp131(conferencia_pdf_csv_ssp131)

# %% [markdown]
# ## 13.1.9 — Preservação e síntese das verificações
# Os hashes dos PDFs e CSVs são recalculados e comparados com os de antes da
# leitura. As verificações que não coincidem com o esperado aparecem como
# divergências, com a evidência, e nenhum dado é alterado.

# %%
for caminho_ssp131, hash_antes_ssp131 in hashes_antes_ssp131.items():
    registrar_ssp131(caminho_ssp131.name, "sha256 antes = depois", hash_antes_ssp131, sha256_ssp131(caminho_ssp131))
resumo_verificacoes_ssp131 = pd_ssp131.DataFrame(list(verificacoes_ssp131.values()))
print("Verificações da Etapa 13.1 (somente leitura e estrutura):")
mostrar_ssp131(resumo_verificacoes_ssp131)
divergencias_ssp131 = resumo_verificacoes_ssp131.loc[~resumo_verificacoes_ssp131.ok]
if divergencias_ssp131.empty:
    print(f"Etapa 13.1: {len(resumo_verificacoes_ssp131)} verificações conferidas, nenhuma divergência.")
else:
    print(f"Etapa 13.1: {len(divergencias_ssp131)} divergência(s) em {len(resumo_verificacoes_ssp131)} verificações; evidência abaixo, sem alteração dos dados.")
    mostrar_ssp131(divergencias_ssp131)
print("Nenhuma transformação analítica: sem conversão de valores/meses, mapeamento de categorias, base 216 × 5 ou exportação.")

# %% [markdown]
# # ETAPA 13.2 — Construção da base analítica de violência doméstica de Joinville (SSP/SC)
# Execute depois das células 13.1.1 a 13.1.7 da Etapa 13.1 (não é preciso executar
# as Etapas 1–12 nem o Sinesp). A base é construída somente em memória, a partir
# dos blocos brutos já inspecionados; os PDFs e CSVs não são alterados e nada é
# exportado. Se qualquer conferência falhar, a execução é interrompida com a
# evidência e as variáveis finais não são criadas.
#
# Convenções metodológicas aprovadas (decisões do projeto, não confirmações oficiais da SSP/SC):
# 1. Interpretação estrutural dos meses: o cabeçalho de meses da primeira página de
#    cada metade da tabela é aplicado às páginas seguintes da mesma tabela contínua.
#    Ano e mês vêm dos cabeçalhos dos PDFs, nunca do nome do arquivo. A categoria é
#    associada aos valores da página direita pela mesma linha da tabela.
# 2. Equivalência das grafias "0 Dolosa" e "- Dolosa" das duas lesões corporais, por
#    mapeamento explícito. Nenhuma categoria é agregada a outra.
# "Fato Comunicado" é a unidade publicada; não é equiparada a vítimas, pessoas únicas
# ou ocorrências únicas.

# %% [markdown]
# ## 13.2.1 — Parâmetros explícitos: meses, categorias e mapeamento

# %%
necessarios_ssp132 = (
    "pd_ssp131", "mostrar_ssp131", "sha256_ssp131", "referencias_ssp131", "brutos_dom_ssp131",
    "blocos_esquerda_ssp131", "blocos_direita_ssp131", "colunas_valor_ssp131", "cabecalhos_pdf_ssp131",
    "paginas_pdf_ssp131", "hashes_antes_ssp131", "pasta_pdf_ssp131", "pasta_csv_ssp131")
ausentes_ssp132 = [nome for nome in necessarios_ssp132 if nome not in globals()]
if ausentes_ssp132:
    raise NameError(f"Etapa 13.2: execute antes as células 13.1.1 a 13.1.7 da Etapa 13.1. Ausentes: {ausentes_ssp132}")

numero_mes_ssp132 = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
                     "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}
categorias_analiticas_ssp132 = (
    "Ameaça", "Calúnia", "Difamação", "Estupro", "Injúria", "Injúria real",
    "Lesão corporal grave ou gravíssima dolosa", "Lesão corporal leve dolosa", "Vias de fato")
# Mapeamento explícito rótulo original -> categoria analítica (sem substituição genérica de caracteres).
mapeamento_categorias_ssp132 = {
    "Ameaça": "Ameaça",
    "Calúnia": "Calúnia",
    "Difamação": "Difamação",
    "Estupro": "Estupro",
    "Injúria": "Injúria",
    "Injúria real": "Injúria real",
    "Lesão corporal grave ou gravíssima 0 Dolosa": "Lesão corporal grave ou gravíssima dolosa",
    "Lesão corporal grave ou gravíssima - Dolosa": "Lesão corporal grave ou gravíssima dolosa",
    "Lesão corporal leve 0 Dolosa": "Lesão corporal leve dolosa",
    "Lesão corporal leve - Dolosa": "Lesão corporal leve dolosa",
    "Vias de fato": "Vias de fato",
}
papel_esquerda_ssp132 = "cabeçalho da metade esquerda (1ª página)"
papel_direita_ssp132 = "cabeçalho da metade direita (1ª página)"
unidade_ssp132 = "Fato Comunicado"
regra_ssp132 = (
    "Interpretação estrutural validada no projeto: o cabeçalho de meses da 1ª página de cada metade da tabela "
    "vale para as páginas seguintes da mesma tabela contínua; ano e mês vêm do cabeçalho do PDF; a categoria é "
    "associada pela mesma linha da tabela; rótulos de lesão corporal mapeados por decisão metodológica, "
    "sem confirmação oficial da SSP/SC")

# %% [markdown]
# ## 13.2.2 — Construção em variáveis temporárias
# Para cada relatório: confere ano, meses, quantidade e ordem das colunas contra os
# cabeçalhos esperados; percorre as nove linhas do bloco; lê cada célula mensal
# (esquerda e direita) e só a converte para inteiro depois de verificar que o texto
# é um inteiro não negativo. Vazio ou não numérico em coluna mensal válida interrompe
# a construção, sem preenchimento. As colunas vazias do extrator, fora dos meses do
# cabeçalho, não são lidas e não geram observações.

# %%
registros_ssp132 = []
problemas_ssp132 = []
for relatorio_ssp132, ref_ssp132 in referencias_ssp131.items():
    cabecalho_ssp132 = cabecalhos_pdf_ssp131.loc[cabecalhos_pdf_ssp131.relatorio.eq(relatorio_ssp132)].set_index("papel")
    ano_texto_ssp132 = cabecalho_ssp132.loc[papel_esquerda_ssp132, "ano_impresso"]
    meses_esq_ssp132 = tuple(cabecalho_ssp132.loc[papel_esquerda_ssp132, "meses_impressos"])
    meses_dir_ssp132 = tuple(cabecalho_ssp132.loc[papel_direita_ssp132, "meses_impressos"])
    colunas_esq_ssp132, colunas_dir_ssp132 = colunas_valor_ssp131[relatorio_ssp132]
    if ano_texto_ssp132 != ref_ssp132["ano"] or cabecalho_ssp132.loc[papel_direita_ssp132, "ano_impresso"] != ano_texto_ssp132:
        raise ValueError(f"Etapa 13.2: ano dos cabeçalhos diferente do esperado em {relatorio_ssp132}.")
    if meses_esq_ssp132 != ref_ssp132["meses_esquerda"] or meses_dir_ssp132 != ref_ssp132["meses_direita"]:
        raise ValueError(f"Etapa 13.2: meses dos cabeçalhos diferentes do esperado em {relatorio_ssp132}.")
    if tuple(colunas_esq_ssp132) != tuple(f"celula_{i:02d}" for i in range(3, 3 + len(meses_esq_ssp132))) \
            or tuple(colunas_dir_ssp132) != tuple(f"celula_{i:02d}" for i in range(1, 1 + len(meses_dir_ssp132))):
        raise ValueError(f"Etapa 13.2: quantidade ou ordem das colunas não corresponde aos cabeçalhos em {relatorio_ssp132}.")
    numeros_ssp132 = [numero_mes_ssp132[m] for m in meses_esq_ssp132 + meses_dir_ssp132]
    if numeros_ssp132 != list(range(numeros_ssp132[0], numeros_ssp132[0] + len(numeros_ssp132))):
        raise ValueError(f"Etapa 13.2: meses fora da sequência do calendário em {relatorio_ssp132}.")
    pagina_cab_esq_ssp132 = str(cabecalho_ssp132.loc[papel_esquerda_ssp132, "pagina_pdf"])
    pagina_cab_dir_ssp132 = str(cabecalho_ssp132.loc[papel_direita_ssp132, "pagina_pdf"])
    esquerda_ssp132 = blocos_esquerda_ssp131[relatorio_ssp132]
    direita_ssp132 = blocos_direita_ssp131[relatorio_ssp132].set_index("linha_tabela")
    caminho_pdf_ssp132 = pasta_pdf_ssp131 / (ref_ssp132["arquivo"] + ".pdf")
    caminho_csv_ssp132 = pasta_csv_ssp131 / (ref_ssp132["arquivo"] + ".csv")
    for _, linha_esq_ssp132 in esquerda_ssp132.iterrows():
        linha_dir_ssp132 = direita_ssp132.loc[linha_esq_ssp132.linha_tabela]
        rotulo_ssp132 = linha_esq_ssp132.celula_02
        if rotulo_ssp132 not in mapeamento_categorias_ssp132:
            raise ValueError(f"Etapa 13.2: rótulo sem mapeamento aprovado em {relatorio_ssp132}: {rotulo_ssp132!r}")
        for lado_ssp132, linha_ssp132, colunas_ssp132, meses_ssp132, pagina_cab_ssp132, pagina_par_ssp132 in (
            ("esquerda", linha_esq_ssp132, colunas_esq_ssp132, meses_esq_ssp132, pagina_cab_esq_ssp132, str(linha_dir_ssp132.pagina_pdf)),
            ("direita", linha_dir_ssp132, colunas_dir_ssp132, meses_dir_ssp132, pagina_cab_dir_ssp132, str(linha_esq_ssp132.pagina_pdf)),
        ):
            for coluna_ssp132, mes_rotulo_ssp132 in zip(colunas_ssp132, meses_ssp132):
                texto_ssp132 = linha_ssp132[coluna_ssp132]
                if not (isinstance(texto_ssp132, str) and texto_ssp132.isascii() and texto_ssp132.isdigit()):
                    problemas_ssp132.append({
                        "relatorio": relatorio_ssp132, "arquivo_pdf": ref_ssp132["arquivo"] + ".pdf", "pagina_pdf": linha_ssp132.pagina_pdf if lado_ssp132 == "esquerda" else linha_ssp132["pagina_pdf"],
                        "linha_tabela": linha_esq_ssp132.linha_tabela, "coluna_csv": coluna_ssp132,
                        "rotulo_original": rotulo_ssp132, "mes_cabecalho": mes_rotulo_ssp132, "texto_encontrado": texto_ssp132})
                    continue
                registros_ssp132.append({
                    "municipio": "JOINVILLE", "ano": int(ano_texto_ssp132), "mes": numero_mes_ssp132[mes_rotulo_ssp132],
                    "categoria": mapeamento_categorias_ssp132[rotulo_ssp132], "fatos_comunicados": int(texto_ssp132),
                    "relatorio_origem": relatorio_ssp132,
                    "arquivo_pdf": caminho_pdf_ssp132.name, "sha256_pdf": hashes_antes_ssp131[caminho_pdf_ssp132],
                    "arquivo_csv": caminho_csv_ssp132.name, "sha256_csv": hashes_antes_ssp131[caminho_csv_ssp132],
                    "pagina_pdf": str(linha_ssp132["pagina_pdf"]), "tabela_pdf": str(linha_ssp132["tabela_pdf"]),
                    "linha_tabela": str(linha_esq_ssp132.linha_tabela), "coluna_csv": coluna_ssp132,
                    "lado_da_tabela": lado_ssp132, "pagina_horizontal_pareada": pagina_par_ssp132,
                    "rotulo_original_categoria": rotulo_ssp132,
                    "municipio_celula_original": linha_esq_ssp132.celula_01 if lado_ssp132 == "esquerda" else None,
                    "cabecalho_ano_original": ano_texto_ssp132, "cabecalho_mes_original": mes_rotulo_ssp132,
                    "pagina_cabecalho_meses": pagina_cab_ssp132,
                    "valor_textual_original": texto_ssp132, "unidade": unidade_ssp132, "regra_interpretacao": regra_ssp132})
if problemas_ssp132:
    print("Etapa 13.2: valor ausente ou não numérico em coluna mensal válida; construção interrompida, sem preenchimento.")
    mostrar_ssp131(pd_ssp131.DataFrame(problemas_ssp132))
    raise ValueError("Etapa 13.2: valores ausentes ou não numéricos em colunas mensais válidas; revisar a evidência acima.")

tabela_temporaria_ssp132 = pd_ssp131.DataFrame(registros_ssp132)
ordem_categoria_ssp132 = {nome: posicao for posicao, nome in enumerate(categorias_analiticas_ssp132)}
tabela_temporaria_ssp132 = tabela_temporaria_ssp132.assign(_ordem=tabela_temporaria_ssp132.categoria.map(ordem_categoria_ssp132)).sort_values(
    ["ano", "mes", "_ordem"], kind="stable").drop(columns="_ordem").reset_index(drop=True)
colunas_base_ssp132 = ["municipio", "ano", "mes", "categoria", "fatos_comunicados"]
tipos_base_ssp132 = {"municipio": "string", "ano": "Int64", "mes": "Int64", "categoria": "string", "fatos_comunicados": "Int64"}
base_temporaria_ssp132 = tabela_temporaria_ssp132[colunas_base_ssp132].astype(tipos_base_ssp132)
colunas_rastreabilidade_ssp132 = [
    "municipio", "ano", "mes", "categoria", "relatorio_origem", "arquivo_pdf", "sha256_pdf", "arquivo_csv", "sha256_csv",
    "pagina_pdf", "tabela_pdf", "linha_tabela", "coluna_csv", "lado_da_tabela", "pagina_horizontal_pareada",
    "rotulo_original_categoria", "municipio_celula_original", "cabecalho_ano_original", "cabecalho_mes_original",
    "pagina_cabecalho_meses", "valor_textual_original", "unidade", "regra_interpretacao"]
tipos_rastreabilidade_ssp132 = {c: "string" for c in colunas_rastreabilidade_ssp132}
tipos_rastreabilidade_ssp132.update({"ano": "Int64", "mes": "Int64"})
rastreabilidade_temporaria_ssp132 = tabela_temporaria_ssp132[colunas_rastreabilidade_ssp132].astype(tipos_rastreabilidade_ssp132)
print(f"Construção temporária concluída: base {base_temporaria_ssp132.shape}, rastreabilidade {rastreabilidade_temporaria_ssp132.shape}. Verificações a seguir.")

# %% [markdown]
# ## 13.2.3 — Verificações mínimas de segurança da construção
# Se qualquer uma falhar, a execução é interrompida e as variáveis finais não são
# criadas. A auditoria integral e a aprovação final ficam para a Etapa 13.3.

# %%
chaves_ssp132 = ["municipio", "ano", "mes", "categoria"]
conjunto_categorias_ssp132 = set(categorias_analiticas_ssp132)
competencias_ssp132 = base_temporaria_ssp132.groupby(["ano", "mes"], sort=True)
contagem_origem_ssp132 = rastreabilidade_temporaria_ssp132.relatorio_origem.value_counts().to_dict()
ligacao_ssp132 = base_temporaria_ssp132.merge(
    rastreabilidade_temporaria_ssp132[chaves_ssp132 + ["valor_textual_original"]],
    on=chaves_ssp132, how="outer", validate="one_to_one", indicator=True)
indices_csv_ssp132 = {
    relatorio: bruto.set_index(["pagina_pdf", "tabela_pdf", "linha_tabela"])
    for relatorio, bruto in brutos_dom_ssp131.items()}
if not all(indice.index.is_unique for indice in indices_csv_ssp132.values()):
    raise ValueError("Etapa 13.2: chave página/tabela/linha não é única nos CSVs brutos.")
textos_csv_ssp132 = [
    indices_csv_ssp132[linha.relatorio_origem].at[(linha.pagina_pdf, linha.tabela_pdf, linha.linha_tabela), linha.coluna_csv]
    for linha in rastreabilidade_temporaria_ssp132.itertuples()]
verificacoes_construcao_ssp132 = [
    ("base com 216 linhas × 5 colunas", (216, 5), base_temporaria_ssp132.shape),
    ("rastreabilidade com 216 linhas", 216, len(rastreabilidade_temporaria_ssp132)),
    ("colunas da base na ordem aprovada", colunas_base_ssp132, base_temporaria_ssp132.columns.tolist()),
    ("tipos da base", ["string", "Int64", "Int64", "string", "Int64"], [str(t) for t in base_temporaria_ssp132.dtypes]),
    ("nove categorias por competência (24 competências)", (24, True),
     (competencias_ssp132.ngroups, bool(competencias_ssp132.categoria.apply(lambda s: set(s) == conjunto_categorias_ssp132 and len(s) == 9).all()))),
    ("12 competências em 2023", list(range(1, 13)), sorted(base_temporaria_ssp132.loc[base_temporaria_ssp132.ano.eq(2023), "mes"].unique().tolist())),
    ("12 competências em 2024", list(range(1, 13)), sorted(base_temporaria_ssp132.loc[base_temporaria_ssp132.ano.eq(2024), "mes"].unique().tolist())),
    ("54 registros do 1º semestre de 2023", 54, contagem_origem_ssp132.get("1º semestre de 2023")),
    ("54 registros do 2º semestre de 2023", 54, contagem_origem_ssp132.get("2º semestre de 2023")),
    ("108 registros de 2024", 108, contagem_origem_ssp132.get("2024")),
    ("chaves duplicadas na base", 0, int(base_temporaria_ssp132.duplicated(chaves_ssp132).sum())),
    ("chaves duplicadas na rastreabilidade", 0, int(rastreabilidade_temporaria_ssp132.duplicated(chaves_ssp132).sum())),
    ("ligação um para um pelas chaves (216 pares)", 216, int(ligacao_ssp132._merge.eq("both").sum())),
    ("valor analítico = texto original convertido", True,
     bool((ligacao_ssp132.fatos_comunicados == ligacao_ssp132.valor_textual_original.map(int)).all())),
    ("texto original da rastreabilidade = célula do CSV bruto (consulta independente)", True,
     textos_csv_ssp132 == rastreabilidade_temporaria_ssp132.valor_textual_original.tolist()),
    ("zeros publicados preservados como zero", int(rastreabilidade_temporaria_ssp132.valor_textual_original.map(int).eq(0).sum()),
     int(base_temporaria_ssp132.fatos_comunicados.eq(0).sum())),
    ("sem ausências nas cinco colunas", 0, int(base_temporaria_ssp132.isna().sum().sum())),
    ("município único JOINVILLE", ["JOINVILLE"], base_temporaria_ssp132.municipio.unique().tolist()),
    ("valores inteiros não negativos", True, bool((base_temporaria_ssp132.fatos_comunicados >= 0).all())),
    ("24 observações por categoria analítica", {c: 24 for c in categorias_analiticas_ssp132},
     base_temporaria_ssp132.categoria.value_counts().reindex(list(categorias_analiticas_ssp132)).to_dict()),
    ("seis arquivos de origem com SHA-256 inalterado", True,
     all(sha256_ssp131(caminho) == hash_antes for caminho, hash_antes in hashes_antes_ssp131.items()) and len(hashes_antes_ssp131) == 6),
]
resumo_verificacoes_construcao_ssp132 = pd_ssp131.DataFrame(
    [{"verificacao": nome, "esperado": esperado, "encontrado": encontrado, "ok": esperado == encontrado}
     for nome, esperado, encontrado in verificacoes_construcao_ssp132])
if not resumo_verificacoes_construcao_ssp132.ok.all():
    print("Etapa 13.2: verificação falhou; as variáveis finais NÃO foram criadas. Evidência:")
    mostrar_ssp131(resumo_verificacoes_construcao_ssp132.loc[~resumo_verificacoes_construcao_ssp132.ok])
    raise ValueError("Etapa 13.2: verificações mínimas não conferiram; interromper e revisar a evidência.")

df_violencia_domestica_joinville_ssp132 = base_temporaria_ssp132.copy()
rastreabilidade_violencia_domestica_ssp132 = rastreabilidade_temporaria_ssp132.copy()
print(f"Etapa 13.2: {len(resumo_verificacoes_construcao_ssp132)} verificações conferidas; base e rastreabilidade criadas somente em memória.")

# %% [markdown]
# ## 13.2.4 — Exibição da base analítica

# %%
print("Primeiras 15 linhas da base:")
mostrar_ssp131(df_violencia_domestica_joinville_ssp132.head(15))
print("Últimas 15 linhas da base:")
mostrar_ssp131(df_violencia_domestica_joinville_ssp132.tail(15))
print("Tipos de dados da base:")
mostrar_ssp131(df_violencia_domestica_joinville_ssp132.dtypes.astype(str).rename("dtype").to_frame())
print("Dimensões:",
      f"base {df_violencia_domestica_joinville_ssp132.shape};",
      f"rastreabilidade {rastreabilidade_violencia_domestica_ssp132.shape}")
print("Quantidade de categorias analíticas:", df_violencia_domestica_joinville_ssp132.categoria.nunique())
print("Cobertura de anos e meses:")
mostrar_ssp131(df_violencia_domestica_joinville_ssp132.groupby("ano").agg(
    competencias=("mes", "nunique"), mes_inicial=("mes", "min"), mes_final=("mes", "max"), linhas=("mes", "size")))

# %% [markdown]
# ## 13.2.5 — Resultado do mapeamento, rastreabilidade e verificações
# O mapeamento é uma decisão metodológica do projeto, não equivalência confirmada
# pela SSP/SC. A amostra da rastreabilidade mostra a primeira e a última observação
# de cada relatório e as lesões corporais leves nos extremos de cada relatório.

# %%
print("Mapeamento aplicado (rótulo original -> categoria analítica):")
mapeamento_aplicado_ssp132 = rastreabilidade_violencia_domestica_ssp132.groupby(
    ["rotulo_original_categoria", "categoria"], sort=False).agg(
    relatorios=("relatorio_origem", lambda s: ", ".join(sorted(set(s)))), observacoes=("mes", "size")).reset_index()
mostrar_ssp131(mapeamento_aplicado_ssp132)
leve_ssp132 = rastreabilidade_violencia_domestica_ssp132.categoria.eq("Lesão corporal leve dolosa")
posicoes_ssp132 = sorted(set(
    rastreabilidade_violencia_domestica_ssp132.groupby("relatorio_origem").head(1).index.tolist()
    + rastreabilidade_violencia_domestica_ssp132.groupby("relatorio_origem").tail(1).index.tolist()
    + rastreabilidade_violencia_domestica_ssp132.loc[leve_ssp132].groupby("relatorio_origem").head(1).index.tolist()
    + rastreabilidade_violencia_domestica_ssp132.loc[leve_ssp132].groupby("relatorio_origem").tail(1).index.tolist()))
print("Amostra da tabela de rastreabilidade:")
mostrar_ssp131(rastreabilidade_violencia_domestica_ssp132.loc[posicoes_ssp132])
print("Resumo das verificações mínimas de construção:")
mostrar_ssp131(resumo_verificacoes_construcao_ssp132)
print("Nenhuma exportação, imputação, deduplicação, soma de categorias, total mensal/anual ou gráfico foi realizado.")

# %% [markdown]
# # ETAPA 13.3 — Validação integral da base analítica de violência doméstica (SSP/SC)
# Execute depois da Etapa 13.2 (que, por sua vez, exige as células 13.1.1 a 13.1.7).
# Não é preciso executar as Etapas 1–12 nem o Sinesp. Esta etapa só lê e audita: não
# altera a base aprovada, os CSVs, os PDFs nem exporta arquivos, e não gera gráficos
# ou estatísticas comparativas.
#
# Como a conferência é independente da construção:
# 1. Os três CSVs são relidos com o módulo `csv` (a Etapa 13.2 usou o Pandas) e cada
#    observação é localizada pelas coordenadas da rastreabilidade (arquivo, página,
#    tabela, linha, coluna). A referência esperada é a célula do CSV, não uma cópia
#    do DataFrame analítico.
# 2. Os PDFs são relidos com o `pypdf`, por linhas de texto (as Etapas 12 e 13.1
#    usaram o `pdfplumber` e suas tabelas). Ano e meses vêm dos cabeçalhos; o mês de
#    cada célula é deduzido da posição da coluna, e não do que a base registra.
# 3. As tabelas do `pdfplumber` da Etapa 13.1 entram como controle adicional.
# Limitação registrada: no `pypdf` a página de continuação tem uma linha de texto por
# linha da tabela, e a página esquerda uma linha por categoria; a correspondência
# com `linha_tabela` é feita pela ordem das linhas (confirmada pelos rótulos), não
# por coordenadas. Isso é conferência direta no PDF, mas não equivale a ler cada
# célula por posição geométrica.
# O mapeamento de "0 Dolosa" e "- Dolosa" é decisão metodológica aprovada pelo
# usuário, não equivalência oficialmente confirmada pela SSP/SC. "Fato Comunicado"
# não é equiparado a vítimas, pessoas únicas ou boletins de ocorrência únicos.

# %% [markdown]
# ## 13.3.1 — Parâmetros próprios e releitura independente dos CSVs

# %%
from pathlib import Path as CaminhoSSP133
import csv as csv_ssp133
import hashlib as hash_ssp133
import re as regex_ssp133

import pandas as pd_ssp133
from pypdf import PdfReader as LeitorSSP133
from IPython.display import display as mostrar_ssp133

necessarios_ssp133 = (
    "df_violencia_domestica_joinville_ssp132", "rastreabilidade_violencia_domestica_ssp132")
ausentes_ssp133 = [nome for nome in necessarios_ssp133 if nome not in globals()]
if ausentes_ssp133:
    raise NameError(f"Etapa 13.3: execute antes a Etapa 13.2. Ausentes: {ausentes_ssp133}")

raiz_ssp133 = CaminhoSSP133(__file__).resolve().parents[2] if "__file__" in globals() else CaminhoSSP133.cwd()
pasta_pdf_ssp133 = raiz_ssp133 / "dados/brutos/ssp_sc/pdfs"
pasta_csv_ssp133 = raiz_ssp133 / "dados/brutos/ssp_sc/csv_extraidos"
base_ssp133 = df_violencia_domestica_joinville_ssp132
rastreio_ssp133 = rastreabilidade_violencia_domestica_ssp132

# Constantes próprias desta etapa (não reutilizam as variáveis da 13.2).
arquivos_ssp133 = {
    "1º semestre de 2023": "2025_06_VIOLENCIA-DOMESTICA-1o-SEMESTRE-2023",
    "2º semestre de 2023": "2025_06_VIOLENCIA-DOMESTICA-2o-SEMESTRE-2023",
    "2024": "2025_06_VIOLENCIA-DOMESTICA-2024-SC",
}
mes_numero_ssp133 = {"jan": 1, "fev": 2, "mar": 3, "abr": 4, "mai": 5, "jun": 6,
                     "jul": 7, "ago": 8, "set": 9, "out": 10, "nov": 11, "dez": 12}
mapeamento_esperado_ssp133 = {
    "Ameaça": "Ameaça", "Calúnia": "Calúnia", "Difamação": "Difamação", "Estupro": "Estupro",
    "Injúria": "Injúria", "Injúria real": "Injúria real", "Vias de fato": "Vias de fato",
    "Lesão corporal grave ou gravíssima 0 Dolosa": "Lesão corporal grave ou gravíssima dolosa",
    "Lesão corporal grave ou gravíssima - Dolosa": "Lesão corporal grave ou gravíssima dolosa",
    "Lesão corporal leve 0 Dolosa": "Lesão corporal leve dolosa",
    "Lesão corporal leve - Dolosa": "Lesão corporal leve dolosa",
}
categorias_esperadas_ssp133 = sorted(set(mapeamento_esperado_ssp133.values()))
chaves_ssp133 = ["municipio", "ano", "mes", "categoria"]

verificacoes_ssp133 = {}


def registrar_ssp133(bloco, verificacao, esperado, encontrado):
    # Registro por chave: executar novamente uma célula substitui, não duplica.
    verificacoes_ssp133[(bloco, verificacao)] = {
        "bloco": bloco, "verificacao": verificacao, "esperado": esperado, "encontrado": encontrado,
        "ok": esperado == encontrado}


def sha256_ssp133(caminho):
    with caminho.open("rb") as fluxo:
        return hash_ssp133.file_digest(fluxo, "sha256").hexdigest()


# Releitura dos CSVs com o módulo csv, em variáveis próprias. Cada linha recebe o
# município de contexto: o último valor não vazio de celula_01 na mesma página e tabela.
csv_linhas_ssp133 = {}
csv_indice_ssp133 = {}
csv_cabecalhos_ssp133 = {}
for relatorio_ssp133, arquivo_ssp133 in arquivos_ssp133.items():
    caminho_ssp133 = pasta_csv_ssp133 / (arquivo_ssp133 + ".csv")
    with caminho_ssp133.open(encoding="utf-8-sig", newline="") as fluxo_ssp133:
        leitor_ssp133 = csv_ssp133.reader(fluxo_ssp133, delimiter=";")
        cabecalho_csv_ssp133 = next(leitor_ssp133)
        linhas_brutas_ssp133 = list(leitor_ssp133)
    registrar_ssp133("CSV", f"todas as linhas com {len(cabecalho_csv_ssp133)} campos ({relatorio_ssp133})",
                     True, all(len(l) == len(cabecalho_csv_ssp133) for l in linhas_brutas_ssp133))
    contexto_ssp133 = {}
    linhas_ssp133 = []
    for linha_bruta_ssp133 in linhas_brutas_ssp133:
        linha_ssp133 = dict(zip(cabecalho_csv_ssp133, linha_bruta_ssp133))
        chave_pagina_ssp133 = (linha_ssp133["pagina_pdf"], linha_ssp133["tabela_pdf"])
        if linha_ssp133["celula_01"] != "":
            contexto_ssp133[chave_pagina_ssp133] = linha_ssp133["celula_01"]
        linha_ssp133["municipio_contexto"] = contexto_ssp133.get(chave_pagina_ssp133, "")
        linhas_ssp133.append(linha_ssp133)
    csv_linhas_ssp133[relatorio_ssp133] = linhas_ssp133
    csv_cabecalhos_ssp133[relatorio_ssp133] = tuple(cabecalho_csv_ssp133)
    csv_indice_ssp133[relatorio_ssp133] = {
        (l["pagina_pdf"], l["tabela_pdf"], l["linha_tabela"]): l for l in linhas_ssp133}
    registrar_ssp133("CSV", f"chave página/tabela/linha única ({relatorio_ssp133})",
                     len(linhas_ssp133), len(csv_indice_ssp133[relatorio_ssp133]))

# Blocos de Joinville descobertos de forma independente no CSV (contexto JOINVILLE e categoria preenchida).
blocos_csv_ssp133 = pd_ssp133.DataFrame([
    {"relatorio": rel, "arquivo_pdf": l["arquivo_pdf"], "pagina_pdf": l["pagina_pdf"], "tabela_pdf": l["tabela_pdf"],
     "linha_tabela": l["linha_tabela"], "celula_01": l["celula_01"], "rotulo_original_csv": l["celula_02"]}
    for rel, linhas in csv_linhas_ssp133.items() for l in linhas
    if l["municipio_contexto"] == "JOINVILLE" and l["celula_02"] != ""])
print("Blocos de Joinville localizados na releitura independente dos CSVs:")
mostrar_ssp133(blocos_csv_ssp133.groupby("relatorio").agg(
    pagina=("pagina_pdf", lambda s: ", ".join(sorted(set(s)))), tabela=("tabela_pdf", "first"),
    linha_inicial=("linha_tabela", lambda s: min(map(int, s))), linha_final=("linha_tabela", lambda s: max(map(int, s))),
    linhas=("linha_tabela", "size")).reindex(list(arquivos_ssp133)))
registrar_ssp133("CSV", "27 linhas de categoria no contexto JOINVILLE (9 por relatório)", 27, len(blocos_csv_ssp133))

# %% [markdown]
# ## 13.3.2 — Releitura independente dos PDFs (pypdf)
# Ano e meses vêm dos cabeçalhos impressos na primeira página de cada metade (página 1
# e página total//2 + 1). Na página esquerda, cada linha de categoria termina com os
# valores dos primeiros meses (linhas só com o nome de um município são ignoradas); na
# página de continuação, cada linha traz os valores dos meses restantes.

# %%
padrao_mes_ssp133 = r"(?:jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez)"
leituras_pdf_ssp133 = {}
linhas_cabecalho_pdf_ssp133 = []
for relatorio_ssp133, arquivo_ssp133 in arquivos_ssp133.items():
    leitor_pdf_ssp133 = LeitorSSP133(pasta_pdf_ssp133 / (arquivo_ssp133 + ".pdf"))
    total_paginas_ssp133 = len(leitor_pdf_ssp133.pages)
    pagina_cab_dir_ssp133 = total_paginas_ssp133 // 2 + 1
    observacoes_relatorio_ssp133 = rastreio_ssp133.loc[rastreio_ssp133.relatorio_origem.eq(relatorio_ssp133)]
    pagina_esq_ssp133 = int(observacoes_relatorio_ssp133.loc[observacoes_relatorio_ssp133.lado_da_tabela.eq("esquerda"), "pagina_pdf"].iloc[0])
    pagina_dir_ssp133 = int(observacoes_relatorio_ssp133.loc[observacoes_relatorio_ssp133.lado_da_tabela.eq("direita"), "pagina_pdf"].iloc[0])
    cabecalhos_ssp133 = {}
    for papel_ssp133, pagina_ssp133 in (("esquerda", 1), ("direita", pagina_cab_dir_ssp133)):
        texto_ssp133 = leitor_pdf_ssp133.pages[pagina_ssp133 - 1].extract_text() or ""
        # No pypdf a ordem do texto difere da visual (o ano pode vir depois das linhas da tabela):
        # procuram-se as linhas de cabeçalho em toda a página, exigindo uma única linha de meses e um único ano.
        linhas_texto_pagina_ssp133 = texto_ssp133.splitlines()
        primeiras_ssp133 = linhas_texto_pagina_ssp133[:4]
        linhas_meses_ssp133 = [l for l in linhas_texto_pagina_ssp133 if regex_ssp133.fullmatch(rf"{padrao_mes_ssp133}(?:\s+{padrao_mes_ssp133})*", l.strip())]
        meses_ssp133 = tuple(m for l in linhas_meses_ssp133 for m in l.split()) if len(linhas_meses_ssp133) == 1 else ()
        anos_ssp133 = sorted({l.strip() for l in linhas_texto_pagina_ssp133 if regex_ssp133.fullmatch(r"20\d{2}", l.strip())})
        cabecalhos_ssp133[papel_ssp133] = (pagina_ssp133, anos_ssp133, meses_ssp133, "Fato Comunicado" in texto_ssp133)
        linhas_cabecalho_pdf_ssp133.append({
            "relatorio": relatorio_ssp133, "metade": papel_ssp133, "pagina_pdf": pagina_ssp133,
            "ano_impresso": anos_ssp133[0] if anos_ssp133 else "", "meses_impressos": meses_ssp133,
            "contem_Fato_Comunicado": "Fato Comunicado" in texto_ssp133,
            "quatro_primeiras_linhas": " | ".join(primeiras_ssp133)})
    texto_esq_ssp133 = leitor_pdf_ssp133.pages[pagina_esq_ssp133 - 1].extract_text().splitlines()
    itens_esq_ssp133 = []
    for linha_texto_ssp133 in texto_esq_ssp133:
        # Rótulo = menor prefixo tal que o restante da linha são só números (valores dos meses).
        correspondencia_ssp133 = regex_ssp133.fullmatch(r"(.*?\S)\s+(\d+(?:\s+\d+)*)", linha_texto_ssp133.strip())
        if correspondencia_ssp133:
            itens_esq_ssp133.append((correspondencia_ssp133.group(1), correspondencia_ssp133.group(2).split()))
    texto_dir_ssp133 = leitor_pdf_ssp133.pages[pagina_dir_ssp133 - 1].extract_text().splitlines()
    linhas_dir_ssp133 = [l.split() for l in texto_dir_ssp133 if regex_ssp133.fullmatch(r"\d+(?:\s+\d+)*", l.strip())]
    leituras_pdf_ssp133[relatorio_ssp133] = {
        "paginas": total_paginas_ssp133, "pagina_esq": pagina_esq_ssp133, "pagina_dir": pagina_dir_ssp133,
        "ano": cabecalhos_ssp133["esquerda"][1], "ano_dir": cabecalhos_ssp133["direita"][1],
        "meses_esq": cabecalhos_ssp133["esquerda"][2], "meses_dir": cabecalhos_ssp133["direita"][2],
        "itens_esq": itens_esq_ssp133, "linhas_dir": linhas_dir_ssp133}
    registrar_ssp133("PDF", f"ano único e igual nas duas metades ({relatorio_ssp133})", True,
                     len(cabecalhos_ssp133["esquerda"][1]) == 1 and cabecalhos_ssp133["esquerda"][1] == cabecalhos_ssp133["direita"][1])
    registrar_ssp133("PDF", f"unidade Fato Comunicado no cabeçalho ({relatorio_ssp133})", True, cabecalhos_ssp133["esquerda"][3])
    registrar_ssp133("PDF", f"linhas de categoria na esquerda = linhas de valores na direita ({relatorio_ssp133})",
                     len(itens_esq_ssp133), len(linhas_dir_ssp133))
cabecalhos_independentes_ssp133 = pd_ssp133.DataFrame(linhas_cabecalho_pdf_ssp133)
print("Cabeçalhos das páginas iniciais das duas metades (leitura pypdf, rótulos literais):")
mostrar_ssp133(cabecalhos_independentes_ssp133)

# %% [markdown]
# ## 13.3.3 — Conferência independente das 216 observações
# Para cada observação da rastreabilidade: o valor é buscado na célula do CSV
# (coordenadas arquivo/página/tabela/linha/coluna), na linha de texto do PDF e na
# tabela do `pdfplumber` da Etapa 13.1 (controle). Ano e mês saem do cabeçalho do PDF
# a partir da posição da coluna. A categoria esperada vem do mapeamento próprio da
# Etapa 13.3. Cada linha da tabela abaixo é uma comparação individual.

# %%
valor_base_ssp133 = {
    (l.municipio, int(l.ano), int(l.mes), l.categoria): l.fatos_comunicados
    for l in base_ssp133.itertuples()}
controle_tabelas_ssp133 = (
    {(l.relatorio, str(l.pagina_pdf), str(l.linha_tabela)): l.celulas_pdf for l in conferencia_pdf_csv_ssp131.itertuples()}
    if "conferencia_pdf_csv_ssp131" in globals() else {})
controle_etapa12_ssp133 = (
    {(str(l.arquivo_pdf), str(l.pagina_pdf), str(l.tabela_pdf), str(l.linha_tabela), str(l.coluna_csv)): l.valor_pdf
     for l in auditoria_celulas_domesticas_ssp12.itertuples()}
    if "auditoria_celulas_domesticas_ssp12" in globals() else {})
linhas_conferencia_ssp133 = []
for linha_ssp133 in rastreio_ssp133.itertuples():
    relatorio_ssp133 = linha_ssp133.relatorio_origem
    leitura_ssp133 = leituras_pdf_ssp133[relatorio_ssp133]
    esquerda_ssp133 = linha_ssp133.lado_da_tabela == "esquerda"
    pagina_esq_obs_ssp133 = linha_ssp133.pagina_pdf if esquerda_ssp133 else linha_ssp133.pagina_horizontal_pareada
    chave_csv_ssp133 = (linha_ssp133.pagina_pdf, linha_ssp133.tabela_pdf, linha_ssp133.linha_tabela)
    linha_csv_ssp133 = csv_indice_ssp133[relatorio_ssp133].get(chave_csv_ssp133)
    linha_csv_esq_ssp133 = csv_indice_ssp133[relatorio_ssp133].get((pagina_esq_obs_ssp133, linha_ssp133.tabela_pdf, linha_ssp133.linha_tabela))
    texto_csv_ssp133 = linha_csv_ssp133.get(linha_ssp133.coluna_csv) if linha_csv_ssp133 else None
    # Posição do mês deduzida da coluna do CSV: a esquerda começa em celula_03 e a direita em celula_01.
    posicao_ssp133 = int(linha_ssp133.coluna_csv[-2:]) - (3 if esquerda_ssp133 else 1)
    meses_cab_ssp133 = leitura_ssp133["meses_esq"] if esquerda_ssp133 else leitura_ssp133["meses_dir"]
    mes_rotulo_ssp133 = meses_cab_ssp133[posicao_ssp133] if 0 <= posicao_ssp133 < len(meses_cab_ssp133) else None
    ano_pdf_ssp133 = int(leitura_ssp133["ano"][0]) if len(leitura_ssp133["ano"]) == 1 else None
    numero_linha_ssp133 = int(linha_ssp133.linha_tabela)
    itens_ssp133 = leitura_ssp133["itens_esq"]
    rotulo_pdf_ssp133, tokens_esq_ssp133 = itens_ssp133[numero_linha_ssp133 - 1] if numero_linha_ssp133 <= len(itens_ssp133) else (None, [])
    tokens_dir_ssp133 = leitura_ssp133["linhas_dir"][numero_linha_ssp133 - 1] if numero_linha_ssp133 <= len(leitura_ssp133["linhas_dir"]) else []
    tokens_ssp133 = tokens_esq_ssp133 if esquerda_ssp133 else tokens_dir_ssp133
    texto_pdf_ssp133 = tokens_ssp133[posicao_ssp133] if 0 <= posicao_ssp133 < len(tokens_ssp133) else None
    celulas_tabela_ssp133 = controle_tabelas_ssp133.get((relatorio_ssp133, linha_ssp133.pagina_pdf, linha_ssp133.linha_tabela))
    texto_tabela_ssp133 = celulas_tabela_ssp133[int(linha_ssp133.coluna_csv[-2:]) - 1] if celulas_tabela_ssp133 is not None else None
    texto_etapa12_ssp133 = controle_etapa12_ssp133.get(
        (linha_ssp133.arquivo_pdf, linha_ssp133.pagina_pdf, linha_ssp133.tabela_pdf, linha_ssp133.linha_tabela, linha_ssp133.coluna_csv))
    mes_independente_ssp133 = mes_numero_ssp133.get(mes_rotulo_ssp133)
    valor_analitico_ssp133 = valor_base_ssp133.get((linha_ssp133.municipio, int(linha_ssp133.ano), int(linha_ssp133.mes), linha_ssp133.categoria))
    texto_inteiro_ssp133 = isinstance(texto_csv_ssp133, str) and texto_csv_ssp133.isascii() and texto_csv_ssp133.isdigit() and str(int(texto_csv_ssp133)) == texto_csv_ssp133
    verificacoes_linha_ssp133 = {
        "csv_linha_existe": linha_csv_ssp133 is not None,
        "csv_pdf_arquivo": bool(linha_csv_ssp133) and linha_csv_ssp133["arquivo_pdf"] == linha_ssp133.arquivo_pdf,
        "municipio_contexto_JOINVILLE": bool(linha_csv_esq_ssp133) and linha_csv_esq_ssp133["municipio_contexto"] == "JOINVILLE" == linha_ssp133.municipio,
        "ano_pdf": ano_pdf_ssp133 is not None and ano_pdf_ssp133 == int(linha_ssp133.ano),
        "mes_pela_coluna_e_cabecalho": mes_independente_ssp133 is not None and mes_independente_ssp133 == int(linha_ssp133.mes) and mes_rotulo_ssp133 == linha_ssp133.cabecalho_mes_original,
        "rotulo_csv": bool(linha_csv_esq_ssp133) and linha_csv_esq_ssp133["celula_02"] == linha_ssp133.rotulo_original_categoria,
        "rotulo_pdf": rotulo_pdf_ssp133 is not None and rotulo_pdf_ssp133.endswith(linha_ssp133.rotulo_original_categoria),
        "categoria_mapeamento": mapeamento_esperado_ssp133.get(linha_ssp133.rotulo_original_categoria) == linha_ssp133.categoria,
        "texto_csv_igual_rastreio": texto_csv_ssp133 == linha_ssp133.valor_textual_original,
        "texto_inteiro_sem_alteracao": bool(texto_inteiro_ssp133),
        "valor_base_igual_csv": bool(texto_inteiro_ssp133) and valor_analitico_ssp133 is not None and not pd_ssp133.isna(valor_analitico_ssp133) and int(valor_analitico_ssp133) == int(texto_csv_ssp133),
        "valor_pdf_texto": texto_pdf_ssp133 is not None and texto_pdf_ssp133 == texto_csv_ssp133,
        "valor_pdf_tabela_controle": texto_tabela_ssp133 is not None and texto_tabela_ssp133 == texto_csv_ssp133,
    }
    if controle_etapa12_ssp133:
        verificacoes_linha_ssp133["valor_pdf_etapa12_controle"] = texto_etapa12_ssp133 is not None and texto_etapa12_ssp133 == texto_csv_ssp133
    linhas_conferencia_ssp133.append({
        "municipio": linha_ssp133.municipio, "ano": int(linha_ssp133.ano), "mes": int(linha_ssp133.mes), "categoria": linha_ssp133.categoria,
        "relatorio_origem": relatorio_ssp133, "arquivo_pdf": linha_ssp133.arquivo_pdf, "pagina_pdf": linha_ssp133.pagina_pdf,
        "tabela_pdf": linha_ssp133.tabela_pdf, "linha_tabela": linha_ssp133.linha_tabela, "coluna_csv": linha_ssp133.coluna_csv,
        "municipio_contexto_csv": linha_csv_esq_ssp133["municipio_contexto"] if linha_csv_esq_ssp133 else None,
        "ano_cabecalho_pdf": ano_pdf_ssp133, "mes_cabecalho_pdf": mes_rotulo_ssp133, "mes_numero_independente": mes_independente_ssp133,
        "rotulo_original_rastreio": linha_ssp133.rotulo_original_categoria,
        "rotulo_original_csv": linha_csv_esq_ssp133["celula_02"] if linha_csv_esq_ssp133 else None,
        "rotulo_original_pdf": rotulo_pdf_ssp133, "categoria_esperada": mapeamento_esperado_ssp133.get(linha_ssp133.rotulo_original_categoria),
        "valor_texto_rastreio": linha_ssp133.valor_textual_original, "valor_texto_csv": texto_csv_ssp133,
        "valor_texto_pdf_pypdf": texto_pdf_ssp133, "valor_texto_pdf_tabela_13_1": texto_tabela_ssp133,
        "valor_inteiro_base": int(valor_analitico_ssp133) if valor_analitico_ssp133 is not None and not pd_ssp133.isna(valor_analitico_ssp133) else None,
        **verificacoes_linha_ssp133, "todas_conferem": all(verificacoes_linha_ssp133.values())})
conferencia_valores_ssp133 = pd_ssp133.DataFrame(linhas_conferencia_ssp133)
colunas_booleanas_ssp133 = [c for c in conferencia_valores_ssp133.columns if conferencia_valores_ssp133[c].dtype == bool]
print(f"{len(conferencia_valores_ssp133)} comparações individuais (uma por observação). Resultado por tipo de conferência:")
mostrar_ssp133(conferencia_valores_ssp133[colunas_booleanas_ssp133].sum().rename("observações conferidas").to_frame().T)
registrar_ssp133("Valores", "216 observações conferidas individualmente", 216, len(conferencia_valores_ssp133))
for coluna_ssp133 in colunas_booleanas_ssp133:
    if coluna_ssp133 != "todas_conferem":
        registrar_ssp133("Valores", f"{coluna_ssp133} em 216/216", 216, int(conferencia_valores_ssp133[coluna_ssp133].sum()))
registrar_ssp133("Valores", "todas as conferências por observação", 216, int(conferencia_valores_ssp133.todas_conferem.sum()))
registrar_ssp133("Valores", "controle da Etapa 12 disponível nesta sessão", bool(controle_etapa12_ssp133), bool(controle_etapa12_ssp133))
print("Controle da Etapa 12 (auditoria_celulas_domesticas_ssp12) na sessão:", bool(controle_etapa12_ssp133),
      "— se False, ele não foi usado; os controles efetivos são pypdf e pdfplumber da 13.1.")
print("Conferência independente dos valores (216 linhas):")
mostrar_ssp133(conferencia_valores_ssp133)

# %% [markdown]
# ## 13.3.4 — Completude: as 216 células esperadas, derivadas do CSV e dos cabeçalhos
# O conjunto esperado de células é derivado sem usar a base: linhas de categoria do
# bloco de Joinville no CSV × posições dos meses nos cabeçalhos do PDF. Ele deve ser
# idêntico ao conjunto de células da rastreabilidade (nenhuma observação descartada,
# duplicada ou acrescentada).

# %%
esperadas_ssp133 = set()
for relatorio_ssp133, arquivo_ssp133 in arquivos_ssp133.items():
    leitura_ssp133 = leituras_pdf_ssp133[relatorio_ssp133]
    bloco_ssp133 = blocos_csv_ssp133.loc[blocos_csv_ssp133.relatorio.eq(relatorio_ssp133)]
    for linha_ssp133 in bloco_ssp133.itertuples():
        for posicao_ssp133 in range(len(leitura_ssp133["meses_esq"])):
            esperadas_ssp133.add((arquivo_ssp133 + ".pdf", linha_ssp133.pagina_pdf, linha_ssp133.tabela_pdf, linha_ssp133.linha_tabela, f"celula_{3 + posicao_ssp133:02d}"))
        for posicao_ssp133 in range(len(leitura_ssp133["meses_dir"])):
            esperadas_ssp133.add((arquivo_ssp133 + ".pdf", str(leitura_ssp133["pagina_dir"]), linha_ssp133.tabela_pdf, linha_ssp133.linha_tabela, f"celula_{1 + posicao_ssp133:02d}"))
presentes_ssp133 = set(zip(rastreio_ssp133.arquivo_pdf, rastreio_ssp133.pagina_pdf, rastreio_ssp133.tabela_pdf, rastreio_ssp133.linha_tabela, rastreio_ssp133.coluna_csv))
registrar_ssp133("Completude", "células esperadas derivadas de CSV e cabeçalhos", 216, len(esperadas_ssp133))
registrar_ssp133("Completude", "células da rastreabilidade distintas (sem duplicar a mesma célula de origem)", 216, len(presentes_ssp133))
registrar_ssp133("Completude", "conjunto esperado = conjunto da rastreabilidade", (set(), set()), (esperadas_ssp133 - presentes_ssp133, presentes_ssp133 - esperadas_ssp133))
print("Células esperadas:", len(esperadas_ssp133), "| células rastreadas:", len(presentes_ssp133),
      "| faltando:", len(esperadas_ssp133 - presentes_ssp133), "| a mais:", len(presentes_ssp133 - esperadas_ssp133))

# %% [markdown]
# ## 13.3.5 — Validação das categorias
# Onze rótulos originais preservados, nove categorias analíticas, 24 observações por
# categoria, nenhuma categoria desconhecida e nenhuma agregação (cada observação vem
# de uma única célula de origem). Mapeamento de "0 Dolosa" e "- Dolosa" = decisão
# metodológica aprovada, não equivalência oficial da SSP/SC.

# %%
rotulos_csv_ssp133 = blocos_csv_ssp133.groupby("rotulo_original_csv").size()
linhas_categorias_ssp133 = []
for rotulo_ssp133, categoria_ssp133 in mapeamento_esperado_ssp133.items():
    parte_ssp133 = rastreio_ssp133.loc[rastreio_ssp133.rotulo_original_categoria.eq(rotulo_ssp133)]
    n_csv_ssp133 = int(rotulos_csv_ssp133.get(rotulo_ssp133, 0))
    linhas_categorias_ssp133.append({
        "rotulo_original": rotulo_ssp133, "categoria_analitica_esperada": categoria_ssp133,
        "categorias_na_rastreabilidade": ", ".join(sorted(parte_ssp133.categoria.unique())),
        "relatorios": ", ".join(sorted(parte_ssp133.relatorio_origem.unique())),
        "linhas_de_categoria_no_csv": n_csv_ssp133, "observacoes": len(parte_ssp133),
        "observacoes_esperadas": n_csv_ssp133 * 0 + int(len(parte_ssp133)),
        "mapeamento_confere": set(parte_ssp133.categoria) == {categoria_ssp133} and n_csv_ssp133 > 0})
conferencia_categorias_ssp133 = pd_ssp133.DataFrame(linhas_categorias_ssp133).drop(columns="observacoes_esperadas")
mostrar_ssp133(conferencia_categorias_ssp133)
contagem_categorias_ssp133 = base_ssp133.categoria.value_counts().sort_index()
print("Observações por categoria analítica:")
mostrar_ssp133(contagem_categorias_ssp133.rename("observacoes").to_frame())
registrar_ssp133("Categorias", "11 rótulos originais distintos na rastreabilidade = 11 no CSV",
                 sorted(mapeamento_esperado_ssp133), sorted(set(rastreio_ssp133.rotulo_original_categoria) & set(rotulos_csv_ssp133.index)))
registrar_ssp133("Categorias", "nenhum rótulo desconhecido (rastreabilidade e CSV)", (set(), set()),
                 (set(rastreio_ssp133.rotulo_original_categoria) - set(mapeamento_esperado_ssp133), set(rotulos_csv_ssp133.index) - set(mapeamento_esperado_ssp133)))
registrar_ssp133("Categorias", "nove categorias analíticas", categorias_esperadas_ssp133, sorted(base_ssp133.categoria.unique()))
registrar_ssp133("Categorias", "24 observações por categoria", {c: 24 for c in categorias_esperadas_ssp133}, contagem_categorias_ssp133.to_dict())
registrar_ssp133("Categorias", "rótulos originais das quatro grafias de lesão preservados (6, 18, 6, 18)",
                 {"Lesão corporal grave ou gravíssima 0 Dolosa": 6, "Lesão corporal grave ou gravíssima - Dolosa": 18,
                  "Lesão corporal leve 0 Dolosa": 6, "Lesão corporal leve - Dolosa": 18},
                 {r: int(rastreio_ssp133.rotulo_original_categoria.eq(r).sum()) for r in (
                     "Lesão corporal grave ou gravíssima 0 Dolosa", "Lesão corporal grave ou gravíssima - Dolosa",
                     "Lesão corporal leve 0 Dolosa", "Lesão corporal leve - Dolosa")})
registrar_ssp133("Categorias", "sem agregação: cada observação tem célula de origem própria e única", 216, len(presentes_ssp133))
registrar_ssp133("Categorias", "mapeamento da rastreabilidade = mapeamento esperado em 216/216", 216,
                 int(conferencia_valores_ssp133.categoria_mapeamento.sum()))

# %% [markdown]
# ## 13.3.6 — Validação temporal
# Doze meses em cada ano, nove categorias por competência, sem competências duplicadas
# nem faltantes. O mês de cada célula foi deduzido do cabeçalho do PDF pela posição da
# coluna (13.3.3), e não do nome do arquivo: os nomes dos arquivos não trazem mês e o
# ano da base é conferido contra o ano impresso no cabeçalho de cada relatório.
# Nenhuma data foi derivada.

# %%
conferencia_temporal_ssp133 = conferencia_valores_ssp133.groupby(["ano", "mes"]).agg(
    categorias=("categoria", "nunique"), observacoes=("categoria", "size"),
    relatorio_origem=("relatorio_origem", lambda s: ", ".join(sorted(set(s)))),
    mes_cabecalho_pdf=("mes_cabecalho_pdf", lambda s: ", ".join(sorted(set(s)))),
    colunas_csv=("coluna_csv", lambda s: ", ".join(sorted(set(s)))),
    ano_cabecalho_pdf=("ano_cabecalho_pdf", lambda s: ", ".join(map(str, sorted(set(s)))))).reset_index()
conferencia_temporal_ssp133["nove_categorias"] = conferencia_temporal_ssp133.categorias.eq(9) & conferencia_temporal_ssp133.observacoes.eq(9)
mostrar_ssp133(conferencia_temporal_ssp133)
grade_ssp133 = {(a, m) for a in (2023, 2024) for m in range(1, 13)}
registrar_ssp133("Temporal", "24 competências esperadas presentes, sem faltantes nem extras", (set(), set()),
                 (grade_ssp133 - set(zip(conferencia_temporal_ssp133.ano, conferencia_temporal_ssp133.mes)),
                  set(zip(conferencia_temporal_ssp133.ano, conferencia_temporal_ssp133.mes)) - grade_ssp133))
registrar_ssp133("Temporal", "12 meses em 2023", list(range(1, 13)), sorted(conferencia_temporal_ssp133.loc[conferencia_temporal_ssp133.ano.eq(2023), "mes"]))
registrar_ssp133("Temporal", "12 meses em 2024", list(range(1, 13)), sorted(conferencia_temporal_ssp133.loc[conferencia_temporal_ssp133.ano.eq(2024), "mes"]))
registrar_ssp133("Temporal", "nenhuma competência duplicada na tabela (24 linhas)", 24, len(conferencia_temporal_ssp133))
registrar_ssp133("Temporal", "nove categorias e nove observações em cada uma das 24 competências", 24, int(conferencia_temporal_ssp133.nove_categorias.sum()))
registrar_ssp133("Temporal", "mês pela posição da coluna e cabeçalho do PDF = mês da base (216/216)", 216, int(conferencia_valores_ssp133.mes_pela_coluna_e_cabecalho.sum()))
registrar_ssp133("Temporal", "ano da base = ano impresso no cabeçalho (216/216)", 216, int(conferencia_valores_ssp133.ano_pdf.sum()))
registrar_ssp133("Temporal", "uma competência vem de um único relatório",
                 24, int(conferencia_temporal_ssp133.relatorio_origem.map(lambda s: "," not in s).sum()))
registrar_ssp133("Temporal", "nomes dos arquivos sem indicação de mês (competência não pode vir deles)", True,
                 all(not regex_ssp133.search(r"jan|fev|mar|abr|mai|jun|jul|ago|set|out|nov|dez", nome.lower()) for nome in arquivos_ssp133.values()))

# %% [markdown]
# ## 13.3.7 — Validação dos valores
# Todos os valores devem ser inteiros não negativos, sem arredondamento nem
# preenchimento. Zeros e ausências são contados na releitura independente do CSV.
# A unidade permanece "Fato Comunicado" (não equivale a vítimas, pessoas únicas ou BOs únicos).

# %%
textos_csv_ssp133 = conferencia_valores_ssp133.valor_texto_csv
zeros_csv_ssp133 = int(textos_csv_ssp133.eq("0").sum())
vazios_csv_ssp133 = int(textos_csv_ssp133.isna().sum() + textos_csv_ssp133.eq("").sum())
resumo_valores_ssp133 = pd_ssp133.DataFrame([
    {"indicador": "observações", "resultado": len(base_ssp133)},
    {"indicador": "zeros publicados no CSV (texto '0')", "resultado": zeros_csv_ssp133},
    {"indicador": "zeros na base analítica", "resultado": int(base_ssp133.fatos_comunicados.eq(0).sum())},
    {"indicador": "células vazias ou não numéricas nas 216 coordenadas", "resultado": vazios_csv_ssp133 + int((~conferencia_valores_ssp133.texto_inteiro_sem_alteracao).sum())},
    {"indicador": "ausências (NA) na base", "resultado": int(base_ssp133.isna().sum().sum())},
    {"indicador": "valor mínimo", "resultado": int(base_ssp133.fatos_comunicados.min())},
    {"indicador": "valor máximo", "resultado": int(base_ssp133.fatos_comunicados.max())},
    {"indicador": "unidade registrada", "resultado": ", ".join(sorted(rastreio_ssp133.unidade.unique()))}])
mostrar_ssp133(resumo_valores_ssp133)
registrar_ssp133("Valores", "27 zeros publicados preservados (CSV e base)", (27, 27), (zeros_csv_ssp133, int(base_ssp133.fatos_comunicados.eq(0).sum())))
registrar_ssp133("Valores", "nenhuma ausência real nas 216 coordenadas", 0, vazios_csv_ssp133 + int((~conferencia_valores_ssp133.texto_inteiro_sem_alteracao).sum()))
registrar_ssp133("Valores", "nenhuma ausência na base (cinco colunas)", 0, int(base_ssp133.isna().sum().sum()))
registrar_ssp133("Valores", "todos inteiros não negativos, sem alteração do texto (sem arredondamento)", 216, int(conferencia_valores_ssp133.texto_inteiro_sem_alteracao.sum()))
registrar_ssp133("Valores", "soma dos valores da base = soma dos valores do CSV (sem preenchimento, descarte ou duplicação)",
                 int(sum(int(t) for t in textos_csv_ssp133)), int(base_ssp133.fatos_comunicados.sum()))
registrar_ssp133("Valores", "unidade única 'Fato Comunicado'", ["Fato Comunicado"], sorted(rastreio_ssp133.unidade.unique()))
registrar_ssp133("Valores", "tipos da base (string, Int64, Int64, string, Int64)", ["string", "Int64", "Int64", "string", "Int64"], [str(t) for t in base_ssp133.dtypes])
registrar_ssp133("Valores", "dimensão da base 216 × 5", (216, 5), base_ssp133.shape)

# %% [markdown]
# ## 13.3.8 — Validação da rastreabilidade e dos hashes
# Ligação um para um, referências existentes e válidas, preservação dos campos
# estruturalmente vazios ou não aplicáveis do município e hashes SHA-256 dos seis
# arquivos originais comparados com o inventário histórico.

# %%
ligacao_ssp133 = base_ssp133[chaves_ssp133].merge(
    rastreio_ssp133[chaves_ssp133], on=chaves_ssp133, how="outer", validate="one_to_one", indicator=True)
paginas_pdf_ssp133 = {rel: len(LeitorSSP133(pasta_pdf_ssp133 / (arq + ".pdf")).pages) for rel, arq in arquivos_ssp133.items()}
manifesto_ssp133 = pd_ssp133.read_csv(raiz_ssp133 / "historico/documentacao_anterior/auditoria/inventario_sha256_pre_reorganizacao.csv", dtype=str, keep_default_na=False)
linhas_hash_ssp133 = []
for relatorio_ssp133, arquivo_ssp133 in arquivos_ssp133.items():
    for caminho_ssp133, coluna_hash_ssp133 in ((pasta_pdf_ssp133 / (arquivo_ssp133 + ".pdf"), "sha256_pdf"), (pasta_csv_ssp133 / (arquivo_ssp133 + ".csv"), "sha256_csv")):
        atual_ssp133 = sha256_ssp133(caminho_ssp133)
        historico_ssp133 = manifesto_ssp133.loc[manifesto_ssp133.caminho_destino_planejado.eq(caminho_ssp133.relative_to(raiz_ssp133).as_posix()), "sha256"].tolist()
        no_rastreio_ssp133 = sorted(rastreio_ssp133.loc[rastreio_ssp133.relatorio_origem.eq(relatorio_ssp133), coluna_hash_ssp133].unique())
        linhas_hash_ssp133.append({
            "arquivo": caminho_ssp133.name, "sha256_atual": atual_ssp133, "sha256_historico": historico_ssp133[0] if len(historico_ssp133) == 1 else None,
            "sha256_na_rastreabilidade": no_rastreio_ssp133[0] if len(no_rastreio_ssp133) == 1 else None,
            "igual_ao_historico": historico_ssp133 == [atual_ssp133], "igual_ao_rastreio": no_rastreio_ssp133 == [atual_ssp133]})
hashes_ssp133 = pd_ssp133.DataFrame(linhas_hash_ssp133)
esquerda_rastreio_ssp133 = rastreio_ssp133.loc[rastreio_ssp133.lado_da_tabela.eq("esquerda")]
direita_rastreio_ssp133 = rastreio_ssp133.loc[rastreio_ssp133.lado_da_tabela.eq("direita")]
referencias_validas_ssp133 = [
    (pasta_pdf_ssp133 / l.arquivo_pdf).is_file() and (pasta_csv_ssp133 / l.arquivo_csv).is_file()
    and 1 <= int(l.pagina_pdf) <= paginas_pdf_ssp133[l.relatorio_origem] and 1 <= int(l.pagina_horizontal_pareada) <= paginas_pdf_ssp133[l.relatorio_origem]
    and l.tabela_pdf == "1" and l.coluna_csv in csv_cabecalhos_ssp133[l.relatorio_origem]
    and (l.pagina_pdf, l.tabela_pdf, l.linha_tabela) in csv_indice_ssp133[l.relatorio_origem]
    for l in rastreio_ssp133.itertuples()]
conferencia_rastreabilidade_ssp133 = pd_ssp133.DataFrame([
    {"verificacao": "chaves distintas na base", "resultado": int(base_ssp133[chaves_ssp133].drop_duplicates().shape[0])},
    {"verificacao": "chaves distintas na rastreabilidade", "resultado": int(rastreio_ssp133[chaves_ssp133].drop_duplicates().shape[0])},
    {"verificacao": "pares ligados um para um", "resultado": int(ligacao_ssp133._merge.eq("both").sum())},
    {"verificacao": "referências (arquivos, página, tabela, linha, coluna) válidas", "resultado": int(sum(referencias_validas_ssp133))},
    {"verificacao": "rótulo original = célula celula_02 do CSV", "resultado": int(conferencia_valores_ssp133.rotulo_csv.sum())},
    {"verificacao": "valor textual = célula do CSV", "resultado": int(conferencia_valores_ssp133.texto_csv_igual_rastreio.sum())},
    {"verificacao": "observações da página esquerda", "resultado": len(esquerda_rastreio_ssp133)},
    {"verificacao": "observações da página direita", "resultado": len(direita_rastreio_ssp133)},
    {"verificacao": "município da célula original = JOINVILLE (só 1ª categoria à esquerda)", "resultado": int(esquerda_rastreio_ssp133.municipio_celula_original.eq("JOINVILLE").sum())},
    {"verificacao": "município da célula original vazio (continuação do bloco, esquerda)", "resultado": int(esquerda_rastreio_ssp133.municipio_celula_original.eq("").sum())},
    {"verificacao": "município não aplicável (<NA>, páginas de continuação)", "resultado": int(direita_rastreio_ssp133.municipio_celula_original.isna().sum())}])
mostrar_ssp133(conferencia_rastreabilidade_ssp133)
print("Hashes SHA-256 dos seis arquivos originais:")
mostrar_ssp133(hashes_ssp133)
registrar_ssp133("Rastreabilidade", "216 chaves distintas na base e na rastreabilidade", (216, 216), (conferencia_rastreabilidade_ssp133.resultado[0], conferencia_rastreabilidade_ssp133.resultado[1]))
registrar_ssp133("Rastreabilidade", "ligação um para um (216 pares, sem órfãos)", (216, 216), (int(ligacao_ssp133._merge.eq("both").sum()), len(ligacao_ssp133)))
registrar_ssp133("Rastreabilidade", "referências válidas em 216/216", 216, int(sum(referencias_validas_ssp133)))
registrar_ssp133("Rastreabilidade", "rótulo original e valor textual iguais ao CSV em 216/216", (216, 216),
                 (int(conferencia_valores_ssp133.rotulo_csv.sum()), int(conferencia_valores_ssp133.texto_csv_igual_rastreio.sum())))
registrar_ssp133("Rastreabilidade", "município original: 6 JOINVILLE + 48 vazios à esquerda; 162 não aplicáveis (<NA>) à direita", (6, 48, 162),
                 (int(esquerda_rastreio_ssp133.municipio_celula_original.eq("JOINVILLE").sum()), int(esquerda_rastreio_ssp133.municipio_celula_original.eq("").sum()),
                  int(direita_rastreio_ssp133.municipio_celula_original.isna().sum())))
registrar_ssp133("Rastreabilidade", "SHA-256 dos seis arquivos = inventário histórico", 6, int(hashes_ssp133.igual_ao_historico.sum()))
registrar_ssp133("Rastreabilidade", "SHA-256 dos seis arquivos = hash registrado na rastreabilidade", 6, int(hashes_ssp133.igual_ao_rastreio.sum()))
registrar_ssp133("Rastreabilidade", "cabeçalho de meses registrado na página 1 ou total//2+1 (216/216)", 216,
                 int(sum(str(l.pagina_cabecalho_meses) == str(1 if l.lado_da_tabela == "esquerda" else paginas_pdf_ssp133[l.relatorio_origem] // 2 + 1) for l in rastreio_ssp133.itertuples())))

# %% [markdown]
# ## 13.3.9 — Resumo das validações, amostras e inconsistências

# %%
resumo_validacoes_ssp133 = pd_ssp133.DataFrame(list(verificacoes_ssp133.values()))
print("Resumo das validações da Etapa 13.3:")
mostrar_ssp133(resumo_validacoes_ssp133.groupby("bloco", sort=False).agg(verificacoes=("ok", "size"), aprovadas=("ok", "sum")))
mostrar_ssp133(resumo_validacoes_ssp133)
print("Amostras da conferência independente (primeiras 3, últimas 3 e lesão leve em cada relatório):")
colunas_amostra_ssp133 = ["relatorio_origem", "ano", "mes", "categoria", "pagina_pdf", "linha_tabela", "coluna_csv",
                          "mes_cabecalho_pdf", "rotulo_original_csv", "valor_texto_csv", "valor_texto_pdf_pypdf",
                          "valor_texto_pdf_tabela_13_1", "valor_inteiro_base", "todas_conferem"]
amostra_ssp133 = pd_ssp133.concat([
    conferencia_valores_ssp133.groupby("relatorio_origem").head(3), conferencia_valores_ssp133.groupby("relatorio_origem").tail(3),
    conferencia_valores_ssp133.loc[conferencia_valores_ssp133.categoria.eq("Lesão corporal leve dolosa")].groupby("relatorio_origem").head(1),
    conferencia_valores_ssp133.loc[conferencia_valores_ssp133.categoria.eq("Lesão corporal leve dolosa")].groupby("relatorio_origem").tail(1)]).drop_duplicates(
    subset=["ano", "mes", "categoria"]).sort_values(["ano", "mes"], kind="stable")
mostrar_ssp133(amostra_ssp133[colunas_amostra_ssp133])
inconsistencias_ssp133 = resumo_validacoes_ssp133.loc[~resumo_validacoes_ssp133.ok]
observacoes_com_divergencia_ssp133 = conferencia_valores_ssp133.loc[~conferencia_valores_ssp133.todas_conferem]
print("Inconsistências encontradas:")
if inconsistencias_ssp133.empty and observacoes_com_divergencia_ssp133.empty:
    print("Nenhuma.")
else:
    mostrar_ssp133(inconsistencias_ssp133)
    mostrar_ssp133(observacoes_com_divergencia_ssp133)
total_ssp133 = len(resumo_validacoes_ssp133)
aprovadas_ssp133 = int(resumo_validacoes_ssp133.ok.sum())
print(f"Etapa 13.3: {aprovadas_ssp133}/{total_ssp133} verificações aprovadas; {int(conferencia_valores_ssp133.todas_conferem.sum())}/216 observações com todas as conferências individuais.")
print("Ressalvas: (1) a correspondência pypdf entre linha de texto e linha_tabela é por ordem das linhas, não por coordenada geométrica; "
      "(2) o mapeamento de '0 Dolosa' e '- Dolosa' e a regra de aplicar o cabeçalho de meses à tabela contínua são decisões metodológicas do projeto, "
      "não confirmações oficiais da SSP/SC; (3) 'Fato Comunicado' não é definido operacionalmente nos PDFs e não equivale a vítimas, pessoas únicas ou BOs únicos.")
print("Nenhum arquivo foi exportado e a base analítica não foi alterada.")

# %% [markdown]
# # ETAPA 14.1 — Análise exploratória e estatísticas descritivas da violência doméstica (SSP/SC)
# Execute depois das Etapas 13.1, 13.2 e 13.3 (não é preciso executar as Etapas 1–12
# nem o Sinesp). Usa somente a base analítica validada
# `df_violencia_domestica_joinville_ssp132` (216 × 5) e a rastreabilidade
# `rastreabilidade_violencia_domestica_ssp132` (216 × 23). Nada é exportado e os
# DataFrames de origem não são modificados: os cálculos usam cópias e, ao final, a
# base e a rastreabilidade são comparadas com retratos feitos no início.
#
# Cuidados de interpretação (valem para todas as tabelas desta etapa):
# - A unidade é "Fato Comunicado", conforme publicada pela SSP/SC. Não é vítima,
#   pessoa única nem boletim de ocorrência único; a definição operacional não consta
#   dos PDFs consultados.
# - As diferenças entre 2023 e 2024 são mudanças nas contagens publicadas; não
#   medem, necessariamente, mudanças na incidência real da violência.
# - Não há total geral por soma das categorias, taxa populacional, dado territorial,
#   dado de vítimas nem inferência causal.
# - As duas grafias "0 Dolosa" e "- Dolosa" das lesões corporais foram unificadas por
#   decisão metodológica do projeto (Etapa 13.2), não por confirmação oficial da SSP/SC.
# - 2023 vem de dois relatórios semestrais e 2024 de um relatório anual.

# %% [markdown]
# ## 14.1.1 — Pré-condições e retratos de preservação

# %%
import math as matematica_eda141
import statistics as estatistica_eda141

import pandas as pd_eda141
from IPython.display import display as mostrar_eda141

necessarios_eda141 = (
    "df_violencia_domestica_joinville_ssp132", "rastreabilidade_violencia_domestica_ssp132", "resumo_validacoes_ssp133")
ausentes_eda141 = [nome for nome in necessarios_eda141 if nome not in globals()]
if ausentes_eda141:
    raise NameError(f"Etapa 14.1: execute antes as Etapas 13.1, 13.2 e 13.3. Ausentes: {ausentes_eda141}")
if resumo_validacoes_ssp133.empty or not resumo_validacoes_ssp133.ok.all():
    raise ValueError("Etapa 14.1: a auditoria da Etapa 13.3 não está integralmente aprovada nesta sessão; interromper.")
print(f"Controle da Etapa 13.3 confirmado nesta sessão: {len(resumo_validacoes_ssp133)} verificações, todas aprovadas.")

categorias_eda141 = (
    "Ameaça", "Calúnia", "Difamação", "Estupro", "Injúria", "Injúria real",
    "Lesão corporal grave ou gravíssima dolosa", "Lesão corporal leve dolosa", "Vias de fato")
anos_eda141 = (2023, 2024)
unidade_eda141 = "Fato Comunicado"
chaves_eda141 = ["municipio", "ano", "mes", "categoria"]

# Retratos profundos para provar, no final, que os DataFrames de origem não mudaram.
retrato_base_eda141 = df_violencia_domestica_joinville_ssp132.copy(deep=True)
retrato_rastreio_eda141 = rastreabilidade_violencia_domestica_ssp132.copy(deep=True)
# Cópia de trabalho: todos os cálculos partem dela.
base_eda141 = df_violencia_domestica_joinville_ssp132.copy()
base_eda141["valor_float"] = base_eda141.fatos_comunicados.astype("float64")
base_eda141["competencia"] = base_eda141.ano.astype(str) + "-" + base_eda141.mes.astype(str).str.zfill(2)
print("Cópia de trabalho criada; a base original permanece intacta.")

# %% [markdown]
# ## 14.1.2 — Inspeção analítica da base

# %%
print("Dimensões da base:", df_violencia_domestica_joinville_ssp132.shape)
print("\nNomes e tipos das colunas:")
mostrar_eda141(df_violencia_domestica_joinville_ssp132.dtypes.astype(str).rename("dtype").to_frame())
print("Categorias identificadas (ordem analítica):")
mostrar_eda141(pd_eda141.DataFrame({"categoria": categorias_eda141,
                                    "presente_na_base": [c in set(base_eda141.categoria) for c in categorias_eda141]}))
print("Competências disponíveis:")
competencias_eda141 = base_eda141.groupby(["ano", "mes"]).agg(
    competencia=("competencia", "first"), categorias=("categoria", "nunique"), observacoes=("categoria", "size")).reset_index()
mostrar_eda141(competencias_eda141)
print("Observações por ano:")
mostrar_eda141(base_eda141.groupby("ano").size().rename("observacoes").to_frame())
print("Observações por categoria:")
mostrar_eda141(base_eda141.groupby("categoria").size().reindex(list(categorias_eda141)).rename("observacoes").to_frame())
print("Zeros publicados por categoria e ano (contagem de meses com valor 0):")
zeros_eda141 = base_eda141.loc[base_eda141.fatos_comunicados.eq(0)]
mostrar_eda141(zeros_eda141.pivot_table(index="categoria", columns="ano", values="mes", aggfunc="size", fill_value=0)
               .reindex(list(categorias_eda141), fill_value=0).assign(total=lambda t: t.sum(axis=1)))
print(f"Zeros publicados identificáveis: {len(zeros_eda141)} linhas (lista completa a seguir):")
mostrar_eda141(zeros_eda141[["ano", "mes", "categoria", "fatos_comunicados"]].reset_index(drop=True))
print("Ausências por coluna:")
mostrar_eda141(df_violencia_domestica_joinville_ssp132.isna().sum().rename("ausencias").to_frame())
print("Duplicidades:",
      f"linhas totalmente duplicadas = {int(df_violencia_domestica_joinville_ssp132.duplicated().sum())};",
      f"chaves (município, ano, mês, categoria) duplicadas = {int(df_violencia_domestica_joinville_ssp132.duplicated(chaves_eda141).sum())}")

# %% [markdown]
# ## 14.1.3 — Estatísticas descritivas por categoria
# Cada categoria é tratada separadamente, sobre os valores mensais publicados. O
# desvio-padrão é descritivo (`ddof=0`). Não há somatório entre categorias.

# %%
def estatisticas_eda141(dados):
    grupos = dados.groupby("categoria")["valor_float"]
    tabela = pd_eda141.DataFrame({
        "soma_publicada": grupos.sum(), "media_mensal": grupos.mean(), "mediana_mensal": grupos.median(),
        "minimo": grupos.min(), "maximo": grupos.max(), "desvio_padrao_ddof0": grupos.std(ddof=0),
        "meses_observados": grupos.count(), "meses_com_zero": grupos.apply(lambda s: int((s == 0).sum()))})
    tabela = tabela.reindex(list(categorias_eda141))
    for coluna in ("soma_publicada", "minimo", "maximo", "meses_observados"):
        tabela[coluna] = tabela[coluna].astype("Int64")
    tabela["unidade"] = unidade_eda141
    return tabela.round({"media_mensal": 4, "mediana_mensal": 4, "desvio_padrao_ddof0": 4})


estatisticas_2023_eda141 = estatisticas_eda141(base_eda141.loc[base_eda141.ano.eq(2023)])
estatisticas_2024_eda141 = estatisticas_eda141(base_eda141.loc[base_eda141.ano.eq(2024)])
estatisticas_24_meses_eda141 = estatisticas_eda141(base_eda141)
for titulo_eda141, tabela_eda141 in (("2023 (12 meses)", estatisticas_2023_eda141), ("2024 (12 meses)", estatisticas_2024_eda141),
                                     ("24 meses (jan/2023 a dez/2024)", estatisticas_24_meses_eda141)):
    print(f"\nEstatísticas descritivas — {titulo_eda141} — unidade: {unidade_eda141}")
    mostrar_eda141(tabela_eda141)

# %% [markdown]
# ## 14.1.4 — Comparação anual 2023 × 2024
# Valor acumulado dos 12 meses de cada ano, diferença absoluta (2024 − 2023) e
# variação percentual, que só existe quando o valor de 2023 é maior que zero; caso
# contrário fica `<NA>` ("não aplicável"), sem substituir a divisão por zero. As
# diferenças são mudanças nas contagens publicadas, não necessariamente na
# incidência real. Em categorias de poucos casos, variações percentuais grandes
# podem refletir diferenças de poucas unidades.

# %%
acumulados_eda141 = base_eda141.groupby(["categoria", "ano"])["fatos_comunicados"].sum().unstack("ano").reindex(list(categorias_eda141))
meses_por_ano_eda141 = base_eda141.groupby(["categoria", "ano"]).size().unstack("ano").reindex(list(categorias_eda141))
comparacao_anual_eda141 = pd_eda141.DataFrame({
    "categoria": list(categorias_eda141),
    "meses_2023": meses_por_ano_eda141[2023].astype("Int64").to_numpy(), "acumulado_2023": acumulados_eda141[2023].astype("Int64").to_numpy(),
    "meses_2024": meses_por_ano_eda141[2024].astype("Int64").to_numpy(), "acumulado_2024": acumulados_eda141[2024].astype("Int64").to_numpy()})
comparacao_anual_eda141["diferenca_absoluta"] = comparacao_anual_eda141.acumulado_2024 - comparacao_anual_eda141.acumulado_2023
comparacao_anual_eda141["variacao_percentual"] = pd_eda141.Series(
    [round((float(d) / float(a)) * 100, 2) if a > 0 else pd_eda141.NA
     for d, a in zip(comparacao_anual_eda141.diferenca_absoluta, comparacao_anual_eda141.acumulado_2023)], dtype="Float64")
comparacao_anual_eda141["observacao_variacao"] = [
    "não aplicável (2023 = 0)" if a == 0 else "calculada sobre 2023" for a in comparacao_anual_eda141.acumulado_2023]
comparacao_anual_eda141["unidade"] = unidade_eda141
print("Comparação anual por categoria (somas de 12 meses de cada ano; sem total entre categorias):")
mostrar_eda141(comparacao_anual_eda141)

# %% [markdown]
# ## 14.1.5 — Análise mensal
# Tabela longa (ano, mês, categoria, fatos comunicados) e visão comparativa mês a mês
# de 2023 e 2024 por categoria. Os extremos preservam todos os meses empatados.

# %%
ordem_categoria_eda141 = {c: i for i, c in enumerate(categorias_eda141)}
mensal_longo_eda141 = base_eda141.assign(_ordem=base_eda141.categoria.map(ordem_categoria_eda141)).sort_values(
    ["_ordem", "ano", "mes"], kind="stable")[["ano", "mes", "categoria", "fatos_comunicados"]].reset_index(drop=True)
print("Tabela mensal (ano, mês, categoria, fatos comunicados) — 216 linhas:")
mostrar_eda141(mensal_longo_eda141)
comparativo_mensal_eda141 = base_eda141.pivot(index=["categoria", "mes"], columns="ano", values="fatos_comunicados")
comparativo_mensal_eda141.columns = [f"fatos_{ano}" for ano in comparativo_mensal_eda141.columns]
comparativo_mensal_eda141 = comparativo_mensal_eda141.reset_index()
comparativo_mensal_eda141 = comparativo_mensal_eda141.assign(_ordem=comparativo_mensal_eda141.categoria.map(ordem_categoria_eda141)).sort_values(
    ["_ordem", "mes"], kind="stable").drop(columns="_ordem").reset_index(drop=True)
comparativo_mensal_eda141["diferenca_2024_menos_2023"] = comparativo_mensal_eda141.fatos_2024 - comparativo_mensal_eda141.fatos_2023
print("Comparativo mensal 2023 × 2024 por categoria (valores publicados; sem causalidade):")
mostrar_eda141(comparativo_mensal_eda141)


def extremos_eda141(dados, rotulo):
    linhas = []
    for categoria in categorias_eda141:
        parte = dados.loc[dados.categoria.eq(categoria)]
        maximo, minimo = parte.fatos_comunicados.max(), parte.fatos_comunicados.min()
        meses_max = parte.loc[parte.fatos_comunicados.eq(maximo), "competencia"].tolist()
        meses_min = parte.loc[parte.fatos_comunicados.eq(minimo), "competencia"].tolist()
        linhas.append({"categoria": categoria, "periodo": rotulo, "valor_maximo": int(maximo), "competencias_do_maximo": ", ".join(meses_max),
                       "empate_no_maximo": len(meses_max) > 1, "valor_minimo": int(minimo), "competencias_do_minimo": ", ".join(meses_min),
                       "empate_no_minimo": len(meses_min) > 1})
    return pd_eda141.DataFrame(linhas)


extremos_eda141_tabela = pd_eda141.concat([
    extremos_eda141(base_eda141.loc[base_eda141.ano.eq(2023)], "2023"),
    extremos_eda141(base_eda141.loc[base_eda141.ano.eq(2024)], "2024"),
    extremos_eda141(base_eda141, "24 meses")], ignore_index=True)
print("Maiores e menores valores mensais por categoria (competência = aaaa-mm; empates preservados, todos os meses listados):")
mostrar_eda141(extremos_eda141_tabela)

# %% [markdown]
# ## 14.1.6 — Validações da Etapa 14.1
# Recalcula tudo por outro caminho (listas e o módulo `statistics`, a partir do texto
# publicado na rastreabilidade) e confere preservação e integridade. Se qualquer
# validação falhar, a execução é interrompida e nenhuma atividade seguinte (README,
# testes, commit) deve ser realizada.

# %%
referencia_eda141 = {
    (int(l.ano), int(l.mes), l.categoria): int(l.valor_textual_original)
    for l in retrato_rastreio_eda141.itertuples()}
verificacoes_eda141 = []


def verificar_eda141(nome, esperado, encontrado):
    verificacoes_eda141.append({"verificacao": nome, "esperado": esperado, "encontrado": encontrado, "ok": esperado == encontrado})


def proximo_eda141(a, b, tolerancia=1e-9):
    return pd_eda141.isna(a) and pd_eda141.isna(b) or (not pd_eda141.isna(a) and not pd_eda141.isna(b) and matematica_eda141.isclose(float(a), float(b), abs_tol=tolerancia))


def serie_eda141(categoria, ano=None):
    return [v for (a, m, c), v in sorted(referencia_eda141.items()) if c == categoria and (ano is None or a == ano)]


# Preservação das 216 observações e da rastreabilidade
verificar_eda141("216 observações preservadas (base = retrato inicial)", True,
                 bool(base_eda141[list(retrato_base_eda141.columns)].equals(retrato_base_eda141)) and len(base_eda141) == 216)
verificar_eda141("chaves da base = chaves da rastreabilidade (216)", set((int(l.ano), int(l.mes), l.categoria) for l in retrato_base_eda141.itertuples()), set(referencia_eda141))
verificar_eda141("valores da base = texto original da rastreabilidade convertido (216)", True,
                 all(referencia_eda141[(int(l.ano), int(l.mes), l.categoria)] == int(l.fatos_comunicados) for l in retrato_base_eda141.itertuples()))
verificar_eda141("rastreabilidade com 216 linhas e 23 colunas", (216, 23), retrato_rastreio_eda141.shape)
verificar_eda141("unidade Fato Comunicado na rastreabilidade", [unidade_eda141], sorted(retrato_rastreio_eda141.unidade.unique()))
verificar_eda141("nove categorias, sem agregação", sorted(categorias_eda141), sorted(base_eda141.categoria.unique()))
verificar_eda141("tabelas por categoria com exatamente nove linhas (sem linha de total)", (9, 9, 9, 9),
                 (len(estatisticas_2023_eda141), len(estatisticas_2024_eda141), len(estatisticas_24_meses_eda141), len(comparacao_anual_eda141)))
verificar_eda141("12 competências por categoria em 2023 e em 2024", True,
                 all(len(serie_eda141(c, a)) == 12 and sorted(m for (aa, m, cc) in referencia_eda141 if cc == c and aa == a) == list(range(1, 13))
                     for c in categorias_eda141 for a in anos_eda141))
verificar_eda141("ausências na base", 0, int(base_eda141[list(retrato_base_eda141.columns)].isna().sum().sum()))
verificar_eda141("27 zeros publicados identificáveis (inspeção e base)", (27, 27),
                 (len(zeros_eda141), sum(1 for v in referencia_eda141.values() if v == 0)))
verificar_eda141("zeros: soma de meses_com_zero de 2023 e 2024 = 27", 27,
                 int(estatisticas_2023_eda141.meses_com_zero.sum() + estatisticas_2024_eda141.meses_com_zero.sum()))
verificar_eda141("zeros dos 24 meses = soma dos zeros anuais", True,
                 bool((estatisticas_24_meses_eda141.meses_com_zero == estatisticas_2023_eda141.meses_com_zero + estatisticas_2024_eda141.meses_com_zero).all()))
verificar_eda141("nenhuma chave duplicada", 0, int(base_eda141.duplicated(chaves_eda141).sum()))
# Estatísticas recalculadas com listas Python e o módulo statistics
vistas_eda141 = {"2023": (estatisticas_2023_eda141, 2023), "2024": (estatisticas_2024_eda141, 2024), "24 meses": (estatisticas_24_meses_eda141, None)}
for nome_vista_eda141, (tabela_eda141, ano_eda141) in vistas_eda141.items():
    resultados_eda141 = []
    for categoria_eda141 in categorias_eda141:
        valores_eda141 = serie_eda141(categoria_eda141, ano_eda141)
        linha_eda141 = tabela_eda141.loc[categoria_eda141]
        resultados_eda141.append(all([
            int(linha_eda141.soma_publicada) == sum(valores_eda141),
            proximo_eda141(linha_eda141.media_mensal, sum(valores_eda141) / len(valores_eda141), 5e-5),
            proximo_eda141(linha_eda141.mediana_mensal, estatistica_eda141.median(valores_eda141), 5e-5),
            int(linha_eda141.minimo) == min(valores_eda141), int(linha_eda141.maximo) == max(valores_eda141),
            proximo_eda141(linha_eda141.desvio_padrao_ddof0, estatistica_eda141.pstdev(valores_eda141), 5e-5),
            int(linha_eda141.meses_observados) == len(valores_eda141),
            int(linha_eda141.meses_com_zero) == sum(1 for v in valores_eda141 if v == 0)]))
    verificar_eda141(f"soma, média, mediana, mínimo, máximo, desvio (ddof=0), meses e zeros corretos — {nome_vista_eda141} (9 categorias)", 9, sum(resultados_eda141))
verificar_eda141("meses observados: 12, 12 e 24 em todas as categorias", (True, True, True),
                 (bool(estatisticas_2023_eda141.meses_observados.eq(12).all()), bool(estatisticas_2024_eda141.meses_observados.eq(12).all()),
                  bool(estatisticas_24_meses_eda141.meses_observados.eq(24).all())))
verificar_eda141("somatórios anuais = soma dos 12 valores mensais (comparação anual)", True,
                 all(int(l.acumulado_2023) == sum(serie_eda141(l.categoria, 2023)) and int(l.acumulado_2024) == sum(serie_eda141(l.categoria, 2024))
                     for l in comparacao_anual_eda141.itertuples()))
verificar_eda141("comparação anual: 12 meses em cada ano, sem misturar períodos", True,
                 bool(comparacao_anual_eda141.meses_2023.eq(12).all() and comparacao_anual_eda141.meses_2024.eq(12).all()))
verificar_eda141("comparação anual: acumulado_2023 + acumulado_2024 = soma_publicada dos 24 meses", True,
                 bool((comparacao_anual_eda141.set_index("categoria").acumulado_2023 + comparacao_anual_eda141.set_index("categoria").acumulado_2024
                       == estatisticas_24_meses_eda141.soma_publicada).all()))
verificar_eda141("diferença absoluta = 2024 − 2023", True,
                 all(int(l.diferenca_absoluta) == int(l.acumulado_2024) - int(l.acumulado_2023) for l in comparacao_anual_eda141.itertuples()))
verificar_eda141("variação percentual: <NA> se e somente se 2023 = 0; senão (2024−2023)/2023×100", True,
                 all((pd_eda141.isna(l.variacao_percentual) if l.acumulado_2023 == 0
                      else (not pd_eda141.isna(l.variacao_percentual) and matematica_eda141.isclose(float(l.variacao_percentual), (int(l.diferenca_absoluta) / int(l.acumulado_2023)) * 100, abs_tol=0.0051)))
                     for l in comparacao_anual_eda141.itertuples()))
verificar_eda141("tabela mensal longa com 216 linhas e valores iguais à referência", True,
                 len(mensal_longo_eda141) == 216 and all(referencia_eda141[(int(l.ano), int(l.mes), l.categoria)] == int(l.fatos_comunicados) for l in mensal_longo_eda141.itertuples()))
verificar_eda141("comparativo mensal: 108 linhas e valores de 2023 e 2024 corretos", True,
                 len(comparativo_mensal_eda141) == 108 and all(
                     referencia_eda141[(2023, int(l.mes), l.categoria)] == int(l.fatos_2023) and referencia_eda141[(2024, int(l.mes), l.categoria)] == int(l.fatos_2024)
                     for l in comparativo_mensal_eda141.itertuples()))
extremos_corretos_eda141 = []
for l in extremos_eda141_tabela.itertuples():
    ano_extremo_eda141 = None if l.periodo == "24 meses" else int(l.periodo)
    itens_eda141 = {f"{a}-{m:02d}": v for (a, m, c), v in referencia_eda141.items() if c == l.categoria and (ano_extremo_eda141 is None or a == ano_extremo_eda141)}
    maior_eda141, menor_eda141 = max(itens_eda141.values()), min(itens_eda141.values())
    extremos_corretos_eda141.append(
        l.valor_maximo == maior_eda141 and l.valor_minimo == menor_eda141
        and sorted(l.competencias_do_maximo.split(", ")) == sorted(k for k, v in itens_eda141.items() if v == maior_eda141)
        and sorted(l.competencias_do_minimo.split(", ")) == sorted(k for k, v in itens_eda141.items() if v == menor_eda141))
verificar_eda141("extremos mensais corretos, com todos os empates preservados (27 linhas)", 27, sum(extremos_corretos_eda141))
# Imutabilidade dos DataFrames de origem
verificar_eda141("df_violencia_domestica_joinville_ssp132 inalterado (comparação exata com o retrato inicial)", True,
                 bool(df_violencia_domestica_joinville_ssp132.equals(retrato_base_eda141)) and df_violencia_domestica_joinville_ssp132.dtypes.equals(retrato_base_eda141.dtypes))
verificar_eda141("rastreabilidade_violencia_domestica_ssp132 inalterada (comparação exata com o retrato inicial)", True,
                 bool(rastreabilidade_violencia_domestica_ssp132.equals(retrato_rastreio_eda141)) and rastreabilidade_violencia_domestica_ssp132.dtypes.equals(retrato_rastreio_eda141.dtypes))
verificar_eda141("colunas da base de origem sem colunas auxiliares", ["municipio", "ano", "mes", "categoria", "fatos_comunicados"], df_violencia_domestica_joinville_ssp132.columns.tolist())
verificacoes_quadro_eda141 = pd_eda141.DataFrame(verificacoes_eda141)
print("Quadro das verificações da Etapa 14.1:")
mostrar_eda141(verificacoes_quadro_eda141)
if not verificacoes_quadro_eda141.ok.all():
    mostrar_eda141(verificacoes_quadro_eda141.loc[~verificacoes_quadro_eda141.ok])
    raise ValueError("Etapa 14.1: validação falhou; interromper as atividades seguintes e NÃO fazer commit.")
print(f"Etapa 14.1: {len(verificacoes_quadro_eda141)} verificações aprovadas; nenhum DataFrame de origem foi modificado; nada foi exportado.")

# %% [markdown]
# ## 14.1.7 — Leitura dos resultados (observado × cuidado metodológico)
# As tabelas abaixo só reordenam resultados já calculados, para leitura. Elas não
# somam categorias nem sugerem causas.

# %%
ranking_eda141 = comparacao_anual_eda141.set_index("categoria")[["acumulado_2023", "acumulado_2024"]].join(
    estatisticas_24_meses_eda141[["soma_publicada", "media_mensal"]]).sort_values("soma_publicada", ascending=False)
print("1) Categorias ordenadas pela contagem publicada nos 24 meses (apenas ordenação; sem total geral):")
mostrar_eda141(ranking_eda141)
variacoes_eda141 = comparacao_anual_eda141.set_index("categoria")[["acumulado_2023", "acumulado_2024", "diferenca_absoluta", "variacao_percentual"]]
print("2) Variações 2023 → 2024 em contagens publicadas, ordenadas pela diferença absoluta:")
mostrar_eda141(variacoes_eda141.sort_values("diferenca_absoluta", key=lambda s: s.abs(), ascending=False))
print("3) Meses com maior e menor valor por categoria nos 24 meses (empates preservados):")
mostrar_eda141(extremos_eda141_tabela.loc[extremos_eda141_tabela.periodo.eq("24 meses"), [
    "categoria", "valor_maximo", "competencias_do_maximo", "valor_minimo", "competencias_do_minimo"]].reset_index(drop=True))
print("Limites: (a) contagens publicadas de 'Fato Comunicado' não são vítimas nem incidência real; (b) variação entre anos não prova mudança de incidência "
      "nem tem causa identificada nestes dados; (c) categorias de poucos casos (Estupro, Injúria real, lesão grave ou gravíssima) variam muito com poucas unidades; "
      "(d) 2023 e 2024 vêm de relatórios diferentes, cuja definição de contagem não está documentada; (e) a unificação das grafias de lesão corporal é decisão do projeto.")

# %% [markdown]
# # ETAPA 14.2 — Visualizações da violência doméstica de Joinville (SSP/SC)
# Execute depois das Etapas 13.1, 13.2, 13.3 e 14.1 (não é preciso executar as Etapas
# 1–12 nem o Sinesp). As seis figuras Matplotlib usam apenas as tabelas validadas das
# Etapas 13.2 e 14.1 (`comparacao_anual_eda141` e `comparativo_mensal_eda141`);
# nenhum valor é recalculado para o gráfico e nada é exportado: as figuras ficam em
# memória. Elas são objetos `Figure` criados fora do `pyplot`, para não serem exibidas
# antes da validação; cada uma é exibida em sua própria célula com `display()`.
#
# Cuidados de leitura (repetidos nas figuras):
# - A unidade é "fatos comunicados" (contagem publicada pela SSP/SC), não vítimas.
# - Cada categoria é plotada separadamente; não há total geral nem soma de categorias.
# - Variações percentuais não são destacadas (categorias de poucos casos são instáveis).
# - Diferenças entre 2023 e 2024 são mudanças nas contagens publicadas, não
#   necessariamente na incidência real, e não têm causa identificada nestes dados.
# - O mapeamento das grafias de lesão corporal e a associação dos meses às páginas
#   seguintes das tabelas são decisões metodológicas do projeto.

# %% [markdown]
# ## 14.2.1 — Pré-condições, retratos de preservação e construção das figuras

# %%
import math as matematica_viz142
import textwrap as quebra_viz142

import pandas as pd_viz142
from IPython.display import display as mostrar_viz142
from matplotlib.figure import Figure as Figura_viz142
from matplotlib.ticker import FuncFormatter as Formatador_viz142
from matplotlib.ticker import MaxNLocator as LocalizadorInteiro_viz142

necessarios_viz142 = (
    "df_violencia_domestica_joinville_ssp132", "rastreabilidade_violencia_domestica_ssp132",
    "comparacao_anual_eda141", "comparativo_mensal_eda141", "verificacoes_quadro_eda141")
ausentes_viz142 = [nome for nome in necessarios_viz142 if nome not in globals()]
if ausentes_viz142:
    raise NameError(f"Etapa 14.2: execute antes as Etapas 13.1 a 14.1. Ausentes: {ausentes_viz142}")
if verificacoes_quadro_eda141.empty or not verificacoes_quadro_eda141.ok.all():
    raise ValueError("Etapa 14.2: as validações da Etapa 14.1 não estão integralmente aprovadas nesta sessão; interromper.")
print(f"Controle da Etapa 14.1 confirmado: {len(verificacoes_quadro_eda141)} verificações aprovadas.")

# Retratos para provar, ao final, que nada de origem foi alterado.
retrato_base_viz142 = df_violencia_domestica_joinville_ssp132.copy(deep=True)
retrato_rastreio_viz142 = rastreabilidade_violencia_domestica_ssp132.copy(deep=True)
retrato_anual_viz142 = comparacao_anual_eda141.copy(deep=True)
retrato_mensal_viz142 = comparativo_mensal_eda141.copy(deep=True)

# Referência independente: texto publicado na rastreabilidade, convertido para inteiro.
referencia_viz142 = {(int(l.ano), int(l.mes), l.categoria): int(l.valor_textual_original) for l in retrato_rastreio_viz142.itertuples()}
anual_ref_viz142 = {(a, c): sum(v for (aa, m, cc), v in referencia_viz142.items() if aa == a and cc == c)
                    for a in (2023, 2024) for c in {c for (_, _, c) in referencia_viz142}}

cores_viz142 = {2023: "#2563A6", 2024: "#D97706"}  # azul e laranja: distinguíveis, inclusive em daltonismo
meses_rotulos_viz142 = ["Jan", "Fev", "Mar", "Abr", "Mai", "Jun", "Jul", "Ago", "Set", "Out", "Nov", "Dez"]
fonte_viz142 = ("Fonte: SSP/SC — relatórios de violência doméstica. Joinville/SC, janeiro/2023 a dezembro/2024. "
                "Unidade: fatos comunicados (contagem publicada).")
nota_viz142 = ("Nota: contagens publicadas, não necessariamente incidência real; sem total entre categorias. "
               "Grafias de lesão corporal unificadas por decisão metodológica do projeto.")


def formatar_inteiro_viz142(valor):
    return f"{int(valor):,}".replace(",", ".")


def rodape_viz142(figura):
    # Duas faixas separadas no rodapé (fonte acima, nota abaixo), sem sobreposição.
    figura.text(0.01, 0.125, "\n".join(quebra_viz142.wrap(fonte_viz142, 125)), fontsize=8.5, color="#333333", ha="left", va="top")
    figura.text(0.01, 0.012, "\n".join(quebra_viz142.wrap(nota_viz142, 125)), fontsize=8.5, color="#333333", ha="left", va="bottom")


def figura_anual_viz142(categorias, titulo):
    """Barras agrupadas 2023 × 2024 (soma dos 12 meses), um eixo único iniciado em zero."""
    dados = comparacao_anual_eda141.set_index("categoria").loc[list(categorias), ["acumulado_2023", "acumulado_2024"]].astype(int)
    figura = Figura_viz142(figsize=(10, 6.6), dpi=120)
    eixo = figura.add_axes([0.10, 0.27, 0.87, 0.53])
    posicoes = list(range(len(categorias)))
    largura = 0.38
    barras = {}
    for deslocamento, ano in ((-largura / 2, 2023), (largura / 2, 2024)):
        barras[ano] = eixo.bar([p + deslocamento for p in posicoes], dados[f"acumulado_{ano}"].tolist(), width=largura,
                               color=cores_viz142[ano], label=str(ano), edgecolor="white")
        eixo.bar_label(barras[ano], labels=[formatar_inteiro_viz142(v) for v in dados[f"acumulado_{ano}"]], padding=3, fontsize=11)
    eixo.set_ylim(0, matematica_viz142.ceil(dados.to_numpy().max() * 1.15))
    eixo.yaxis.set_major_locator(LocalizadorInteiro_viz142(integer=True))
    eixo.yaxis.set_major_formatter(Formatador_viz142(lambda valor, _: formatar_inteiro_viz142(valor)))
    eixo.set_xticks(posicoes)
    eixo.set_xticklabels(["\n".join(quebra_viz142.wrap(c, 20)) for c in categorias], fontsize=11)
    eixo.set_ylabel("Fatos comunicados (soma dos 12 meses do ano)", fontsize=11)
    eixo.set_xlabel("Categoria", fontsize=11)
    eixo.set_title("Soma dos 12 meses de cada ano, por categoria (mesmo eixo vertical, iniciado em zero)", fontsize=10.5, color="#333333")
    eixo.grid(axis="y", color="#DDDDDD", linewidth=0.8)
    eixo.set_axisbelow(True)
    eixo.spines[["top", "right"]].set_visible(False)
    eixo.legend(title="Ano", fontsize=11, title_fontsize=11, loc="upper right", frameon=False)
    figura.suptitle("\n".join(quebra_viz142.wrap(titulo, 75)), fontsize=13, fontweight="bold", x=0.01, ha="left", y=0.985)
    rodape_viz142(figura)
    return figura


def figura_mensal_viz142(categoria):
    """Linhas de 2023 e 2024 por mês para uma única categoria, eixo vertical iniciado em zero."""
    dados = comparativo_mensal_eda141.loc[comparativo_mensal_eda141.categoria.eq(categoria)].sort_values("mes")
    figura = Figura_viz142(figsize=(10, 6.6), dpi=120)
    eixo = figura.add_axes([0.10, 0.27, 0.87, 0.53])
    for ano in (2023, 2024):
        eixo.plot(dados.mes.astype(int).tolist(), dados[f"fatos_{ano}"].astype(int).tolist(), marker="o", linewidth=2.2, markersize=6,
                  color=cores_viz142[ano], label=str(ano))
    eixo.set_xticks(range(1, 13))
    eixo.set_xticklabels(meses_rotulos_viz142, fontsize=11)
    eixo.set_ylim(0, matematica_viz142.ceil(max(dados.fatos_2023.astype(int).max(), dados.fatos_2024.astype(int).max()) * 1.12))
    eixo.yaxis.set_major_locator(LocalizadorInteiro_viz142(integer=True))
    eixo.yaxis.set_major_formatter(Formatador_viz142(lambda valor, _: formatar_inteiro_viz142(valor)))
    eixo.set_ylabel("Fatos comunicados por mês", fontsize=11)
    eixo.set_xlabel("Mês", fontsize=11)
    eixo.set_title(f"Categoria: {categoria} — 2023 × 2024, janeiro a dezembro (eixo vertical iniciado em zero)", fontsize=10.5, color="#333333")
    eixo.grid(axis="y", color="#DDDDDD", linewidth=0.8)
    eixo.set_axisbelow(True)
    eixo.spines[["top", "right"]].set_visible(False)
    eixo.legend(title="Ano", fontsize=11, title_fontsize=11, loc="lower right", frameon=False)
    figura.suptitle("\n".join(quebra_viz142.wrap(f"Violência doméstica em Joinville/SC — {categoria}: fatos comunicados por mês", 75)),
                    fontsize=13, fontweight="bold", x=0.01, ha="left", y=0.985)
    rodape_viz142(figura)
    return figura


especificacoes_viz142 = {
    "fig_anual_maiores_ssp142": ("anual", ("Ameaça", "Lesão corporal leve dolosa", "Injúria"),
                                 "Violência doméstica em Joinville/SC — categorias de maiores contagens publicadas: 2023 × 2024"),
    "fig_anual_intermediarias_ssp142": ("anual", ("Vias de fato", "Difamação", "Calúnia"),
                                        "Violência doméstica em Joinville/SC — categorias de contagens intermediárias: 2023 × 2024"),
    "fig_anual_menores_ssp142": ("anual", ("Estupro", "Lesão corporal grave ou gravíssima dolosa", "Injúria real"),
                                 "Violência doméstica em Joinville/SC — categorias de menores contagens publicadas: 2023 × 2024"),
    "fig_mensal_ameaca_ssp142": ("mensal", ("Ameaça",), None),
    "fig_mensal_lesao_leve_ssp142": ("mensal", ("Lesão corporal leve dolosa",), None),
    "fig_mensal_injuria_ssp142": ("mensal", ("Injúria",), None),
}
figuras_viz142 = {}
for nome_viz142, (tipo_viz142, categorias_viz142, titulo_viz142) in especificacoes_viz142.items():
    figuras_viz142[nome_viz142] = (figura_anual_viz142(categorias_viz142, titulo_viz142) if tipo_viz142 == "anual"
                                   else figura_mensal_viz142(categorias_viz142[0]))
print("Seis figuras construídas em memória (ainda não exibidas):", ", ".join(figuras_viz142))

# %% [markdown]
# ## 14.2.2 — Validação das figuras antes de exibi-las
# Os valores são lidos dos próprios elementos gráficos (alturas das barras, rótulos
# sobre as barras e coordenadas das linhas) e comparados com a tabela da Etapa 14.1 e
# com a soma/valor obtidos do texto publicado na rastreabilidade. Também são conferidos
# títulos, legendas, eixos, fonte e ausência de total geral. Se algo divergir, a
# execução é interrompida e as figuras não devem ser aprovadas.

# %%
linhas_valores_viz142 = []
linhas_conferencia_viz142 = []


def registrar_viz142(figura, verificacao, esperado, encontrado):
    linhas_conferencia_viz142.append({"figura": figura, "verificacao": verificacao, "esperado": esperado, "encontrado": encontrado,
                                      "ok": esperado == encontrado})


def textos_da_figura_viz142(figura):
    eixo = figura.axes[0]
    textos = [t.get_text() for t in figura.texts] + [figura.get_suptitle(), eixo.get_title(), eixo.get_xlabel(), eixo.get_ylabel()]
    legenda = eixo.get_legend()
    if legenda is not None:
        textos += [legenda.get_title().get_text()] + [t.get_text() for t in legenda.get_texts()]
    return [t for t in textos if t]


for nome_viz142, (tipo_viz142, categorias_viz142, _) in especificacoes_viz142.items():
    figura_viz142 = figuras_viz142[nome_viz142]
    eixo_viz142 = figura_viz142.axes[0]
    textos_viz142 = " | ".join(textos_da_figura_viz142(figura_viz142))
    registrar_viz142(nome_viz142, "um único eixo (mesma escala vertical para 2023 e 2024)", 1, len(figura_viz142.axes))
    registrar_viz142(nome_viz142, "eixo vertical começa em zero e comporta o maior valor", True,
                     eixo_viz142.get_ylim()[0] == 0)
    registrar_viz142(nome_viz142, "título, eixos, fonte SSP/SC, Joinville/SC, período e unidade presentes", True,
                     all(trecho in textos_viz142 for trecho in ("Joinville/SC", "SSP/SC", "janeiro/2023 a dezembro/2024", "fatos comunicados", "Fatos comunicados", "2023", "2024")))
    registrar_viz142(nome_viz142, "sem a palavra 'vítima' e sem 'total' nos textos", True, "vítima" not in textos_viz142.lower() and "total geral" not in textos_viz142.lower())
    registrar_viz142(nome_viz142, "legenda com 2023 e 2024", ["2023", "2024"], [t.get_text() for t in eixo_viz142.get_legend().get_texts()])
    registrar_viz142(nome_viz142, "cores de 2023 e 2024 diferentes", 2, len({cores_viz142[2023], cores_viz142[2024]}))
    if tipo_viz142 == "anual":
        rotulos_eixo_viz142 = [t.get_text().replace("\n", " ") for t in eixo_viz142.get_xticklabels()]
        registrar_viz142(nome_viz142, "categorias no eixo horizontal (ordem da especificação)", list(categorias_viz142), rotulos_eixo_viz142)
        registrar_viz142(nome_viz142, "duas séries de barras (2023 e 2024) com uma barra por categoria", (3, 3), (len(eixo_viz142.containers[0]), len(eixo_viz142.containers[1])))
        registrar_viz142(nome_viz142, "total de barras = 2 × categorias", 6, len(eixo_viz142.patches))
        topo_viz142 = eixo_viz142.get_ylim()[1]
        for contador_viz142, ano_viz142 in enumerate((2023, 2024)):
            contêiner_viz142 = eixo_viz142.containers[contador_viz142]
            rotulos_valor_viz142 = [t.get_text() for t in eixo_viz142.texts][contador_viz142 * 3:(contador_viz142 + 1) * 3]
            for posicao_viz142, categoria_viz142 in enumerate(categorias_viz142):
                plotado_viz142 = contêiner_viz142[posicao_viz142].get_height()
                tabela_viz142 = int(comparacao_anual_eda141.set_index("categoria").loc[categoria_viz142, f"acumulado_{ano_viz142}"])
                independente_viz142 = anual_ref_viz142[(ano_viz142, categoria_viz142)]
                linhas_valores_viz142.append({
                    "figura": nome_viz142, "categoria": categoria_viz142, "ano": ano_viz142, "mes": pd_viz142.NA, "valor_plotado": int(plotado_viz142),
                    "rotulo_sobre_a_barra": rotulos_valor_viz142[posicao_viz142], "valor_tabela_14_1": tabela_viz142,
                    "valor_rastreabilidade_soma_12_meses": independente_viz142,
                    "conferem": int(plotado_viz142) == tabela_viz142 == independente_viz142 == int(rotulos_valor_viz142[posicao_viz142].replace(".", "")) and plotado_viz142 == int(plotado_viz142)})
        registrar_viz142(nome_viz142, "valor máximo cabe no eixo (topo > maior barra)", True, topo_viz142 > max(p.get_height() for p in eixo_viz142.patches))
    else:
        categoria_viz142 = categorias_viz142[0]
        linhas_eixo_viz142 = eixo_viz142.lines
        registrar_viz142(nome_viz142, "duas linhas (2023 e 2024), cada uma com 12 pontos", [(12, 12)], [(len(linhas_eixo_viz142[0].get_ydata()), len(linhas_eixo_viz142[1].get_ydata()))])
        registrar_viz142(nome_viz142, "eixo horizontal de janeiro a dezembro", meses_rotulos_viz142, [t.get_text() for t in eixo_viz142.get_xticklabels()])
        registrar_viz142(nome_viz142, "meses plotados 1 a 12 em ambas as linhas", True,
                         all(list(l.get_xdata()) == list(range(1, 13)) for l in linhas_eixo_viz142))
        registrar_viz142(nome_viz142, "uma única categoria no gráfico", [categoria_viz142], [categoria_viz142])
        registrar_viz142(nome_viz142, "título menciona a categoria", True, categoria_viz142 in figura_viz142.get_suptitle() and categoria_viz142 in eixo_viz142.get_title())
        maior_viz142 = max(max(l.get_ydata()) for l in linhas_eixo_viz142)
        registrar_viz142(nome_viz142, "valor máximo cabe no eixo (topo > maior valor)", True, eixo_viz142.get_ylim()[1] > maior_viz142)
        for linha_viz142, ano_viz142 in zip(linhas_eixo_viz142, (2023, 2024)):
            registrar_viz142(nome_viz142, f"rótulo da linha de {ano_viz142}", str(ano_viz142), linha_viz142.get_label())
            for mes_viz142, plotado_viz142 in zip(linha_viz142.get_xdata(), linha_viz142.get_ydata()):
                tabela_viz142 = int(comparativo_mensal_eda141.loc[comparativo_mensal_eda141.categoria.eq(categoria_viz142) & comparativo_mensal_eda141.mes.eq(int(mes_viz142)), f"fatos_{ano_viz142}"].iloc[0])
                independente_viz142 = referencia_viz142[(ano_viz142, int(mes_viz142), categoria_viz142)]
                linhas_valores_viz142.append({
                    "figura": nome_viz142, "categoria": categoria_viz142, "ano": ano_viz142, "mes": int(mes_viz142), "valor_plotado": int(plotado_viz142),
                    "rotulo_sobre_a_barra": pd_viz142.NA, "valor_tabela_14_1": tabela_viz142, "valor_rastreabilidade_soma_12_meses": independente_viz142,
                    "conferem": int(plotado_viz142) == tabela_viz142 == independente_viz142 and plotado_viz142 == int(plotado_viz142)})
valores_plotados_viz142 = pd_viz142.DataFrame(linhas_valores_viz142)
valores_plotados_viz142 = valores_plotados_viz142.rename(columns={"valor_rastreabilidade_soma_12_meses": "valor_rastreabilidade (soma de 12 meses nas figuras anuais)"})
# Cada categoria mantém os próprios valores: nenhum valor aparece em mais de uma categoria por engano e não há categoria "Total".
categorias_plotadas_viz142 = set(valores_plotados_viz142.categoria)
registrar_viz142("todas", "nove categorias plotadas, cada uma com os valores próprios", 9, len(categorias_plotadas_viz142 & {c for (_, _, c) in referencia_viz142}))
registrar_viz142("todas", "nenhuma categoria agregada ou 'Total' plotada", True, all("total" not in c.lower() for c in categorias_plotadas_viz142))
registrar_viz142("todas", "valores plotados: 18 anuais + 72 mensais, todos conferidos", (18, 72, 90), (
    int(valores_plotados_viz142.mes.isna().sum()), int(valores_plotados_viz142.mes.notna().sum()), int(valores_plotados_viz142.conferem.sum())))
registrar_viz142("todas", "12 meses de cada ano em cada série mensal (3 figuras × 2 anos × 12)", 72, int(valores_plotados_viz142.mes.notna().sum()))
registrar_viz142("todas", "base analítica inalterada", True, bool(df_violencia_domestica_joinville_ssp132.equals(retrato_base_viz142)) and df_violencia_domestica_joinville_ssp132.dtypes.equals(retrato_base_viz142.dtypes))
registrar_viz142("todas", "rastreabilidade inalterada", True, bool(rastreabilidade_violencia_domestica_ssp132.equals(retrato_rastreio_viz142)))
registrar_viz142("todas", "tabelas da Etapa 14.1 inalteradas", True, bool(comparacao_anual_eda141.equals(retrato_anual_viz142)) and bool(comparativo_mensal_eda141.equals(retrato_mensal_viz142)))
conferencia_visualizacoes_viz142 = pd_viz142.DataFrame(linhas_conferencia_viz142)
print("Valores lidos dos elementos gráficos e comparados com a Etapa 14.1 e com a rastreabilidade (90 valores):")
mostrar_viz142(valores_plotados_viz142)
print("Conferência das visualizações:")
mostrar_viz142(conferencia_visualizacoes_viz142)
if not conferencia_visualizacoes_viz142.ok.all() or not valores_plotados_viz142.conferem.all():
    mostrar_viz142(conferencia_visualizacoes_viz142.loc[~conferencia_visualizacoes_viz142.ok])
    mostrar_viz142(valores_plotados_viz142.loc[~valores_plotados_viz142.conferem])
    raise ValueError("Etapa 14.2: conferência das figuras falhou; não aprovar nem integrar as visualizações.")
print(f"Etapa 14.2: {len(conferencia_visualizacoes_viz142)} verificações e {len(valores_plotados_viz142)} valores plotados conferidos; "
      "dados de origem inalterados; nenhuma imagem exportada.")

# %% [markdown]
# ## 14.2.3 — Figura 1: categorias de maiores contagens (Ameaça, Lesão corporal leve dolosa, Injúria)

# %%
mostrar_viz142(figuras_viz142["fig_anual_maiores_ssp142"])

# %% [markdown]
# ## 14.2.4 — Figura 2: categorias de contagens intermediárias (Vias de fato, Difamação, Calúnia)

# %%
mostrar_viz142(figuras_viz142["fig_anual_intermediarias_ssp142"])

# %% [markdown]
# ## 14.2.5 — Figura 3: categorias de menores contagens (Estupro, Lesão corporal grave ou gravíssima dolosa, Injúria real)

# %%
mostrar_viz142(figuras_viz142["fig_anual_menores_ssp142"])

# %% [markdown]
# ## 14.2.6 — Figura 4: Ameaça, mês a mês, 2023 × 2024

# %%
mostrar_viz142(figuras_viz142["fig_mensal_ameaca_ssp142"])

# %% [markdown]
# ## 14.2.7 — Figura 5: Lesão corporal leve dolosa, mês a mês, 2023 × 2024

# %%
mostrar_viz142(figuras_viz142["fig_mensal_lesao_leve_ssp142"])

# %% [markdown]
# ## 14.2.8 — Figura 6: Injúria, mês a mês, 2023 × 2024

# %%
mostrar_viz142(figuras_viz142["fig_mensal_injuria_ssp142"])
