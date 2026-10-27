import csv
import hashlib
from pathlib import Path

from pypdf import PdfReader
import pytest

from scripts.utils import project_path


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def test_pre_migration_backup_matches_inventory():
    backup_root = project_path("backup_pre_migracao_20261016")
    if not backup_root.is_dir():
        pytest.skip("Cópia local pré-migração está fora de clones e ignorada pelo Git")
    inventory = project_path("documentacao", "auditoria", "inventario_pre_migracao.csv")
    with inventory.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 82
    for row in rows:
        backup = backup_root.joinpath(*Path(row["caminho_relativo"]).parts)
        assert backup.is_file()
        assert backup.stat().st_size == int(row["tamanho_bytes"])
        assert sha256(backup) == row["sha256_original"] == row["sha256_backup"]


def test_expected_source_files_are_in_new_layout():
    raw = project_path("dados", "brutos")
    assert len(list((raw / "ssp_sc" / "pdfs").glob("*.pdf"))) == 45
    assert len(list((raw / "ssp_sc" / "csv_extraidos").glob("*.csv"))) == 45
    assert len(list((raw / "sinesp").glob("bancovde-*.xlsx"))) == 4
    assert len(list((raw / "sinesp" / "pendente_validacao_origem").glob("*.xlsx"))) == 3


def test_historical_sinesp_sources_have_official_provenance_evidence():
    audit = project_path("documentacao", "auditoria", "sinesp_origem_2020_2022.csv")
    with audit.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert {int(row["ano"]) for row in rows} == {2020, 2021, 2022}
    for row in rows:
        source = project_path(*Path(row["caminho_local"]).parts)
        assert row["status"] == "VALIDADO"
        assert row["colunas"] == "14"
        assert row["url_download_oficial"].startswith("https://www.gov.br/mj/")
        assert source.is_file()
        assert source.stat().st_size == int(row["bytes_local"])
        assert sha256(source) == row["sha256_local"]


def test_raw_sources_still_match_the_migration_integrity_map():
    mapping = project_path("documentacao", "auditoria", "inventario_sha256_pre_reorganizacao.csv")
    with mapping.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 158
    data_matches = 0
    output_matches = 0
    archived_readmes = {
        "documentacao/analises/graficos/README.md",
        "documentacao/analises/tabelas/README.md",
    }
    archive_root = project_path("backup_pre_migracao_20261016", "ajustes_academicos_finais_20261016", "readmes_arquivados")
    archived_matches = 0
    # Documentos realocados na Etapa 21 (reorganização da documentação); a existência continua obrigatória no novo caminho.
    realocados_etapa21 = {
        "documentacao/ESCOPO_PROJETO.md": "documentacao/analises/ESCOPO_PROJETO.md",
        "documentacao/fontes/ORIGEM_SINESP.md": "documentacao/analises/fontes/ORIGEM_SINESP.md",
        "documentacao/fontes/SINESP_JOINVILLE.md": "documentacao/analises/fontes/SINESP_JOINVILLE.md",
        "documentacao/fontes/SSP_SC.md": "historico/ssp_sc/documentacao/SSP_SC.md",
        "documentacao/fontes/CSV_COMPLEMENTARES.md": "historico/ssp_sc/documentacao/CSV_COMPLEMENTARES.md",
        "documentacao/auditoria/ssp_csv_extraction_etapa_1_1.md": "historico/ssp_sc/documentacao/ssp_csv_extraction_etapa_1_1.md",
        "documentacao/analises/relatorio_pipeline_analitico.md": "historico/documentacao_anterior/relatorio_pipeline_analitico.md",
        **{
            f"documentacao/auditoria/{nome}": f"historico/documentacao_anterior/{nome}"
            for nome in (
                "inspecao_etapa1_leitura_bruta.md", "levantamento_2026.md", "MIGRACAO_E_AUDITORIA.md",
                "mapa_migracao.csv", "sinesp_python_js_comparison.csv",
            )
        },
    }
    for row in rows:
        destination = row["caminho_destino_planejado"]
        source = project_path(*Path(realocados_etapa21.get(destination, destination)).parts)
        if destination in archived_readmes:
            assert not source.exists()
            if archive_root.is_dir():
                archived = archive_root.joinpath(*Path(destination).parts)
                assert archived.is_file()
                assert archived.stat().st_size == int(row["tamanho_bytes"])
                assert sha256(archived) == row["sha256"]
                archived_matches += 1
            continue
        assert source.is_file()
        current_hash = sha256(source)
        if row["caminho_original"].startswith("data/"):
            assert source.stat().st_size == int(row["tamanho_bytes"])
            assert current_hash == row["sha256"]
            data_matches += 1
        elif row["caminho_original"].startswith("outputs/") and not row["caminho_original"].endswith(("manifesto_analise.json", "validacao_csv_parquet.csv")):
            assert source.stat().st_size == int(row["tamanho_bytes"])
            assert current_hash == row["sha256"]
            output_matches += 1
    assert data_matches == 125
    assert output_matches == 14
    assert archived_matches == (2 if archive_root.is_dir() else 0)


def test_ssp_outputs_keep_traceability_and_valid_pages():
    raw_pdfs = {path.name: path for path in project_path("dados", "brutos", "ssp_sc", "pdfs").glob("*.pdf")}
    csv_files = sorted(project_path("dados", "processados", "ssp_sc").glob("ssp_joinville_*.csv"))
    assert len(csv_files) == 4
    page_counts = {}
    expected_rows = {2023: 185, 2024: 192, 2025: 84, 2026: 49}
    for csv_path in csv_files:
        year = int(csv_path.stem[-4:])
        with csv_path.open(encoding="utf-8-sig", newline="") as stream:
            reader = csv.DictReader(stream)
            rows = list(reader)
        assert len(rows) == expected_rows[year]
        assert {"municipio", "indicador", "indicador_normalizado", "periodo_referencia", "arquivo_pdf", "pagina_pdf"}.issubset(reader.fieldnames)
        for row in rows:
            assert row["municipio"] == "JOINVILLE"
            assert row["indicador"]
            assert row["periodo_referencia"]
            assert row["arquivo_pdf"] in raw_pdfs
            pdf_name = row["arquivo_pdf"]
            if pdf_name not in page_counts:
                page_counts[pdf_name] = len(PdfReader(str(raw_pdfs[pdf_name])).pages)
            for page in row["pagina_pdf"].split(";"):
                assert 1 <= int(page) <= page_counts[pdf_name]


def test_sinesp_processed_workbooks_have_expected_city_rows():
    from openpyxl import load_workbook

    expected = {2020: 132, 2021: 132, 2022: 132, 2023: 132, 2024: 132, 2025: 144, 2026: 96}
    for year, count in expected.items():
        path = project_path("dados", "processados", "sinesp", f"bancovde-{year}-joinville-sc.xlsx")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            worksheet = workbook.active
            rows = worksheet.iter_rows(values_only=True)
            fields = list(next(rows))
            assert len(fields) == 14
            data = list(rows)
            assert len(data) == count
            assert all(row[0] == "SC" and row[1] == "JOINVILLE" for row in data)
        finally:
            workbook.close()


def test_violence_domestica_normalized_labels_preserve_extracted_values():
    rows = []
    for path in sorted(project_path("dados", "processados", "ssp_sc").glob("ssp_joinville_*.csv")):
        with path.open(encoding="utf-8-sig", newline="") as stream:
            rows.extend(csv.DictReader(stream))
    domestic = [row for row in rows if row["periodo_tipo"] == "mensal"]
    assert domestic
    assert all(row["indicador"] for row in domestic)
    assert all(row["indicador_normalizado"] for row in domestic)
    assert all(row["indicador_normalizado"] == "Lesão corporal leve - Dolosa"
               for row in domestic if row["indicador"] == "Lesão corporal leve 0 Dolosa")
    assert all(row["indicador_normalizado"] == "Lesão corporal grave ou gravíssima - Dolosa"
               for row in domestic if row["indicador"] == "Lesão corporal grave ou gravíssima 0 Dolosa")
    key_counts = {}
    for row in domestic:
        key = (row["ano_referencia"], row["periodo_referencia"], row["indicador_normalizado"])
        key_counts[key] = key_counts.get(key, 0) + 1
    assert max(key_counts.values()) == 1


def test_eda_exports_csv_parquet_and_source_control_totals():
    import pandas as pd

    root = project_path("dados", "processados", "analises")
    for stem in ("sinesp_joinville", "ssp_joinville", "prefeitura_furto_roubo_joinville"):
        csv_path = root / f"{stem}.csv"
        parquet_path = root / f"{stem}.parquet"
        csv_frame = pd.read_csv(csv_path, encoding="utf-8-sig")
        parquet_frame = pd.read_parquet(parquet_path)
        assert len(csv_frame) == len(parquet_frame)
        assert list(csv_frame.columns) == list(parquet_frame.columns)
    homicide = pd.read_csv(project_path("documentacao", "analises", "tabelas", "comparacao_ssp_homicidio_parcial.csv"))
    assert homicide["ano"].tolist() == [2023, 2024, 2025, 2026]
    assert homicide["vitimas_homicidio"].tolist() == [36, 20, 17, 22]
    cross_source = pd.read_csv(project_path("documentacao", "analises", "tabelas", "comparacao_fontes_homicidio.csv"))
    row_2025 = cross_source[cross_source["ano"] == 2025].iloc[0]
    assert row_2025["sinesp_vitimas_homicidio"] == 29
    assert row_2025["ssp_vitimas_homicidio_ano_completo"] == 32
    row_2026 = cross_source[cross_source["ano"] == 2026].iloc[0]
    assert not row_2026["comparacao_anual_valida"]
    assert row_2026["sinesp_periodo"] == "janeiro–agosto"
    assert row_2026["ssp_periodo_parcial"] == "01/01 - 31/07"
    city = pd.read_csv(project_path("documentacao", "analises", "tabelas", "estatisticas_prefeitura_furto_roubo.csv"))
    robbery = city[city["indicador"] == "Roubo"].sort_values("ano")
    assert robbery["quantidade"].tolist() == [933, 593, 476]
    for chart in project_path("documentacao", "analises", "graficos").glob("*.png"):
        assert chart.stat().st_size > 1000


def test_paths_are_root_relative():
    assert project_path("dados", "brutos").is_absolute()
    assert project_path("dados", "brutos").is_dir()
