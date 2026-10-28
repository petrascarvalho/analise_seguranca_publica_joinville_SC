from __future__ import annotations

import argparse
import csv
import re
from datetime import date, datetime
from pathlib import Path
from typing import Iterable
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from lxml import etree

from src.pipeline.utils import project_path


HEADERS = [
    "uf", "municipio", "evento", "data_referencia", "agente", "arma",
    "faixa_etaria", "feminino", "masculino", "nao_informado",
    "total_vitima", "total", "total_peso", "abrangencia",
]
YEARS = tuple(range(2020, 2027))
NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
CITY_CRIME_VALUES = {
    2023: {"Furto": 8865, "Roubo": 933},
    2024: {"Furto": 8730, "Roubo": 593},
    2025: {"Furto": 8372, "Roubo": 476},
}
CITY_CRIME_SOURCE = "Prefeitura de Joinville / SEPROT"
CITY_CRIME_URL = (
    "https://www.joinville.sc.gov.br/noticias/indices-de-criminalidade-"
    "tem-reducao-em-joinville-pelo-terceiro-ano-consecutivo/"
)

SSP_INDICATOR_LABELS = {
    "Lesão corporal grave ou gravíssima 0 Dolosa": "Lesão corporal grave ou gravíssima - Dolosa",
    "Lesão corporal leve 0 Dolosa": "Lesão corporal leve - Dolosa",
}


def normalize_ssp_indicator(value: str) -> str:
    """Canonicaliza apenas duas falhas de extração observadas no texto dos PDFs.

    O campo original ``indicador`` é mantido; esta função fornece o rótulo de
    análise sem alterar os valores ou os documentos-fonte.
    """
    return SSP_INDICATOR_LABELS.get(value, value)


def sinesp_source_path(year: int) -> Path:
    preferred = project_path("dados", "brutos", "sinesp", f"bancovde-{year}.xlsx")
    if preferred.exists():
        return preferred
    pending = project_path(
        "dados", "brutos", "sinesp", "pendente_validacao_origem", f"bancovde-{year}.xlsx"
    )
    if pending.exists():
        return pending
    raise FileNotFoundError(f"XLSX Sinesp de {year} não localizado: {preferred}")


def _excel_row(row: Iterable[object], worksheet) -> list[object]:
    values = list(row)
    if len(values) != len(HEADERS):
        raise ValueError(f"Esperadas {len(HEADERS)} colunas, recebidas {len(values)}")
    formatted = []
    for index, value in enumerate(values):
        if index == 3:
            cell = WriteOnlyCell(worksheet, value=value)
            cell.number_format = "yyyy-mm-dd"
            formatted.append(cell)
        else:
            formatted.append(value)
    return formatted


def _cell_value(cell, shared_strings: list[str]):
    kind = cell.get("t")
    if kind == "inlineStr":
        return "".join(cell.itertext())
    node = cell.find(NS + "v")
    if node is None:
        return None
    raw = node.text
    if raw is None:
        return None
    if kind == "s":
        return shared_strings[int(raw)]
    if kind in ("str", "e"):
        return raw
    if kind == "b":
        return raw == "1"
    number = float(raw)
    return int(number) if number.is_integer() else number


def _parse_sheet_row(raw_row: bytes):
    # Atributos de extensão como x14ac:dyDescent dependem de declarações
    # herdadas do XML da planilha; o fragmento isolado não precisa deles.
    raw_row = re.sub(rb'\s+x14ac:[\w]+="[^"]*"', b"", raw_row)
    return etree.fromstring(
        raw_row.replace(
            b"<row ",
            b'<row xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" ',
            1,
        )
    )


def _shared_strings(archive: ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    root = etree.fromstring(archive.read("xl/sharedStrings.xml"))
    return ["".join(node.itertext()) for node in root.findall(NS + "si")]


def _joinville_rows(source: Path) -> list[list[object]]:
    """Lê apenas as linhas Joinville no XML do XLSX para reduzir tempo e memória."""
    rows = []
    with ZipFile(source) as archive:
        shared_strings = _shared_strings(archive)
        joinville_shared_ids = [
            str(index).encode("ascii")
            for index, value in enumerate(shared_strings)
            if value == "JOINVILLE"
        ]
        xml = archive.read("xl/worksheets/sheet1.xml")
    for match in re.finditer(rb"<row\b[^>]*\br=\"\d+\"[^>]*>.*?</row>", xml, re.DOTALL):
        raw_row = match.group()
        has_inline_city = b"JOINVILLE" in raw_row
        has_shared_city = any(
            re.search(rb'<c\b[^>]*\br="B\d+"[^>]*>.*?<v>' + index + rb"</v>.*?</c>", raw_row, re.DOTALL)
            for index in joinville_shared_ids
        )
        if not has_inline_city and not has_shared_city:
            continue
        element = _parse_sheet_row(raw_row)
        cells = {
            cell.get("r").rstrip("0123456789"): _cell_value(cell, shared_strings)
            for cell in element.findall(NS + "c")
        }
        row = [cells.get(chr(65 + column)) for column in range(len(HEADERS))]
        if row[0] == "SC" and row[1] == "JOINVILLE":
            rows.append(row)
    return rows


def transform_sinesp_year(year: int) -> tuple[Path, int]:
    """Filtra o XLSX nacional sem alterar valores ou ordem das linhas."""
    if year not in YEARS:
        raise ValueError(f"Ano fora da série Sinesp validada: {year}")
    source = sinesp_source_path(year)
    target_dir = project_path("dados", "processados", "sinesp")
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"bancovde-{year}-joinville-sc.xlsx"

    with ZipFile(source) as archive:
        shared_strings = _shared_strings(archive)
        workbook_xml = etree.fromstring(archive.read("xl/workbook.xml"))
        first_sheet_name = workbook_xml.find(
            ".//" + NS + "sheet"
        ).get("name")
        sheet_xml = archive.read("xl/worksheets/sheet1.xml")
        header_match = re.search(rb"<row\b[^>]*\br=\"1\"[^>]*>.*?</row>", sheet_xml, re.DOTALL)
        if header_match is None:
            raise ValueError(f"Cabeçalho ausente em {source.name}")
        header_element = _parse_sheet_row(header_match.group())
        header_cells = {
            cell.get("r").rstrip("0123456789"): _cell_value(cell, shared_strings)
            for cell in header_element.findall(NS + "c")
        }
        header = [header_cells.get(chr(65 + column)) for column in range(len(HEADERS))]
    if header != HEADERS or first_sheet_name not in {str(year), ""}:
        raise ValueError(f"Esquema/aba Sinesp inesperado em {source.name}: {header}, {first_sheet_name}")

    rows = _joinville_rows(source)
    output = Workbook(write_only=True)
    output_sheet = output.create_sheet(str(year))
    output_sheet.freeze_panes = "A2"
    output_sheet.append(HEADERS)
    for row in rows:
        output_sheet.append(_excel_row(row, output_sheet))
    if not rows:
        raise ValueError(f"Nenhuma linha de SC/JOINVILLE encontrada em {source.name}")
    output.save(target)
    count = len(rows)
    return target, count


def export_femicide_csv(workbook_path: Path) -> tuple[Path, int]:
    year = int(workbook_path.stem.split("-")[1])
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        headers = list(next(rows))
        selected = []
        for row in rows:
            if row[0] == "SC" and row[1] == "JOINVILLE" and row[2] in {
                "Feminicídio", "Tentativa de feminicídio"
            }:
                selected.append([
                    value.date().isoformat() if isinstance(value, datetime)
                    else value.isoformat() if isinstance(value, date)
                    else value
                    for value in row
                ])
    finally:
        workbook.close()

    target = workbook_path.parent / f"feminicidio_sinesp_joinville_{year}.csv"
    with target.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(selected)
    return target, len(selected)


def export_municipal_theft_robbery() -> list[tuple[Path, int]]:
    """Transcreve os totais anuais publicados pela Prefeitura, sem inferir 2026."""
    output_dir = project_path("dados", "processados", "prefeitura")
    output_dir.mkdir(parents=True, exist_ok=True)
    results = []
    for year, indicators in CITY_CRIME_VALUES.items():
        target = output_dir / f"furto_roubo_joinville_{year}.csv"
        with target.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream)
            writer.writerow([
                "uf", "municipio", "ano", "indicador", "periodo",
                "quantidade", "fonte", "url_fonte",
            ])
            for indicator, value in indicators.items():
                writer.writerow([
                    "SC", "JOINVILLE", year, indicator, "janeiro-dezembro",
                    value, CITY_CRIME_SOURCE, CITY_CRIME_URL,
                ])
        results.append((target, len(indicators)))
    return results


def run_sinesp(years: Iterable[int] = YEARS) -> list[tuple[Path, int]]:
    outputs = []
    for year in years:
        workbook_path, count = transform_sinesp_year(year)
        outputs.append((workbook_path, count))
        csv_path, csv_count = export_femicide_csv(workbook_path)
        outputs.append((csv_path, csv_count))
    return outputs


def main() -> None:
    parser = argparse.ArgumentParser(description="Transformações locais do projeto.")
    parser.add_argument("--sinesp", action="store_true", help="Filtra Sinesp 2020-2026")
    parser.add_argument("--furto-roubo", action="store_true", help="Exporta a série da Prefeitura")
    args = parser.parse_args()
    if not args.sinesp and not args.furto_roubo:
        parser.error("Informe --sinesp, --furto-roubo ou ambos")
    if args.sinesp:
        for path, count in run_sinesp():
            print(f"{path.relative_to(project_path())}: {count} registros")
    if args.furto_roubo:
        for path, count in export_municipal_theft_robbery():
            print(f"{path.relative_to(project_path())}: {count} registros")


if __name__ == "__main__":
    main()
